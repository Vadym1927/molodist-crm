"""Фінанси → Планування ціни реалізації: курс, валютна закупівля, РРЦ, рекомендована ціна, Excel."""
from fastapi.testclient import TestClient

from app.config import settings
from app.main import app
from app.models import SessionLocal, Product
from app.services import pricing as pr


def test_pricing_flow(monkeypatch):
    settings.admin_password = "owner-pass-1"
    with SessionLocal() as db:
        p = Product(name="Medik8 C-Tetra 30 мл", name_norm="medik8 c-tetra", brand="MediK8", warehouse="vitrina", price=2600)
        db.add(p); db.commit(); pid = p.id
    assert pr.round_price(1234, 50) == 1250 and pr.margin(2000, 1200) == 40.0 and pr.markup(2000, 1000) == 100.0
    assert round(pr.price_for_margin(1000, 60)) == 2500
    with TestClient(app) as c:
        c.post("/login", data={"login": settings.owner_login, "password": "owner-pass-1"})
        c.post("/finance/rates", data={"usd": "41", "eur": "48,5", "round": "10"})
        monkeypatch.setattr(pr, "nbu_rates", lambda: {"USD": 41.2, "EUR": 47.9})
        c.post("/finance/rates", data={"nbu": "1"})
        with SessionLocal() as db:
            rt = pr.rates(db)
            assert rt["EUR"] == 47.9 and rt["source"] == "НБУ" and rt["round"] == 10
        page = c.get("/finance/pricing?brand=MediK8").text
        assert "Medik8 C-Tetra" in page and "Ціна при маржі 47%" in page and "Рекомендована ціна продажу" in page
        r = c.post("/finance/pricing", data={"pid": [str(pid)], f"cur_{pid}": "EUR", f"cfx_{pid}": "25", f"rrp_{pid}": "60",
                                             f"plan_{pid}": "2900"}, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"].endswith(":+1")
        with SessionLocal() as db:
            p = db.get(Product, pid)
            assert (p.cost_currency, p.cost_fx, p.rrp_fx, p.planned_price) == ("EUR", 25, 60, 2900) and p.planned_at
        x = c.get("/finance/pricing.xlsx?brand=MediK8")
        assert x.status_code == 200 and x.content[:2] == b"PK"
        from openpyxl import load_workbook
        import io
        ws = load_workbook(io.BytesIO(x.content)).active
        row = [cell.value for cell in ws[3]]
        assert row[1] == "Medik8 C-Tetra 30 мл" and row[6] == round(25 * 47.9, 2) and row[11] == 2870  # 60×47.9 → до 10
        assert "Планування ціни реалізації" in c.get("/finance").text
