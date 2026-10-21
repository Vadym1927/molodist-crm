"""Експорт, імпорт, перенос з CleverBox, резервна копія."""
import io
import zipfile
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.config import settings
from app.main import app
from app.models import SessionLocal, Patient, BeautyPlan, Visit, ProductSale
from app.services.kpi import compute_plan

PDF = Path(__file__).parent / "sample_plan.pdf"
pytestmark = pytest.mark.skipif(not PDF.exists(), reason="потрібен tests/sample_plan.pdf")


def login(c, lg, pw):
    assert c.post("/login", data={"login": lg, "password": pw}, follow_redirects=False).status_code == 303


@pytest.fixture
def c(tmp_path):
    settings.admin_password = "owner-pass-1"
    settings.files_dir = str(tmp_path / "files")
    with TestClient(app) as client:
        login(client, settings.owner_login, "owner-pass-1")
        yield client


def sheets(resp):
    return load_workbook(io.BytesIO(resp.content)).sheetnames


def test_export_import_backup(c):
    c.post("/upload", files={"files": ("plan.pdf", PDF.read_bytes(), "application/pdf")})
    c.post(f"/patients/1/photos", files={"files": ("a.jpg", b"\xff\xd8fakejpeg", "image/jpeg")}, data={"caption": "до"})

    # --- експорт
    r = c.get("/export/report.xlsx")
    assert r.status_code == 200 and sheets(r) == ["Плани", "Позиції", "KPI лікарів", "Прогноз"]
    z = zipfile.ZipFile(io.BytesIO(c.get("/export/report.csv.zip").content))
    assert "Плани.csv" in z.namelist() and z.read("Плани.csv").decode("utf-8-sig").startswith("ID плану;")
    files = zipfile.ZipFile(io.BytesIO(c.get("/export/files.zip").content)).namelist()
    assert any(n.startswith("Олена Даниленко/") and n.endswith(".jpg") for n in files)
    assert "Друк / Зберегти PDF" in c.get("/patients/1/print").text

    # адмін: без прогнозу і KPI
    c.post("/access/users", data={"login": "adm", "name": "Адмін", "role": "admin", "password": "password-123", "active": "1"})
    c.get("/logout"); login(c, "adm", "password-123")
    assert sheets(c.get("/export/report.xlsx")) == ["Плани", "Позиції"]
    assert c.get("/backup/download").status_code == 403
    c.get("/logout"); login(c, settings.owner_login, "owner-pass-1")

    # --- імпорт шаблону (3 аркуші: пацієнти, плани, історія)
    tpl = c.get("/import/template.xlsx").content
    page = c.post("/import/table", files={"file": ("tpl.xlsx", tpl, "application/octet-stream")}).text
    assert "Результат імпорту" in page and "бюті-плани" in page
    with SessionLocal() as db:
        assert db.query(Patient).filter(Patient.full_name == "Приклад Пацієнтка").count() == 1   # без дублів між аркушами
        imp = db.query(BeautyPlan).filter(BeautyPlan.title == "Імпорт").one()
        ps = compute_plan(db, imp)
        assert ps.proc_done == 1 and ps.prod_bought == 1          # історія з шаблону зарахувалась
    # повторний імпорт не дублює історію/план
    c.post("/import/table", files={"file": ("tpl.xlsx", tpl, "application/octet-stream")})
    with SessionLocal() as db:
        assert db.query(BeautyPlan).filter(BeautyPlan.title == "Імпорт").count() == 1
        n_visits = db.query(Visit).filter(Visit.source == "import").count()
        assert n_visits == 1

    # CSV з іншої CRM: cp1251, «;», російські заголовки
    csv_bytes = "ФИО;Телефон;Дата;Услуга;Статус;Сумма\nИрина Тест;0939998877;01.09.2026 10:00;Чистка лица;Выполнен;1500\n" \
                "Ирина Тест;0939998877;05.09.2026;Крем SPF 50 мл;оплачено;900\n".encode("cp1251")
    c.post("/import/table", files={"file": ("old_crm.csv", csv_bytes, "text/csv")}, data={"kind": "history"})
    with SessionLocal() as db:
        ir = db.query(Patient).filter(Patient.phone == "0939998877").one()
        assert len(ir.visits) == 1 and ir.visits[0].status == "done" and len(ir.sales) == 1

    # ZIP з PDF
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zz:
        zz.writestr("plans/a.pdf", PDF.read_bytes())
    assert "PDF Observ" in c.post("/import/pdf-zip", files={"file": ("plans.zip", buf.getvalue(), "application/zip")}).text

    # --- резервна копія → зміни → відновлення
    backup = c.get("/backup/download").content
    with SessionLocal() as db:
        n_pat = db.query(Patient).count()
        db.query(ProductSale).delete(); db.commit()
        db.add(Patient(full_name="Зайвий")); db.commit()
    assert c.post("/backup/restore", files={"file": ("b.zip", backup, "application/zip")},
                  data={"confirm": "ні"}).status_code == 400
    r = c.post("/backup/restore", files={"file": ("b.zip", backup, "application/zip")},
               data={"confirm": "ВІДНОВИТИ"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    with SessionLocal() as db:
        assert db.query(Patient).count() == n_pat and db.query(ProductSale).count() >= 1
        assert not db.query(Patient).filter(Patient.full_name == "Зайвий").count()
    login(c, settings.owner_login, "owner-pass-1")
    assert c.get("/patients/1/files.zip").status_code == 200


def test_migrate_cleverbox(monkeypatch):
    from app.integrations import cleverbox as cbm
    from tests.test_cleverbox_sync import handler
    settings.cleverbox_token, settings.cleverbox_min_interval_sec = "test", 0
    orig = cbm.CleverBox.__init__

    def init(self, *a, **k):
        orig(self, *a, **k)
        self.http = httpx.Client(transport=httpx.MockTransport(handler), headers={"token": "test"})
    monkeypatch.setattr(cbm.CleverBox, "__init__", init)

    class J:
        def log(self, m): pass
    rep = cbm.migrate(J(), months=1, all_clients=True)
    assert rep["patients_created"] == 1 and rep["services"] == 2
    with SessionLocal() as db:
        p = db.query(Patient).filter(Patient.cleverbox_id == 777).one()
        assert p.phone == "0671112233" and len(p.visits) >= 1 and len(p.sales) == 1
