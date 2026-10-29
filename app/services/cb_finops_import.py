"""Журнал фіноперацій CleverBox → CbFinOp: імпорт Excel-вивантаження («Фінанси → Фіноперації → Інструменти → Excel»)
і рядки з API (якщо CleverBox віддає метод фіноперацій). Повторний імпорт оновлює операції за № (без дублів)."""
from __future__ import annotations

import re
import zlib
from datetime import datetime, timedelta

from sqlalchemy.orm import Session

from ..models import CbFinOp

KIND_WORDS = (("сторно", "storno"), ("перемiщ", "transfer"), ("переміщ", "transfer"), ("перемещ", "transfer"),
              ("витрат", "expense"), ("расход", "expense"), ("дохід", "income"), ("доход", "income"), ("прихід", "income"))
COLS = {"date": ("дата",), "time": ("час", "время"), "type": ("тип",), "num": ("n", "№", "номер", "id", "№ документа"),
        "who": ("від кого/ кому", "від кого/кому", "від кого / кому", "від кого", "контрагент", "клієнт / постачальник"),
        "basis": ("підстава", "основание"), "article": ("стаття", "статья"), "filial": ("філія", "филиал", "локація"),
        "account": ("рахунок", "счет", "каса"), "vid": ("вид",), "author": ("автор",), "amount": ("сума", "сумма"),
        "comment": ("коментар", "комментарий", "примітка")}


def _read(path: str) -> list[list]:
    if path.lower().endswith(".xls"):
        import xlrd
        bk = xlrd.open_workbook(path, ignore_workbook_corruption=True)
        sh = bk.sheet_by_index(0)
        rows = []
        for i in range(sh.nrows):
            r = []
            for j, c in enumerate(sh.row(i)):
                if c.ctype == xlrd.XL_CELL_DATE:
                    r.append(xlrd.xldate.xldate_as_datetime(c.value, bk.datemode))
                else:
                    r.append(c.value)
            rows.append(r)
        return rows
    from openpyxl import load_workbook
    ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower())


def is_journal(path: str) -> bool:
    try:
        rows = _read(path)[:12]
    except Exception:  # noqa: BLE001
        return False
    return any({"стаття", "сума"} <= {_norm(x) for x in r} or ({"статья", "сумма"} <= {_norm(x) for x in r}) for r in rows)


def _dt(v, t=None) -> datetime | None:
    if isinstance(v, datetime):
        d = v
    elif isinstance(v, (int, float)) and v > 20000:
        d = datetime(1899, 12, 30) + timedelta(days=float(v))
    else:
        s = str(v or "").strip()
        d = None
        for f in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                d = datetime.strptime(s[:19], f)
                break
            except ValueError:
                continue
        if not d:
            m = re.search(r"(\d{2})\.(\d{2})\.(\d{4})(?:\s+(\d{1,2}):(\d{2})(?::(\d{2}))?)?", s)
            if not m:
                return None
            d = datetime(int(m[3]), int(m[2]), int(m[1]), int(m[4] or 0), int(m[5] or 0), int(m[6] or 0))
    if t not in (None, "") and d.hour == 0 and d.minute == 0:
        m = re.match(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", str(t))
        if m:
            d = d.replace(hour=int(m[1]), minute=int(m[2]), second=int(m[3] or 0))
    return d


def _num(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v or "")).replace(",", ".")
    if s.count(".") > 1:
        s = s.replace(".", "", s.count(".") - 1)
    try:
        return float(s)
    except ValueError:
        return None


def kind_of(text: str, amount: float | None = None) -> str:
    t = (text or "").lower()
    for w, k in KIND_WORDS:
        if w in t:
            return k
    return "expense" if (amount or 0) < 0 else "income"


def upsert(db: Session, rec: dict, source: str) -> bool:
    """Одна операція (з API). True — нова."""
    oid = _key(rec)
    x = db.get(CbFinOp, oid)
    new = x is None
    if new:
        x = CbFinOp(id=oid)
        db.add(x)
    _fill(x, rec, source)
    return new


def parse_file(path: str) -> tuple[list[dict], int]:
    rows = _read(path)
    hi = next((i for i, r in enumerate(rows[:12]) if {"сума", "стаття"} <= {_norm(x) for x in r}
               or {"сумма", "статья"} <= {_norm(x) for x in r}), None)
    if hi is None:
        raise ValueError("Це не журнал фіноперацій CleverBox: потрібні колонки «Стаття» і «Сума».")
    head = [_norm(h) for h in rows[hi]]
    idx = {}
    for key, names in COLS.items():
        for n in names:
            if n in head:
                idx[key] = head.index(n)
                break
    get = lambda r, k: r[idx[k]] if k in idx and idx[k] < len(r) else None   # noqa: E731
    out, skipped = [], 0
    for r in rows[hi + 1:]:
        if not any(v not in (None, "") for v in r):
            continue
        at = _dt(get(r, "date"), get(r, "time"))
        amt = _num(get(r, "amount"))
        if not at or amt is None or amt == 0:
            skipped += 1
            continue
        ttxt = str(get(r, "type") or "")
        num = get(r, "num")
        m = re.search(r"#\s*(\d+)", ttxt) or re.search(r"(\d{3,})", str(num or ""))
        who = str(get(r, "who") or "")
        basis = get(r, "basis")
        if not basis and "підстава" in who.lower():
            parts = re.split(r"(?i)підстава:?", who, maxsplit=1)
            who, basis = parts[0], parts[1]
        out.append({"at": at, "kind": kind_of(ttxt, amt), "amount": amt, "num": int(m[1]) if m else None, "who": who.strip(),
                    "basis": (str(basis).strip() if basis else None), "article": get(r, "article"), "filial": get(r, "filial"),
                    "account": get(r, "account"), "vid": get(r, "vid"), "author": get(r, "author"), "comment": get(r, "comment")})
    return out, skipped


FIELDS = (("who", "counterparty"), ("basis", "basis"), ("article", "article"), ("filial", "filial"),
          ("account", "account"), ("vid", "vid"), ("author", "author"), ("comment", "comment"))


def _key(rec: dict) -> int:
    if rec.get("num"):
        return int(rec["num"])
    key = f"{rec['at']:%Y%m%d%H%M%S}|{rec.get('amount')}|{rec.get('account')}|{rec.get('article')}|{rec.get('who')}"
    return -(zlib.crc32(key.encode()) & 0x7FFFFFFF)


def _fill(x: CbFinOp, rec: dict, source: str) -> None:
    x.at, x.kind, x.amount = rec["at"], rec["kind"], round(abs(rec.get("amount") or 0), 2)
    for k, attr in FIELDS:
        v = rec.get(k)
        setattr(x, attr, (str(v).strip()[:300] if v not in (None, "") else None))
    x.source, x.updated_at = source, datetime.utcnow()


def import_file(db: Session, path: str, job=None) -> dict:
    """Швидкий імпорт: існуючі операції вантажаться пачками, нові додаються пачками по 2000."""
    from sqlalchemy import select
    log = job.log if job else (lambda m: None)
    recs, skipped = parse_file(path)
    log(f"Прочитано {len(recs)} операцій, обробляю…")
    by_id: dict[int, dict] = {}
    for r in recs:
        by_id[_key(r)] = r                                   # дублікати в самому файлі — останній рядок
    ids = list(by_id)
    created = updated = 0
    for i in range(0, len(ids), 2000):
        chunk = ids[i:i + 2000]
        have = {x.id: x for x in db.scalars(select(CbFinOp).where(CbFinOp.id.in_(chunk)))}
        for oid in chunk:
            x = have.get(oid)
            if x is None:
                x = CbFinOp(id=oid)
                db.add(x)
                created += 1
            else:
                updated += 1
            _fill(x, by_id[oid], "import")
        db.commit()
        log(f"Збережено {min(i + 2000, len(ids))} з {len(ids)}")
    return {"created": created, "updated": updated, "skipped": skipped, "rows": len(recs)}


def from_api_row(r: dict) -> dict | None:
    """Рядок методу фіноперацій API CleverBox → rec (назви полів невідомі заздалегідь — беремо перші наявні)."""
    from ..integrations.cb_finance import AMOUNT_KEYS, DATE_KEYS, TYPE_KEYS, CASHBOX_KEYS, _first, _flat, _text, classify
    from ..integrations.cleverbox import parse_dt
    at = parse_dt(str(_first(r, DATE_KEYS) or ""))
    c = classify(r)
    if not at or not c:
        return None
    k, amt, _ = c
    num = _first(r, ("id", "operation_id", "fin_id", "number", "num"))
    pick = lambda *ks: _flat(_first(r, ks))                                         # noqa: E731
    return {"at": at, "kind": {"in": "income", "out": "expense", "transfer": "transfer"}[k], "amount": amt,
            "num": int(num) if str(num or "").isdigit() else None,
            "who": pick("client_name", "supplier_name", "counterparty", "contragent", "staff_name", "who", "name"),
            "basis": pick("basis", "reason", "document", "doc_name", "sale_id", "check_id"),
            "article": pick("article_name", "article", "category_name", "category", "type_name"),
            "filial": pick("filial_name", "filial", "salon_name"), "account": _text(r, CASHBOX_KEYS)[:200] or None,
            "vid": pick("payment_name", "pay_type", "cashbox_to_name", "to_cashbox"), "author": pick("author_name", "author", "user_name"),
            "comment": pick("comment", "note", "description")}
