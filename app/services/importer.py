"""Імпорт даних з Excel / CSV (вивантаження з будь-якої CRM, таблиці, старі звіти) і пакетів PDF.

Три типи таблиць (тип визначається автоматично за колонками або обирається вручну):
  patients — довідник пацієнтів: ПІБ, телефон, лікар, ID у CleverBox, нотатки
  plans    — бюті-плани построчно: пацієнт, [дата], [лікар], позиція, [тип], кількість, ціна
             (рядки з однаковим пацієнтом+датою або «№ плану» збираються в один план)
  history  — історія візитів і покупок: пацієнт, дата, послуга/товар, [статус], [сума], [тип]
             (потрібна, щоб перенести реалізацію з іншої CRM — CRM порахує, що вже виконано)
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from datetime import datetime, date
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Patient, Visit, ProductSale
from ..parsers.observ import ParsedPlan, PlanItem, norm
from .plans import import_plan, find_patient, get_or_create_doctor

# поле -> синоніми заголовків (порівняння без регістру/пунктуації; підтримуються вивантаження з різних CRM)
FIELDS: dict[str, list[str]] = {
    "patient":      ["пацієнт", "пацієнтка", "клієнт", "піб", "імя", "ім я", "прізвище імя", "name", "full name",
                     "client", "client name", "patient", "фио", "клиент", "имя", "full_name", "client_name"],
    "phone":        ["телефон", "тел", "мобільний", "phone", "mobile", "номер телефону", "телефон клієнта"],
    "email":        ["email", "e mail", "пошта", "електронна пошта"],
    "doctor":       ["лікар", "спеціаліст", "майстер", "косметолог", "doctor", "master", "specialist", "staff",
                     "врач", "мастер", "співробітник", "виконавець", "user_name", "staff_name"],
    "date":         ["дата", "date", "дата візиту", "дата плану", "дата консультації", "дата продажу", "дата і час",
                     "date_start", "datetime", "дата время"],
    "kind":         ["тип", "kind", "type", "вид", "item_type"],
    "item":         ["позиція", "послуга", "процедура", "товар", "засіб", "service", "product", "item",
                     "назва", "найменування", "услуга", "service_name", "product_name", "item_name"],
    "qty":          ["кількість", "к сть", "кол во", "qty", "quantity", "количество", "шт"],
    "price":        ["ціна", "price", "вартість", "цена", "ціна за од", "unit_price"],
    "amount":       ["сума", "total", "amount", "сумма", "оплачено", "до сплати", "total_amount", "payd_amount"],
    "status":       ["статус", "status", "стан", "_status"],
    "cleverbox_id": ["id клієнта", "client_id", "cleverbox_id", "id cleverbox", "код клієнта"],
    "plan_key":     ["план", "№ плану", "номер плану", "plan", "plan_id", "plan id"],
    "notes":        ["коментар", "нотатки", "примітка", "comment", "notes", "note", "примечание", "комментарий"],
    "interval":     ["інтервал", "інтервал днів", "interval_days"],
}

PRODUCT_HINT = re.compile(r"\b(\d+\s*(мл|ml|г|g))\b|крем|гель|тонік|tonic|сироватк|serum|spf|спф|маска|пінка|"
                          r"шампун|лосьйон|cream|gel|патч|товар|засіб", re.I)
DONE_RE = re.compile(r"викон|заверш|оплач|прийш|надано|done|paid|complete|выполн|пришел", re.I)
CANCEL_RE = re.compile(r"скас|відмін|не\s*прийш|cancel|отмен|не\s*пришел", re.I)


def _key(h) -> str:
    return re.sub(r"[^\wʼ']+", " ", str(h or "")).replace("ʼ", " ").replace("'", " ").strip().lower()


SHEET_TYPES = {"пацієнти": "patients", "patients": "patients", "клієнти": "patients", "clients": "patients",
               "плани": "plans", "plans": "plans", "бюті плани": "plans",
               "історія": "history", "history": "history", "візити": "history", "продажі": "history"}


def read_sheets(data: bytes, filename: str) -> list[tuple[str, list[list]]]:
    """[(назва аркуша, рядки)]; для CSV — один «аркуш»."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
            if len(rows) >= 2:
                out.append((ws.title, rows))
        return out
    text = None
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Не вдалося прочитати CSV (кодування)")
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
    return [(Path(filename).stem, [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)])]


def sheet_type(title: str) -> str | None:
    return SHEET_TYPES.get(_key(title))


def detect_columns(header: list) -> dict[str, int]:
    found: dict[str, int] = {}
    keys = [_key(h) for h in header]
    for field, syns in FIELDS.items():
        syn_keys = [_key(s) for s in syns]
        for i, k in enumerate(keys):
            if i in found.values():
                continue
            if k in syn_keys:
                found[field] = i
                break
        else:
            for i, k in enumerate(keys):          # часткове входження («Дата візиту клієнта»)
                if i not in found.values() and k and any(s and (s == k or f" {s} " in f" {k} ") for s in syn_keys):
                    found[field] = i
                    break
    return found


def guess_type(cols: dict) -> str:
    if "item" in cols and ("qty" in cols or "price" in cols) and "status" not in cols:
        return "plans"
    if "item" in cols and "date" in cols:
        return "history"
    return "patients"


def _val(row, cols, f):
    i = cols.get(f)
    if i is None or i >= len(row):
        return None
    v = row[i]
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _num(v, default=0.0):
    if v in (None, ""):
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v)).replace(",", ".")
    if s.count(".") > 1:
        s = s.replace(".", "", s.count(".") - 1)
    try:
        return float(s)
    except ValueError:
        return default


def _date(v) -> datetime | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = str(v).strip()
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
                "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s[:19], fmt)
        except ValueError:
            continue
    return None


def _kind(v, item: str) -> str:
    if v:
        k = str(v).lower()
        if any(x in k for x in ("товар", "засіб", "догляд", "product", "goods", "косметика")):
            return "product"
        if any(x in k for x in ("послуга", "процедура", "service", "procedure", "услуга")):
            return "procedure"
    return "product" if PRODUCT_HINT.search(item or "") else "procedure"


def _patient(db: Session, row, cols, create: bool = True) -> Patient | None:
    name, phone = _val(row, cols, "patient"), _val(row, cols, "phone")
    cb = _val(row, cols, "cleverbox_id")
    cb = int(_num(cb)) if cb not in (None, "") and _num(cb, -1) > 0 else None
    p = db.scalar(select(Patient).where(Patient.cleverbox_id == cb)) if cb else None
    if not p and (name or phone):
        p = find_patient(db, str(name or ""), str(phone) if phone else None)
    if not p and create and name:
        p = Patient(full_name=str(name))
        db.add(p)
        db.flush()
    if p:
        if phone and not p.phone:
            p.phone = str(phone)
        if cb and not p.cleverbox_id and not db.scalar(select(Patient).where(Patient.cleverbox_id == cb)):
            p.cleverbox_id = cb
    return p


def import_rows(db: Session, rows: list[list], kind: str | None = None, cols: dict | None = None) -> dict:
    if len(rows) < 2:
        raise ValueError("Таблиця порожня")
    header, body = rows[0], rows[1:]
    cols = cols or detect_columns(header)
    kind = kind or guess_type(cols)
    if "patient" not in cols and "phone" not in cols and "cleverbox_id" not in cols:
        raise ValueError("Не знайшов колонку з пацієнтом (ПІБ / телефон / ID клієнта)")
    rep = {"type": kind, "rows": len(body), "patients_created": 0, "plans": 0, "visits": 0, "sales": 0,
           "skipped": 0, "errors": []}
    before = db.query(Patient).count()

    if kind == "patients":
        for n, r in enumerate(body, 2):
            p = _patient(db, r, cols)
            if not p:
                rep["skipped"] += 1
                continue
            doc = _val(r, cols, "doctor")
            if doc and not p.doctor_id:
                p.doctor_id = get_or_create_doctor(db, str(doc)).id
            note = _val(r, cols, "notes")
            if note and not p.notes:
                p.notes = str(note)

    elif kind == "plans":
        groups: dict[tuple, dict] = {}
        for n, r in enumerate(body, 2):
            item = _val(r, cols, "item")
            p = _patient(db, r, cols)
            if not p or not item:
                rep["skipped"] += 1
                continue
            dt = _date(_val(r, cols, "date"))
            key = (p.id, str(_val(r, cols, "plan_key") or (dt.date() if dt else "")))
            g = groups.setdefault(key, {"patient": p, "date": dt, "doctor": _val(r, cols, "doctor"), "items": []})
            qty = int(_num(_val(r, cols, "qty"), 1)) or 1
            price = _num(_val(r, cols, "price"))
            if not price and _val(r, cols, "amount") is not None:
                price = _num(_val(r, cols, "amount")) / qty
            g["items"].append(PlanItem(kind=_kind(_val(r, cols, "kind"), str(item)), name=str(item), qty=qty,
                                       unit_price=round(price, 2), total=round(price * qty, 2)))
        for g in groups.values():
            parsed = ParsedPlan(patient_name=g["patient"].full_name, doctor_name=str(g["doctor"]) if g["doctor"] else None,
                                session_at=g["date"] or datetime.now(), complaints=None, comments=None,
                                routine_comment=None, items=g["items"], total=sum(i.total for i in g["items"]),
                                plan_title="Імпорт")
            import_plan(db, parsed, source_file="import", patient=g["patient"])
            rep["plans"] += 1

    elif kind == "history":
        for n, r in enumerate(body, 2):
            item, dt = _val(r, cols, "item"), _date(_val(r, cols, "date"))
            p = _patient(db, r, cols)
            if not p or not item or not dt:
                rep["skipped"] += 1
                if len(rep["errors"]) < 20:
                    rep["errors"].append(f"рядок {n}: немає пацієнта, назви або дати")
                continue
            status_raw = str(_val(r, cols, "status") or "")
            if CANCEL_RE.search(status_raw):
                rep["skipped"] += 1
                continue
            ext = hashlib.sha1(f"{p.id}|{dt.isoformat()}|{item}|{n}".encode()).hexdigest()[:20]
            amount = _num(_val(r, cols, "amount") or _val(r, cols, "price"), 0) or None
            if _kind(_val(r, cols, "kind"), str(item)) == "product":
                if not db.scalar(select(ProductSale).where(ProductSale.external_id == "imp-" + ext)):
                    db.add(ProductSale(external_id="imp-" + ext, patient_id=p.id, product_name=str(item),
                                       quantity=int(_num(_val(r, cols, "qty"), 1)) or 1, amount=amount or 0, sold_at=dt))
                    rep["sales"] += 1
            else:
                status = "done" if (DONE_RE.search(status_raw) or (not status_raw and dt <= datetime.now())) else "booked"
                if not db.scalar(select(Visit).where(Visit.source == "import", Visit.external_id == ext)):
                    db.add(Visit(source="import", external_id=ext, patient_id=p.id, service_name=str(item),
                                 master_name=str(_val(r, cols, "doctor") or "") or None, starts_at=dt, status=status,
                                 status_label=status_raw or None, amount=amount))
                    rep["visits"] += 1
    else:
        raise ValueError(f"Невідомий тип імпорту: {kind}")

    db.commit()
    rep["patients_created"] = db.query(Patient).count() - before
    rep["columns"] = {f: str(header[i]) for f, i in cols.items()}
    return rep


def import_file(db: Session, data: bytes, filename: str, kind: str | None = None) -> list[dict]:
    """Імпорт усього файлу: кожен аркуш xlsx окремо (тип — з назви аркуша, з параметра або автоматично)."""
    reports = []
    for title, rows in read_sheets(data, filename):
        k = kind or sheet_type(title)
        try:
            rep = import_rows(db, rows, k)
        except ValueError as e:
            rep = {"type": k or "?", "error": str(e)}
            db.rollback()
        rep["sheet"] = title
        reports.append(rep)
    if not reports:
        raise ValueError("У файлі немає даних")
    return reports


def import_pdf_zip(db: Session, data: bytes) -> dict:
    from .ingest import ingest_pdf
    rep = {"pdf": 0, "plans": 0, "errors": []}
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        for info in z.infolist():
            if info.is_dir() or not info.filename.lower().endswith(".pdf"):
                continue
            rep["pdf"] += 1
            plan, _, err = ingest_pdf(db, z.read(info), Path(info.filename).name)
            if plan:
                rep["plans"] += 1
            else:
                rep["errors"].append(f"{info.filename}: {err}")
    return rep


def template_xlsx() -> bytes:
    from .exporter import to_xlsx
    return to_xlsx({
        "Пацієнти": [["Пацієнт", "Телефон", "Лікар", "ID клієнта", "Коментар"],
                     ["Приклад Пацієнтка", "0670000000", "Дарʼя Колган", "", ""]],
        "Плани": [["Пацієнт", "Телефон", "Дата", "Лікар", "Тип", "Позиція", "Кількість", "Ціна"],
                  ["Приклад Пацієнтка", "0670000000", "10.10.2026", "Дарʼя Колган", "процедура", "IPL-терапія обличчя", 4, 3900],
                  ["Приклад Пацієнтка", "0670000000", "10.10.2026", "Дарʼя Колган", "догляд", "AG SKIN Azelorosa Rosy Cream Gel 50 мл", 1, 1700]],
        "Історія": [["Пацієнт", "Телефон", "Дата", "Послуга", "Тип", "Статус", "Сума", "Лікар"],
                    ["Приклад Пацієнтка", "0670000000", "13.10.2026 12:00", "IPL-терапія обличчя", "процедура", "виконано", 3900, "Дарʼя Колган"],
                    ["Приклад Пацієнтка", "0670000000", "10.10.2026", "AG SKIN Azelorosa Rosy Cream Gel 50 мл", "товар", "оплачено", 1700, ""]],
    })
