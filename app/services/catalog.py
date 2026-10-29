"""Розділи «Продукти» і «Склад»: каталог засобів, конверсія «призначено → куплено», потреба за планами vs залишки."""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Product, PlanItem, StockLevel, BeautyPlan
from ..parsers.observ import norm
from .kpi import name_match, all_statuses, PlanStatus


# Склади (розділи) товарів. Ключі стабільні, назви — для інтерфейсу.
WAREHOUSES = {"vitrina": "Вітрина", "inject": "Ін'єкційна", "classic": "Класична", "single": "Одноразка", "equip": "Обладнання"}


def warehouse_of(category: str | None) -> str | None:
    """Категорія / напрямок CleverBox → склад CRM."""
    c = (category or "").lower()
    if not c:
        return None
    if "вітрин" in c or "витрин" in c or "магазин" in c:
        return "vitrina"
    if "ін'єкц" in c or "інєкц" in c or "ін’єкц" in c or "инъек" in c or "inject" in c:
        return "inject"
    if "класич" in c or "классич" in c:
        return "classic"
    if "однораз" in c or "расход" in c or "витрат" in c:
        return "single"
    if "обладн" in c or "апарат" in c or "оборуд" in c:
        return "equip"
    return None


def home_care(products):
    """Лише засоби для продажу (вітрина) — з ними зіставляються засоби з бюті-планів."""
    return [p for p in products if p.warehouse in (None, "vitrina") and p.for_sale is not False]


def match_product(db: Session, name: str, products: list[Product] | None = None) -> Product | None:
    products = products if products is not None else db.scalars(select(Product)).all()
    key = norm(name)
    best, score = None, 0.0
    for p in products:
        if p.name_norm == key:
            return p
        s = name_match(name, p.name)
        if s > score:
            best, score = p, s
    return best if score >= 0.7 else None


def link_item(db: Session, it: PlanItem, products: list[Product]) -> Product:
    prod = match_product(db, it.name, home_care(products))
    if not prod:                                        # засіб є в плані, але ще немає в каталозі
        prod = Product(name=it.name, name_norm=norm(it.name), volume=it.volume, price=it.unit_price, url=it.url)
        db.add(prod)
        db.flush()
        products.append(prod)
    prod.volume = prod.volume or it.volume
    prod.url = prod.url or it.url
    it.url = it.url or prod.url                          # посилання з каталогу — у нагадування і скрипти
    if not prod.price:
        prod.price = it.unit_price
    it.product_id = prod.id
    return prod


def link_plan(db: Session, plan: BeautyPlan):
    products = db.scalars(select(Product)).all()
    for it in plan.items:
        if it.kind == "product" and not it.product_id:
            prod = link_item(db, it, list(products))
            if prod.supply_days:                        # тривалість засобу, задана в каталозі, має пріоритет
                it.supply_days = prod.supply_days


def link_all_items(db: Session):
    products = list(db.scalars(select(Product)))
    for it in db.scalars(select(PlanItem).where(PlanItem.kind == "product", PlanItem.product_id.is_(None))):
        link_item(db, it, products)
    # засоби, створені з плану, які тепер зіставились з товаром CleverBox — зливаємо
    cb = [p for p in home_care(products) if p.cb_product_id]
    for p in [p for p in products if not p.cb_product_id]:
        twin = match_product(db, p.name, cb)
        if twin:
            for it in db.scalars(select(PlanItem).where(PlanItem.product_id == p.id)):
                it.product_id = twin.id
            twin.supply_days = twin.supply_days or p.supply_days
            twin.url = twin.url or p.url
            twin.volume = twin.volume or p.volume
            db.delete(p)
    db.commit()


@dataclass
class ProductRow:
    product: Product
    prescribed: int = 0           # призначено в планах (шт)
    bought: int = 0               # куплено пацієнтами з планів
    revenue: float = 0.0
    demand_open: int = 0          # призначено в активних планах, але ще не куплено
    rebuy_30: int = 0             # повторні покупки, що «настають» у найближчі 30 днів
    patients: set = field(default_factory=set)
    stock: float | None = None
    cost_avg: float | None = None      # середня ціна закупівлі (зважена за кількістю)
    cost_last: float | None = None     # ціна останньої закупівлі
    last_bought: date | None = None
    last_supplier: str | None = None
    expiry: date | None = None         # найближчий термін придатності
    purchases: int = 0
    sold_90: float = 0.0               # продано за 90 днів (усім пацієнтам, з CleverBox)
    sold_90_sum: float = 0.0

    @property
    def days_left(self):
        """На скільки днів вистачить залишку за темпом продажів останніх 90 днів."""
        if self.stock is None or not self.sold_90:
            return None
        return int(self.stock / (self.sold_90 / 90))

    @property
    def sale_price(self):
        return self.product.price or self.product.price_manual

    @property
    def margin(self):
        """Націнка, % від середньої закупівлі."""
        c = self.cost_avg or self.cost_last
        return round(100 * (self.sale_price - c) / c) if c and self.sale_price else None

    @property
    def days_to_expiry(self):
        return (self.expiry - date.today()).days if self.expiry else None

    last_source: str | None = None     # «CleverBox», якщо остання ціна — поточна собівартість з CleverBox

    @property
    def stock_value(self):
        """Залишок — в одиницях обліку (мл, шт), ціна закупівлі — за одиницю приходу (флакон)."""
        c = self.cost_avg or self.cost_last
        k = self.product.coef if self.product.coef and self.product.coef > 0 else 1
        return (self.stock or 0) * c / k if c and self.stock else 0.0

    @property
    def conversion(self):
        return round(100 * self.bought / self.prescribed, 1) if self.prescribed else 0.0

    @property
    def need_30(self):
        return self.demand_open + self.rebuy_30

    @property
    def deficit(self):
        if self.stock is None:
            return None
        return max(0, int(round(self.need_30 + (self.product.min_stock or 0) - self.stock)))


def stock_by_product(db: Session, products: list[Product]) -> dict[int, float]:
    levels = db.scalars(select(StockLevel)).all()
    out: dict[int, float] = defaultdict(float)
    by_cb = {p.cb_product_id: p.id for p in products if p.cb_product_id}
    for lv in levels:
        pid = by_cb.get(lv.cb_product_id)
        if pid is None:
            m = match_product(db, lv.name, products)
            pid = m.id if m else None
        if pid is not None:
            out[pid] += lv.qty
    return out


def product_rows(db: Session, statuses: list[PlanStatus] | None = None, now: datetime | None = None) -> list[ProductRow]:
    now = now or datetime.now()
    statuses = statuses if statuses is not None else all_statuses(db)
    products = db.scalars(select(Product).order_by(Product.name)).all()
    rows = {p.id: ProductRow(p) for p in products}
    for ps in statuses:
        if ps.plan.status == "cancelled":
            continue
        for st in ps.items:
            it = st.item
            if it.kind != "product" or it.product_id not in rows:
                continue
            r = rows[it.product_id]
            r.prescribed += it.qty
            r.bought += min(st.done, it.qty)
            r.revenue += min(st.done, it.qty) * it.unit_price
            r.patients.add(ps.plan.patient_id)
            if ps.plan.status == "active" and st.done < it.qty:
                r.demand_open += it.qty - st.done
            if st.last_done and it.supply_days and st.done >= it.qty:
                due = st.last_done + timedelta(days=it.supply_days)
                if now - timedelta(days=30) <= due <= now + timedelta(days=30):
                    r.rebuy_30 += 1
    from .stock import ledger_active, balance_by_product
    if ledger_active(db):                        # облік у CRM — залишок із журналу рухів
        stock = balance_by_product(db)
        for pid, r in rows.items():
            r.stock = round(stock.get(pid, 0.0), 3)
    elif db.scalar(select(StockLevel).limit(1)) is not None:
        stock = stock_by_product(db, products)
        for pid, r in rows.items():
            r.stock = stock.get(pid, 0.0)
    apply_purchases(db, rows)
    if ledger_active(db):                        # термін — найближча партія, що реально лежить на складі
        from .stock import batch_balances
        near: dict[int, date] = {}
        for (pid, _w, e), q in batch_balances(db).items():
            if e and q > 1e-6 and (pid not in near or e < near[pid]):
                near[pid] = e
        for pid, r in rows.items():
            if pid in near:
                r.expiry = min(near[pid], r.product.expiry_date) if r.product.expiry_date else near[pid]
    apply_sales(db, rows, now)
    return sorted(rows.values(), key=lambda r: (-r.prescribed, r.product.name))


def apply_purchases(db: Session, rows: dict) -> None:
    """Середня/остання ціна закупівлі і найближчий термін придатності з партій."""
    from ..models import ProductPurchase
    today = date.today()
    agg: dict[int, dict] = {}
    for pp in db.scalars(select(ProductPurchase).order_by(ProductPurchase.bought_on, ProductPurchase.id)):
        a = agg.setdefault(pp.product_id, {"q": 0.0, "sum": 0.0, "n": 0, "last": None, "batches": []})
        if pp.unit_cost and pp.qty:
            a["q"] += pp.qty
            a["sum"] += pp.qty * pp.unit_cost
        a["n"] += 1
        a["last"] = pp
        a["batches"].append(pp)
    for pid, r in rows.items():
        p = r.product
        a = agg.get(pid)
        exp = []
        k = p.coef if p.coef and p.coef > 0 else 1     # закупівлі зберігаються за одиницю обліку → показуємо за одиницю приходу
        if a:
            r.purchases = a["n"]
            r.cost_avg = round(a["sum"] / a["q"] * k, 2) if a["q"] else None
            last = a["last"]
            r.cost_last = round(last.unit_cost * k, 2) if last.unit_cost else None
            r.last_bought, r.last_supplier = last.bought_on, last.supplier
            # FIFO: на складі лежать найсвіжіші партії — термін беремо лише з них
            if r.stock is None:
                live = [b for b in a["batches"] if b.bought_on >= today - timedelta(days=730)]
            else:
                live, left = [], r.stock
                for b in reversed(a["batches"]):
                    if left <= 0:
                        break
                    live.append(b)
                    left -= b.qty or 0
            exp = [b.expiry_date for b in live if b.expiry_date]
        if not r.cost_avg and not p.cost and p.stock_cost:
            r.cost_avg = p.stock_cost                # собівартість залишку зі звіту «Залишки по складах»
        if not r.cost_avg and p.cost:
            r.cost_avg = p.cost                      # собівартість з CleverBox, якщо закупівель не внесено
        # поточна собівартість CleverBox свіжіша за давні закупівлі (напр. імпорт руху за 2022)
        if p.cost and (not r.last_bought or r.last_bought < today - timedelta(days=180)):
            r.cost_last, r.last_bought, r.last_source = p.cost, None, "CleverBox"
        if p.expiry_date:
            exp.append(p.expiry_date)
        r.expiry = min(exp) if exp else None


def apply_sales(db: Session, rows: dict, now: datetime) -> None:
    from sqlalchemy import func
    from ..models import ProductSale
    by_cb = {r.product.cb_product_id: r for r in rows.values() if r.product.cb_product_id}
    for pid, q, sm in db.execute(select(ProductSale.product_id, func.sum(ProductSale.quantity), func.sum(ProductSale.amount))
                                 .where(ProductSale.sold_at >= now - timedelta(days=90), ProductSale.product_id.is_not(None))
                                 .group_by(ProductSale.product_id)):
        r = by_cb.get(pid)
        if r:
            r.sold_90, r.sold_90_sum = float(q or 0), float(sm or 0)


def brand_groups(rows: list) -> list[tuple[str, list]]:
    groups: dict[str, list] = defaultdict(list)
    for r in rows:
        groups[r.product.brand or ""].append(r)
    return sorted(groups.items(), key=lambda g: (g[0] == "", g[0].lower()))


@dataclass
class ProcedureRow:
    name: str
    planned: int = 0
    done: int = 0
    booked: int = 0
    planned_sum: float = 0.0
    realized_sum: float = 0.0
    plans: int = 0
    prices: list = field(default_factory=list)

    @property
    def conversion(self):
        return round(100 * self.done / self.planned, 1) if self.planned else 0.0

    @property
    def avg_price(self):
        return round(sum(self.prices) / len(self.prices)) if self.prices else 0


def procedure_rows(statuses: list[PlanStatus]) -> list[ProcedureRow]:
    rows: dict[str, ProcedureRow] = {}
    for ps in statuses:
        if ps.plan.status == "cancelled":
            continue
        for st in ps.items:
            it = st.item
            if it.kind != "procedure":
                continue
            r = rows.setdefault(norm(it.name), ProcedureRow(it.name))
            r.planned += it.qty
            r.done += min(st.done, it.qty)
            r.booked += st.booked
            r.planned_sum += it.total
            r.realized_sum += st.realized_sum
            r.plans += 1
            r.prices.append(it.unit_price)
    return sorted(rows.values(), key=lambda r: -r.planned_sum)


# ---------------- Імпорт товарів з Excel-вивантаження CleverBox ----------------
IMPORT_COLS = {"id": ("id",), "name": ("назва",), "sku": ("артикул",), "barcode": ("штрихкод",),
               "unit": ("од. списання", "од списання", "одиниця"), "category": ("категорія",), "direction": ("напрямок",),
               "brand": ("бренд",), "supplier": ("постачальник",), "min_stock": ("крит. залишок", "критичний залишок"),
               "for_sale": ("для продажу",), "cost": ("ціна закупівлі", "собівартість"), "price": ("ціна продажу",),
               "qty": ("кількість", "залишок"), "coef": ("коефiцiєнт", "коефіцієнт", "коеф")}


def _cell_num(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace("\xa0", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def read_table(path: str) -> list[dict]:
    """XLS (CleverBox) або XLSX → список рядків {назва колонки: значення}."""
    if path.lower().endswith(".xls"):
        import xlrd
        sh = xlrd.open_workbook(path, ignore_workbook_corruption=True).sheet_by_index(0)
        rows = [sh.row_values(i) for i in range(sh.nrows)]
    else:
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        return []
    head = [str(h or "").strip().lower() for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if any(x not in (None, "") for x in r)]


def detect_export(path: str) -> str:
    """Який це файл CleverBox: products (Склад → Товари) | supply (Склад → Поставки) | turnover (Рух товарів) |
    balances (Залишки по складах) | unknown."""
    data = read_table(path)
    if not data:
        return "unknown"
    h = set(data[0].keys())
    if "n" in h and "постачальник" in h:
        return "supply"
    if any(k.startswith("залишок од") for k in h) and "назва" in h:
        return "balances"
    if "товар" in h and any(k.startswith("надходження") for k in h):
        return "turnover"
    if "назва" in h and any(k.startswith("категорія") or k.startswith("ціна") for k in h):
        return "products"
    return "unknown"


WRONG_FILE = {
    "products": "Це вивантаження «Товари» — його імпорт: Продукти → «Імпорт товарів з CleverBox (Excel)».",
    "supply": "Це вивантаження «Поставки» — його імпорт: Склад → Постачальники → «Імпорт поставок з CleverBox».",
    "turnover": "Це звіт «Рух товарів» — його імпорт: Продукти → «Ціни закупівлі з руху товарів CleverBox» (вкажіть період).",
    "balances": "Це звіт «Залишки по складах» — його імпорт: Склад → Залишки → «Імпорт залишків з CleverBox».",
    "unknown": "Не вдалося розпізнати файл. Потрібне вивантаження CleverBox у форматі Excel (.xls/.xlsx).",
}


def import_products(db: Session, path: str) -> dict:
    """Оновлює каталог з файлу: склад, бренд, постачальник, артикул, штрихкод, одиниця, ціни, крит. залишок, кількість.
    Зіставлення — за ID CleverBox, потім за назвою. Нові товари створюються."""
    data = read_table(path)
    if not data:
        return {"error": "Файл порожній або не розпізнано"}
    head = list(data[0].keys())

    def col(key):
        for h in head:
            if any(h == a or h.startswith(a) for a in IMPORT_COLS[key]):
                return h
        return None
    cmap = {k: col(k) for k in IMPORT_COLS}
    if not cmap["name"]:
        return {"error": "Немає колонки «Назва»"}
    products = db.scalars(select(Product)).all()
    by_cb = {p.cb_product_id: p for p in products if p.cb_product_id}
    by_name = {p.name_norm: p for p in products}
    brand_case: dict[str, str] = {}
    rep = {"rows": len(data), "created": 0, "updated": 0, "by_warehouse": defaultdict(int), "stock": 0}
    stock_rows = []
    for r in data:
        g = lambda k: (r.get(cmap[k]) if cmap[k] else None)
        name = str(g("name") or "").strip()
        if not name:
            continue
        cid = _cell_num(g("id"))
        cid = int(cid) if cid else None
        p = by_cb.get(cid) if cid else None
        p = p or by_name.get(norm(name))
        if not p:
            p = Product(name=name, name_norm=norm(name), cb_product_id=cid)
            db.add(p)
            rep["created"] += 1
        else:
            rep["updated"] += 1
        if cid and not p.cb_product_id:
            p.cb_product_id = cid
        wh = warehouse_of(str(g("category") or "")) or warehouse_of(str(g("direction") or ""))
        if wh:
            p.warehouse = wh
        rep["by_warehouse"][p.warehouse or "none"] += 1
        brand = str(g("brand") or "").strip()
        if brand:
            p.brand = brand_case.setdefault(brand.lower(), brand)[:120]
        sup = str(g("supplier") or "").strip()
        if sup:
            p.supplier = sup[:200]
        for k, attr in (("sku", "sku"), ("barcode", "barcode"), ("unit", "unit")):
            v = g(k)
            if v not in (None, "") and str(v).strip() not in ("0", "0.0"):
                setattr(p, attr, str(v).strip().rstrip(".")[:100] if k == "unit" else str(v).strip()[:100])
        if g("for_sale") is not None and cmap["for_sale"]:
            p.for_sale = str(g("for_sale")).strip() in ("1", "1.0", "так", "true", "True")
        kf = _cell_num(g("coef"))
        if kf and kf > 0:
            p.coef = kf
        price, cost, ms = _cell_num(g("price")), _cell_num(g("cost")), _cell_num(g("min_stock"))
        if price:
            p.price = price
        if cost:
            p.cost = cost
        if ms is not None and ms > 0:
            p.min_stock = int(ms)
        q = _cell_num(g("qty"))
        if q is not None and cmap["qty"]:
            stock_rows.append((p, name, max(q, 0.0)))
    db.flush()
    if stock_rows:
        ids = [p.cb_product_id for p, _, _ in stock_rows if p.cb_product_id]
        if ids:
            db.execute(StockLevel.__table__.delete().where(StockLevel.cb_product_id.in_(ids)))
        for p, name, q in stock_rows:
            if p.cb_product_id:
                db.add(StockLevel(cb_product_id=p.cb_product_id, name=name, qty=q,
                                  storage=WAREHOUSES.get(p.warehouse or "", None), unit=p.unit))
                rep["stock"] += 1
    db.commit()
    rep["by_warehouse"] = dict(rep["by_warehouse"])
    rep["brands"] = len({p.brand for p in db.scalars(select(Product)) if p.brand})
    return rep


# ---------------- Оборотна відомість CleverBox (рух товарів за період) → закупівлі ----------------

def import_turnover(db: Session, path: str, d_from: date, d_to: date) -> dict:
    """Звіт «Рух товарів» (Товар, артикул, …, надходження к-ть, Прихід сума, витрата…, залишок) за період.
    По кожному товару з надходженням створює закупівлю «за період» (дата = кінець періоду) з ціною за одиницю обліку
    = сума приходу / к-ть надходження. Повторний імпорт того самого періоду замінює попередній."""
    from ..models import ProductPurchase
    data = read_table(path)
    if not data or "товар" not in data[0] or not any(k.startswith("надходження") for k in data[0]):
        return {"error": "Це не звіт «Рух товарів» з CleverBox (немає колонок Товар / надходження)"}
    head = list(data[0].keys())
    k_qty = next((h for h in head if h.startswith("надходження") and "обліку" in h), None) or \
        next(h for h in head if h.startswith("надходження"))
    k_sum = next((h for h in head if h.startswith("прихід")), None)
    k_out = next((h for h in head if h.startswith("витрата") and "обліку" in h), None)
    tag = f"turnover:{d_from.isoformat()}:{d_to.isoformat()}"
    db.execute(ProductPurchase.__table__.delete().where(ProductPurchase.note == tag))
    products = db.scalars(select(Product)).all()
    by_sku = {p.sku.strip().lower(): p for p in products if p.sku}
    by_name = {p.name_norm: p for p in products}
    rep = {"rows": 0, "with_income": 0, "matched": 0, "created": 0, "not_found": []}
    for r in data:
        name = str(r.get("товар") or "").strip()
        if not name:
            continue
        rep["rows"] += 1
        q, sm = _cell_num(r.get(k_qty)), abs(_cell_num(r.get(k_sum)) or 0) if k_sum else 0
        if not q or q <= 0 or not sm:
            continue
        rep["with_income"] += 1
        sku = str(r.get("артикул") or "").strip().lower()
        p = (by_sku.get(sku) if sku and sku not in ("0", "0.0") else None) or by_name.get(norm(name)) or match_product(db, name, products)
        if not p:
            p = Product(name=name, name_norm=norm(name), sku=(str(r.get("артикул")).strip() or None) if sku else None,
                        unit=str(r.get("одиниця обліку") or "").strip().rstrip(".")[:20] or None)
            db.add(p)
            db.flush()
            products.append(p)
            by_name[p.name_norm] = p
            rep["created"] += 1
            rep["not_found"].append(name[:80])
        else:
            rep["matched"] += 1
        kf = _cell_num(r.get("коеф."))
        if kf and kf > 0 and not p.coef:
            p.coef = kf
        db.add(ProductPurchase(product_id=p.id, bought_on=d_to, qty=q, unit_cost=round(sm / q, 4),
                               note=tag, doc_no=f"рух {d_from:%d.%m.%Y}–{d_to:%d.%m.%Y}"))
    db.commit()
    rep["not_found"] = rep["not_found"][:20]
    return rep


def purchase_history(p, purchases):
    """Історія поставок товару для картки: ціна й кількість — в одиницях приходу (флакон), а не обліку (мл).
    purchases — від найновіших до найстаріших. Повертає (рядки, підсумки)."""
    k = p.coef if p.coef and p.coef > 0 else 1
    rows = []
    for pp in purchases:
        note = pp.note or ""
        if note.startswith("stockdoc:"):
            src, link = "Поставка", "/stock/docs/" + note.split(":", 1)[1]
        elif note.startswith("turnover:"):
            src, link = "Рух CleverBox", None
        else:
            src, link = "Вручну", None
        period = None
        if note.startswith("turnover:"):
            try:
                a, b = note.split(":")[1:3]
                period = f"{date.fromisoformat(a):%d.%m.%y}–{date.fromisoformat(b):%d.%m.%y}"
            except (ValueError, IndexError):
                period = None
        rows.append({"id": pp.id, "date": pp.bought_on, "period": period, "supplier": pp.supplier,
                     "qty": round((pp.qty or 0) / k, 3), "price": round((pp.unit_cost or 0) * k, 2),
                     "sum": round((pp.qty or 0) * (pp.unit_cost or 0), 2), "expiry": pp.expiry_date,
                     "doc_no": pp.doc_no, "src": src, "link": link, "delta": None,
                     "deletable": not note.startswith("stockdoc:")})
    for i, r in enumerate(rows):                      # зміна ціни до попередньої (старшої) поставки
        prev = next((x for x in rows[i + 1:] if x["price"]), None)
        if prev and r["price"]:
            r["delta"] = round(100 * (r["price"] - prev["price"]) / prev["price"], 1)
    priced = [r for r in rows if r["price"]]
    stats = None
    if priced:
        tq = sum(r["qty"] for r in rows)
        ts = sum(r["sum"] for r in rows)
        first, last = priced[-1], priced[0]
        stats = {"n": len(rows), "qty": round(tq, 3), "sum": round(ts, 2),
                 "avg": round(ts / tq, 2) if tq else None,
                 "min": min(priced, key=lambda r: r["price"]), "max": max(priced, key=lambda r: r["price"]),
                 "first": first, "last": last,
                 "change": round(100 * (last["price"] - first["price"]) / first["price"], 1) if len(priced) > 1 else None,
                 "chart": list(reversed(priced))[-24:]}
        stats["top"] = max(r["price"] for r in stats["chart"])
    return rows, stats


def parse_balances(db: Session, path: str) -> dict:
    """Звіт CleverBox «Залишки по складах»: рядки-заголовки складів (лише назва), далі товари з
    «Собівартість» (вартість усього залишку), «Залишок од. обліку». Зіставляє товари (артикул → назва),
    відсутні створює, оновлює собівартість залишку. Повертає знімок {date, items:[[pid, wh, qty, cost_total]], …}."""
    data = read_table(path)
    if not data or not any(k.startswith("залишок од") for k in data[0]):
        raise ValueError("Це не звіт «Залишки по складах» з CleverBox (немає колонки «Залишок од. обліку»)")
    head = list(data[0].keys())
    k_qty = next(h for h in head if h.startswith("залишок од") and "обл" in h) if any(
        h.startswith("залишок од") and "обл" in h for h in head) else next(h for h in head if h.startswith("залишок од"))
    k_cost = next((h for h in head if h.startswith("собівартість")), None)
    products = db.scalars(select(Product)).all()
    by_sku = {p.sku.strip().lower(): p for p in products if p.sku}
    by_name = {p.name_norm: p for p in products}
    group = None
    items: dict[tuple[int, str], list] = {}
    rep = {"rows": 0, "matched": 0, "created": 0, "negative": [], "unknown_wh": 0, "by_wh": {}}
    for r in data:
        name = str(r.get("назва") or "").strip()
        if not name or name.lower().startswith("total"):
            continue
        if r.get(k_qty) in (None, "") and not str(r.get("категорія") or "").strip():
            group = name                                  # заголовок складу
            continue
        rep["rows"] += 1
        wh = warehouse_of(group) or warehouse_of(r.get("категорія")) or warehouse_of(r.get("напрямок"))
        if not wh:
            rep["unknown_wh"] += 1
            wh = "vitrina"
        q = _cell_num(r.get(k_qty)) or 0.0
        cost = _cell_num(r.get(k_cost)) if k_cost else None
        sku_raw = str(r.get("артикул") or "").strip()
        sku = sku_raw.lower() if sku_raw not in ("", "0", "0.0") else ""
        p = (by_sku.get(sku) if sku else None) or by_name.get(norm(name)) or match_product(db, name, products)
        if not p:
            p = Product(name=name, name_norm=norm(name), sku=sku_raw or None, warehouse=wh,
                        brand=str(r.get("бренд") or "").strip() or None,
                        unit=str(r.get("од. обліку") or "").strip().rstrip(".")[:20] or None)
            db.add(p)
            db.flush()
            products.append(p)
            by_name[p.name_norm] = p
            rep["created"] += 1
        else:
            rep["matched"] += 1
        if not p.brand and str(r.get("бренд") or "").strip():
            p.brand = str(r.get("бренд")).strip()
        if q < 0:
            rep["negative"].append(f"{name[:70]} ({WAREHOUSES[wh]}: {q:g})")
        it = items.setdefault((p.id, wh), [p.id, wh, 0.0, 0.0])
        it[2] += q
        it[3] += cost or 0.0
    # собівартість залишку за одиницю приходу (усі склади разом, лише додатні залишки)
    tot: dict[int, list] = {}
    for pid, wh, q, c in items.values():
        if q > 0 and c > 0:
            t = tot.setdefault(pid, [0.0, 0.0])
            t[0] += q
            t[1] += c
        b = rep["by_wh"].setdefault(wh, {"n": 0, "qty_pos": 0, "value": 0.0})
        if q > 0:
            b["n"] += 1
            b["value"] = round(b["value"] + (c or 0), 2)
    pmap = {p.id: p for p in products}
    today = date.today()
    for pid, (q, c) in tot.items():
        p = pmap[pid]
        k = p.coef if p.coef and p.coef > 0 else 1
        p.stock_cost, p.stock_cost_at = round(c / q * k, 2), today
    db.commit()
    rep["items"] = [[pid, wh, round(q, 4), round(c, 2)] for pid, wh, q, c in items.values()]
    rep["date"] = today.isoformat()
    rep["value"] = round(sum(b["value"] for b in rep["by_wh"].values()), 2)
    return rep
