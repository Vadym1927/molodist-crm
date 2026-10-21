"""Експорт: звіти Excel / CSV, архів файлів пацієнтів, повна резервна копія CRM (для переїзду на інший сервер)."""
from __future__ import annotations

import csv
import io
import json
import re
import tempfile
import zipfile
from datetime import datetime, date
from pathlib import Path

from sqlalchemy import select, delete, text, DateTime, Date
from sqlalchemy.orm import Session

from ..config import settings
from ..models import Base, Patient, PatientFile, engine
from .kpi import all_statuses

BACKUP_VERSION = 1


# ---------------- звіти ----------------

def report_tables(db: Session, user, date_from: date | None = None, date_to: date | None = None) -> dict:
    from ..integrations.sheets import build_tables
    from ..security import can
    doc = user.doctor_id if user.role == "doctor" else None
    sts = all_statuses(db, date_from, date_to, doctor_id=doc)
    tables = build_tables(db, statuses=sts, finance=can(user, "finance"))
    if not (can(user, "finance") or can(user, "kpi_own")):
        tables.pop("KPI лікарів", None)
    return tables


MONEY_COLS = {"Сума плану", "Реалізовано", "Ціна", "Сума", "Сума планів", "Записано (сума)", "Середній план",
              "Підтверджено (записи)", "Очікується за планами", "Повторний догляд", "Разом"}
PCT_COLS = {"% реалізації", "% активних із записом", "Конверсія в старт"}


def to_xlsx(tables: dict[str, list[list]], title: str = "MOLODIST CRM") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1D1B19")
    for name, rows in tables.items():
        ws = wb.create_sheet(name[:31])
        for r in rows:
            ws.append(r)
        if not rows:
            continue
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for idx, h in enumerate(rows[0], 1):
            letter = get_column_letter(idx)
            fmt = '#,##0 "₴"' if h in MONEY_COLS else ("0.0%" if h in PCT_COLS else None)
            width = max(len(str(h)), *(len(str(r[idx - 1])) for r in rows[1:] if idx - 1 < len(r))) if len(rows) > 1 else len(str(h))
            ws.column_dimensions[letter].width = min(max(10, width + 2), 60)
            if fmt:
                for cell in ws[letter][1:]:
                    cell.number_format = fmt
    wb.properties.title = title
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_csv_zip(tables: dict[str, list[list]]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, rows in tables.items():
            s = io.StringIO()
            csv.writer(s, delimiter=";").writerows(rows)       # «;» — щоб Excel з українською локаллю відкривав коректно
            z.writestr(f"{name}.csv", "﻿" + s.getvalue())
    return buf.getvalue()


# ---------------- файли пацієнтів ----------------

def _safe(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", s).strip()[:80] or "file"


def files_zip(db: Session, patients: list[Patient]) -> str:
    """Архів фото і PDF: <Пацієнт>/<дата>_<файл>. Повертає шлях до тимчасового zip."""
    tmp = tempfile.NamedTemporaryFile(prefix="crm_files_", suffix=".zip", delete=False)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_STORED) as z:     # фото вже стиснуті
        for p in patients:
            for f in p.files:
                src = Path(f.path)
                if src.exists():
                    z.write(src, f"{_safe(p.full_name)}/{f.created_at:%Y-%m-%d}_{_safe(src.name.split('_', 1)[-1])}")
    tmp.close()
    return tmp.name


# ---------------- повна резервна копія ----------------

def _ser(v):
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


def backup_zip(db: Session, with_files: bool = True) -> str:
    tmp = tempfile.NamedTemporaryFile(prefix="crm_backup_", suffix=".zip", delete=False)
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        meta = {"version": BACKUP_VERSION, "created_at": datetime.now().isoformat(timespec="seconds"),
                "tables": {}}
        for t in Base.metadata.sorted_tables:
            rows = [{k: _ser(v) for k, v in r._mapping.items()} for r in db.execute(t.select())]
            meta["tables"][t.name] = len(rows)
            z.writestr(f"db/{t.name}.json", json.dumps(rows, ensure_ascii=False))
        if with_files:
            base = Path(settings.files_dir)
            for f in base.rglob("*"):
                if f.is_file():
                    z.write(f, "files/" + str(f.relative_to(base)))
        z.writestr("backup.json", json.dumps(meta, ensure_ascii=False, indent=2))
    tmp.close()
    return tmp.name


def restore_backup(db: Session, zip_path: str) -> dict:
    """Повністю замінює дані CRM вмістом резервної копії (переїзд на новий сервер)."""
    with zipfile.ZipFile(zip_path) as z:
        names = set(z.namelist())
        if "backup.json" not in names:
            raise ValueError("Це не резервна копія MOLODIST CRM (немає backup.json)")
        meta = json.loads(z.read("backup.json"))
        tables = Base.metadata.sorted_tables
        for t in reversed(tables):
            db.execute(delete(t))
        counts = {}
        for t in tables:
            fn = f"db/{t.name}.json"
            if fn not in names:
                continue
            rows = json.loads(z.read(fn))
            for r in rows:
                for col in t.columns:
                    v = r.get(col.name)
                    if isinstance(v, str) and isinstance(col.type, (DateTime, Date)):
                        r[col.name] = datetime.fromisoformat(v) if isinstance(col.type, DateTime) else date.fromisoformat(v)
                r_clean = {k: v for k, v in r.items() if k in t.c}
                db.execute(t.insert().values(**r_clean))
            counts[t.name] = len(rows)
        db.commit()
        # файли
        base = Path(settings.files_dir)
        nfiles = 0
        for n in names:
            if n.startswith("files/") and not n.endswith("/"):
                rel = Path(n[6:])
                if ".." in rel.parts:
                    continue
                dst = base / rel
                dst.parent.mkdir(parents=True, exist_ok=True)
                dst.write_bytes(z.read(n))
                nfiles += 1
    # шляхи файлів у БД -> на новий files_dir
    for f in db.scalars(select(PatientFile)):
        p = Path(f.path)
        parts = p.parts
        if "files" in parts:
            rel = Path(*parts[parts.index("files") + 1:])
            f.path = str(base / rel)
    db.commit()
    _fix_sequences(db)
    return {"backup_created": meta.get("created_at"), "rows": counts, "files": nfiles}


def _fix_sequences(db: Session):
    if engine.dialect.name != "postgresql":
        return
    for t in Base.metadata.sorted_tables:
        if "id" in t.c and t.c.id.autoincrement and str(t.c.id.type).upper().startswith("INT"):
            db.execute(text(f"SELECT setval(pg_get_serial_sequence('{t.name}', 'id'), "
                            f"COALESCE((SELECT MAX(id) FROM {t.name}), 0) + 1, false)"))
    db.commit()
