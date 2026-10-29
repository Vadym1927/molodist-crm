"""Звіти: виручка, лікарі, послуги, товари, пацієнти, бюті-плани, поставки.

Кожен звіт повертає Report — однакову структуру (плитки з порівнянням з попереднім періодом, колонки, рядки,
підсумок, міні-графік), яку малює один шаблон reports.html і вивантажує один Excel-експорт.
Джерела: чеки CleverBox (Check — фактична виручка), записи (Visit), бюті-плани, поставки (SupplierInvoice)."""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from sqlalchemy import select, func
from sqlalchemy.orm import Session

from ..models import Check, Visit, BeautyPlan, Patient, Product, SupplierInvoice
from ..parsers.observ import norm

REPORTS = {"cashflow": "Cashflow", "fop": "Ліміти ФОП", "expiry": "Терміни придатності", "revenue": "Виручка", "doctors": "Лікарі", "services": "Послуги", "products": "Товари",
           "patients": "Пацієнти", "plans": "Бюті-плани", "certificates": "Сертифікати й абонементи", "supply": "Поставки"}


@dataclass
class Col:
    key: str
    label: str
    kind: str = "text"          # text | uah | int | num | pct | date
    total: bool = False         # підсумовувати в рядку «Разом»


@dataclass
class Report:
    key: str
    title: str
    hint: str = ""
    tiles: list = field(default_factory=list)       # (назва, значення, попереднє, формат, підпис)
    cols: list[Col] = field(default_factory=list)
    rows: list[dict] = field(default_factory=list)
    total: dict | None = None
    chart: list = field(default_factory=list)       # (підпис, значення)
    chart_title: str = ""
    empty: str = "За період немає даних."

    def finish(self):
        if any(c.total for c in self.cols) and self.rows:
            self.total = {c.key: sum(r.get(c.key) or 0 for r in self.rows) for c in self.cols if c.total}
        return self


def _rng(d_from: date, d_to: date):
    return datetime.combine(d_from, datetime.min.time()), datetime.combine(d_to, datetime.max.time())


def prev_period(d_from: date, d_to: date) -> tuple[date, date]:
    """Попередній період тієї ж довжини (для місяця — попередній календарний місяць)."""
    if d_from.day == 1 and (d_to + timedelta(days=1)).day == 1:
        p_to = d_from - timedelta(days=1)
        months = (d_to.year - d_from.year) * 12 + d_to.month - d_from.month + 1
        y, m = p_to.year, p_to.month - months + 1
        while m < 1:
            y, m = y - 1, m + 12
        return date(y, m, 1), p_to
    n = (d_to - d_from).days + 1
    return d_from - timedelta(days=n), d_from - timedelta(days=1)


def _checks(db: Session, d_from: date, d_to: date):
    a, b = _rng(d_from, d_to)
    return db.scalars(select(Check).where(Check.deleted.is_(False), Check.date >= a, Check.date <= b)).all()


def _items(c: Check):
    for it in c.items or []:
        amt = float(it.get("payd_amount") or it.get("total") or 0)
        yield (it.get("item_type") or "other", (it.get("item_name") or it.get("name") or "—").strip(),
               float(it.get("quantity") or 1), amt, (it.get("initiator_name") or c.author or "—").strip())


def _money(checks):
    t = {"revenue": 0.0, "checks": 0, "services": 0.0, "products": 0.0, "other": 0.0, "discount": 0.0, "patients": set()}
    for c in checks:
        t["revenue"] += c.amount or 0
        t["discount"] += c.discount or 0
        t["checks"] += 1
        if c.patient_id:
            t["patients"].add(c.patient_id)
        for kind, _, _, amt, _ in _items(c):
            t["services" if kind == "service" else "products" if kind == "product" else "other"] += amt
    t["patients"] = len(t["patients"])
    t["avg"] = round(t["revenue"] / t["checks"]) if t["checks"] else 0
    return t


# ---------------- Виручка ----------------

def revenue(db: Session, d_from: date, d_to: date) -> Report:
    cur, prev = _checks(db, d_from, d_to), _checks(db, *prev_period(d_from, d_to))
    t, p = _money(cur), _money(prev)
    r = Report("revenue", "Виручка", "За чеками CleverBox. Порівняння — з попереднім періодом такої ж довжини.")
    r.tiles = [("Виручка", t["revenue"], p["revenue"], "uah", f"{t['checks']} чеків"),
               ("Середній чек", t["avg"], p["avg"], "uah", ""),
               ("Послуги", t["services"], p["services"], "uah", ""),
               ("Товари", t["products"], p["products"], "uah", ""),
               ("Пацієнтів", t["patients"], p["patients"], "int", "з оплатою")]
    by_month = (d_to - d_from).days > 62
    buckets: dict = defaultdict(list)
    for c in cur:
        buckets[(c.date.year, c.date.month) if by_month else c.date.date()].append(c)
    keys = []
    if by_month:
        y, m = d_from.year, d_from.month
        while (y, m) <= (d_to.year, d_to.month):
            keys.append((y, m))
            y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    else:
        d = d_from
        while d <= d_to:
            keys.append(d)
            d += timedelta(days=1)
    r.cols = [Col("period", "Місяць" if by_month else "Дата"), Col("checks", "Чеків", "int", True),
              Col("patients", "Пацієнтів", "int"), Col("services", "Послуги", "uah", True),
              Col("products", "Товари", "uah", True), Col("other", "Інше", "uah", True),
              Col("discount", "Знижки", "uah", True), Col("revenue", "Виручка", "uah", True), Col("avg", "Сер. чек", "uah")]
    for k in keys:
        m = _money(buckets.get(k, []))
        label = f"{k[1]:02d}.{k[0]}" if by_month else k.strftime("%d.%m %a").replace("Mon", "пн").replace("Tue", "вт") \
            .replace("Wed", "ср").replace("Thu", "чт").replace("Fri", "пт").replace("Sat", "сб").replace("Sun", "нд")
        if m["checks"] or by_month:
            r.rows.append({"period": label, **{x: m[x] for x in ("checks", "patients", "services", "products", "other",
                                                                 "discount", "revenue", "avg")}})
        r.chart.append((f"{k[1]:02d}.{str(k[0])[2:]}" if by_month else k.strftime("%d"), m["revenue"]))
    r.chart_title = "Виручка по місяцях" if by_month else "Виручка по днях"
    r.finish()
    if r.total:
        r.total["avg"] = round(r.total["revenue"] / r.total["checks"]) if r.total.get("checks") else 0
        r.total["patients"] = t["patients"]
    return r


# ---------------- Лікарі ----------------

def doctors(db: Session, d_from: date, d_to: date) -> Report:
    a, b = _rng(d_from, d_to)
    agg = defaultdict(lambda: {"services": 0.0, "proc": 0.0, "products": 0.0, "patients": set(), "done": 0,
                               "booked": 0, "cancelled": 0, "plans": 0, "plan_sum": 0.0})
    for c in _checks(db, d_from, d_to):
        for kind, _, q, amt, staff in _items(c):
            x = agg[staff]
            if kind == "service":
                x["services"] += amt
                x["proc"] += q
            elif kind == "product":
                x["products"] += amt
            if c.patient_id and kind in ("service", "product"):
                x["patients"].add(c.patient_id)
    seen = defaultdict(set)          # візит = пацієнт × день у лікаря (надані послуги й записи CleverBox)
    for v in db.scalars(select(Visit).where(Visit.starts_at >= a, Visit.starts_at <= b)):
        m = (v.master_name or "—").strip()
        if v.status == "cancelled":
            agg[m]["cancelled"] += 1
        elif v.status == "done" or (v.source != "cb_service_sale" and v.starts_at < datetime.now()):
            seen[m].add((v.patient_id, v.starts_at.date()))
        else:
            agg[m]["booked"] += 1
    for m, days in seen.items():
        agg[m]["done"] = len(days)
    for pl in db.scalars(select(BeautyPlan).where(BeautyPlan.session_at >= a, BeautyPlan.session_at <= b,
                                                  BeautyPlan.status != "cancelled")):
        x = agg[pl.doctor.name if pl.doctor else "Без лікаря"]
        x["plans"] += 1
        x["plan_sum"] += pl.total_planned or sum(i.total for i in pl.items)
    r = Report("doctors", "Лікарі", "Послуги й товари — за ініціатором у чеках CleverBox; записи — за майстром у записі; "
               "бюті-плани — за датою консультації.")
    r.cols = [Col("name", "Лікар / спеціаліст"), Col("services", "Послуги", "uah", True), Col("proc", "Процедур", "num", True),
              Col("products", "Догляд", "uah", True), Col("total", "Разом", "uah", True), Col("patients", "Пацієнтів", "int"),
              Col("per_patient", "На пацієнта", "uah"), Col("done", "Візитів відбулось", "int", True),
              Col("cancelled", "Скасовано", "int", True), Col("cancel_pct", "% скасувань", "pct"),
              Col("plans", "Бюті-планів", "int", True), Col("plan_sum", "Сума планів", "uah", True)]
    for name, x in agg.items():
        tot = x["services"] + x["products"]
        if not (tot or x["done"] or x["booked"] or x["plans"]):
            continue
        n_p = len(x["patients"])
        vis = x["done"] + x["cancelled"]
        r.rows.append({"name": name, "services": x["services"], "proc": x["proc"], "products": x["products"], "total": tot,
                       "patients": n_p, "per_patient": round(tot / n_p) if n_p else 0, "done": x["done"],
                       "cancelled": x["cancelled"], "cancel_pct": round(100 * x["cancelled"] / vis, 1) if vis else 0,
                       "plans": x["plans"], "plan_sum": x["plan_sum"]})
    r.rows.sort(key=lambda z: -z["total"])
    r.chart = [(z["name"][:18], z["total"]) for z in r.rows[:12] if z["total"]]
    r.chart_title = "Виручка за лікарями"
    top = r.rows[0] if r.rows else None
    r.finish()
    r.tiles = [("Лікарів з виручкою", sum(1 for z in r.rows if z["total"]), None, "int", ""),
               ("Найбільша виручка", top["total"] if top else 0, None, "uah", top["name"] if top else ""),
               ("Бюті-планів", r.total["plans"] if r.total else 0, None, "int", ""),
               ("Візитів відбулось", r.total["done"] if r.total else 0, None, "int", "")]
    return r


# ---------------- Послуги / Товари (позиції чеків) ----------------

def _positions(db, d_from, d_to, kind):
    agg = defaultdict(lambda: {"qty": 0.0, "sum": 0.0, "patients": set(), "staff": defaultdict(float)})
    for c in _checks(db, d_from, d_to):
        for k, name, q, amt, staff in _items(c):
            if k != kind:
                continue
            x = agg[name]
            x["qty"] += q
            x["sum"] += amt
            x["staff"][staff] += amt
            if c.patient_id:
                x["patients"].add(c.patient_id)
    return agg


def services(db: Session, d_from: date, d_to: date) -> Report:
    agg = _positions(db, d_from, d_to, "service")
    prev = _positions(db, *prev_period(d_from, d_to), "service")
    total = sum(x["sum"] for x in agg.values()) or 1
    from .cb_services_import import service_index
    idx = service_index(db)
    r = Report("services", "Послуги", "Надані й оплачені послуги за чеками CleverBox. Напрямок — з довідника "
               "«Послуги» CleverBox (Дані → Імпорт).")
    r.cols = [Col("name", "Послуга"), Col("direction", "Напрямок"), Col("qty", "Кількість", "num", True), Col("sum", "Виручка", "uah", True),
              Col("avg", "Сер. ціна", "uah"), Col("share", "Частка", "pct"), Col("patients", "Пацієнтів", "int"),
              Col("prev", "Попередній період", "uah", True), Col("staff", "Хто виконував")]
    for name, x in agg.items():
        sv = idx.get(norm(name))
        r.rows.append({"name": name, "direction": (sv.direction if sv else None) or "", "qty": x["qty"], "sum": x["sum"], "avg": round(x["sum"] / x["qty"]) if x["qty"] else 0,
                       "share": round(100 * x["sum"] / total, 1), "patients": len(x["patients"]),
                       "prev": prev[name]["sum"] if name in prev else 0,
                       "staff": ", ".join(s for s, _ in sorted(x["staff"].items(), key=lambda kv: -kv[1])[:3])})
    r.rows.sort(key=lambda z: -z["sum"])
    by_dir = defaultdict(float)
    for z in r.rows:
        by_dir[z["direction"] or "Без напрямку"] += z["sum"]
    if len(by_dir) > 1:
        r.chart = sorted(by_dir.items(), key=lambda kv: -kv[1])[:12]
        r.chart_title = "Виручка за напрямками"
    else:
        r.chart = [(z["name"][:18], z["sum"]) for z in r.rows[:10]]
        r.chart_title = "Топ-10 послуг за виручкою"
    r.finish()
    pt = sum(x["sum"] for x in prev.values())
    r.tiles = [("Виручка послуг", total if agg else 0, pt, "uah", ""),
               ("Послуг надано", sum(x["qty"] for x in agg.values()), sum(x["qty"] for x in prev.values()), "num", ""),
               ("Різних послуг", len(agg), len(prev), "int", "")]
    return r


def products(db: Session, d_from: date, d_to: date) -> Report:
    from .catalog import product_rows
    agg = _positions(db, d_from, d_to, "product")
    prev = _positions(db, *prev_period(d_from, d_to), "product")
    cat = {norm(pr.product.name): pr for pr in product_rows(db, [])}
    r = Report("products", "Товари", "Продажі товарів за чеками CleverBox. Собівартість — середня ціна закупівлі "
               "(або остання / CleverBox), прибуток = виручка − собівартість.")
    r.cols = [Col("name", "Товар"), Col("brand", "Бренд"), Col("qty", "К-сть", "num", True), Col("sum", "Виручка", "uah", True),
              Col("cost", "Собівартість", "uah", True), Col("profit", "Валовий прибуток", "uah", True),
              Col("margin", "Маржа", "pct"), Col("patients", "Пацієнтів", "int"), Col("prev", "Попередній період", "uah", True)]
    for name, x in agg.items():
        pr = cat.get(norm(name))
        unit = (pr.cost_avg or pr.cost_last) if pr else None
        cost = round(unit * x["qty"], 2) if unit else 0
        profit = x["sum"] - cost if unit else 0
        r.rows.append({"name": name, "brand": (pr.product.brand if pr else "") or "", "qty": x["qty"], "sum": x["sum"],
                       "cost": cost, "profit": profit, "margin": round(100 * profit / x["sum"], 1) if unit and x["sum"] else None,
                       "patients": len(x["patients"]), "prev": prev[name]["sum"] if name in prev else 0,
                       "_href": f"/products/{pr.product.id}" if pr else None})
    r.rows.sort(key=lambda z: -z["sum"])
    by_brand = defaultdict(float)
    for z in r.rows:
        by_brand[z["brand"] or "Без бренду"] += z["sum"]
    r.chart = sorted(by_brand.items(), key=lambda kv: -kv[1])[:10]
    r.chart_title = "Виручка за брендами"
    r.finish()
    t = r.total or {}
    known = [z for z in r.rows if z["margin"] is not None]
    ks = sum(z["sum"] for z in known)
    r.tiles = [("Виручка товарів", t.get("sum", 0), sum(x["sum"] for x in prev.values()), "uah", ""),
               ("Продано, шт", t.get("qty", 0), sum(x["qty"] for x in prev.values()), "num", ""),
               ("Валовий прибуток", t.get("profit", 0), None, "uah",
                f"маржа {round(100 * sum(z['profit'] for z in known) / ks, 1)}%" if ks else ""),
               ("Без собівартості", len(r.rows) - len(known), None, "int", "позицій — внесіть закупівлі")]
    return r


# ---------------- Пацієнти ----------------

def patients(db: Session, d_from: date, d_to: date) -> Report:
    a, b = _rng(d_from, d_to)
    first = dict(db.execute(select(Check.patient_id, func.min(Check.date)).where(
        Check.deleted.is_(False), Check.patient_id.is_not(None)).group_by(Check.patient_id)).all())
    with_plan = {pid for (pid,) in db.execute(select(BeautyPlan.patient_id).where(BeautyPlan.status != "cancelled"))}
    agg = defaultdict(lambda: {"sum": 0.0, "checks": 0, "services": 0.0, "products": 0.0, "last": None})
    for c in _checks(db, d_from, d_to):
        if not c.patient_id:
            continue
        x = agg[c.patient_id]
        x["sum"] += c.amount or 0
        x["checks"] += 1
        x["last"] = max(x["last"] or c.date, c.date)
        for kind, _, _, amt, _ in _items(c):
            if kind == "service":
                x["services"] += amt
            elif kind == "product":
                x["products"] += amt
    new = {pid for pid in agg if first.get(pid) and first[pid] >= a}
    pa, pb = prev_period(d_from, d_to)
    prev_ids = {c.patient_id for c in _checks(db, pa, pb) if c.patient_id}
    names = {p.id: p.full_name for p in db.scalars(select(Patient).where(Patient.id.in_(list(agg) or [0])))}
    r = Report("patients", "Пацієнти", "Пацієнти з оплатою за період. Нова — перша оплата в CleverBox припала на цей період.")
    r.cols = [Col("name", "Пацієнт"), Col("type", "Тип"), Col("plan", "Бюті-план"), Col("checks", "Чеків", "int", True),
              Col("services", "Послуги", "uah", True), Col("products", "Догляд", "uah", True), Col("sum", "Разом", "uah", True),
              Col("last", "Остання оплата", "date")]
    for pid, x in agg.items():
        r.rows.append({"name": names.get(pid, f"#{pid}"), "type": "нова" if pid in new else "повторна",
                       "plan": "так" if pid in with_plan else "—", "checks": x["checks"], "services": x["services"],
                       "products": x["products"], "sum": x["sum"], "last": x["last"], "_href": f"/patients/{pid}"})
    r.rows.sort(key=lambda z: -z["sum"])
    n = len(agg)
    tot = sum(x["sum"] for x in agg.values())
    plan_sum = sum(x["sum"] for pid, x in agg.items() if pid in with_plan)
    r.tiles = [("Пацієнтів з оплатою", n, len(prev_ids), "int", ""),
               ("Нових", len(new), None, "int", f"{round(100 * len(new) / n)}%" if n else ""),
               ("Повторних", n - len(new), None, "int", f"повернулись з минулого періоду: {len(prev_ids & set(agg))}"),
               ("Виручка на пацієнта", round(tot / n) if n else 0, None, "uah", ""),
               ("Пацієнти з бюті-планом", plan_sum, None, "uah", f"{round(100 * plan_sum / tot)}% виручки" if tot else "")]
    r.chart = [("нові", sum(agg[p]["sum"] for p in new)), ("повторні", tot - sum(agg[p]["sum"] for p in new))]
    r.chart_title = "Виручка: нові / повторні"
    return r.finish()


# ---------------- Бюті-плани ----------------

def plans(db: Session, d_from: date, d_to: date) -> Report:
    from .kpi import all_statuses, kpi_summary
    total, by_doc = kpi_summary(all_statuses(db, d_from, d_to))
    pt, _ = kpi_summary(all_statuses(db, *prev_period(d_from, d_to)))
    r = Report("plans", "Бюті-плани", "Плани з датою консультації в періоді. Реалізація — виконані процедури й куплений "
               "догляд за даними CleverBox (на сьогодні).")
    r.tiles = [("Бюті-планів", total.plans, pt.plans, "int", ""),
               ("Сума планів", total.planned_sum, pt.planned_sum, "uah", f"сер. {total.avg_plan:,.0f} ₴".replace(",", " ")),
               ("Реалізовано", total.realized_sum, pt.realized_sum, "uah", f"{total.realization_pct}%"),
               ("Почали виконувати", total.started_plans, pt.started_plans, "int", f"конверсія {total.conversion_pct}%"),
               ("Записані далі", total.booked_plans, None, "int", f"{total.booked_pct}% активних")]
    r.cols = [Col("name", "Лікар"), Col("plans", "Планів", "int", True), Col("planned", "Сума планів", "uah", True),
              Col("avg", "Сер. план", "uah"), Col("realized", "Реалізовано", "uah", True), Col("real_pct", "% реалізації", "pct"),
              Col("conv", "Конверсія", "pct"), Col("booked", "Записані", "int", True), Col("red", "Без руху", "int", True),
              Col("proc", "Процедур виконано", "text")]
    for k in by_doc:
        r.rows.append({"name": k.label, "plans": k.plans, "planned": k.planned_sum, "avg": k.avg_plan,
                       "realized": k.realized_sum, "real_pct": k.realization_pct, "conv": k.conversion_pct,
                       "booked": k.booked_plans, "red": k.red, "proc": f"{k.proc_done} з {k.proc_planned}"})
    r.chart = [(z["name"][:18], z["planned"]) for z in r.rows]
    r.chart_title = "Сума планів за лікарями"
    return r.finish()


# ---------------- Поставки ----------------

def supply(db: Session, d_from: date, d_to: date) -> Report:
    a, b = _rng(d_from, d_to)
    inv = db.scalars(select(SupplierInvoice).where(SupplierInvoice.inv_date >= a, SupplierInvoice.inv_date <= b)).all()
    pa, pb = _rng(*prev_period(d_from, d_to))
    prev = db.scalars(select(SupplierInvoice).where(SupplierInvoice.inv_date >= pa, SupplierInvoice.inv_date <= pb)).all()
    agg = defaultdict(lambda: {"n": 0, "sum": 0.0, "paid": 0.0, "last": None, "wh": set()})
    for i in inv:
        x = agg[i.supplier or "—"]
        x["n"] += 1
        x["sum"] += i.amount or 0
        x["paid"] += i.paid or 0
        x["last"] = max(x["last"] or i.inv_date, i.inv_date)
        if i.warehouse_name:
            x["wh"].add(i.warehouse_name.strip())
    r = Report("supply", "Поставки", "З імпорту «Поставки» CleverBox (Склад → Постачальники). Борг — за поставками періоду.",
               empty="Поставок за період немає — завантажте вивантаження «Поставки» в Склад → Постачальники.")
    r.cols = [Col("name", "Постачальник"), Col("n", "Поставок", "int", True), Col("sum", "Сума", "uah", True),
              Col("paid", "Оплачено", "uah", True), Col("debt", "Борг", "uah", True), Col("share", "Частка", "pct"),
              Col("last", "Остання", "date"), Col("wh", "Склади")]
    tot = sum(x["sum"] for x in agg.values()) or 1
    for name, x in agg.items():
        r.rows.append({"name": name, "n": x["n"], "sum": x["sum"], "paid": x["paid"], "debt": max(x["sum"] - x["paid"], 0),
                       "share": round(100 * x["sum"] / tot, 1), "last": x["last"], "wh": ", ".join(sorted(x["wh"]))})
    r.rows.sort(key=lambda z: -z["sum"])
    r.chart = [(z["name"][:14], z["sum"]) for z in r.rows[:10]]
    r.chart_title = "Закупівлі за постачальниками"
    r.finish()
    t = r.total or {}
    r.tiles = [("Закуплено", t.get("sum", 0), sum(i.amount or 0 for i in prev), "uah", f"{len(inv)} поставок"),
               ("Оплачено", t.get("paid", 0), None, "uah", ""),
               ("Борг за період", t.get("debt", 0), None, "uah", ""),
               ("Постачальників", len(agg), len({i.supplier for i in prev}), "int", "")]
    return r


# ---------------- Сертифікати ----------------

def certificates(db: Session, d_from: date, d_to: date) -> Report:
    from ..models import GiftCertificate
    from .certificates import STATES
    today = date.today()
    allc = db.scalars(select(GiftCertificate)).all()
    a, b = _rng(d_from, d_to)
    pa, pb = _rng(*prev_period(d_from, d_to))
    sold = [c for c in allc if c.sold_at and a <= c.sold_at <= b and c.state(today) != "unpaid"]
    sold_prev = [c for c in allc if c.sold_at and pa <= c.sold_at <= pb and c.state(today) != "unpaid"]
    live = [c for c in allc if c.state(today) in ("active", "expiring")]
    burned = [c for c in allc if c.state(today) == "expired" and c.valid_to and d_from <= c.valid_to <= d_to]
    r = Report("certificates", "Сертифікати", "З вивантаження CleverBox «Сертифікати» (Дані → Імпорт). Таблиця — "
               "сертифікати, продані за період, плюс усі активні зараз. «Згоріло» — залишок на сертифікатах, "
               "термін яких минув у періоді.", empty="Сертифікатів ще немає — завантажте вивантаження «Сертифікати» з CleverBox.")
    r.tiles = [("Продано", sum(c.paid for c in sold), sum(c.paid for c in sold_prev), "uah", f"{len(sold)} шт"),
               ("Активні зараз", sum(c.balance for c in live), None, "uah",
                f"{len(live)} шт — зобов'язання клініки"),
               ("Згорають за 14 дн", sum(c.balance for c in live if c.state(today) == "expiring"), None, "uah",
                f"{sum(1 for c in live if c.state(today) == 'expiring')} шт — задачі адмінам створено"),
               ("Згоріло з залишком", sum(c.balance for c in burned), None, "uah", f"{len(burned)} шт у періоді")]
    from ..models import Subscription
    subs = [x for x in db.scalars(select(Subscription)) if x.state(today) in ("active", "expiring")]
    if subs:
        r.tiles.append(("Абонементи: невикористано", sum(x.value_left for x in subs), None, "uah",
                        f"{len(subs)} активних абонементів"))
    r.cols = [Col("code", "№"), Col("type", "Сертифікат"), Col("client", "Пацієнт"), Col("seller", "Продав(ла)"),
              Col("sold", "Продано", "date"), Col("nominal", "Номінал", "uah", True), Col("balance", "Залишок", "uah", True),
              Col("used", "Використано", "uah", True), Col("valid", "Діє до", "date"), Col("state", "Стан")]
    seen = set()
    for c in sorted(sold + live, key=lambda c: c.sold_at or datetime.min, reverse=True):
        if c.id in seen:
            continue
        seen.add(c.id)
        r.rows.append({"code": c.code or str(c.id), "type": c.type_name or "", "client": c.client_name or "—",
                       "seller": c.seller or "", "sold": c.sold_at, "nominal": c.nominal, "balance": c.balance,
                       "used": max(c.nominal - c.balance, 0), "valid": c.valid_to, "state": STATES[c.state(today)],
                       "_href": f"/patients/{c.patient_id}" if c.patient_id else None})
    by_seller = defaultdict(float)
    for c in sold:
        by_seller[c.seller or "—"] += c.paid
    r.chart = sorted(by_seller.items(), key=lambda kv: -kv[1])
    r.chart_title = "Продажі сертифікатів за адміністраторами"
    return r.finish()


BUILDERS = {"revenue": revenue, "doctors": doctors, "services": services, "products": products,
            "patients": patients, "plans": plans, "certificates": certificates, "supply": supply}


def build(db: Session, key: str, d_from: date, d_to: date) -> Report:
    return BUILDERS.get(key, revenue)(db, d_from, d_to)


def to_xlsx(rep: Report, d_from: date, d_to: date) -> bytes:
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = rep.title[:30]
    ws.append([f"{rep.title} · {d_from:%d.%m.%Y} — {d_to:%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append([c.label for c in rep.cols])
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3B3050")
    for row in rep.rows:
        ws.append([_xl(row.get(c.key), c.kind) for c in rep.cols])
    if rep.total:
        ws.append(["Разом"] + [_xl(rep.total.get(c.key), c.kind) if c.key in rep.total else None for c in rep.cols[1:]])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for i, c in enumerate(rep.cols, 1):
        ws.column_dimensions[ws.cell(3, i).column_letter].width = 42 if i == 1 else (24 if c.kind == "text" else 15)
        if c.kind in ("uah", "num"):
            for (cell,) in ws.iter_rows(min_row=4, min_col=i, max_col=i):
                cell.number_format = "#,##0" if c.kind == "uah" else "#,##0.##"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xl(v, kind):
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    return v
