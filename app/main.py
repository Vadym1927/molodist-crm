"""MOLODIST Beauty Plan CRM — веб-інтерфейс (з ролями), REST API, планувальник синхронізації та Telegram-бот."""
from __future__ import annotations

import asyncio
from collections import defaultdict
import json
import re
import logging
import mimetypes
import secrets
from contextlib import asynccontextmanager
from datetime import datetime, date
import httpx
from pathlib import Path

from fastapi import FastAPI, Request, Form, UploadFile, File, HTTPException, Depends
from fastapi.openapi.docs import get_swagger_ui_html
from fastapi.responses import HTMLResponse, Response, RedirectResponse, FileResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from .config import settings
from .models import (init_db, SessionLocal, Patient, Doctor, BeautyPlan, PlanItem, PatientFile, Alias,
                     CbService, Visit, SyncState, User, ApiKey, Webhook, AuditLog, ROLES, CashDay)
from .parsers.observ import norm
from .security import (can, hash_password, verify_password, ensure_owner, db_dep, current_user, require,
                       patient_visible, plan_visible, new_api_key, audit)
from .services import ingest, reminders
from .services.kpi import compute_plan, all_statuses, kpi_summary, forecast, name_match
from .services.serialize import plan_dict
from .integrations import webhooks
from . import api as api_module

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(name)s %(levelname)s %(message)s")
log = logging.getLogger("crm")
BASE = Path(__file__).parent
templates = Jinja2Templates(directory=str(BASE / "templates"))
templates.env.filters["uah"] = lambda v: f"{(v or 0):,.2f}".replace(",", " ").replace(".", ",") + " ₴"   # копійка в копійку
templates.env.filters["dt"] = lambda v, f="%d.%m.%Y": v.strftime(f) if v else "—"


def _local(v):
    """UTC (datetime.utcnow у БД) → київський час для показу."""
    from datetime import timezone
    from zoneinfo import ZoneInfo
    return v.replace(tzinfo=timezone.utc).astimezone(ZoneInfo(settings.timezone)).replace(tzinfo=None) if v else v


templates.env.filters["local"] = _local
FIN_TABS = [("/finance", "Огляд"), ("/sales", "Каса CRM"), ("/finance/ops", "Фіноперації"), ("/finance/payroll", "Нарахування зарплати"),
            ("/finance/settle", "Взаєморозрахунки по з/п"), ("/finance/day", "Продажі за день"), ("/finance/sales", "Продажі"),
            ("/finance/balances", "Баланс пацієнтів"), ("/finance/bonuses", "Рух бонусів"), ("/finance/pricing", "Планування ціни реалізації")]
templates.env.globals.update(
    can=can, ROLES=ROLES, EVENTS=webhooks.EVENTS, FIN_TABS=FIN_TABS,
    FLAG={"green": ("записана", "ok"), "yellow": ("без запису", "warn"),
          "red": ("прострочено", "bad"), "done": ("завершено", "done")})

def _pay_accounts() -> list[dict]:
    """Готові підключення платіжних систем — для вибору ФОП у формі посилання на оплату."""
    from .services import paylinks as _pls
    try:
        with SessionLocal() as _db:
            t_ = _pls.public(_db)
            return [{"id": a_["id"], "name": a_["name"], "default": int(t_.get("default") or 0) == int(a_["id"])} for a_ in t_["accounts"] if a_["ready"]]
    except Exception:  # noqa: BLE001
        return []


templates.env.globals["pay_accounts"] = _pay_accounts

state = {"bot": None, "last_sync": None, "last_error": None}


def _set_sync_state(key: str, value: str | None) -> None:
    with SessionLocal() as db:
        item = db.get(SyncState, key) or SyncState(key=key)
        item.value = value
        db.merge(item)
        db.commit()


def auto_sync_enabled() -> bool:
    """Auto-sync stays enabled unless an owner explicitly turns it off."""
    with SessionLocal() as db:
        item = db.get(SyncState, "cleverbox_auto_sync")
        return item is None or item.value != "0"


def ensure_sync_schedule() -> None:
    """Keep one automatic CleverBox job, including after a token is set in the UI."""
    sched = state.get("scheduler")
    if not sched:
        return
    if settings.cleverbox_token and auto_sync_enabled():
        sched.add_job(job_sync, "cron", hour="*/2", minute=5,
                      id="cleverbox_sync", replace_existing=True)
    elif sched.get_job("cleverbox_sync"):
        sched.remove_job("cleverbox_sync")


# ---------------- фонові задачі ----------------

def run_sync() -> dict:
    from .integrations.cleverbox import full_sync
    from .integrations.cb_webhooks import SYNC_LOCK
    with SYNC_LOCK, SessionLocal() as db:
        rep = full_sync(db)
    state["last_sync"] = (datetime.now(), rep)
    from .services.tasks import generate_auto
    from .services import stock as stock_svc
    with SessionLocal() as db:
        rep["tasks"] = generate_auto(db)
        try:
            rep["stock_auto"] = stock_svc.auto_moves(db)      # продажі й процедури списують склад
        except Exception as e:      # noqa: BLE001
            log.warning("stock auto moves failed: %s", e)
        try:
            from .services import cashflow as _cf
            rep["cashflow_days"] = _cf.auto_fill(db)        # нові дні — надходження з чеків до наступного звіту
        except Exception as e:      # noqa: BLE001
            log.warning("cashflow auto fill failed: %s", e)
    webhooks.fire("sync.done", {"report": {k: v for k, v in rep.items() if k != "patients"},
                                "patients_linked": rep.get("patients", {}).get("linked")})
    if settings.google_sheet_id and settings.google_service_account_file:
        try:
            run_sheets()
        except Exception as e:      # noqa: BLE001
            log.warning("Sheets export failed: %s", e)
    return rep


def run_sheets() -> str:
    from .integrations.sheets import export
    with SessionLocal() as db:
        return export(db)


async def job_sync():
    from .services import alerts
    try:
        rep = await asyncio.to_thread(run_sync)
        state["last_error"] = None
        _set_sync_state("cleverbox_last_success", datetime.now().isoformat(timespec="seconds"))
        _set_sync_state("cleverbox_last_error", "")
        partial = [f"{k}: {v}" for k, v in rep.items() if isinstance(v, str) and v.startswith("помилка")]
        if partial:      # основний синк пройшов, але окремий модуль (склад/чеки/календар) — ні
            await asyncio.to_thread(alerts.failure, "sync_part", "; ".join(partial),
                                    "Можливо, модуль вимкнений у тарифі CleverBox або немає прав у токена.")
        else:
            await asyncio.to_thread(alerts.success, "sync_part")
        await asyncio.to_thread(alerts.success, "sync")
    except Exception as e:          # noqa: BLE001
        state["last_error"] = f"{datetime.now():%d.%m %H:%M} {e}"
        _set_sync_state("cleverbox_last_error", state["last_error"])
        log.exception("sync failed")
        await asyncio.to_thread(alerts.failure, "sync", str(e), alerts.hint_for(str(e)))


async def job_campaigns():
    """Запуск запланованих розсилок і продовження перерваних (після рестарту / тихих годин)."""
    from .services import campaigns as camp, jobs
    from .models import Campaign
    try:
        with SessionLocal() as db:
            camp.run_due(db)
            ids = [c.id for c in db.scalars(select(Campaign).where(Campaign.status == "sending"))]
        for cid in ids:
            jobs.start(f"Розсилка #{cid}", camp.deliver, campaign_id=cid)
    except Exception:              # noqa: BLE001
        log.exception("campaigns job failed")


async def job_message_statuses():
    from .services import campaigns as camp
    def _refresh():
        with SessionLocal() as db:
            return camp.refresh_statuses(db)
    try:
        await asyncio.to_thread(_refresh)
    except Exception:              # noqa: BLE001
        log.exception("statuses refresh failed")



async def job_shop():
    """molodist.store (Shop-Express): щоранку — посилання, фото й ціни товарів у картки CRM."""
    from .integrations import shop_express as sx
    def run():
        with SessionLocal() as db:
            if sx.creds(db).get("login"):
                return sx.sync(db)
    try:
        await asyncio.to_thread(run)
    except Exception as e:  # noqa: BLE001
        log.warning("shop-express sync: %s", e)


async def job_expiry():
    """Терміни придатності: щодня — нові межі 30/7/0 днів, у вибраний день — підсумок по складах."""
    from .services import expiry as ex
    def run():
        with SessionLocal() as db:
            return ex.daily(db)
    try:
        await asyncio.to_thread(run)
    except Exception as e:  # noqa: BLE001
        log.warning("expiry notify: %s", e)


async def job_loyalty():
    """Щодня: бонуси до дня народження, згоряння бонусів з минулим терміном."""
    from .services import loyalty as ly
    def run():
        with SessionLocal() as db:
            if ly.settings(db).get("enabled"):
                return {"birthday": ly.birthday_bonuses(db), "expired": ly.expire(db)}
    try:
        await asyncio.to_thread(run)
    except Exception as e:  # noqa: BLE001
        log.warning("loyalty: %s", e)


async def job_rematch():
    """Перерозносити платежі, що вже є в CRM (без запиту в monobank): нові правила, частково рознесені."""
    from .services import payments as pay
    def run():
        with SessionLocal() as db:
            return pay.rematch_all(db)
    try:
        n = await asyncio.to_thread(run)
        log.info("rematch: %s", n)
    except Exception as e:  # noqa: BLE001
        log.warning("rematch: %s", e)


async def job_mono():
    """monobank: щогодини — списання за 2 доби, рознесення оплат постачальникам."""
    from .integrations import monobank
    def run():
        with SessionLocal() as db:
            if monobank.saved(db):
                return monobank.sync(db, days=2)
    try:
        await asyncio.to_thread(run)
    except Exception as e:  # noqa: BLE001
        log.warning("monobank sync: %s", e)


async def job_checkbox():
    """Оборот ФОП з Checkbox (Z-звіти) — поточний і попередній місяць."""
    from .integrations import checkbox as chb
    if not any(a["configured"] for a in chb.accounts()):
        return
    def run():
        with SessionLocal() as db:
            return chb.sync(db)
    try:
        await asyncio.to_thread(run)
    except Exception as e:  # noqa: BLE001
        log.warning("checkbox sync: %s", e)

async def job_backup():
    from .services import backups
    try:
        await asyncio.to_thread(backups.run_backup, "auto")
    except Exception:              # noqa: BLE001
        pass                        # сповіщення вже надіслано в run_backup


async def job_digest():
    from .services.tasks import generate_auto
    from .services import alerts
    try:
        with SessionLocal() as db:
            generate_auto(db)
            rems = reminders.collect(db)
        webhooks.fire("reminders.daily", {"items": [{"kind": r.kind, "patient_id": r.patient_id, "text_html": r.text}
                                                    for r in rems]})
        if state["bot"]:
            from .integrations.telegram_bot import send_digest
            await send_digest(state["bot"])
        await asyncio.to_thread(alerts.success, "digest")
    except Exception as e:          # noqa: BLE001
        log.exception("digest failed")
        await asyncio.to_thread(alerts.failure, "digest", str(e))


@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    with SessionLocal() as db:
        ensure_owner(db)
        from .integrations.cleverbox import load_saved_token
        load_saved_token(db)
        from .services import clinic as _clinic
        _clinic.apply(_clinic.get(db))                      # часовий пояс з «Налаштування → Загальні»
        from .services import roles as _roles
        _roles.apply(db)                                    # свої ролі й права з «Права доступу»
        from .services import channels as _channels
        _channels.apply(db)
        from .services import cashboxes as _cbx
        _cbx.apply_seed(db)                                 # залишки кас на 30.10.2026 (один раз)                                 # канали (TurboSMS, бот) з «Налаштування → Канали»
        from .services.scripts import seed_defaults
        seed_defaults(db)
    tasks, sched = [], None
    try:
        from apscheduler.schedulers.asyncio import AsyncIOScheduler
        sched = AsyncIOScheduler(timezone=settings.timezone)
        state["scheduler"] = sched
        ensure_sync_schedule()
        sched.add_job(job_digest, "cron", day_of_week="mon-sat", hour=settings.remind_hour, minute=0)
        sched.add_job(job_backup, "cron", hour=settings.backup_hour, minute=30)
        sched.add_job(job_campaigns, "interval", minutes=5)
        sched.add_job(job_message_statuses, "interval", minutes=30)
        sched.add_job(job_checkbox, "cron", hour="8,14,23", minute=50)
        sched.add_job(job_mono, "cron", minute=20)
        sched.add_job(job_rematch, "date", run_date=datetime.now() + timedelta(seconds=40))   # після оновлення — одразу
        sched.add_job(job_expiry, "cron", hour=9, minute=15)
        sched.add_job(job_loyalty, "cron", hour=9, minute=5)
        sched.add_job(job_shop, "cron", hour=6, minute=40)
        sched.start()
    except Exception as e:          # noqa: BLE001
        log.warning("scheduler not started: %s", e)
    if settings.telegram_bot_token:
        from .integrations.telegram_bot import make_bot
        bot, dp = make_bot()
        state["bot"] = bot
        tasks.append(asyncio.create_task(dp.start_polling(bot, handle_signals=False)))
    yield
    for t in tasks:
        t.cancel()
    if sched:
        sched.shutdown(wait=False)
    state["scheduler"] = None


app = FastAPI(title="MOLODIST Beauty Plan CRM API", version="1.0", lifespan=lifespan,
              docs_url=None, redoc_url=None, openapi_url=None,
              description=api_module.__doc__)
class _DropEmptyQuery:
    """Порожні параметри з форм фільтрів (?doctor=&page=) — як відсутні, а не помилка 422."""
    def __init__(self, app_):
        self.app = app_

    async def __call__(self, scope, receive, send):
        if scope["type"] == "http" and scope.get("query_string") and not scope["path"].startswith("/api/"):
            from urllib.parse import parse_qsl, urlencode
            pairs = parse_qsl(scope["query_string"].decode("latin-1"), keep_blank_values=True)
            kept = [(k, v) for k, v in pairs if v != ""]
            if len(kept) != len(pairs):
                scope = dict(scope, query_string=urlencode(kept).encode("latin-1"))
        await self.app(scope, receive, send)


app.add_middleware(_DropEmptyQuery)
app.add_middleware(SessionMiddleware, secret_key=settings.secret_key, max_age=60 * 60 * 24 * 14,
                   https_only=settings.public_url.startswith("https"))
app.include_router(api_module.router)
from fastapi.staticfiles import StaticFiles   # noqa: E402
app.mount("/static", StaticFiles(directory=str(BASE / "static")), name="static")


@app.exception_handler(HTTPException)
async def http_exc(request: Request, exc: HTTPException):
    if exc.status_code == 303:
        return RedirectResponse(exc.headers["Location"], 303)
    if request.url.path.startswith("/api/"):
        return JSONResponse({"detail": exc.detail}, status_code=exc.status_code, headers=exc.headers)
    user = getattr(request.state, "user", None)
    return templates.TemplateResponse(request, "error.html", {"request": request, "user": user, "nav": "",
                                      "code": exc.status_code, "detail": exc.detail}, status_code=exc.status_code)


def render(request: Request, name: str, **ctx):
    user = getattr(request.state, "user", None)
    counts = None
    if user:
        from .services.tasks import counts_for
        with SessionLocal() as db:
            counts = counts_for(db, user)
    return templates.TemplateResponse(request, name, {"request": request, "nav": name.split(".")[0],
                                                      "user": user, "task_counts": counts, **ctx})


def doctor_scope(user: User) -> int | None:
    """Лікар бачить лише свої плани; власник/адмін — усі."""
    return user.doctor_id if user.role == "doctor" else None


# ---------------- вхід ----------------

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    return render(request, "login.html")


@app.post("/login")
def login(request: Request, login: str = Form(...), password: str = Form(...), db: Session = Depends(db_dep)):
    from .security import login_limiter, ip_limiter, client_ip
    ip, lg = client_ip(request), login.strip().lower()
    left = max(login_limiter.blocked_for(f"login:{lg}"), ip_limiter.blocked_for(f"ip:{ip}"))
    if left:
        return render(request, "login.html", login=login,
                      error=f"Забагато невдалих спроб. Вхід тимчасово заблоковано — спробуйте через {left} хв.")
    u = db.scalar(select(User).where(func.lower(User.login) == lg))
    if not u or not u.active or not verify_password(password, u.password_hash):
        just_blocked = login_limiter.fail(f"login:{lg}") | ip_limiter.fail(f"ip:{ip}")
        audit(db, "login.failed", user=u if u else None, login=lg[:64], ip=ip)
        if just_blocked:
            audit(db, "login.blocked", login=lg[:64], ip=ip, minutes=settings.login_block_minutes)
            import threading
            from .services import alerts
            threading.Thread(target=alerts.send, daemon=True, args=(
                f"🔐 <b>Безпека: блокування входу</b>\n{settings.login_max_attempts} невдалих спроб — логін «{lg[:40]}», "
                f"IP {ip}. Вхід заблоковано на {settings.login_block_minutes} хв.",)).start()
            return render(request, "login.html", login=login,
                          error=f"Забагато невдалих спроб. Вхід заблоковано на {settings.login_block_minutes} хв.")
        return render(request, "login.html", error="Невірний логін або пароль", login=login)
    login_limiter.reset(f"login:{lg}")
    request.session.clear()
    request.session["uid"] = u.id
    u.last_login = datetime.utcnow()
    db.commit()
    audit(db, "login", user=u)
    return RedirectResponse("/", 303)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", 303)


@app.get("/me", response_class=HTMLResponse)
def me(request: Request, user: User = Depends(current_user)):
    return render(request, "me.html")


@app.post("/me")
def me_save(request: Request, current: str = Form(""), new: str = Form(""), telegram_chat_id: str = Form(""),
            user: User = Depends(current_user), db: Session = Depends(db_dep)):
    u = db.get(User, user.id)
    u.telegram_chat_id = telegram_chat_id.strip() or None
    msg = "Збережено"
    if new:
        if len(new) < 8:
            msg = "Новий пароль — щонайменше 8 символів"
        elif not verify_password(current, u.password_hash):
            msg = "Поточний пароль невірний"
        else:
            u.password_hash = hash_password(new)
            msg = "Пароль змінено"
            audit(db, "user.password_change", "user", u.id, user=u)
    db.commit()
    request.state.user = u
    return render(request, "me.html", msg=msg)


# ---------------- дашборд ----------------

@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request, df: str | None = None, dt_: str | None = None,
              user: User = Depends(current_user), db: Session = Depends(db_dep)):
    d_from = date.fromisoformat(df) if df else None
    d_to = date.fromisoformat(dt_) if dt_ else None
    doc = doctor_scope(user)
    if user.role == "doctor" and not doc:
        return render(request, "error.html", code=403,
                      detail="Ваш обліковий запис ще не привʼязаний до лікаря. Зверніться до власника.")
    sts = all_statuses(db, d_from, d_to, doctor_id=doc)
    total, by_doc = kpi_summary(sts)
    ctx = dict(total=total, by_doc=by_doc, df=df, dt_=dt_)
    if can(user, "finance"):
        fc = forecast(all_statuses(db))
        ctx.update(fc=fc, fc_max=max([r["total"] for r in fc] + [1]))
    if can(user, "reminders") or user.role == "doctor":
        allowed_patients = {ps.plan.patient_id for ps in sts} if doc else None
        rems = [r for r in reminders.collect(db) if allowed_patients is None or r.patient_id in allowed_patients]
        ctx.update(not_booked=[r for r in rems if r.kind == "not_booked"],
                   care=[r for r in rems if r.kind == "care_ending"])
    last = db.get(SyncState, "last_sales_sync")
    ctx.update(last_sync=last.value if last else None, last_error=state["last_error"],
               sheets=bool(settings.google_sheet_id), cb=bool(settings.cleverbox_token),
               unassigned=db.scalar(select(func.count()).select_from(PatientFile)
                                    .where(PatientFile.patient_id.is_(None))) if can(user, "files_assign") else 0)
    return render(request, "dashboard.html", **ctx)


# ---------------- пацієнти ----------------

def _patient_or_403(db, user, pid) -> Patient:
    p = db.get(Patient, pid)
    if not p:
        raise HTTPException(404, "Пацієнта не знайдено")
    if not patient_visible(user, p):
        raise HTTPException(403, "Це не ваш пацієнт")
    return p


@app.get("/patients/{pid}", response_class=HTMLResponse)
def patient_card(request: Request, pid: int, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    statuses = [compute_plan(db, pl) for pl in p.plans if plan_visible(user, pl)]
    completed_visits = [v for v in p.visits if v.status == "done"]
    planned_visits = [v for v in p.visits if v.status == "booked" and v.starts_at >= datetime.now()]
    dated_events = [v.starts_at for v in p.visits if v.starts_at]
    service_revenue = sum(v.amount or 0 for v in completed_visits)
    product_revenue = sum(s.amount or 0 for s in p.sales)
    from .models import Task as _Task
    from .services.tasks import visible_query as _vq
    ptasks = db.scalars(_vq(user).where(_Task.patient_id == p.id).order_by(_Task.status, _Task.due_date)).all()
    from .services import scripts as _sc
    from .integrations import messaging as _msg
    from .models import MessageLog as _ML
    msgs = db.scalars(select(_ML).where(_ML.patient_id == p.id).order_by(_ML.at.desc()).limit(30)).all()
    kind = "not_booked" if statuses and statuses[0].next_procedure else "follow_up"
    from .services import patient_match as _pm
    dupes = _pm.suggestions(db, p) if not p.cleverbox_id and can(user, "patients_all") else []
    from .services import certificates as _cert
    from .models import PayLink as _PL
    paylinks = db.scalars(select(_PL).where(_PL.patient_id == p.id).order_by(_PL.id.desc()).limit(10)).all() if can(user, "sales") else []
    return render(request, "patient.html", p=p, statuses=statuses, ptasks=ptasks, msgs=msgs, dupes=dupes, paylinks=paylinks,
                  certs=_cert.for_patient(db, p.id), CERT_STATES=_cert.STATES, today=date.today(),
                  subs=_cert.subs_for_patient(db, p.id), SUB_STATES=_cert.SUB_STATES,
                  CB_LABELS=CB_LABELS, loyal=_loyal_ctx(db, p),
                  pscripts=_sc.for_patient(db, p, kind, user.name) if can(user, "message") else [],
                  avail=_msg.available_channels(p),
                  users=db.scalars(select(User).where(User.active.is_(True)).order_by(User.name)).all(),
                  doctors=db.scalars(select(Doctor).order_by(Doctor.name)).all(),
                  photos=[f for f in p.files if f.kind == "photo"],
                  docs=[f for f in p.files if f.kind != "photo"], now=datetime.now(),
                  patient_kpi=dict(done=len(completed_visits), planned=len(planned_visits),
                                   first=min(dated_events) if dated_events else None,
                                   next=min((v.starts_at for v in planned_visits), default=None),
                                   service_revenue=service_revenue, product_revenue=product_revenue,
                                   revenue=service_revenue + product_revenue))


def _loyal_ctx(db: Session, p: Patient) -> dict:
    from .services import loyalty as ly
    from .models import BonusTxn
    return {"s": ly.settings(db), "balance": ly.balance(db, p.id), "KL": ly.KIND_LABELS,
            "txns": db.scalars(select(BonusTxn).where(BonusTxn.patient_id == p.id).order_by(BonusTxn.at.desc()).limit(30)).all()}


CB_LABELS = {"sex": "Стать", "gender": "Стать", "address": "Адреса", "city": "Місто", "source": "Джерело",
             "source_name": "Джерело", "groups": "Групи", "categories": "Категорії", "category_name": "Категорія",
             "tags": "Теги", "balance": "Баланс", "deposit": "Депозит", "bonus": "Бонуси", "discount": "Знижка, %",
             "card": "Картка лояльності", "card_number": "Картка лояльності", "comment": "Коментар", "note": "Нотатка",
             "notes": "Нотатки", "instagram": "Instagram", "telegram": "Telegram", "viber": "Viber",
             "status": "Статус", "vip": "VIP", "visits": "Візитів", "last_visit": "Останній візит", "ltv": "LTV",
             "total_paid": "Сплачено всього", "average_check": "Середній чек", "master": "Майстер", "master_name": "Майстер",
             "filial": "Філія", "filial_name": "Філія", "blacklist": "Чорний список", "sms": "SMS-розсилка",
             "notification": "Сповіщення", "fields": "Додаткові поля", "custom_fields": "Додаткові поля"}


@app.post("/patients/{pid:int}/merge")
def patient_merge(pid: int, target: int = Form(...), user: User = Depends(require("patients_all")),
                  db: Session = Depends(db_dep)):
    """Картка-дубль (напр. створена з PDF без CleverBox) → у картку CleverBox; дубль видаляється."""
    from .services import patient_match as _pm
    src, dst = db.get(Patient, pid), db.get(Patient, target)
    if not src or not dst:
        raise HTTPException(404)
    name = src.full_name
    moved = _pm.merge(db, src, dst)
    audit(db, "patient.merge", "patient", dst.id, user=user, merged_from=pid, merged_name=name, moved=moved)
    return RedirectResponse(f"/patients/{dst.id}?merged=1", 303)


@app.post("/patients/{pid:int}")
def patient_update(pid: int, full_name: str = Form(...), phone: str = Form(""), cleverbox_id: str = Form(""),
                   doctor_id: str = Form(""), notes: str = Form(""), email: str = Form(""),
                   gender: str = Form(""), address: str = Form(""), source: str = Form(""),
                   communication_channel: str = Form(""), client_groups: str = Form(""),
                   cb_balance: str = Form(""), cb_bonus: str = Form(""),
                   user: User = Depends(require("patient_edit")), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    p.full_name, p.phone, p.notes = full_name.strip(), phone.strip() or None, notes.strip() or None
    p.cleverbox_id = int(cleverbox_id) if cleverbox_id.strip().isdigit() else None
    p.email, p.gender, p.address = email.strip() or None, gender.strip() or None, address.strip() or None
    p.source, p.communication_channel = source.strip() or None, communication_channel.strip() or None
    p.client_groups = [x.strip() for x in client_groups.split(",") if x.strip()] or None
    def money(value: str):
        try:
            return float(value.strip().replace(" ", "").replace(",", ".")) if value.strip() else None
        except ValueError:
            return None
    p.cb_balance, p.cb_bonus = money(cb_balance), money(cb_bonus)
    if user.role != "doctor":
        p.doctor_id = int(doctor_id) if doctor_id.isdigit() else None
    db.commit()
    audit(db, "patient.update", "patient", p.id, user=user)
    return RedirectResponse(f"/patients/{pid}", 303)


@app.post("/patients")
def patient_create(full_name: str = Form(...), phone: str = Form(""), doctor_id: str = Form(""),
                   notes: str = Form(""), user: User = Depends(require("patient_edit")),
                   db: Session = Depends(db_dep)):
    """Ручне створення картки без прив'язки до CleverBox."""
    name = full_name.strip()
    if not name:
        raise HTTPException(400, "Вкажіть ПІБ пацієнта")
    assigned_doctor = user.doctor_id if user.role == "doctor" else (int(doctor_id) if doctor_id.isdigit() else None)
    p = Patient(full_name=name, phone=phone.strip() or None, doctor_id=assigned_doctor, notes=notes.strip() or None,
                source="manual")
    db.add(p)
    db.commit()
    audit(db, "patient.create_manual", "patient", p.id, user=user)
    return RedirectResponse(f"/patients/{p.id}", 303)


@app.post("/patients/test")
def patient_create_test(user: User = Depends(require("patient_edit")), db: Session = Depends(db_dep)):
    """Одна демонстраційна картка для перевірки інтерфейсу, без персональних даних."""
    p = db.scalar(select(Patient).where(Patient.source == "manual_test"))
    if not p:
        doctor_id = user.doctor_id if user.role == "doctor" else None
        p = Patient(full_name="Тестова картка пацієнта", phone=None, doctor_id=doctor_id,
                    notes="Демонстраційна картка. Не є даними реального пацієнта.", source="manual_test")
        db.add(p)
        db.commit()
        audit(db, "patient.create_test", "patient", p.id, user=user)
    return RedirectResponse(f"/patients/{p.id}", 303)


@app.post("/plans/{plan_id}/items")
async def plan_items_update(request: Request, plan_id: int, user: User = Depends(require("plan_edit")),
                            db: Session = Depends(db_dep)):
    form = await request.form()
    plan = db.get(BeautyPlan, plan_id)
    if not plan or not plan_visible(user, plan):
        raise HTTPException(403, "Немає доступу до цього плану")
    changes = {}
    for it in plan.items:
        for f in ("interval_days", "supply_days", "manual_done", "qty"):
            v = form.get(f"{f}_{it.id}")
            if v is not None and str(v).strip().isdigit() and getattr(it, f) != int(v):
                changes[f"{it.name[:40]}.{f}"] = [getattr(it, f), int(v)]
                setattr(it, f, int(v))
        it.total = it.qty * it.unit_price
    plan.total_planned = sum(i.total for i in plan.items)
    if form.get("status") in ("active", "completed", "cancelled") and form["status"] != plan.status:
        changes["status"] = [plan.status, form["status"]]
        plan.status = form["status"]
    if user.role != "doctor" and str(form.get("doctor_id", "")).isdigit():
        plan.doctor_id = int(form["doctor_id"])
    db.commit()
    if changes:
        audit(db, "plan.update", "plan", plan.id, user=user, changes=changes)
        webhooks.fire("plan.updated", plan_dict(compute_plan(db, plan)))
    return RedirectResponse(f"/patients/{plan.patient_id}#plan{plan.id}", 303)


@app.post("/plans/{plan_id}/delete")
def plan_delete(plan_id: int, del_password: str = Form(""), user: User = Depends(require("plan_delete")),
                db: Session = Depends(db_dep)):
    plan = db.get(BeautyPlan, plan_id)
    pid = plan.patient_id
    err = _delpw(db, user, del_password, "plan", "plan", plan_id)
    if err:
        return RedirectResponse(_with_err(f"/patients/{pid}#plan{plan_id}", err), 303)
    audit(db, "plan.delete", "plan", plan.id, user=user, patient=plan.patient.full_name, total=plan.total_planned)
    db.delete(plan)
    db.commit()
    return RedirectResponse(f"/patients/{pid}", 303)


# ---------------- файли ----------------

@app.post("/upload")
async def upload(request: Request, files: list[UploadFile] = File(...), user: User = Depends(require("upload")),
                 db: Session = Depends(db_dep)):
    results = []
    for f in files:
        data = await f.read()
        if (f.filename or "").lower().endswith(".pdf"):
            plan, _, err = ingest.ingest_pdf(db, data, f.filename)
            results.append((f.filename, plan, err))
            if plan:
                audit(db, "plan.upload", "plan", plan.id, user=user, file=f.filename)
        else:
            ff = ingest.ingest_photo(db, data, f.filename or "photo.jpg", None)
            results.append((f.filename, None, None if ff.patient_id else "фото без пацієнта → «Нерозібрані»"))
    if len(results) == 1 and results[0][1]:
        return RedirectResponse(f"/patients/{results[0][1].patient_id}", 303)
    return render(request, "upload_result.html", results=results)


@app.post("/patients/{pid}/photos")
async def upload_photos(pid: int, files: list[UploadFile] = File(...), caption: str = Form(""), stage: str = Form(""),
                        zone: str = Form(""), user: User = Depends(require("upload")), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    for f in files:
        pf = ingest.ingest_photo(db, await f.read(), f.filename or "photo.jpg", caption or None, patient=p)
        if stage in ("before", "after", "progress") or zone.strip():
            pf.stage, pf.zone, pf.taken_at = (stage if stage in ("before", "after", "progress") else None), zone.strip()[:80] or None, date.today()
            db.commit()
    return RedirectResponse(f"/patients/{pid}#files", 303)


# ---------------- Звернення (заявки) ----------------
def _appeal_filters(request: Request) -> dict:
    qp = request.query_params
    f = {k: qp.get(k, "") for k in ("source", "status", "cause", "resp", "mark", "q", "due", "view")}
    f["d_from"], f["d_to"] = _day(qp.get("df", "")), _day(qp.get("dt_", ""))
    if not qp.get("df") and not qp.get("all"):
        f["d_from"] = date.today() - timedelta(days=60)
    return f


@app.get("/appeals", response_class=HTMLResponse)
def appeals_page(request: Request, user: User = Depends(require("appeals")), db: Session = Depends(db_dep)):
    from .services import appeals as ap
    from .models import Appeal
    f = _appeal_filters(request)
    rows = ap.listing(db, f)
    d = ap.dirs(db)
    users = {u_.id: u_.name for u_ in db.scalars(select(User).where(User.active.is_(True)))}
    pats = {p_.id: p_ for p_ in db.scalars(select(Patient).where(Patient.id.in_([a.patient_id for a in rows if a.patient_id] or [0])))}
    due_n = len([a for a in db.scalars(select(Appeal).where(Appeal.deleted_at.is_(None), Appeal.closed_at.is_(None),
                                                            Appeal.next_contact <= datetime.now().replace(hour=23, minute=59)))])
    return render(request, "appeals.html", nav="appeals", rows=rows, d=d, f=f, st=ap.stats(db, rows), users=users, pats=pats,
                  ap=ap, due_n=due_n, now=datetime.now(), open_=int(request.query_params.get("open") or 0) or None)


@app.get("/appeals/{aid:int}", response_class=HTMLResponse)
def appeal_card(request: Request, aid: int, user: User = Depends(require("appeals")), db: Session = Depends(db_dep)):
    from .services import appeals as ap
    from .models import Appeal, AppealNote
    a = db.get(Appeal, aid)
    if not a or a.deleted_at:
        raise HTTPException(404)
    notes = db.scalars(select(AppealNote).where(AppealNote.appeal_id == aid).order_by(AppealNote.at.desc())).all()
    users = {u_.id: u_.name for u_ in db.scalars(select(User))}
    other = db.scalars(select(Appeal).where(Appeal.id != aid, Appeal.deleted_at.is_(None), Appeal.phone.is_not(None))).all()
    pk = ap.phone_key(a.phone)
    history = [x for x in other if pk and ap.phone_key(x.phone) == pk][:20]
    return templates.TemplateResponse(request, "_appeal_card.html", {"request": request, "user": user, "a": a, "notes": notes, "users": users,
                                      "d": ap.dirs(db), "ap": ap, "p": db.get(Patient, a.patient_id) if a.patient_id else None,
                                      "history": history, "match": None if a.patient_id else ap.find_patient(db, a.phone)})


@app.post("/appeals")
async def appeal_save(request: Request, user: User = Depends(require("appeals")), db: Session = Depends(db_dep)):
    from .services import appeals as ap
    from .models import Appeal
    f = await request.form()
    back = str(f.get("back") or "/appeals")
    back = back if back.startswith("/appeals") or back.startswith("/patients/") else "/appeals"
    nc = str(f.get("next_contact") or "")
    try:
        next_contact = datetime.fromisoformat(nc) if nc else None
    except ValueError:
        next_contact = None
    data = {k: str(f.get(k) or "") for k in ("name", "phone", "source", "cause", "interest", "comment", "status", "utm")}
    data["marks"] = [str(x) for x in f.getlist("marks")]
    data["responsible_id"] = int(f["responsible_id"]) if str(f.get("responsible_id") or "").isdigit() else None
    data["next_contact"], data["amount"] = next_contact, _num(str(f.get("amount") or ""))
    aid = str(f.get("aid") or "")
    if aid.isdigit():
        a = db.get(Appeal, int(aid))
        if not a:
            raise HTTPException(404)
        changes = []
        for k in ("name", "phone", "source", "cause", "interest", "comment", "utm"):
            v = data[k].strip() or None
            if (getattr(a, k) or None) != v:
                changes.append(k)
                setattr(a, k, v)
        if (a.marks or []) != data["marks"]:
            a.marks = data["marks"] or None
        a.responsible_id, a.next_contact, a.amount = data["responsible_id"], next_contact, data["amount"]
        a.updated_at = datetime.utcnow()
        if not a.patient_id:
            m = ap.find_patient(db, a.phone)
            a.patient_id = m.id if m else None
        db.commit()
        if data["status"] and data["status"] != a.status:
            ap.set_status(db, a, data["status"], user.id, str(f.get("deny_reason") or "") or None)
        elif str(f.get("deny_reason") or "") and ap.status_type(ap.dirs(db), a.status) == "fail":
            a.deny_reason = str(f.get("deny_reason"))[:120]
            db.commit()
        if changes:
            ap.note(db, a, "Змінено: " + ", ".join(changes), user.id, "system")
            db.commit()
        audit(db, "appeal.edit", "appeal", a.id, user=user, status=a.status)
    else:
        if not (data["name"].strip() or data["phone"].strip()):
            return RedirectResponse(_with_msg(back, "Вкажіть імʼя або телефон", True), 303)
        a = ap.create(db, data, user.id)
        if data["status"] and ap.status_type(ap.dirs(db), data["status"]) == "fail" and f.get("deny_reason"):
            a.deny_reason = str(f.get("deny_reason"))[:120]
            db.commit()
        audit(db, "appeal.create", "appeal", a.id, user=user, source=a.source)
    if a.next_contact and not a.closed_at:
        _appeal_task(db, a, user)
    return RedirectResponse(back if "open=" in back or back.startswith("/patients/") else _with_msg(back, f"Звернення №{a.id} збережено"), 303)


def _appeal_task(db: Session, a, user: User) -> None:
    """Наступний контакт → задача в «Задачах» (одна на звернення, оновлюється)."""
    t = db.scalar(select(Task).where(Task.source_key == f"appeal:{a.id}"))
    t = t or Task(source_key=f"appeal:{a.id}", kind="follow_up", created_by=user.id)
    t.title = f"Звернення №{a.id}: звʼязатися — {a.name or a.phone or ''}"[:300]
    t.description = "\n".join(x for x in [a.interest and f"Цікавить: {a.interest}", a.comment, f"/appeals?open={a.id}"] if x)
    t.due_date, t.patient_id, t.assignee_id = a.next_contact.date(), a.patient_id, a.responsible_id
    t.assignee_role = None if a.responsible_id else "admin"
    t.status = "open"
    db.add(t)
    db.commit()


@app.post("/appeals/{aid:int}/note")
def appeal_note(aid: int, text: str = Form(...), user: User = Depends(require("appeals")), db: Session = Depends(db_dep)):
    from .services import appeals as ap
    from .models import Appeal
    a = db.get(Appeal, aid)
    if a and text.strip():
        ap.note(db, a, text.strip(), user.id)
        a.updated_at = datetime.utcnow()
        db.commit()
    return RedirectResponse(f"/appeals?open={aid}", 303)


@app.post("/appeals/{aid:int}/status")
def appeal_status(aid: int, status: str = Form(...), deny_reason: str = Form(""), back: str = Form("/appeals"),
                  user: User = Depends(require("appeals")), db: Session = Depends(db_dep)):
    from .services import appeals as ap
    from .models import Appeal
    a = db.get(Appeal, aid)
    if a:
        ap.set_status(db, a, status, user.id, deny_reason or None)
        if a.closed_at:
            t = db.scalar(select(Task).where(Task.source_key == f"appeal:{aid}", Task.status == "open"))
            if t:
                t.status, t.done_at, t.result = "done", datetime.utcnow(), f"Звернення: {a.status}"
                db.commit()
        audit(db, "appeal.status", "appeal", aid, user=user, status=status, deny=deny_reason or None)
    return RedirectResponse(back if back.startswith("/appeals") else "/appeals", 303)


@app.post("/appeals/{aid:int}/patient")
def appeal_to_patient(aid: int, user: User = Depends(require("appeals")), db: Session = Depends(db_dep)):
    from .services import appeals as ap
    from .models import Appeal
    a = db.get(Appeal, aid) or (_ for _ in ()).throw(HTTPException(404))
    created = not a.patient_id and not ap.find_patient(db, a.phone)
    p = ap.to_patient(db, a)
    ap.note(db, a, ("Створено картку пацієнта" if created else "Привʼязано до картки пацієнта") + f" №{p.id}", user.id, "system")
    db.commit()
    audit(db, "appeal.patient", "appeal", aid, user=user, patient=p.id, created=created)
    return RedirectResponse(f"/patients/{p.id}", 303)


@app.post("/appeals/{aid:int}/delete")
def appeal_delete(aid: int, del_password: str = Form(""), user: User = Depends(require("appeals")), db: Session = Depends(db_dep)):
    from .models import Appeal
    err = _delpw(db, user, del_password, "appeal", "appeal", aid)
    if err:
        return RedirectResponse(_with_err("/appeals", err), 303)
    a = db.get(Appeal, aid)
    if a:
        a.deleted_at = datetime.utcnow()
        db.commit()
        audit(db, "appeal.delete", "appeal", aid, user=user)
    return RedirectResponse(_with_msg("/appeals", f"Звернення №{aid} видалено"), 303)


@app.get("/settings/appeals", response_class=HTMLResponse)
def settings_appeals(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import appeals as ap
    return render(request, "appeal_dirs.html", nav="settings", d=ap.dirs(db), ap=ap, open_=request.query_params.get("open", ""))


@app.post("/settings/appeals/{what}")
def settings_appeals_save(what: str, action: str = Form(...), idx: str = Form(""), name: str = Form(""), type_: str = Form("work", alias="type"),
                          color: str = Form(""), del_password: str = Form(""), user: User = Depends(require("users")),
                          db: Session = Depends(db_dep)):
    import re as _re
    from .services import appeals as ap
    if what not in ap.DIR_LABELS:
        raise HTTPException(404)
    d = ap.dirs(db)
    items = d[what]
    i = int(idx) if idx.isdigit() and int(idx) < len(items) else None
    nm = name.strip()[:80]
    is_st = what == "statuses"
    if action == "add" and nm:
        items.append({"name": nm, "type": type_ if type_ in ap.STATUS_TYPES else "work", "color": color if _re.fullmatch(r"#[0-9a-fA-F]{6}", color or "") else "#888888"} if is_st else nm)
    elif i is not None and action == "save" and nm:
        if is_st:
            items[i].update(name=nm, type=type_ if type_ in ap.STATUS_TYPES else items[i]["type"],
                            color=color if _re.fullmatch(r"#[0-9a-fA-F]{6}", color or "") else items[i].get("color"))
        else:
            items[i] = nm
    elif i is not None and action in ("up", "down"):
        j = i - 1 if action == "up" else i + 1
        if 0 <= j < len(items):
            items[i], items[j] = items[j], items[i]
    elif i is not None and action == "delete":
        err = _delpw(db, user, del_password, f"appeal_{what}", None, None)
        if err:
            return RedirectResponse(_with_err(f"/settings/appeals?open={what}", err), 303)
        if is_st and len(items) <= 1:
            return RedirectResponse(_with_msg(f"/settings/appeals?open={what}", "Має лишитися хоча б один статус", True), 303)
        items.pop(i)
    ap.save_dirs(db, d)
    audit(db, "settings.appeals", user=user, what=what, op=action, name=nm or None)
    return RedirectResponse(f"/settings/appeals?open={what}#{what}", 303)


# ---------------- Запис у CRM ----------------
@app.get("/lookup/patients")
def lookup_patients(q: str = "", user: User = Depends(current_user), db: Session = Depends(db_dep)):
    import re as _re
    q = q.strip()
    digits = _re.sub(r"\D", "", q)
    is_phone = bool(digits) and not _re.search(r"[^\d\s+()\-]", q)          # шукаємо за номером з першої цифри
    if not is_phone and len(q) < 2:
        return JSONResponse([])
    ql = q.lower()
    d9 = digits.lstrip("0")[-9:] if len(digits) >= 10 else digits                # 0981234567 / 380981234567 → однаково
    out = []
    for p_ in db.scalars(select(Patient).order_by(Patient.full_name)):
        ph = _re.sub(r"\D", "", p_.phone or "")
        if (is_phone and ph and (digits in ph or (d9 and d9 in ph))) or (not is_phone and ql in (p_.full_name or "").lower()):
            if patient_visible(user, p_):
                out.append({"id": p_.id, "name": p_.full_name, "phone": p_.phone or "", "bonus": None})
        if len(out) >= 15:
            break
    return JSONResponse(out)


@app.get("/lookup/services")
def lookup_services(q: str = "", user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .models import CbService
    ql = q.strip().lower()
    rows = [s for s in db.scalars(select(CbService).order_by(CbService.name)) if not ql or ql in (s.name or "").lower()
            or ql in (s.category or "").lower()]
    return JSONResponse([{"id": s.id, "name": s.name, "price": s.price or 0, "minutes": s.time_min or 30, "category": s.category or ""}
                         for s in rows[:25]])


def _appt_ctx(db: Session) -> dict:
    from .services import clinic as cl, appointments as aps
    data = cl.get(db)
    loc = next((x for x in data["locations"] if x.get("active")), {"places": [], "resources": []})
    return {"doctors": spec_svc.calendar_doctors(db), "places": [p_["name"] for p_ in loc.get("places", []) if p_.get("active", True)],
            "resources": [r_["name"] for r_ in loc.get("resources", []) if r_.get("active", True)], "STATUSES": aps.STATUSES,
            "COLORS": aps.COLORS, "cb_on": bool(settings.cleverbox_token or db.get(SyncState, "cleverbox_token_enc"))}


@app.get("/appointments/new", response_class=HTMLResponse)
def appointment_new(request: Request, doc: int = 0, start: str = "", patient: int = 0, appeal: int = 0,
                    user: User = Depends(require("book")), db: Session = Depends(db_dep)):
    try:
        st_ = datetime.fromisoformat(start) if start else None
    except ValueError:
        st_ = None
    if not st_:
        n_ = datetime.now() + timedelta(hours=1)
        st_ = n_.replace(minute=0 if n_.minute < 30 else 30, second=0, microsecond=0)
    pt = db.get(Patient, patient) if patient else None
    ap_ = None
    if appeal:
        from .models import Appeal
        ap_ = db.get(Appeal, appeal)
        pt = pt or (db.get(Patient, ap_.patient_id) if ap_ and ap_.patient_id else None)
    return templates.TemplateResponse(request, "_appt_form.html", {"request": request, "user": user, "a": None, "doc": doc, "st": st_,
                                      "pt": pt, "appeal": ap_, **_appt_ctx(db)})


@app.get("/appointments/{aid:int}", response_class=HTMLResponse)
def appointment_card(request: Request, aid: int, edit: int = 0, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .models import Appointment
    a = db.get(Appointment, aid) or (_ for _ in ()).throw(HTTPException(404))
    _patient_or_403(db, user, a.patient_id)
    return templates.TemplateResponse(request, "_appt_form.html", {"request": request, "user": user, "a": a, "doc": a.doctor_id,
                                      "st": a.start, "pt": a.patient, "appeal": None, "view_only": not edit, **_appt_ctx(db),
                                      "doctor": db.get(Doctor, a.doctor_id) if a.doctor_id else None})


@app.post("/appointments")
async def appointment_save(request: Request, user: User = Depends(require("book")), db: Session = Depends(db_dep)):
    from .services import appointments as aps
    from .services.booking import BookingError
    from .models import Appointment, Appeal
    f = await request.json()
    pid = f.get("patient_id")
    if not pid and (f.get("new_name") or "").strip():
        from .services.appeals import find_patient
        pt = find_patient(db, f.get("new_phone")) if f.get("new_phone") else None
        if not pt:
            pt = Patient(full_name=f["new_name"].strip()[:200], phone=(f.get("new_phone") or "").strip()[:64] or None, source="Запис у CRM")
            db.add(pt)
            db.commit()
        pid = pt.id
    if not pid:
        return JSONResponse({"error": "Оберіть пацієнта або введіть імʼя нового"}, status_code=400)
    p_ = _patient_or_403(db, user, int(pid))
    try:
        start = datetime.fromisoformat(f.get("start") or "")
    except ValueError:
        return JSONResponse({"error": "Вкажіть дату й час"}, status_code=400)
    svc = []
    for s in f.get("services") or []:
        try:
            svc.append({"id": int(s["id"]), "name": str(s.get("name") or "")[:300], "category": s.get("category"),
                        "price": round(float(s.get("price") or 0), 2), "minutes": max(5, int(s.get("minutes") or 30))})
        except (KeyError, TypeError, ValueError):
            continue
    data = {"patient_id": p_.id, "doctor_id": int(f["doctor_id"]) if str(f.get("doctor_id") or "").isdigit() else None,
            "services": svc, "start": start, "minutes": int(f.get("minutes") or 0) or None, "place": f.get("place"),
            "resource": f.get("resource"), "comment": (f.get("comment") or "").strip()[:2000],
            "appeal_id": int(f["appeal_id"]) if str(f.get("appeal_id") or "").isdigit() else None}
    appt = db.get(Appointment, int(f["id"])) if str(f.get("id") or "").isdigit() else None
    mins = data["minutes"] or sum(s["minutes"] for s in svc) or 30
    warn = aps.conflicts(db, data["doctor_id"], start, start + timedelta(minutes=mins), data["place"], data["resource"],
                         appt.id if appt else None, p_.id)
    if warn and not f.get("force"):
        return JSONResponse({"warnings": warn})
    a = aps.save(db, data, user.id, appt)
    msg = None
    if f.get("to_cb") and not a.cb_record_id:
        try:
            rec = await asyncio.to_thread(aps.mirror_to_cleverbox, db, a)
            msg = f"Записано й у CleverBox (№{rec})"
        except (BookingError, Exception) as e:  # noqa: BLE001
            msg = f"Запис у CRM є, але в CleverBox не створено: {str(e)[:200]}"
    if a.appeal_id:
        from .services import appeals as apl
        ap_ = db.get(Appeal, a.appeal_id)
        if ap_:
            ok = next((s_["name"] for s_ in apl.dirs(db)["statuses"] if s_["type"] == "success"), None)
            if ok:
                apl.set_status(db, ap_, ok, user.id)
            ap_.patient_id = ap_.patient_id or p_.id
            apl.note(db, ap_, f"Записано на {a.start:%d.%m.%Y %H:%M}", user.id, "system")
            db.commit()
    audit(db, "appointment.save" if appt else "appointment.create", "patient", p_.id, user=user, appt=a.id,
          start=a.start.isoformat(timespec="minutes"), doctor=a.doctor_id, forced=bool(warn), cb=a.cb_record_id)
    return JSONResponse({"ok": True, "id": a.id, "msg": msg, "day": a.start.date().isoformat()})


@app.post("/appointments/{aid:int}/status")
def appointment_status(aid: int, status: str = Form(...), reason: str = Form(""), back: str = Form(""),
                       user: User = Depends(require("book")), db: Session = Depends(db_dep)):
    from .services import appointments as aps
    from .models import Appointment
    a = db.get(Appointment, aid) or (_ for _ in ()).throw(HTTPException(404))
    if status == "cancelled" and len(reason.strip()) < 3:
        return RedirectResponse(_with_msg(back or f"/calendar?start={a.start.date()}", "Вкажіть причину скасування", True), 303)
    aps.set_status(db, a, status, reason.strip()[:200] or None)
    audit(db, "appointment.status", "patient", a.patient_id, user=user, appt=aid, status=status, reason=reason or None)
    b = back if back.startswith("/calendar") or back.startswith("/patients/") else f"/calendar?start={a.start.date()}"
    return RedirectResponse(b, 303)


# ---------------- Каса: продаж у CRM ----------------
@app.get("/pos/visits")
def pos_visits(d: str = "", user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import pos
    day_ = _day(d) or date.today()
    rows = [r for r in pos.visits(db, day_) if not r["patient_id"] or patient_visible(user, db.get(Patient, r["patient_id"]))]
    return JSONResponse({"day": day_.isoformat(), "rows": rows})


@app.get("/pos/patient/{pid:int}")
def pos_patient(pid: int, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    """Картка пацієнта в касі: мітка, бонуси, баланс, персональна знижка."""
    from .services import loyalty as ly, calendar as _cal
    from sqlalchemy import func as _f
    p_ = _patient_or_403(db, user, pid)
    n = db.scalar(select(_f.count(_f.distinct(_f.date(Visit.starts_at)))).where(Visit.patient_id == p_.id, Visit.status.in_(("attended", "done")))) or 0
    from .services import certificates as _cert
    certs = [c_ for c_ in _cert.for_patient(db, p_.id) if c_.state() in ("active", "expiring")]
    subs = [s_ for s_ in _cert.subs_for_patient(db, p_.id) if s_.state() in ("active", "expiring")]
    labels = []
    for g_ in (p_.client_groups or []):
        nm = g_.get("name") or g_.get("title") if isinstance(g_, dict) else str(g_)
        if nm and nm not in labels:
            labels.append(str(nm)[:40])
    if db.scalar(select(User.id).where(User.active.is_(True), User.name == p_.full_name)) or \
            db.scalar(select(Doctor.id).where(Doctor.active.is_(True), Doctor.name == p_.full_name)):
        labels.insert(0, "Співробітник") if "Співробітник" not in labels else None
    return JSONResponse({"id": p_.id, "name": p_.full_name, "phone": p_.phone or "", "tag": _cal.tag_for(n), "visits": n,
                         "bonus": round(ly.balance(db, p_.id), 2), "balance": p_.cb_balance, "discount": p_.discount_pct,
                         "labels": labels[:6],
                         "certs": {"n": len(certs), "sum": round(sum(c_.balance or 0 for c_ in certs), 2),
                                   "items": [{"id": c_.id, "code": c_.code or str(c_.id), "balance": c_.balance or 0} for c_ in certs]},
                         "subs": {"n": len(subs), "names": [(s_.type_name or s_.code or "")[:40] for s_ in subs][:3],
                                  "items": [{"id": s_.id, "name": (s_.type_name or s_.code or f"№{s_.id}")[:80], "balance": s_.balance,
                                             "visits": s_.visits_total is not None, "left": s_.visits_left, "services": (s_.services or "")[:300]}
                                            for s_ in subs]}})


@app.post("/pos/patient")
async def pos_patient_new(request: Request, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services.appeals import find_patient
    f = await request.json()
    name, phone = str(f.get("name") or "").strip()[:200], str(f.get("phone") or "").strip()[:64]
    if len(name) < 2:
        return JSONResponse({"error": "Вкажіть імʼя пацієнта"}, status_code=400)
    p_ = find_patient(db, phone) if phone else None
    if not p_:
        p_ = Patient(full_name=name, phone=phone or None, source="Каса CRM")
        db.add(p_)
        db.commit()
        audit(db, "patient.create", "patient", p_.id, user=user, via="pos")
    return JSONResponse({"id": p_.id, "name": p_.full_name, "phone": p_.phone or "", "existed": p_.full_name != name})


@app.get("/pos/services")
def pos_services(group: str = "category", key: str | None = None, q: str = "", user: User = Depends(require("sales")),
                 db: Session = Depends(db_dep)):
    from .services import pos
    return JSONResponse(pos.services(db, group, key, q))


@app.get("/pos/products")
def pos_products(group: str = "warehouse", key: str | None = None, q: str = "", stock: int = 0, user: User = Depends(require("sales")),
                 db: Session = Depends(db_dep)):
    from .services import pos
    return JSONResponse(pos.products(db, group, key, q, bool(stock)))


@app.get("/pos/lots/{pid:int}")
def pos_lots(pid: int, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import stock as st
    from .models import Product as _P
    p_ = db.get(_P, pid) or (_ for _ in ()).throw(HTTPException(404))
    wh = p_.warehouse or "vitrina"
    lots = []
    if st.ledger_active(db):
        lots = sorted([{"expiry": e.isoformat(), "qty": round(qq, 3)} for (pid_, w_, e), qq in st.batch_balances(db, p_.id, wh).items()
                       if e and qq > 1e-6], key=lambda x: x["expiry"])
    return JSONResponse(lots)


@app.get("/lookup/sale-products")
def lookup_sale_products(q: str = "", user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import stock as st
    q = (q or "").strip()
    if len(q) < 2:
        return JSONResponse([])
    words, ql = set(norm(q).split()), q.lower()
    led = st.ledger_active(db)
    out = []
    for r in catalog.product_rows(db, []):
        p_ = r.product
        if p_.active is False:
            continue
        exact = ql in ((p_.sku or "").lower(), (p_.barcode or "").lower())
        if not exact and not words <= set(norm(" ".join(filter(None, [p_.name, p_.sku, p_.brand, p_.barcode]))).split()):
            continue
        wh = p_.warehouse or "vitrina"
        lots = []
        if led:
            lots = [{"expiry": e.isoformat(), "qty": round(qq, 3)} for (pid_, w_, e), qq in st.batch_balances(db, p_.id, wh).items() if e and qq > 1e-6]
            lots.sort(key=lambda x: x["expiry"])
        out.append((0 if exact else (0 if wh == "vitrina" else 1), {"id": p_.id, "name": p_.name, "brand": p_.brand or "", "sku": p_.sku or "", "photo": p_.photo_url or "", "category": p_.brand or p_.category or "",
                    "price": p_.price or p_.price_manual or p_.shop_price or 0, "stock": r.stock, "warehouse": wh,
                    "wh": catalog.WAREHOUSES.get(wh), "lots": lots}))
    out.sort(key=lambda x: (x[0], x[1]["name"].lower()))
    return JSONResponse([x[1] for x in out[:25]])


@app.get("/sales", response_class=HTMLResponse)
def sales_list(request: Request, df: str = "", dt_: str = "", q: str = "", d: str = "", user: User = Depends(require("sales")),
               db: Session = Depends(db_dep)):
    from .models import Sale
    if not (df or dt_ or q):                                     # каса за добу (як «Продажі за добу» у CleverBox)
        from .services import daybook, clinic as _cl
        day_ = _day(d) or date.today()
        return render(request, "fin_cash_day.html", nav="finance", day=day_, v=daybook.day(db, day_), timedelta=timedelta,
                      today=date.today(), locs=_cl.get(db).get("locations") or [])
    d_from, d_to = _period(df, dt_)
    a_, b_ = datetime.combine(d_from, datetime.min.time()), datetime.combine(d_to, datetime.max.time())
    rows = db.scalars(select(Sale).where(Sale.at >= a_, Sale.at <= b_).order_by(Sale.at.desc()).limit(2000)).all()
    if q.strip():
        ql = q.strip().lower()
        rows = [s_ for s_ in rows if ql == str(s_.id) or (s_.patient and ql in s_.patient.full_name.lower()) or any(ql in i.name.lower() for i in s_.items)]
    paid = [s_ for s_ in rows if s_.status == "paid"]
    by_m: dict = {}
    for s_ in paid:
        for p_ in s_.payments:
            by_m[p_.method] = round(by_m.get(p_.method, 0) + p_.amount, 2)
    users = {u_.id: u_.name for u_ in db.scalars(select(User))}
    return render(request, "sales.html", nav="finance", rows=rows, d_from=d_from, d_to=d_to, q=q, users=users, by_m=by_m,
                  tot={"n": len(paid), "total": round(sum(s_.total for s_ in paid), 2), "discount": round(sum(s_.discount for s_ in paid), 2),
                       "refunded": round(sum(s_.total for s_ in rows if s_.status == "refunded"), 2)})


@app.get("/sales/new", response_class=HTMLResponse)
def sale_new(request: Request, appt: int = 0, patient: int = 0, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .models import Appointment, DiscountRule, GiftCertificate
    from .services import sales as sl
    a_ = db.get(Appointment, appt) if appt else None
    pt = a_.patient if a_ else (db.get(Patient, patient) if patient else None)
    items = []
    if a_:
        items = [{"kind": "service", "ref_id": s_.get("id"), "name": s_.get("name"), "category": s_.get("category") or "", "qty": 1,
                  "price": float(s_.get("price") or 0), "doctor_id": a_.doctor_id} for s_ in a_.services or []]
    certs = db.scalars(select(GiftCertificate).where(GiftCertificate.balance > 0)).all()
    from .services import fiscal as _fis
    _fc = _fis.cfg(db)
    from .services import clinic as _cl
    _locs = [l_ for l_ in (_cl.get(db).get("locations") or []) if l_.get("active", True)]
    return render(request, "sale_new.html", nav="finance", appt=a_, pt=pt, items=items, methods=sl.methods(db), today=date.today().isoformat(),
                  loc_name=_locs[0]["name"] if _locs else "MOLODIST",
                  mkinds={m_: sl.method_kind(m_) for m_ in sl.methods(db)}, rros=[r_ for r_ in _fis.rros(db) if r_["on"]],
                  link_for=_pos_link_map(sl.methods(db), sl), rro_m=_fis.rro_methods(db),
                  pl_term={t_["method"]: t_["name"] for t_ in _fis.terminals(_fc)["items"] if t_.get("on") and t_.get("method")},
                  fauto=bool(_fc.get("auto")), fdef=_fc.get("default_rro"),
                  doctors=[d_ for d_ in db.scalars(select(Doctor).where(Doctor.active.is_(True)).order_by(Doctor.name)) if d_.sale_initiator is not False],
                  rules=db.scalars(select(DiscountRule).where(DiscountRule.active.is_(True),
                                                                                                      DiscountRule.auto.is_(False))).all(),
                  certs=[{"id": c_.id, "code": c_.code or str(c_.id), "balance": c_.balance, "patient_id": c_.patient_id,
                          "name": c_.client_name or ""} for c_ in certs if c_.state() in ("active", "expiring")][:500],
                  default_doc=user.doctor_id)


def _pos_link_map(methods: list[str], sl) -> dict:
    """Термінал ФОП X → «Посилання · monobank ФОП X» (за прізвищем ФОП); непривʼязані посилання — під ключем ''."""
    import re as _re
    links = [m_ for m_ in methods if sl.method_kind(m_) == "link"]
    out, used = {}, set()
    for m_ in methods:
        if sl.method_kind(m_) != "card":
            continue
        mm = _re.search(r"ФОП\s+(\S+)", m_)
        if mm:
            hit = [l_ for l_ in links if mm.group(1).lower() in l_.lower()]
            if hit:
                out[m_] = hit[0]
                used.add(hit[0])
    out[""] = [l_ for l_ in links if l_ not in used]
    return out


def _sale_payload(f: dict) -> dict:
    items = []
    for it in f.get("items") or []:
        try:
            items.append({"kind": it.get("kind") if it.get("kind") in ("service", "product", "certificate", "other") else "other",
                          "ref_id": int(it["ref_id"]) if str(it.get("ref_id") or "").isdigit() else None,
                          "name": str(it.get("name") or "")[:300], "category": str(it.get("category") or ""),
                          "qty": max(0.001, float(it.get("qty") or 1)), "price": max(0.0, float(it.get("price") or 0)),
                          "doctor_id": int(it["doctor_id"]) if str(it.get("doctor_id") or "").isdigit() else None,
                          "warehouse": it.get("warehouse") or None, "expiry": _day(str(it.get("expiry") or ""))})
        except (TypeError, ValueError):
            continue
    return {"patient_id": int(f["patient_id"]) if str(f.get("patient_id") or "").isdigit() else None,
            "appointment_id": int(f["appointment_id"]) if str(f.get("appointment_id") or "").isdigit() else None,
            "items": items, "rule_ids": [int(x) for x in f.get("rule_ids") or [] if str(x).isdigit()],
            "manual_discount": float(f.get("manual_discount") or 0), "bonus_pay": float(f.get("bonus_pay") or 0),
            "fiscal_rro": (int(f["fiscal_rro"]) if str(f.get("fiscal_rro")).isdigit() else None) if "fiscal_rro" in f else None,
            "payments": f.get("payments") or [], "comment": str(f.get("comment") or "")}


@app.post("/sales/quote")
async def sale_quote(request: Request, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import sales as sl
    d_ = _sale_payload(await request.json())
    pt = db.get(Patient, d_["patient_id"]) if d_["patient_id"] else None
    return JSONResponse(sl.quote(db, d_["items"], pt, d_["rule_ids"], d_["manual_discount"], d_["bonus_pay"]))


async def _sale_links(request: Request, db: Session, s_, user: User) -> list[str]:
    """Посилання на оплату для чека (оплати «Посилання на оплату · …» без створеного посилання) + надсилання пацієнту."""
    from .services import paylinks as _pls, sales as sl
    from .models import PayLink
    from .integrations.messaging import send_to_patient
    base = str(request.base_url).rstrip("/").replace("http://", "https://")
    have = {(p_.account_id, round(p_.amount, 2)) for p_ in db.scalars(select(PayLink).where(PayLink.sale_id == s_.id,
                                                                                        PayLink.status.in_(("created", "paid"))))}
    basket = [{"name": i.name, "qty": i.qty, "total": i.total, "code": f"{i.kind[:1]}{i.ref_id or i.id}"} for i in s_.items]
    msgs = []
    for p_ in s_.payments:
        if not p_.method.startswith(sl.PAYLINK):
            continue
        acc = _pls.account_by_name(_pls._raw(db), p_.method.split(" · ", 1)[-1])
        if acc and (acc["id"], round(p_.amount, 2)) in have:
            continue
        try:
            pl = await asyncio.to_thread(_pls.create, db, p_.amount, f"Чек Ч-{s_.id} · MOLODIST", base, s_.patient_id, s_.appointment_id,
                                         user.id, None, acc["id"] if acc else None, basket)
        except _pls.PayError as e:
            msgs.append(f"посилання не створено: {e}")
            continue
        pl.sale_id = s_.id
        db.commit()
        if s_.patient:
            res = await asyncio.to_thread(send_to_patient, s_.patient, _pls.message(db, pl))
            pl.info = {**(pl.info or {}), "sent": {"ok": res.ok, "channel": res.channel, "at": datetime.now().strftime("%d.%m.%Y %H:%M")}}
            db.commit()
            msgs.append(f"посилання на {pl.amount:.2f} ₴ " + (f"надіслано ({res.channel})" if res.ok else f"створено, не надіслано: {res.error}"))
    return msgs


@app.post("/sales/{sid:int}/paylink")
async def sale_paylink_retry(request: Request, sid: int, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    if s_.status != "awaiting":
        return RedirectResponse(_with_msg(f"/sales/{sid}", "Чек не очікує оплати", True), 303)
    msgs = await _sale_links(request, db, s_, user)
    bad = any("не створено" in m_ for m_ in msgs)
    return RedirectResponse(_with_msg(f"/sales/{sid}", "; ".join(msgs) or "Посилання вже створено", bad), 303)


@app.get("/sales/{sid:int}/pay-status")
async def sale_pay_status(sid: int, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    """Для сторінки чека: перепитати monobank і повернути статуси посилань і чека."""
    from .models import Sale, PayLink
    from .services import paylinks as _pls
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    for pl in db.scalars(select(PayLink).where(PayLink.sale_id == sid, PayLink.status == "created")).all():
        try:
            await asyncio.to_thread(_pls.check_status, db, pl)
        except Exception:  # noqa: BLE001
            pass
    for pl in db.scalars(select(PayLink).where(PayLink.sale_id == sid, PayLink.status == "paid")).all():
        if not (s_.fiscal or {}).get("fiscal_code"):
            try:
                await asyncio.to_thread(_pls.mono_fiscal, db, pl)
            except Exception:  # noqa: BLE001
                pass
    db.refresh(s_)
    return JSONResponse({"sale": s_.status, "fiscal": bool((s_.fiscal or {}).get("fiscal_code")),
                         "links": [{"id": p_.id, "status": p_.status} for p_ in db.scalars(select(PayLink).where(PayLink.sale_id == sid))]})


@app.get("/paylinks/{lid:int}/fiscal.pdf")
async def paylink_fiscal_pdf(lid: int, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    """PDF фіскального чека monobank (base64 з /invoice/fiscal-checks)."""
    import base64 as _b64
    from fastapi.responses import Response
    from .models import PayLink
    from .services import paylinks as _pls
    pl = db.get(PayLink, lid) or (_ for _ in ()).throw(HTTPException(404))
    checks = await asyncio.to_thread(_pls.mono_fiscal, db, pl, None, True)
    f_ = next((c_.get("file") for c_ in checks if c_.get("file")), None)
    if not f_:
        raise HTTPException(404, "Фіскальний чек ще не готовий")
    return Response(_b64.b64decode(f_), media_type="application/pdf", headers={"Content-Disposition": f'inline; filename="check-{lid}.pdf"'})


@app.post("/sales")
async def sale_create(request: Request, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import sales as sl
    d_ = _sale_payload(await request.json())
    if d_["patient_id"]:
        _patient_or_403(db, user, d_["patient_id"])
    if any(not it_.get("doctor_id") for it_ in d_["items"]):
        return JSONResponse({"error": "Оберіть спеціаліста / продавця (ініціатора продажу) для кожної позиції"}, status_code=400)
    d_["manual_discount"] = 0                                   # знижка — лише зі списку в «Налаштування → Знижки»
    from .services import stock as _st0
    if _st0.ledger_active(db):                                  # кілька термінів дії товару — термін обовʼязковий
        from .models import Product as _P0
        for it_ in d_["items"]:
            if it_.get("kind") == "product" and it_.get("ref_id") and not it_.get("expiry"):
                p0_ = db.get(_P0, int(it_["ref_id"]))
                if p0_ and len([1 for (_a, _b, e_), q_ in _st0.batch_balances(db, p0_.id, it_.get("warehouse") or p0_.warehouse or "vitrina").items()
                                if e_ and q_ > 1e-6]) > 1:
                    return JSONResponse({"error": f"«{p0_.name}»: оберіть термін дії для списання"}, status_code=400)
    from .services import fiscal as _fis0, sales as _sl0
    _allowed = _fis0.rro_methods(db).get(d_.get("fiscal_rro") or 0) if d_.get("fiscal_rro") else None
    if _allowed:
        bad_ = [p_.get("method") for p_ in d_["payments"] if _sl0.method_kind(str(p_.get("method") or "")) in ("cash", "card")
                and p_.get("method") not in _allowed]
        if bad_:
            return JSONResponse({"error": "Ця каса РРО не приймає: " + ", ".join(bad_)}, status_code=400)
    _fc0 = _fis0.cfg(db)
    term_wait = {}                                              # оплата на терміналі PayLink: чек закриється лише після оплати
    for p_ in d_["payments"]:
        m_ = str(p_.get("method") or "")
        if float(p_.get("amount") or 0) <= 0 or _sl0.method_kind(m_) != "card" or m_.startswith(_sl0.PAYLINK):
            continue
        t_ = _fis0.terminal_for(_fc0, m_)
        if not t_:
            return JSONResponse({"error": f"«{m_}»: термінал не підключено до CRM — без відповіді термінала чек не проводиться. "
                                          "Додайте термінал (ID з PayLink) у «Налаштування → Checkbox → Термінали» і привʼяжіть до цього способу оплати."},
                                status_code=400)
        if True:
                w_ = term_wait.setdefault(p_["method"], {"device": t_["device_id"], "name": t_["name"], "amount": 0.0})
                w_["amount"] = round(w_["amount"] + float(p_["amount"]), 2)
    d_["await_terminal"] = bool(term_wait)
    try:
        s_ = sl.create(db, d_, user.id)
    except sl.SaleError as e:
        db.rollback()
        return JSONResponse({"error": str(e)}, status_code=400)
    if term_wait:
        s_.fiscal = {**(s_.fiscal or {}), "term_wait": term_wait, "rro_pending": d_.get("fiscal_rro")}
        db.commit()
    audit(db, "sale.create", "sale", s_.id, user=user, total=s_.total, patient=s_.patient_id, appt=s_.appointment_id)
    link_msgs = await _sale_links(request, db, s_, user) if s_.status == "awaiting" and not term_wait else []
    if term_wait and any(str(p_.get("method") or "").startswith(sl.PAYLINK) for p_ in d_["payments"]):
        link_msgs = await _sale_links(request, db, s_, user)
    from .services import fiscal as _fis
    rro_ = d_.get("fiscal_rro")
    if rro_ and s_.status == "paid":                              # вибрана фіскальна каса в POS
        import threading
        threading.Thread(target=_fis.auto, args=(s_.id, rro_), daemon=True).start()
    elif _fis.cfg(db).get("auto") and rro_ != 0 and s_.status == "paid":
        import threading
        threading.Thread(target=_fis.auto, args=(s_.id,), daemon=True).start()
    return JSONResponse({"ok": True, "id": s_.id, "msg": "; ".join(link_msgs),
                         "term": list(term_wait),
                         "title": "Чек створено — оплата на терміналі" if term_wait else ("Чек створено — очікує оплати" if s_.status == "awaiting" else "Продаж проведено")})


@app.get("/sales/{sid:int}", response_class=HTMLResponse)
def sale_view(request: Request, sid: int, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    docs = {d_.id: d_.name for d_ in db.scalars(select(Doctor))}
    from .services import fiscal as _fis
    return render(request, "sale_view.html", nav="finance", s=s_, docs=docs, author=db.get(User, s_.created_by) if s_.created_by else None,
                  print_=bool(request.query_params.get("print")), rros=[r for r in _fis.rros(db) if r["on"]], fcfg=_fis.cfg(db),
                  links=db.scalars(select(__import__("app.models", fromlist=["PayLink"]).PayLink).where(
                      __import__("app.models", fromlist=["PayLink"]).PayLink.sale_id == s_.id)).all())


def _term_has_data(r) -> bool:
    return isinstance(r, dict) and any(r.get(k) for k in ("rrn", "auth_code", "card_mask", "receipt_no"))


def _term_done(db: Session, s_, method: str, result: dict, user) -> dict:
    """Оплата на терміналі пройшла: дані картки в чек → чек «оплачено» → фіскалізація в Checkbox (каса з POS)."""
    from .services import fiscal as _fis
    from .models import PayLink
    f = dict(s_.fiscal or {})
    wait = dict(f.get("term_wait") or {})
    wait.pop(method, None)
    f["term_wait"] = wait
    f["term"] = {**(f.get("term") or {}), method: {k: result.get(k) for k in _fis.TERM_FIELDS if result.get(k) not in (None, "")}}
    s_.fiscal = f
    db.commit()
    open_links = db.query(PayLink).filter(PayLink.sale_id == s_.id, PayLink.status == "created").count()
    if s_.status == "awaiting" and not wait and not open_links:
        s_.status = "paid"
        db.commit()
        audit(db, "sale.terminal_paid", "sale", s_.id, user=user, method=method, rrn=result.get("rrn"))
        rro_ = f.get("rro_pending")
        if rro_ and not f.get("fiscal_code"):
            try:
                _fis.fiscalize(db, s_, rro_)
                _fis.send_receipt(db, s_)
            except (_fis.FiscalError, httpx.HTTPError) as e:
                f2 = dict(s_.fiscal or {})
                f2["error"] = str(e)[:200]
                s_.fiscal = f2
                db.commit()
    return {"sale": s_.status, "fiscal": (s_.fiscal or {}).get("fiscal_code"), "error": (s_.fiscal or {}).get("error")}


@app.post("/sales/{sid:int}/terminal/start")
async def sale_terminal_start(sid: int, request: Request, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    """Передати суму на термінал. PayLink на компʼютері каси → повертаємо адресу для браузера; інакше платить сервер."""
    from .services import fiscal as _fis
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    js = await request.json()
    method = str(js.get("method") or "")
    w = ((s_.fiscal or {}).get("term_wait") or {}).get(method)
    if s_.status != "awaiting" or not w:
        return JSONResponse({"error": "Цей чек не очікує оплати на терміналі"}, status_code=400)
    c = _fis.cfg(db)
    now_ = datetime.now()
    started = w.get("started")
    if started and not js.get("force") and (now_ - datetime.fromisoformat(started)).total_seconds() < 150:
        return JSONResponse({"already": True})
    f = dict(s_.fiscal)
    f["term_wait"] = {**f["term_wait"], method: {**w, "started": now_.isoformat(timespec="seconds")}}
    s_.fiscal = f
    db.commit()
    audit(db, "sale.terminal_start", "sale", sid, user=user, method=method, amount=w["amount"])
    if _fis.paylink_local(c):
        return JSONResponse({"mode": "browser", "url": f"{_fis.paylink_url(c)}/api/pos/{w['device']}/purchase", "amount": int(round(w["amount"] * 100))})
    try:
        res = await asyncio.to_thread(_fis.purchase, c, w["device"], w["amount"])
    except httpx.HTTPError as e:
        return JSONResponse({"mode": "server", "ok": False, "error": f"PayLink недоступний: {e}"[:200]})
    if res.get("success") and not res.get("error") and _term_has_data(res.get("result")):
        out = await asyncio.to_thread(_term_done, db, s_, method, res.get("result") or {}, user)
        return JSONResponse({"mode": "server", "ok": True, **out})
    err_ = str(res.get("error") or ("термінал не повернув даних оплати" if res.get("success") else "оплату не проведено"))[:200]
    audit(db, "sale.terminal_fail", "sale", sid, user=user, method=method, error=err_)
    return JSONResponse({"mode": "server", "ok": False, "error": err_})


@app.post("/sales/{sid:int}/terminal/result")
async def sale_terminal_result(sid: int, request: Request, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    """Результат PayLink від браузера (або «Оплата пройшла» вручну) → чек оплачено + фіскалізація."""
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    js = await request.json()
    method = str(js.get("method") or "")
    if s_.status != "awaiting" or method not in ((s_.fiscal or {}).get("term_wait") or {}):
        return JSONResponse({"error": "Цей чек не очікує оплати на терміналі", "sale": s_.status}, status_code=400)
    r = js.get("result") if isinstance(js.get("result"), dict) else {}
    if not js.get("ok") or not _term_has_data(r):               # лише підтверджена терміналом оплата з даними картки
        audit(db, "sale.terminal_fail", "sale", sid, user=user, method=method, error=str(js.get("error") or "немає даних оплати")[:200])
        return JSONResponse({"ok": False, "error": str(js.get("error") or "Термінал не повернув даних оплати — чек не закрито")[:200]})
    out = await asyncio.to_thread(_term_done, db, s_, method, r, user)
    return JSONResponse({"ok": True, **out})


@app.post("/sales/{sid:int}/fiscalize")
async def sale_fiscalize(sid: int, rro: int = Form(0), ret: str = Form(""), user: User = Depends(require("sales")),
                         db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    is_ret = bool(ret)
    if is_ret and not can(user, "sale_refund"):
        raise HTTPException(403)
    try:
        info = await asyncio.to_thread(_fis.fiscalize, db, s_, rro or None, None, 1.0, is_ret)
    except (_fis.FiscalError, httpx.HTTPError) as e:
        return RedirectResponse(_with_msg(f"/sales/{sid}", f"Checkbox: {e}", True), 303)
    audit(db, "sale.fiscal_return" if is_ret else "sale.fiscalize", "sale", sid, user=user, receipt=info.get("id"), code=info.get("fiscal_code"))
    if not is_ret and _fis.cfg(db)["send"].get("on"):
        await asyncio.to_thread(_fis.send_receipt, db, s_)
    return RedirectResponse(_with_msg(f"/sales/{sid}", ("Чек повернення" if is_ret else "Фіскальний чек") + f" № {info.get('fiscal_code') or info.get('id')}"), 303)


@app.post("/sales/{sid:int}/send-receipt")
async def sale_send_receipt(sid: int, ch: str = Form(""), user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    c = _fis.cfg(db)
    c["send"]["on"] = True
    err = None
    if not (s_.patient and s_.fiscal and s_.fiscal.get("url")):
        err = "Немає пацієнта або фіскального чека"
    else:
        from .integrations.messaging import send_to_patient
        text = (c["send"].get("text") or _fis.DEFAULT["send"]["text"]).replace("{url}", s_.fiscal["url"]).replace("{sum}", f"{s_.total:.2f}") \
            .replace("{name}", (s_.patient.full_name or "").split(" ")[0])
        prefer_ = [ch] if ch in ("telegram", "viber", "sms") else (c["send"].get("channels") or None)
        res = await asyncio.to_thread(send_to_patient, s_.patient, text, prefer_)
        err = None if res.ok else (res.error or "не надіслано")
        f = dict(s_.fiscal)
        f["sent"] = {"ok": res.ok, "channel": res.channel, "at": datetime.now().strftime("%d.%m.%Y %H:%M")}
        s_.fiscal = f
        db.commit()
    audit(db, "sale.send_receipt", "sale", sid, user=user, ok=err is None)
    lbl_ = {"telegram": " в Telegram", "viber": " у Viber", "sms": " в SMS"}.get(ch, "")
    return RedirectResponse(_with_msg(f"/sales/{sid}", err or f"Чек надіслано пацієнту{lbl_}", bool(err)), 303)


@app.get("/sales/{sid:int}/fiscal.png")
def sale_fiscal_png(sid: int, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    """Зображення фіскального чека Checkbox (GET /receipts/{id}/png)."""
    from .services import fiscal as _fis
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    f = s_.fiscal or {}
    if not f.get("id") or f.get("source") == "monobank":
        raise HTTPException(404)
    png = _fis.receipt_png(db, f)
    if not png:
        raise HTTPException(404)
    return Response(png, media_type="image/png", headers={"Cache-Control": "private, max-age=86400"})


def _safe_back(back: str, default: str) -> str:
    return back if back.startswith("/") and not back.startswith("//") else default


@app.post("/sales/{sid:int}/cancel")
def sale_cancel(sid: int, reason: str = Form(""), back: str = Form(""), user: User = Depends(require("sales")),
                      db: Session = Depends(db_dep)):
    """Скасувати чек, що очікує оплати (посилання на оплату інвалідуються)."""
    from .services import sales as sl
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    back = _safe_back(back, f"/sales/{sid}")
    try:
        sl.cancel(db, s_, reason.strip() or "скасовано в касі", user.id)
    except sl.SaleError as e:
        return RedirectResponse(_with_msg(back, str(e), True), 303)
    audit(db, "sale.cancel", "sale", sid, user=user, total=s_.total, reason=reason)
    return RedirectResponse(_with_msg(back, f"Чек Ч-{sid} скасовано"), 303)


@app.post("/sales/{sid:int}/archive")
def sale_archive(sid: int, on: str = Form("1"), back: str = Form(""), user: User = Depends(require("sales")),
                 db: Session = Depends(db_dep)):
    """Видалити скасований чек з каси в архів (on=1) або повернути з архіву (on=0)."""
    from .services import sales as sl
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    back = _safe_back(back, "/sales/archive")
    try:
        if on == "1" and s_.status == "awaiting":
            sl.cancel(db, s_, "видалено з каси", user.id)
            audit(db, "sale.cancel", "sale", sid, user=user, total=s_.total, reason="видалено з каси")
        sl.archive(db, s_, on == "1")
    except sl.SaleError as e:
        return RedirectResponse(_with_msg(back, str(e), True), 303)
    audit(db, "sale.archive" if on == "1" else "sale.unarchive", "sale", sid, user=user)
    return RedirectResponse(_with_msg(back, f"Чек Ч-{sid} " + ("перенесено в архів" if on == "1" else "повернуто з архіву")), 303)


@app.get("/sales/archive", response_class=HTMLResponse)
def sales_archive(request: Request, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .models import Sale
    rows = db.scalars(select(Sale).where((Sale.archived_at.is_not(None)) | (Sale.status == "cancelled"))
                      .order_by(Sale.at.desc()).limit(1000)).all()
    users = {u_.id: u_.name for u_ in db.scalars(select(User))}
    return render(request, "sales_archive.html", nav="finance", rows=rows, users=users)


@app.post("/sales/{sid:int}/refund")
def sale_refund(sid: int, reason: str = Form(""), del_password: str = Form(""), user: User = Depends(require("sale_refund")),
                db: Session = Depends(db_dep)):
    from .services import sales as sl
    from .models import Sale
    s_ = db.get(Sale, sid) or (_ for _ in ()).throw(HTTPException(404))
    if len(reason.strip()) < 3:
        return RedirectResponse(_with_msg(f"/sales/{sid}", "Вкажіть причину повернення", True), 303)
    err = _delpw(db, user, del_password, "sale_refund", "sale", sid)
    if err:
        return RedirectResponse(_with_err(f"/sales/{sid}", err), 303)
    try:
        sl.refund(db, s_, reason.strip(), user.id)
    except sl.SaleError as e:
        return RedirectResponse(_with_msg(f"/sales/{sid}", str(e), True), 303)
    audit(db, "sale.refund", "sale", sid, user=user, total=s_.total, reason=reason)
    from .services import fiscal as _fis
    if s_.fiscal and s_.fiscal.get("fiscal_code") and not s_.fiscal.get("return") and _fis.cfg(db).get("auto"):
        import threading
        threading.Thread(target=_fis.auto_return, args=(sid,), daemon=True).start()
    return RedirectResponse(_with_msg(f"/sales/{sid}", "Чек повернено"), 303)


# ---------------- Знижки й бонусна програма ----------------
@app.get("/settings/loyalty", response_class=HTMLResponse)
def settings_loyalty(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import loyalty as ly
    from .models import DiscountRule, BonusTxn, CbService
    rules = db.scalars(select(DiscountRule).order_by(DiscountRule.active.desc(), DiscountRule.name)).all()
    cats = sorted({c for (c,) in db.execute(select(CbService.category).distinct()) if c} |
                  {c for (c,) in db.execute(select(Product.brand).distinct()) if c})
    bal = ly.balances(db)
    calc = None
    qp = request.query_params
    if qp.get("calc"):
        items = []
        for t, pr, cat in (("service", qp.get("sp"), qp.get("sc")), ("product", qp.get("pp"), qp.get("pc"))):
            v = _num(pr or "")
            if v:
                items.append({"type": t, "price": v, "qty": 1, "category": cat or ""})
        pt = db.get(Patient, int(qp["pid"])) if (qp.get("pid") or "").isdigit() else None
        if items:
            dsc = ly.discount_for(db, items, pt, _day(qp.get("day", "")) or date.today())
            for it, dl in zip(items, dsc["lines"]):
                it["net"] = it["price"] - dl
            calc = {**dsc, "bonus": ly.bonus_accrual(db, items), "patient": pt,
                    "max_pay": ly.max_bonus_pay(db, dsc["subtotal"] - dsc["total"], pt.id) if pt else None}
    return render(request, "loyalty.html", nav="settings", s=ly.settings(db), rules=rules, cats=cats, ly=ly, calc=calc,
                  n_holders=sum(1 for v in bal.values() if v > 0.005), total_bonus=round(sum(v for v in bal.values() if v > 0), 2),
                  n_txns=db.scalar(select(func.count(BonusTxn.id))) or 0,
                  n_cb=db.scalar(select(func.count(Patient.id)).where(Patient.cb_bonus > 0)) or 0,
                  personal=db.scalars(select(Patient).where(Patient.discount_pct > 0).order_by(Patient.full_name)).all(),
                  edit=next((r for r in rules if str(r.id) == qp.get("edit", "")), None), WD=["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Нд"])


@app.post("/settings/loyalty/bonus")
async def settings_loyalty_bonus(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import loyalty as ly
    f = await request.form()
    s = ly.settings(db)
    s["enabled"] = bool(f.get("enabled"))
    s["accrue_on_bonus_paid"] = bool(f.get("accrue_on_bonus_paid"))
    for k in ("pct_services", "pct_products", "max_pay_pct", "welcome", "birthday", "min_check"):
        v = _num(str(f.get(k) or ""))
        s[k] = max(0.0, v) if v is not None else 0.0
    v = _num(str(f.get("expire_days") or ""))
    s["expire_days"] = int(v) if v and v > 0 else 0
    ly.save_settings(db, s)
    audit(db, "loyalty.settings", user=user, **s)
    return RedirectResponse(f"/settings/loyalty?msg={_q('Бонусну програму збережено')}", 303)


@app.post("/settings/loyalty/discount")
async def settings_loyalty_discount(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .models import DiscountRule
    f = await request.form()
    rid = str(f.get("rid") or "")
    r = db.get(DiscountRule, int(rid)) if rid.isdigit() else DiscountRule()
    name = str(f.get("name") or "").strip()[:150]
    val = _num(str(f.get("value") or ""))
    if not name or not val or val <= 0:
        return RedirectResponse(_with_msg("/settings/loyalty", "Вкажіть назву і розмір знижки", True), 303)
    r.name, r.kind = name, "amount" if f.get("kind") == "amount" else "percent"
    r.value = min(val, 100) if r.kind == "percent" else val
    r.applies_to = str(f.get("applies_to")) if f.get("applies_to") in ("all", "services", "products") else "all"
    r.categories = [str(x) for x in f.getlist("categories") if str(x)] or None
    r.date_from, r.date_to = _day(str(f.get("date_from") or "")), _day(str(f.get("date_to") or ""))
    r.weekdays = [int(x) for x in f.getlist("weekdays") if str(x).isdigit()] or None
    r.min_sum = _num(str(f.get("min_sum") or ""))
    r.segment = str(f.get("segment") or "") or None
    r.auto, r.combine, r.active = bool(f.get("auto")), bool(f.get("combine")), bool(f.get("active"))
    r.note = str(f.get("note") or "").strip()[:300] or None
    db.add(r)
    db.commit()
    audit(db, "loyalty.discount", "discount", r.id, user=user, name=r.name, value=r.value, kind=r.kind)
    return RedirectResponse(f"/settings/loyalty?msg={_q('Знижку збережено')}#disc", 303)


@app.post("/settings/loyalty/discount/{rid:int}/delete")
def settings_loyalty_discount_delete(rid: int, del_password: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .models import DiscountRule
    err = _delpw(db, user, del_password, "discount", "discount", rid)
    if err:
        return RedirectResponse(_with_err("/settings/loyalty#disc", err), 303)
    r = db.get(DiscountRule, rid)
    if r:
        audit(db, "loyalty.discount_delete", "discount", rid, user=user, name=r.name)
        db.delete(r)
        db.commit()
    return RedirectResponse("/settings/loyalty#disc", 303)


@app.post("/settings/loyalty/import")
def settings_loyalty_import(user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import loyalty as ly
    n = ly.import_cleverbox(db, user.id)
    audit(db, "loyalty.import", user=user, n=n)
    return RedirectResponse(f"/settings/loyalty?msg={_q(f'Перенесено залишки бонусів: {n} пацієнтів')}", 303)


@app.post("/patients/{pid:int}/bonus")
def patient_bonus(pid: int, amount: str = Form(...), sign: str = Form("+"), comment: str = Form(""),
                  user: User = Depends(require("loyalty_patient")), db: Session = Depends(db_dep)):
    from .services import loyalty as ly
    _patient_or_403(db, user, pid)
    v = _num(amount)
    if v and v > 0:
        try:
            ly.add(db, pid, -v if sign == "-" else v, "manual", comment.strip()[:300] or None, user_id=user.id)
        except ValueError as e:
            return RedirectResponse(_with_msg(f"/patients/{pid}#loyal", str(e), True), 303)
        audit(db, "loyalty.manual", "patient", pid, user=user, amount=-v if sign == "-" else v, comment=comment or None)
    return RedirectResponse(f"/patients/{pid}#loyal", 303)


@app.post("/patients/{pid:int}/discount")
def patient_discount(pid: int, discount_pct: str = Form(""), user: User = Depends(require("loyalty_patient")), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    v = _num(discount_pct)
    p.discount_pct = min(max(v, 0), 100) if v else None
    db.commit()
    audit(db, "loyalty.personal", "patient", pid, user=user, pct=p.discount_pct)
    return RedirectResponse(f"/patients/{pid}#loyal", 303)


# ---------------- Медична картка й фото до/після ----------------
MED_KINDS = {"consultation": "Консультація", "procedure": "Процедура", "control": "Контрольний огляд"}
MED_FIELDS = ("title", "complaints", "anamnesis", "objective", "diagnosis", "protocol", "recommendations")
PHOTO_STAGES = {"before": "До", "after": "Після", "progress": "У процесі"}


def _med_patient(db, user, pid) -> Patient:
    if not can(user, "medical"):
        raise HTTPException(403, "Медична картка доступна лише лікарям і власнику")
    return _patient_or_403(db, user, pid)


@app.get("/patients/{pid}/medcard", response_class=HTMLResponse)
def medcard_page(request: Request, pid: int, edit: int = 0, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .models import MedRecord, MedTemplate
    p = _med_patient(db, user, pid)
    recs = db.scalars(select(MedRecord).where(MedRecord.patient_id == pid, MedRecord.deleted_at.is_(None))
                      .order_by(MedRecord.at.desc())).all()
    photos = [f for f in p.files if f.kind == "photo"]
    by_rec: dict[int, list] = {}
    for f in photos:
        if f.record_id:
            by_rec.setdefault(f.record_id, []).append(f)
    docs = {d_.id: d_ for d_ in db.scalars(select(Doctor))}
    inject = [x.name for x in db.scalars(select(Product).where(Product.warehouse.in_(("inject", "classic", "single"))).order_by(Product.name)).all()][:1500]
    return render(request, "medcard.html", nav="patient", p=p, recs=recs, by_rec=by_rec, docs=docs, KINDS=MED_KINDS, STAGES=PHOTO_STAGES,
                  doctors=db.scalars(select(Doctor).where(Doctor.active.is_(True)).order_by(Doctor.name)).all(),
                  templates_=db.scalars(select(MedTemplate).order_by(MedTemplate.name)).all(), products=inject,
                  edit=next((r for r in recs if r.id == edit), None), now=datetime.now(), photos=photos,
                  my_doctor=user.doctor_id)


@app.post("/patients/{pid}/medcard/summary")
def medcard_summary(pid: int, med_allergies: str = Form(""), med_conditions: str = Form(""), med_contra: str = Form(""),
                    phototype: str = Form(""), user: User = Depends(current_user), db: Session = Depends(db_dep)):
    p = _med_patient(db, user, pid)
    p.med_allergies, p.med_conditions = med_allergies.strip() or None, med_conditions.strip() or None
    p.med_contra, p.phototype = med_contra.strip() or None, phototype.strip()[:10] or None
    db.commit()
    audit(db, "medcard.summary", "patient", pid, user=user)
    return RedirectResponse(f"/patients/{pid}/medcard?msg={_q('Збережено')}", 303)


@app.post("/patients/{pid}/medcard/record")
async def medcard_record(request: Request, pid: int, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .models import MedRecord, MedTemplate
    p = _med_patient(db, user, pid)
    f = await request.form()
    rid = str(f.get("rid") or "")
    r = db.get(MedRecord, int(rid)) if rid.isdigit() else None
    if r and r.patient_id != pid:
        raise HTTPException(404)
    new = r is None
    r = r or MedRecord(patient_id=pid, created_by=user.id)
    at_s = str(f.get("at") or "")
    try:
        r.at = datetime.fromisoformat(at_s) if at_s else (r.at or datetime.now())
    except ValueError:
        r.at = r.at or datetime.now()
    r.kind = str(f.get("kind") or "procedure") if str(f.get("kind")) in MED_KINDS else "procedure"
    did = str(f.get("doctor_id") or "")
    r.doctor_id = int(did) if did.isdigit() else (user.doctor_id or None)
    for k in MED_FIELDS:
        v = str(f.get(k) or "").strip()
        setattr(r, k, (v[:200] if k == "title" else v) or None)
    drugs = []
    for i in range(12):
        nm = str(f.get(f"drug_{i}") or "").strip()
        if nm:
            drugs.append({k: str(f.get(f"{k}_{i}") or "").strip()[:120] for k in ("batch", "expiry", "dose", "zone")} | {"name": nm[:200]})
    r.drugs = drugs or None
    r.next_visit = _day(str(f.get("next_visit") or ""))
    r.consent = bool(f.get("consent"))
    if not new:
        r.updated_at = datetime.utcnow()
    db.add(r)
    db.commit()
    for up in f.getlist("photos"):
        if getattr(up, "filename", None):
            pf = ingest.ingest_photo(db, await up.read(), up.filename or "photo.jpg", str(f.get("photo_caption") or "") or None, patient=p)
            pf.stage = str(f.get("photo_stage") or "") if str(f.get("photo_stage") or "") in PHOTO_STAGES else None
            pf.zone = str(f.get("photo_zone") or "").strip()[:80] or None
            pf.taken_at, pf.record_id = r.at.date(), r.id
            db.commit()
    tname = str(f.get("save_template") or "").strip()
    if tname:
        db.add(MedTemplate(name=tname[:150], kind=r.kind, created_by=user.id,
                           fields={k: getattr(r, k) for k in MED_FIELDS if k not in ("anamnesis",) and getattr(r, k)}))
        db.commit()
    audit(db, "medcard.record_add" if new else "medcard.record_edit", "patient", pid, user=user, record=r.id, kind=r.kind)
    return RedirectResponse(f"/patients/{pid}/medcard?msg={_q('Запис збережено')}#r{r.id}", 303)


@app.post("/patients/{pid}/medcard/record/{rid:int}/delete")
def medcard_record_delete(pid: int, rid: int, del_password: str = Form(""), user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .models import MedRecord
    _med_patient(db, user, pid)
    err = _delpw(db, user, del_password, "med_record", "patient", pid)
    if err:
        return RedirectResponse(_with_err(f"/patients/{pid}/medcard", err), 303)
    r = db.get(MedRecord, rid)
    if r and r.patient_id == pid:
        r.deleted_at = datetime.utcnow()
        db.commit()
        audit(db, "medcard.record_delete", "patient", pid, user=user, record=rid)
    return RedirectResponse(f"/patients/{pid}/medcard?msg={_q('Запис видалено')}", 303)


@app.post("/medcard/templates/{tid:int}/delete")
def medcard_template_delete(tid: int, back: str = Form("/"), del_password: str = Form(""), user: User = Depends(require("medical")),
                            db: Session = Depends(db_dep)):
    from .models import MedTemplate
    back = back if back.startswith("/patients/") else "/"
    err = _delpw(db, user, del_password, "med_template", None, tid)
    if err:
        return RedirectResponse(_with_err(back, err), 303)
    t = db.get(MedTemplate, tid)
    if t:
        db.delete(t)
        db.commit()
    return RedirectResponse(back, 303)


@app.get("/patients/{pid}/medcard/print", response_class=HTMLResponse)
def medcard_print(request: Request, pid: int, rid: int = 0, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .models import MedRecord
    p = _med_patient(db, user, pid)
    q = select(MedRecord).where(MedRecord.patient_id == pid, MedRecord.deleted_at.is_(None))
    if rid:
        q = q.where(MedRecord.id == rid)
    recs = db.scalars(q.order_by(MedRecord.at)).all()
    audit(db, "medcard.print", "patient", pid, user=user, record=rid or None)
    return render(request, "medcard_print.html", nav="print", p=p, recs=recs, KINDS=MED_KINDS,
                  docs={d_.id: d_ for d_ in db.scalars(select(Doctor))}, now=datetime.now())


@app.get("/patients/{pid}/photos", response_class=HTMLResponse)
def patient_photos(request: Request, pid: int, a: int = 0, b: int = 0, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    photos = sorted([f for f in p.files if f.kind == "photo"], key=lambda f: (f.taken_at or f.created_at.date(), f.id))
    zones: dict[str, list] = {}
    for f in photos:
        zones.setdefault(f.zone or "Без зони", []).append(f)
    fa = next((f for f in photos if f.id == a), None) or next((f for f in photos if f.stage == "before"), photos[0] if photos else None)
    fb = next((f for f in photos if f.id == b), None) or next((f for f in reversed(photos) if f.stage == "after" and f is not fa),
                                                              photos[-1] if len(photos) > 1 else None)
    return render(request, "photos.html", nav="patient", p=p, photos=photos, zones=zones, fa=fa, fb=fb, STAGES=PHOTO_STAGES)


@app.post("/files/{fid:int}/meta")
def file_meta(fid: int, stage: str = Form(""), zone: str = Form(""), taken_at: str = Form(""), caption: str = Form(""),
              user: User = Depends(require("upload")), db: Session = Depends(db_dep)):
    f = db.get(PatientFile, fid)
    if not f or not f.patient_id:
        raise HTTPException(404)
    _patient_or_403(db, user, f.patient_id)
    f.stage = stage if stage in PHOTO_STAGES else None
    f.zone, f.taken_at, f.caption = zone.strip()[:80] or None, _day(taken_at), caption.strip() or f.caption
    db.commit()
    return RedirectResponse(f"/patients/{f.patient_id}/photos?a={fid}", 303)


@app.get("/files/{fid}")
def file_get(fid: int, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    f = db.get(PatientFile, fid)
    if not f or not Path(f.path).exists():
        raise HTTPException(404, "Файл не знайдено")
    if f.patient_id:
        _patient_or_403(db, user, f.patient_id)
    elif not can(user, "files_assign"):
        raise HTTPException(403, "Немає доступу")
    return FileResponse(f.path, media_type=mimetypes.guess_type(f.path)[0] or "application/octet-stream")


@app.get("/unassigned", response_class=HTMLResponse)
def unassigned(request: Request, user: User = Depends(require("files_assign")), db: Session = Depends(db_dep)):
    files = db.scalars(select(PatientFile).where(PatientFile.patient_id.is_(None))
                       .order_by(PatientFile.created_at.desc())).all()
    return render(request, "unassigned.html", files=files,
                  patients=db.scalars(select(Patient).order_by(Patient.full_name)).all())


@app.post("/files/{fid}/assign")
def file_assign(fid: int, patient_id: int = Form(...), user: User = Depends(require("files_assign")),
                db: Session = Depends(db_dep)):
    f = db.get(PatientFile, fid)
    f.patient_id = patient_id
    db.commit()
    audit(db, "file.assign", "patient", patient_id, user=user, file_id=fid)
    return RedirectResponse("/unassigned", 303)


# ---------------- зіставлення з CleverBox ----------------

@app.get("/mapping", response_class=HTMLResponse)
def mapping(request: Request, user: User = Depends(require("mapping")), db: Session = Depends(db_dep)):
    names: dict[tuple[str, str], str] = {}
    for it in db.scalars(select(PlanItem).where(PlanItem.kind == "procedure")):
        names.setdefault((it.kind, norm(it.name)), it.name)
    services = db.scalars(select(CbService).order_by(CbService.category, CbService.name)).all()
    aliases = {(a.kind, a.plan_name_norm): a for a in db.scalars(select(Alias))}
    seen_names = sorted({v.service_name for v in db.scalars(select(Visit)) if v.service_name})
    rows = [(key, name, aliases.get(key), [s for s in seen_names if name_match(name, s) >= 0.6])
            for key, name in sorted(names.items(), key=lambda kv: kv[1])]
    return render(request, "mapping.html", rows=rows, services=services)


@app.post("/mapping")
async def mapping_save(request: Request, user: User = Depends(require("mapping")), db: Session = Depends(db_dep)):
    form = await request.form()
    kind, key = form["kind"], form["key"]
    ids = [int(x) for x in form.getlist("cb_ids") if str(x).isdigit()]
    al = db.scalar(select(Alias).where(Alias.kind == kind, Alias.plan_name_norm == key))
    if not ids:
        if al:
            db.delete(al)
    else:
        al = al or Alias(kind=kind, plan_name_norm=key)
        al.cb_ids = ids
        al.cb_names = [db.get(CbService, i).name for i in ids if db.get(CbService, i)]
        db.add(al)
    db.commit()
    audit(db, "mapping.update", user=user, key=key, cb_ids=ids)
    return RedirectResponse("/mapping", 303)


# ---------------- доступи (лише власник) ----------------

@app.get("/access", response_class=HTMLResponse)
def access(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    return render(request, "access.html", tab="users",
                  users=db.scalars(select(User).order_by(User.role.desc(), User.name)).all(),
                  doctors=db.scalars(select(Doctor).order_by(Doctor.name)).all())


@app.post("/access/users")
def user_save(request: Request, uid: str = Form(""), login: str = Form(...), name: str = Form(...),
              role: str = Form(...), doctor_id: str = Form(""), password: str = Form(""), active: str = Form(""),
              delete_pw: str = Form(""), delete_pw_clear: str = Form(""),
              user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    if role not in ROLES:
        raise HTTPException(400, "Невідома роль")
    u = db.get(User, int(uid)) if uid.isdigit() else None
    if u is None:
        if db.scalar(select(User).where(func.lower(User.login) == login.strip().lower())):
            raise HTTPException(400, f"Логін «{login}» уже зайнятий")
        if len(password) < 8:
            raise HTTPException(400, "Пароль — щонайменше 8 символів")
        u = User(login=login.strip(), name=name.strip(), role=role, password_hash=hash_password(password))
        db.add(u)
        action = "user.create"
    else:
        action = "user.update"
        if u.id == user.id and (role != "owner" or not active):
            raise HTTPException(400, "Не можна зняти з себе роль власника або деактивувати себе")
        u.login, u.name, u.role = login.strip(), name.strip(), role
        u.active = bool(active)
        if password:
            if len(password) < 8:
                raise HTTPException(400, "Пароль — щонайменше 8 символів")
            u.password_hash = hash_password(password)
    u.doctor_id = int(doctor_id) if role == "doctor" and doctor_id.isdigit() else None
    if delete_pw:
        if len(delete_pw) < 4:
            raise HTTPException(400, "Пароль для видалення — щонайменше 4 символи")
        if verify_password(delete_pw, u.password_hash or ""):
            raise HTTPException(400, "Пароль для видалення має відрізнятися від пароля входу")
        u.delete_pw_hash = hash_password(delete_pw)
    elif delete_pw_clear:
        u.delete_pw_hash = None
    db.commit()
    audit(db, action, "user", u.id, user=user, login=u.login, role=u.role, active=u.active,
          password_changed=bool(password), delete_pw=("set" if delete_pw else "cleared" if delete_pw_clear else None))
    return RedirectResponse("/access", 303)


@app.get("/access/roles", response_class=HTMLResponse)
def access_roles(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import roles as rl
    return render(request, "roles.html", nav="access", tab="roles", rs=rl.roles(db), rl=rl, counts=rl.users_by_role(db),
                  edit=request.query_params.get("edit", ""), view=request.query_params.get("view", ""),
                  copy=request.query_params.get("copy", ""))


@app.post("/access/roles")
async def access_roles_save(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import roles as rl
    f = await request.form()
    key = str(f.get("key") or "") or None
    if key == "owner":
        return RedirectResponse("/access/roles", 303)
    name = str(f.get("name") or "").strip()
    if not name:
        return RedirectResponse(_with_msg("/access/roles", "Вкажіть назву ролі", True), 303)
    k = rl.save_role(db, key, name, [str(x) for x in f.getlist("perms")], bool(f.get("hidden")))
    audit(db, "roles.save", user=user, role=k, name=name, perms=sorted(str(x) for x in f.getlist("perms")))
    return RedirectResponse(_with_msg("/access/roles", f"Роль «{name}» збережено") + f"#r-{k}", 303)


@app.post("/access/roles/{key}/reset")
def access_roles_reset(key: str, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import roles as rl
    rl.reset_role(db, key)
    audit(db, "roles.reset", user=user, role=key)
    return RedirectResponse(_with_msg("/access/roles", "Права ролі скинуто до заводських"), 303)


@app.post("/access/roles/{key}/delete")
def access_roles_delete(key: str, del_password: str = Form(""), user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import roles as rl
    err = _delpw(db, user, del_password, "role", None, None)
    if err:
        return RedirectResponse(_with_err("/access/roles", err), 303)
    e = rl.delete_role(db, key)
    if e:
        return RedirectResponse(_with_msg("/access/roles", e, True), 303)
    audit(db, "roles.delete", user=user, role=key)
    return RedirectResponse(_with_msg("/access/roles", "Роль видалено"), 303)


@app.get("/access/api", response_class=HTMLResponse)
def access_api(request: Request, user: User = Depends(require("api")), db: Session = Depends(db_dep)):
    return render(request, "access.html", tab="api",
                  keys=db.scalars(select(ApiKey).order_by(ApiKey.created_at.desc())).all(),
                  hooks=db.scalars(select(Webhook).order_by(Webhook.created_at.desc())).all(),
                  new_key=request.session.pop("new_key", None), new_secret=request.session.pop("new_secret", None))


@app.post("/access/api/keys")
def key_create(request: Request, name: str = Form(...), scopes: list[str] = Form([]),
               user: User = Depends(require("api")), db: Session = Depends(db_dep)):
    raw, prefix, h = new_api_key()
    k = ApiKey(name=name.strip(), prefix=prefix, key_hash=h,
               scopes=[s for s in scopes if s in ("read", "write", "medical", "bot")] or ["read"], created_by=user.id)
    db.add(k)
    db.commit()
    audit(db, "api_key.create", "api_key", k.id, user=user, name=k.name, scopes=k.scopes)
    request.session["new_key"] = raw          # показуємо один раз
    return RedirectResponse("/access/api", 303)


@app.post("/access/api/keys/{kid}/revoke")
def key_revoke(kid: int, user: User = Depends(require("api")), db: Session = Depends(db_dep)):
    k = db.get(ApiKey, kid)
    k.active = False
    db.commit()
    audit(db, "api_key.revoke", "api_key", k.id, user=user, name=k.name)
    return RedirectResponse("/access/api", 303)


@app.post("/access/api/hooks")
def hook_create(request: Request, url: str = Form(...), events: list[str] = Form([]),
                user: User = Depends(require("api")), db: Session = Depends(db_dep)):
    if not url.startswith("https://"):
        raise HTTPException(400, "Вебхук має бути на https://")
    secret = secrets.token_hex(24)
    h = Webhook(url=url.strip(), events=[e for e in events if e in webhooks.EVENTS], secret=secret)
    db.add(h)
    db.commit()
    audit(db, "webhook.create", "webhook", h.id, user=user, url=h.url, events=h.events)
    request.session["new_secret"] = secret
    return RedirectResponse("/access/api", 303)


@app.post("/access/api/hooks/{hid}/toggle")
def hook_toggle(hid: int, user: User = Depends(require("api")), db: Session = Depends(db_dep)):
    h = db.get(Webhook, hid)
    h.active = not h.active
    db.commit()
    audit(db, "webhook." + ("enable" if h.active else "disable"), "webhook", h.id, user=user)
    return RedirectResponse("/access/api", 303)


@app.post("/access/api/hooks/{hid}/test")
def hook_test(hid: int, user: User = Depends(require("api")), db: Session = Depends(db_dep)):
    h = db.get(Webhook, hid)
    import json
    body = json.dumps({"event": "test", "sent_at": datetime.now().isoformat(timespec="seconds"),
                       "data": {"message": "Тестова подія MOLODIST CRM"}}, ensure_ascii=False).encode()
    webhooks._deliver(h.id, h.url, h.secret, body)
    return RedirectResponse("/access/api", 303)


@app.get("/access/audit", response_class=HTMLResponse)
def access_audit(request: Request, user: User = Depends(require("audit")), db: Session = Depends(db_dep)):
    return render(request, "access.html", tab="audit",
                  log=db.scalars(select(AuditLog).order_by(AuditLog.at.desc()).limit(300)).all())


@app.get("/api/openapi.json", include_in_schema=False)
def openapi_json(user: User = Depends(require("api"))):
    return app.openapi()


@app.get("/api/docs", include_in_schema=False)
def api_docs(user: User = Depends(require("api"))):
    return get_swagger_ui_html(openapi_url="/api/openapi.json", title="MOLODIST CRM API")


# ---------------- дії ----------------

@app.post("/actions/sync")
async def action_sync(user: User = Depends(require("sync")), db: Session = Depends(db_dep)):
    await job_sync()
    audit(db, "sync.manual", user=user)
    return RedirectResponse("/", 303)


@app.post("/actions/sheets")
async def action_sheets(user: User = Depends(require("sheets"))):
    try:
        url = await asyncio.to_thread(run_sheets)
        return RedirectResponse(url, 303)
    except Exception as e:          # noqa: BLE001
        state["last_error"] = f"Sheets: {e}"
        return RedirectResponse("/", 303)


@app.post("/actions/digest")
async def action_digest(user: User = Depends(require("send_digest"))):
    await job_digest()
    return RedirectResponse("/", 303)


# ---------------- дані: експорт / імпорт / перенос / резервна копія ----------------

import os
import tempfile
from starlette.background import BackgroundTask
from fastapi.responses import Response
from .services import exporter, importer, jobs

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LAST_IMPORT: dict[int, list[dict]] = {}       # результати останнього імпорту по користувачу (не в cookie — може бути великим)


def _attach(name: str) -> dict:
    from urllib.parse import quote
    return {"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"}


def _tmp_response(path: str, name: str, media: str = "application/zip"):
    return FileResponse(path, media_type=media, headers=_attach(name), background=BackgroundTask(os.unlink, path))


@app.get("/data", response_class=HTMLResponse)
def data_page(request: Request, user: User = Depends(require("export_reports")), db: Session = Depends(db_dep)):
    last = db.get(SyncState, "last_migration")
    cb_success = db.get(SyncState, "cleverbox_last_success")
    cb_error = db.get(SyncState, "cleverbox_last_error")
    job = state.get("scheduler").get_job("cleverbox_sync") if state.get("scheduler") else None
    from .services import backups as bk
    return render(request, "data.html", jobs=jobs.recent(), cb=bool(settings.cleverbox_token),
                  backups=bk.list_backups() if can(user, "backup") else [], settings=settings,
                  last_migration=last.value if last else None, reports=LAST_IMPORT.pop(user.id, None),
                  cb_last_success=cb_success.value if cb_success else None,
                  cb_last_error=cb_error.value if cb_error else None,
                  cb_next_sync=job.next_run_time if job else None,
                  cb_auto_sync=auto_sync_enabled(),
                  cb_notice=request.query_params.get("cb_notice"),
                  wh=cb_webhooks.status(db) if user.role == "owner" else None,
                  wh_url=cb_webhooks.hook_url(db) if user.role == "owner" else None,
                  wh_subs=cb_webhooks.subscriptions(db) if user.role == "owner" and settings.cleverbox_token
                  and request.query_params.get("wh") else None)


@app.get("/export/report.xlsx")
def export_report(df: str | None = None, dt_: str | None = None, user: User = Depends(require("export_reports")),
                  db: Session = Depends(db_dep)):
    tables = exporter.report_tables(db, user, date.fromisoformat(df) if df else None,
                                    date.fromisoformat(dt_) if dt_ else None)
    audit(db, "export.report", user=user, format="xlsx", df=df, dt=dt_)
    return Response(exporter.to_xlsx(tables), media_type=XLSX,
                    headers=_attach(f"MOLODIST_звіт_{datetime.now():%Y-%m-%d}.xlsx"))


@app.get("/export/report.csv.zip")
def export_report_csv(df: str | None = None, dt_: str | None = None, user: User = Depends(require("export_reports")),
                      db: Session = Depends(db_dep)):
    tables = exporter.report_tables(db, user, date.fromisoformat(df) if df else None,
                                    date.fromisoformat(dt_) if dt_ else None)
    audit(db, "export.report", user=user, format="csv", df=df, dt=dt_)
    return Response(exporter.to_csv_zip(tables), media_type="application/zip",
                    headers=_attach(f"MOLODIST_звіт_{datetime.now():%Y-%m-%d}_csv.zip"))


@app.get("/export/files.zip")
def export_files(user: User = Depends(require("export_files")), db: Session = Depends(db_dep)):
    pats = [p for p in db.scalars(select(Patient)) if patient_visible(user, p) and p.files]
    audit(db, "export.files", user=user, patients=len(pats))
    return _tmp_response(exporter.files_zip(db, pats), f"MOLODIST_файли_{datetime.now():%Y-%m-%d}.zip")


@app.get("/patients/{pid}/files.zip")
def export_patient_files(pid: int, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    audit(db, "export.files", "patient", pid, user=user)
    return _tmp_response(exporter.files_zip(db, [p]), f"{p.full_name}_файли.zip")


@app.get("/patients/{pid}/print", response_class=HTMLResponse)
def patient_print(request: Request, pid: int, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    statuses = [compute_plan(db, pl) for pl in p.plans if plan_visible(user, pl)]
    return render(request, "print.html", p=p, statuses=statuses, now=datetime.now())


@app.get("/import/template.xlsx")
def import_template(user: User = Depends(require("import"))):
    return Response(importer.template_xlsx(), media_type=XLSX, headers=_attach("MOLODIST_шаблон_імпорту.xlsx"))


def _cash_imp():
    from .services import cashflow
    return cashflow


def _cert_imp():
    from .services import certificates
    return certificates


def _svc_imp():
    from .services import cb_services_import
    return cb_services_import


@app.post("/import/table")
async def import_table(request: Request, file: UploadFile = File(...), kind: str = Form(""),
                       user: User = Depends(require("import")), db: Session = Depends(db_dep)):
    data = await file.read()
    try:
        from .services import cb_clients_import as cbc
        sheets = importer.read_sheets(data, file.filename or "data.xlsx")
        if sheets and sheets[0][1] and cbc.is_clients_export(sheets[0][1][0]):     # вивантаження «Клієнти» CleverBox
            rep = await asyncio.to_thread(cbc.import_clients, db, sheets[0][1][1:], sheets[0][1][0])
            rep["sheet"] = file.filename
            reports = [rep]
        elif sheets and sheets[0][1] and _cash_imp().is_profit_export(sheets[0][1][0]):   # «Прибуток» → cashflow
            rep = await asyncio.to_thread(_cash_imp().import_profit, db, sheets[0][1][1:], sheets[0][1][0])
            _cash_imp().auto_fill(db)
            rep["sheet"] = file.filename
            reports = [rep]
        elif sheets and sheets[0][1] and _cert_imp().is_certificates_export(sheets[0][1][0]):   # «Сертифікати»
            rep = await asyncio.to_thread(_cert_imp().import_certificates, db, sheets[0][1][1:], sheets[0][1][0])
            rep["sheet"] = file.filename
            reports = [rep]
        elif sheets and sheets[0][1] and _cert_imp().is_subscriptions_export(sheets[0][1][0]):   # «Абонементи»
            rep = await asyncio.to_thread(_cert_imp().import_subscriptions, db, sheets[0][1][1:], sheets[0][1][0])
            rep["sheet"] = file.filename
            reports = [rep]
        elif sheets and sheets[0][1] and _svc_imp().is_services_export(sheets[0][1][0]):   # «Послуги» CleverBox
            rep = await asyncio.to_thread(_svc_imp().import_services, db, sheets[0][1][1:], sheets[0][1][0])
            rep["sheet"] = file.filename
            reports = [rep]
        else:
            reports = importer.import_file(db, data, file.filename or "data.xlsx", kind or None)
    except Exception as e:          # noqa: BLE001
        reports = [{"sheet": file.filename, "error": str(e)}]
    audit(db, "import.table", user=user, file=file.filename,
          summary=[{k: r.get(k) for k in ("sheet", "type", "plans", "visits", "sales", "patients_created", "error")}
                   for r in reports])
    LAST_IMPORT[user.id] = reports
    return RedirectResponse("/data#import", 303)


@app.post("/import/pdf-zip")
async def import_pdf_zip(request: Request, file: UploadFile = File(...), user: User = Depends(require("import")),
                         db: Session = Depends(db_dep)):
    try:
        rep = importer.import_pdf_zip(db, await file.read())
    except Exception as e:          # noqa: BLE001
        rep = {"error": str(e)}
    rep["sheet"] = file.filename
    rep["type"] = "PDF Observ"
    audit(db, "import.pdf_zip", user=user, file=file.filename, plans=rep.get("plans"))
    LAST_IMPORT[user.id] = [rep]
    return RedirectResponse("/data#import", 303)


@app.post("/migrate/cleverbox")
def migrate_cleverbox(months: int = Form(24), all_clients: str = Form(""), all_history: str = Form(""),
                      user: User = Depends(require("migrate")), db: Session = Depends(db_dep)):
    from .integrations.cleverbox import migrate
    full_history = all_history == "1"
    job = jobs.start("Перенос з CleverBox", migrate, months=max(1, min(months, 120)),
                     all_clients=bool(all_clients), start_date=date(2000, 1, 1) if full_history else None)
    audit(db, "migrate.cleverbox", user=user, months="all" if full_history else months,
          all_clients=bool(all_clients), job=job.id)
    return RedirectResponse("/data#migrate", 303)


@app.post("/settings/cleverbox-token")
async def cleverbox_token_save(token: str = Form(...), user: User = Depends(require("migrate")),
                               db: Session = Depends(db_dep)):
    """Validate first; never display or audit the supplied secret."""
    token = token.strip()
    if not token:
        return RedirectResponse("/data?cb_notice=Введіть+токен#migrate", 303)
    try:
        from .integrations.cleverbox import CleverBox, save_token
        await asyncio.to_thread(lambda: CleverBox(token=token).filials())
        save_token(db, token)
        ensure_sync_schedule()
        audit(db, "cleverbox.token_saved", user=user)
        return RedirectResponse("/data?cb_notice=Токен+перевірено+і+збережено#migrate", 303)
    except Exception as e:  # noqa: BLE001
        log.warning("CleverBox token validation failed: %s", type(e).__name__)
        audit(db, "cleverbox.token_rejected", user=user)
        return RedirectResponse("/data?cb_notice=Токен+не+пройшов+перевірку#migrate", 303)


@app.post("/actions/cleverbox-sync")
async def cleverbox_manual_sync(user: User = Depends(require("sync")), db: Session = Depends(db_dep)):
    if not settings.cleverbox_token:
        return RedirectResponse("/data?cb_notice=Спершу+збережіть+токен#migrate", 303)
    await job_sync()
    audit(db, "cleverbox.sync.manual", user=user)
    return RedirectResponse("/data#migrate", 303)


@app.post("/settings/cleverbox-auto-sync")
def cleverbox_auto_sync_save(enabled: str = Form("0"), user: User = Depends(require("migrate")),
                             db: Session = Depends(db_dep)):
    value = "1" if enabled == "1" else "0"
    item = db.get(SyncState, "cleverbox_auto_sync") or SyncState(key="cleverbox_auto_sync")
    item.value = value
    db.merge(item)
    db.commit()
    ensure_sync_schedule()
    audit(db, "cleverbox.auto_sync", user=user, enabled=value == "1")
    notice = "Автооновлення+увімкнено" if value == "1" else "Автооновлення+вимкнено"
    return RedirectResponse(f"/data?cb_notice={notice}#migrate", 303)


@app.get("/jobs/{jid}")
def job_status(jid: str, user: User = Depends(require("import"))):
    j = jobs.JOBS.get(jid)
    if not j:
        raise HTTPException(404, "Задачу не знайдено")
    return {"id": j.id, "title": j.title, "status": j.status, "progress": j.progress[-30:],
            "result": j.result, "error": j.error}


@app.get("/backup/download")
def backup_download(with_files: int = 1, user: User = Depends(require("backup")), db: Session = Depends(db_dep)):
    audit(db, "backup.download", user=user, with_files=bool(with_files))
    return _tmp_response(exporter.backup_zip(db, bool(with_files)),
                         f"MOLODIST_CRM_backup_{datetime.now():%Y-%m-%d_%H%M}.zip")


@app.post("/backup/now")
async def backup_now(request: Request, user: User = Depends(require("backup")), db: Session = Depends(db_dep)):
    from .services import backups
    try:
        rep = await asyncio.to_thread(backups.run_backup, "manual")
        audit(db, "backup.create", user=user, **{k: v for k, v in rep.items() if k != "trigger"})
    except Exception as e:          # noqa: BLE001
        state["last_error"] = f"Бекап: {e}"
    return RedirectResponse("/data#backup", 303)


@app.get("/backup/file/{name}")
def backup_file(name: str, user: User = Depends(require("backup")), db: Session = Depends(db_dep)):
    from .services import backups
    try:
        p = backups.backup_path(name)
    except FileNotFoundError:
        raise HTTPException(404, "Копію не знайдено")
    audit(db, "backup.download", user=user, file=name)
    return FileResponse(p, media_type="application/zip", headers=_attach(name))


@app.post("/backup/restore")
async def backup_restore(request: Request, file: UploadFile = File(...), confirm: str = Form(""),
                         password: str = Form(""), user: User = Depends(require("backup")),
                         db: Session = Depends(db_dep)):
    from .services import backups
    if confirm.strip().upper() != "ВІДНОВИТИ":
        raise HTTPException(400, "Для підтвердження введіть слово ВІДНОВИТИ")
    tmp = tempfile.NamedTemporaryFile(suffix=".zip", delete=False)
    while chunk := await file.read(1 << 20):
        tmp.write(chunk)
    tmp.close()
    plain = tmp.name
    try:
        if backups.is_encrypted(tmp.name):           # копія з Telegram (.zip.enc)
            if not password:
                raise HTTPException(400, "Копія зашифрована — введіть пароль (BACKUP_PASSWORD)")
            plain = tmp.name + ".dec"
            try:
                backups.decrypt_file(tmp.name, plain, password)
            except ValueError as e:
                raise HTTPException(400, str(e))
        try:
            rep = exporter.restore_backup(db, plain)
        except ValueError as e:
            raise HTTPException(400, str(e))
    finally:
        for f in {tmp.name, plain}:
            if os.path.exists(f):
                os.unlink(f)
    request.session.clear()            # користувачі могли змінитися — входимо заново
    log.warning("Backup restored: %s", rep)
    with SessionLocal() as db2:
        audit(db2, "backup.restore", details_rows=sum(rep["rows"].values()), files=rep["files"],
              backup_created=rep["backup_created"])
    return RedirectResponse("/login", 303)


# ---------------- Клієнти / Календар / Задачі / Продукти / Склад / Фінанси ----------------

from datetime import timedelta
from .models import Task, Product, StockLevel, Check, StockDoc
from .services import clients as clients_svc, tasks as tasks_svc, calendar as cal_svc, catalog, finance as fin_svc


@app.get("/clients", response_class=HTMLResponse)
def clients_page(request: Request, seg: str = "all", q: str = "", doctor: int | None = None, page: int = 1,
                 user: User = Depends(current_user), db: Session = Depends(db_dep)):
    words = set(norm(q).split()) if q else set()
    if user.role == "doctor":
        doctor = user.doctor_id
    pats = []
    from sqlalchemy.orm import selectinload
    for p in db.scalars(select(Patient).options(selectinload(Patient.plans)).order_by(Patient.full_name)):
        if not patient_visible(user, p):
            continue
        if words and not (words <= set(norm(p.full_name).split()) or (p.phone and q.strip() in (p.phone or ""))):
            continue
        if doctor and p.doctor_id != doctor and not any(pl.doctor_id == doctor for pl in p.plans):
            continue
        pats.append(p)
    rows = clients_svc.client_rows(db, pats)
    counts = clients_svc.segment_counts(rows)
    rows = [r for r in rows if seg in r.segment_flags]
    order = {"red": 0, "yellow": 1, "green": 2, "done": 3}
    rows.sort(key=lambda r: (order.get(r.plan_status.flag, 4) if r.plan_status else 5,
                             -(r.last_visit.timestamp() if r.last_visit else 0)))
    per = 100
    total = len(rows)
    rows = rows[(page - 1) * per: page * per]
    return render(request, "clients.html", rows=rows, counts=counts, seg=seg, q=q, doctor=doctor, page=page,
                  pages=max(1, (total + per - 1) // per), total=total, SEGMENTS=clients_svc.SEGMENTS,
                  doctors=db.scalars(select(Doctor).order_by(Doctor.name)).all(),
                  unassigned=db.scalar(select(func.count()).select_from(PatientFile)
                                       .where(PatientFile.patient_id.is_(None))) if can(user, "files_assign") else 0)


@app.get("/calendar", response_class=HTMLResponse)
def calendar_page(request: Request, start: str | None = None, doc: str = "", view: str = "day",
                  user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .services import specialists as spec_svc
    today = date.today()
    try:
        s = date.fromisoformat(start) if start else today
    except ValueError:
        s = today
    view = view if view in ("day", "week", "month") else "day"
    doc_id = int(doc) if doc.isdigit() else None
    ctx = dict(view=view, today=today, doc=doc_id, specialists=spec_svc.calendar_doctors(db),
               color=spec_svc.color, initials=spec_svc.initials)
    if view == "month":
        first = s.replace(day=1)
        nxt = (first + timedelta(days=32)).replace(day=1)
        prv = (first - timedelta(days=1)).replace(day=1)
        return render(request, "calendar.html", start=first, end=nxt - timedelta(days=1), prev=prv, next=nxt,
                      month=cal_svc.month_summary(db, user, first, doc_id), ev={}, cols={}, days=0, **ctx)
    days = 1 if view == "day" else 7
    if view == "week":
        s = s - timedelta(days=s.weekday())
    ev = cal_svc.events(db, user, s, days, doc_id)
    cols = cal_svc.columns(db, ev, fixed=(view == "day" and not doc_id))
    if doc_id:                                   # фільтр по спеціалісту — лише його колонка
        cols = {d: [c for c in cs if c.doctor and c.doctor.id == doc_id] or
                   [cal_svc.Column(str(doc_id), db.get(Doctor, doc_id), [])] for d, cs in cols.items()}
    from .services import clinic as _cl
    _wh = _cl.hours_for(db, s) or (settings.work_hour_from, settings.work_hour_to)      # режим роботи з «Налаштування → Локації»
    grid = cal_svc.time_grid(s, cols[s], _wh[0], _wh[1]) if view == "day" else None
    n_ev = sum(1 for e in ev.get(s, []) if e.kind in ("visit", "attended")) if view == "day" else 0
    from .services import schedule as _sch
    ctx["shifts"] = _sch.for_day(db, s) if view == "day" else {}
    ctx["now_h"] = min(max(datetime.now().hour + 1, 8), 20)
    return render(request, "calendar.html", ev=ev, cols=cols, grid=grid, start=s, days=days, n_ev=n_ev,
                  prev=s - timedelta(days=days), next=s + timedelta(days=days), end=s + timedelta(days=days - 1),
                  month=None, **ctx)


@app.get("/tasks", response_class=HTMLResponse)
def tasks_page(request: Request, f: str = "open", user: User = Depends(current_user), db: Session = Depends(db_dep)):
    q = tasks_svc.visible_query(user)
    today = date.today()
    if f == "open":
        q = q.where(Task.status == "open")
    elif f == "mine":
        q = q.where(Task.status == "open", Task.assignee_id == user.id)
    elif f == "overdue":
        q = q.where(Task.status == "open", Task.due_date < today)
    elif f == "done":
        q = q.where(Task.status != "open")
    tasks = db.scalars(q.order_by(Task.status, Task.due_date.is_(None), Task.due_date, Task.id.desc()).limit(500)).all()
    users = db.scalars(select(User).where(User.active.is_(True)).order_by(User.name)).all()
    pats = [p for p in db.scalars(select(Patient).order_by(Patient.full_name)) if patient_visible(user, p)]
    from .services import scripts as _sc
    from .integrations import messaging as _msg
    scripts = {}
    if can(user, "message"):
        for t in [t for t in tasks if t.status == "open" and t.patient_id][:100]:
            scripts[t.id] = _sc.for_task(db, t, user.name)
    return render(request, "tasks.html", tasks=tasks, f=f, today=today, users=users, patients=pats,
                  KINDS=tasks_svc.KINDS, PRIORITY=tasks_svc.PRIORITY, scripts=scripts,
                  tg_on=_msg.telegram_available(), ts_on=_msg.turbosms_available(),
                  avail={t.id: _msg.available_channels(t.patient) for t in tasks if t.patient})


@app.post("/tasks")
def task_create(title: str = Form(...), description: str = Form(""), patient_id: str = Form(""),
                assignee: str = Form(""), due_date: str = Form(""), priority: str = Form("normal"),
                back: str = Form("/tasks"), user: User = Depends(current_user), db: Session = Depends(db_dep)):
    t = Task(title=title.strip(), description=description.strip() or None, kind="manual", created_by=user.id,
             priority=priority if priority in tasks_svc.PRIORITY else "normal",
             due_date=date.fromisoformat(due_date) if due_date else None)
    if patient_id.isdigit():
        _patient_or_403(db, user, int(patient_id))
        t.patient_id = int(patient_id)
    if assignee.isdigit():
        t.assignee_id = int(assignee)
    elif assignee in ("admin", "owner"):
        t.assignee_role = assignee
    else:
        t.assignee_id = user.id
    db.add(t)
    db.commit()
    audit(db, "task.create", "task", t.id, user=user, title=t.title)
    _notify_assignee(db, t)
    return RedirectResponse(back if back.startswith("/") else "/tasks", 303)


def _notify_assignee(db, t: Task):
    """Особисте сповіщення в Telegram, якщо виконавець вказав свій chat ID у «Мій кабінет»."""
    if not state["bot"] or not t.assignee_id:
        return
    u = db.get(User, t.assignee_id)
    if u and u.telegram_chat_id:
        from html import escape
        txt = f"📝 Нова задача: <b>{escape(t.title)}</b>" + (f"\nТермін: {t.due_date:%d.%m}" if t.due_date else "") + \
              f'\n<a href="{settings.public_url.rstrip("/")}/tasks">Задачі</a>'
        try:
            asyncio.get_event_loop().create_task(state["bot"].send_message(u.telegram_chat_id, txt))
        except Exception:      # noqa: BLE001
            pass


@app.post("/tasks/{tid}")
def task_update(tid: int, status: str = Form(...), result: str = Form(""), back: str = Form("/tasks"),
                user: User = Depends(current_user), db: Session = Depends(db_dep)):
    t = db.get(Task, tid)
    if not t or t not in db.scalars(tasks_svc.visible_query(user).where(Task.id == tid)).all():
        raise HTTPException(403, "Немає доступу до задачі")
    if status in ("open", "done", "cancelled"):
        t.status = status
        t.done_at = datetime.utcnow() if status != "open" else None
        t.result = result.strip() or t.result
        if status == "open" and not t.assignee_id and not t.assignee_role:
            t.assignee_id = user.id
    elif status == "take":
        t.assignee_id, t.assignee_role = user.id, None
    db.commit()
    audit(db, "task." + status, "task", t.id, user=user)
    return RedirectResponse(back if back.startswith("/") else "/tasks", 303)


@app.post("/tasks-generate")
def tasks_generate(user: User = Depends(require("reminders")), db: Session = Depends(db_dep)):
    rep = tasks_svc.generate_auto(db)
    audit(db, "task.generate", user=user, **rep)
    return RedirectResponse("/tasks", 303)


PRODUCT_FILTERS = {
    "low": "Закінчуються", "out": "Немає в наявності", "expiring": "Термін ≤ 90 дн", "expired": "Прострочені",
    "prescribed": "Призначені в планах", "selling": "Продаються (90 дн)", "nosales": "Без продажів 90 дн",
    "lowmargin": "Націнка < 40%", "nocost": "Без ціни закупівлі", "nobrand": "Без бренду",
}
PRODUCT_SORTS = {"brand": "Бренд, назва", "stock": "Залишок ↑", "expiry": "Термін ↑", "margin": "Націнка ↓",
                 "sold": "Продажі 90 дн ↓", "revenue": "Виручка з планів ↓", "value": "Вартість залишку ↓"}


def _product_flags(r) -> set[str]:
    f = set()
    if r.stock is not None and r.stock <= 0:
        f.add("out")
    if r.stock is not None and (r.deficit or (r.product.min_stock and 0 < r.stock <= r.product.min_stock)):
        f.add("low")
    d = r.days_to_expiry
    if d is not None and d < 0:
        f.add("expired")
    elif d is not None and d <= 90:
        f.add("expiring")
    if r.prescribed:
        f.add("prescribed")
    f.add("selling" if r.sold_90 else "nosales")
    if r.margin is not None and r.margin < 40:
        f.add("lowmargin")
    if not (r.cost_avg or r.cost_last):
        f.add("nocost")
    if not r.product.brand:
        f.add("nobrand")
    return f


@app.get("/products", response_class=HTMLResponse)
def products_page(request: Request, tab: str = "care", q: str = "", brand: str = "", supplier: str = "",
                  f: str = "", sort: str = "brand", wh: str = "vitrina", user: User = Depends(require("products")),
                  db: Session = Depends(db_dep)):
    sts = all_statuses(db)
    if tab == "procedures":
        return render(request, "products.html", tab=tab, procedures=catalog.procedure_rows(sts))
    from .models import ProductPurchase
    every = [r for r in catalog.product_rows(db, sts) if r.product.active is not False]
    wh = wh if wh in catalog.WAREHOUSES or wh == "all" else "vitrina"
    wh_counts = {k: 0 for k in catalog.WAREHOUSES}
    for r in every:
        wh_counts[r.product.warehouse or "vitrina"] = wh_counts.get(r.product.warehouse or "vitrina", 0) + 1
    rows = every if wh == "all" else [r for r in every if (r.product.warehouse or "vitrina") == wh]
    all_rows = rows
    counts = {k: 0 for k in PRODUCT_FILTERS}
    flags = {}
    for r in rows:
        flags[r.product.id] = _product_flags(r)
        for k in flags[r.product.id]:
            if k in counts:
                counts[k] += 1
    if q:
        words = set(norm(q).split())
        rows = [r for r in rows if words <= set(norm(" ".join(filter(None, [r.product.name, r.product.sku, r.product.brand]))).split())
                or (r.product.sku and q.strip().lower() == r.product.sku.lower())]
    if brand:
        rows = [r for r in rows if (r.product.brand or "") == ("" if brand == "—" else brand)]
    if supplier:
        rows = [r for r in rows if (r.product.supplier or r.last_supplier or "") == supplier]
    if f in PRODUCT_FILTERS:
        rows = [r for r in rows if f in flags[r.product.id]]
    keyf = {"stock": lambda r: (r.stock if r.stock is not None else 1e9),
            "expiry": lambda r: (r.expiry or date(2999, 1, 1)),
            "margin": lambda r: -(r.margin or -1e9), "sold": lambda r: -r.sold_90,
            "revenue": lambda r: -r.revenue, "value": lambda r: -r.stock_value}.get(sort)
    if keyf:
        rows = sorted(rows, key=keyf)
    groups = catalog.brand_groups(rows) if sort == "brand" else [("", rows)]
    brands = sorted({r.product.brand for r in all_rows if r.product.brand}, key=str.lower)
    suppliers = sorted({x for r in all_rows for x in (r.product.supplier, r.last_supplier) if x}, key=str.lower)
    tot = dict(n=len(rows), stock=sum(r.stock or 0 for r in rows), value=sum(r.stock_value for r in rows),
               sold=sum(r.sold_90_sum for r in rows))
    supplier_rows = None
    if tab == "suppliers":
        agg: dict[str, dict] = {}
        for pp in db.scalars(select(ProductPurchase).order_by(ProductPurchase.bought_on)):
            a = agg.setdefault(pp.supplier or "Без постачальника", {"n": 0, "sum": 0.0, "last": None, "products": set()})
            a["n"] += 1; a["sum"] += (pp.qty or 0) * (pp.unit_cost or 0); a["last"] = pp.bought_on; a["products"].add(pp.product_id)
        for r in every:
            if r.product.supplier:
                agg.setdefault(r.product.supplier, {"n": 0, "sum": 0.0, "last": None, "products": set()})["products"].add(r.product.id)
        supplier_rows = sorted(agg.items(), key=lambda x: x[0].lower())
    return render(request, "products.html", tab=tab, groups=groups, flags=flags, counts=counts, FILTERS=PRODUCT_FILTERS,
                  SORTS=PRODUCT_SORTS, q=q, brand=brand, supplier=supplier, f=f, sort=sort, brands=brands,
                  suppliers=suppliers, tot=tot, supplier_rows=supplier_rows, today=date.today(), wh=wh,
                  WAREHOUSES=catalog.WAREHOUSES, wh_counts=wh_counts, total_all=len(every),
                  imported=LAST_IMPORT.pop(f"products-{user.id}", None))


def _num(v: str):
    try:
        return float(v.replace(" ", "").replace(",", ".")) if v and v.strip() else None
    except ValueError:
        return None


def _day(v: str):
    try:
        return date.fromisoformat(v) if v else None
    except ValueError:
        return None


def _exp_day(v) -> date | None:
    """Термін придатності з тексту: 2027-12-31 · 31.12.2027 · 31.12.27 · 12.2027 · 12/27 · 31122027 (місяць-рік → останній день)."""
    import calendar as _cal
    s = re.sub(r"[/\-\s,]+", ".", str(v or "").strip()).strip(".")
    if not s:
        return None
    pats = [(r"(\d{4})\.(\d{1,2})\.(\d{1,2})", "ymd"), (r"(\d{1,2})\.(\d{1,2})\.(\d{2}|\d{4})", "dmy"), (r"(\d{1,2})\.(\d{2}|\d{4})", "my"),
            (r"(\d{2})(\d{2})(\d{4})", "dmy"), (r"(\d{2})(\d{2})(\d{2})", "dmy"), (r"(\d{2})(\d{2})", "my")]
    for rx, kind in pats:
        m_ = re.fullmatch(rx, s)
        if not m_:
            continue
        g = [int(x) for x in m_.groups()]
        y, mo, d_ = (g[0], g[1], g[2]) if kind == "ymd" else ((g[2], g[1], g[0]) if kind == "dmy" else (g[1], g[0], 0))
        y = y + 2000 if y < 100 else y
        if not (1 <= mo <= 12 and 2000 <= y <= 2100):
            return None
        last = _cal.monthrange(y, mo)[1]
        d_ = d_ or last
        return date(y, mo, d_) if 1 <= d_ <= last else None
    return None


@app.get("/products/{pid:int}", response_class=HTMLResponse)
def product_card(pid: int, request: Request, user: User = Depends(require("products")), db: Session = Depends(db_dep)):
    from .models import ProductPurchase, ProductSale
    prod = db.get(Product, pid)
    if not prod:
        raise HTTPException(404)
    row = next(r for r in catalog.product_rows(db) if r.product.id == pid)
    purchases = db.scalars(select(ProductPurchase).where(ProductPurchase.product_id == pid)
                           .order_by(ProductPurchase.bought_on.desc(), ProductPurchase.id.desc())).all()
    sales = []
    if prod.cb_product_id:
        sales = db.scalars(select(ProductSale).where(ProductSale.product_id == prod.cb_product_id)
                           .order_by(ProductSale.sold_at.desc()).limit(50)).all()
    items = db.scalars(select(PlanItem).where(PlanItem.product_id == pid)).all()
    plans = [(it, it.plan) for it in items if it.plan and it.plan.status != "cancelled"]
    # продажі по місяцях (6 міс) для міні-графіка
    months = []
    today = date.today()
    for k in range(5, -1, -1):
        m0 = (today.replace(day=1) - timedelta(days=31 * k)).replace(day=1)
        months.append([m0, 0])
    if prod.cb_product_id:
        for (q, at) in db.execute(select(ProductSale.quantity, ProductSale.sold_at).where(
                ProductSale.product_id == prod.cb_product_id, ProductSale.sold_at >= datetime.combine(months[0][0], datetime.min.time()))):
            for mm in months:
                if (at.year, at.month) == (mm[0].year, mm[0].month):
                    mm[1] += q or 0
    stock_rows = db.scalars(select(StockLevel).where(StockLevel.cb_product_id == prod.cb_product_id)).all() if prod.cb_product_id else []
    buys, buy_stats = catalog.purchase_history(prod, purchases)
    from .services import stock as _st
    from .models import StockMove as _SM
    ledger_on = _st.ledger_active(db)
    batches = _st.batches_view(db, pid) if ledger_on else []
    wh_tot: dict[str, float] = {}
    for b in batches:
        wh_tot[b["wh"]] = round(wh_tot.get(b["wh"], 0) + b["qty"], 4)
    out_moves = db.scalars(select(_SM).where(_SM.product_id == pid, _SM.qty < 0, _SM.kind.in_(("sale", "procedure", "writeoff")))
                           .order_by(_SM.at.desc()).limit(30)).all() if ledger_on else []
    batch_opts: dict[str, list] = {}
    for b in batches:
        batch_opts.setdefault(b["wh"], []).append(b["expiry"])
    for m in out_moves:                              # поточний термін руху теж має бути в списку
        if m.expiry_date not in batch_opts.setdefault(m.warehouse, []):
            batch_opts[m.warehouse].append(m.expiry_date)
    return render(request, "product.html", WAREHOUSES=catalog.WAREHOUSES, p=prod, r=row, purchases=purchases, buys=buys, buy_stats=buy_stats,
                  batches=batches, out_moves=out_moves, batch_opts=batch_opts, wh_tot=wh_tot, ledger_on=ledger_on, sales=sales, plans=plans, months=months,
                  m_max=max([m[1] for m in months] + [1]), stock_rows=stock_rows, today=today,
                  brands=sorted({x for (x,) in db.execute(select(Product.brand).where(Product.brand.is_not(None)).distinct())}),
                  suppliers=sorted({x for (x,) in db.execute(select(Product.supplier).where(Product.supplier.is_not(None)).distinct())}))


@app.post("/products/{pid:int}")
def product_update(pid: int, supply_days: str = Form(""), min_stock: str = Form(""), url: str = Form(""),
                   apply: str = Form(""), brand: str = Form(None), supplier: str = Form(None), sku: str = Form(None),
                   expiry_date: str = Form(None), price_manual: str = Form(None), back: str = Form(""),
                   warehouse: str = Form(None),
                   user: User = Depends(require("products")), db: Session = Depends(db_dep)):
    prod = db.get(Product, pid) or (_ for _ in ()).throw(HTTPException(404))
    prod.supply_days = int(supply_days) if supply_days.strip().isdigit() else None
    prod.min_stock = int(min_stock) if min_stock.strip().isdigit() else None
    prod.url = url.strip() or None
    if brand is not None:
        prod.brand = brand.strip()[:120] or None
    if supplier is not None:
        prod.supplier = supplier.strip()[:200] or None
    if sku is not None:
        prod.sku = sku.strip()[:100] or None
    if expiry_date is not None:
        prod.expiry_date = _day(expiry_date)
    if price_manual is not None:
        prod.price_manual = _num(price_manual)
    if warehouse is not None and warehouse in catalog.WAREHOUSES:
        prod.warehouse = warehouse
    if prod.url:                                         # посилання потрапляє в засоби бюті-планів без власного
        for it in db.scalars(select(PlanItem).where(PlanItem.product_id == prod.id, PlanItem.url.is_(None))):
            it.url = prod.url
    n = 0
    if apply and prod.supply_days:
        for it in db.scalars(select(PlanItem).where(PlanItem.product_id == prod.id)):
            if it.plan.status == "active":
                it.supply_days = prod.supply_days
                n += 1
    db.commit()
    audit(db, "product.update", "product", prod.id, user=user, supply_days=prod.supply_days,
          min_stock=prod.min_stock, applied_to_items=n)
    return RedirectResponse(back if back.startswith("/products") else f"/products/{pid}?saved=1", 303)


@app.post("/products/{pid:int}/batches")
async def product_batches_set(pid: int, request: Request, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Ручне внесення термінів: розподіл залишку складу на кілька партій (термін + к-сть)."""
    from .services import stock as st
    f = await request.form()
    wh = f.get("wh") if f.get("wh") in catalog.WAREHOUSES else None
    if not wh or not st.ledger_active(db):
        return RedirectResponse(f"/products/{pid}?err=Облік+складу+ще+не+розпочато#batches", 303)
    prod = db.get(Product, pid)
    k = prod.coef if prod and prod.coef and prod.coef > 0 else 1
    targets = []
    for e, q in zip(f.getlist("exp"), f.getlist("qty")):
        qq = _num(q)
        if qq and qq > 0:
            targets.append((_exp_day(e), round(qq * (k if f.get("in_packs") else 1), 4)))
    try:
        rep = st.set_batches(db, pid, wh, targets, note=f"терміни вручну · {user.name}")
    except st.StockError as e:
        return RedirectResponse(f"/products/{pid}?err={_q(str(e))}#batches", 303)
    audit(db, "stock.batches_manual", "product", pid, user=user, wh=wh, batches=len(targets), changed=rep["changed"])
    return RedirectResponse(f"/products/{pid}?saved=1#batches", 303)


@app.post("/stock/moves/{mid:int}/batch")
def stock_move_batch(mid: int, expiry: str = Form(""), back: str = Form(""), user: User = Depends(require("stock")),
                     db: Session = Depends(db_dep)):
    """Який термін продали / списали — для автоматичних продажів CleverBox і процедур."""
    from .services import stock as st
    m = st.reassign_batch(db, mid, _day(expiry))
    if m:
        audit(db, "stock.batch", user=user, move=mid, expiry=expiry or None)
    return RedirectResponse(back if back.startswith("/") else f"/products/{m.product_id if m else ''}#batches", 303)


@app.post("/products/{pid:int}/purchases")
def product_purchase_add(pid: int, bought_on: str = Form(""), qty: str = Form("1"), unit_cost: str = Form(""),
                         supplier: str = Form(""), expiry_date: str = Form(""), doc_no: str = Form(""),
                         user: User = Depends(require("products")), db: Session = Depends(db_dep)):
    from .models import ProductPurchase
    prod = db.get(Product, pid) or (_ for _ in ()).throw(HTTPException(404))
    q, c = _num(qty), _num(unit_cost)
    if not q or q <= 0 or c is None or c < 0:
        return RedirectResponse(f"/products/{pid}?err=Вкажіть+кількість+і+ціну+закупівлі#buy", 303)
    k = prod.coef if prod.coef and prod.coef > 0 else 1   # у формі — одиниці приходу, зберігаємо за одиницю обліку
    pp = ProductPurchase(product_id=pid, bought_on=_day(bought_on) or date.today(), qty=q * k, unit_cost=c / k,
                         supplier=supplier.strip()[:200] or prod.supplier, expiry_date=_day(expiry_date),
                         doc_no=doc_no.strip()[:100] or None)
    db.add(pp)
    if pp.supplier and not prod.supplier:
        prod.supplier = pp.supplier
    db.commit()
    audit(db, "product.purchase", "product", pid, user=user, qty=q, unit_cost=c)
    return RedirectResponse(f"/products/{pid}?saved=1#buy", 303)


@app.post("/products/{pid:int}/purchases/{ppid:int}/delete")
def product_purchase_delete(pid: int, ppid: int, del_password: str = Form(""), user: User = Depends(require("products")),
                            db: Session = Depends(db_dep)):
    from .models import ProductPurchase
    err = _delpw(db, user, del_password, "purchase", "product", pid)
    if err:
        return RedirectResponse(_with_err(f"/products/{pid}#buy", err), 303)
    pp = db.get(ProductPurchase, ppid)
    if pp and pp.product_id == pid:
        db.delete(pp)
        db.commit()
        audit(db, "product.purchase_delete", "product", pid, user=user, purchase=ppid)
    return RedirectResponse(f"/products/{pid}#buy", 303)


@app.post("/products/import")
async def products_import(file: UploadFile = File(...), user: User = Depends(require("import")), db: Session = Depends(db_dep)):
    """Excel-вивантаження товарів з CleverBox (Склад → Товари → Експорт): склади, бренди, постачальники, ціни, залишки."""
    import tempfile
    suffix = ".xls" if (file.filename or "").lower().endswith(".xls") else ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
    try:
        kind = await asyncio.to_thread(catalog.detect_export, tmp.name)
        if kind == "supply":                         # переплутали форму — поставки теж приймаємо
            from .services import stock as st_
            rep = await asyncio.to_thread(st_.import_supplier_invoices, db, tmp.name)
            if not rep.get("error"):
                rep = {"error": f"Це був файл «Поставки» — імпортовано в Склад → Постачальники ({rep['rows']} поставок)."}
        elif kind == "balances":                     # «Залишки по складах» — теж приймаємо
            b = await _import_balances_file(db, tmp.name)
            rep = {"error": f"Це був звіт «Залишки по складах» — {b['rows']} позицій завантажено, собівартість оновлено. "
                            "Далі: Склад → Залишки (старт обліку або звірка)."}
        elif kind != "products":
            rep = {"error": catalog.WRONG_FILE[kind]}
        else:
            rep = await asyncio.to_thread(catalog.import_products, db, tmp.name)
    except Exception as e:  # noqa: BLE001
        rep = {"error": f"Не вдалося прочитати файл: {type(e).__name__}"}
    finally:
        os.unlink(tmp.name)
    audit(db, "products.import", user=user, **{k: v for k, v in rep.items() if k != "by_warehouse"})
    LAST_IMPORT[f"products-{user.id}"] = rep
    return RedirectResponse("/products?wh=all", 303)


@app.post("/products/import-turnover")
async def products_import_turnover(file: UploadFile = File(...), d_from: str = Form(...), d_to: str = Form(...),
                                   user: User = Depends(require("import")), db: Session = Depends(db_dep)):
    """Звіт «Рух товарів» CleverBox за період → ціни закупівлі (середня за період, остання — з найсвіжішого періоду)."""
    import tempfile
    a, b = _day(d_from), _day(d_to)
    if not a or not b or a > b:
        LAST_IMPORT[f"products-{user.id}"] = {"error": "Вкажіть період звіту (з … по …)"}
        return RedirectResponse("/products?wh=all", 303)
    suffix = ".xls" if (file.filename or "").lower().endswith(".xls") else ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
    try:
        kind = await asyncio.to_thread(catalog.detect_export, tmp.name)
        if kind != "turnover":
            raise ValueError(catalog.WRONG_FILE[kind])
        rep = await asyncio.to_thread(catalog.import_turnover, db, tmp.name, a, b)
    except ValueError as e:
        rep = {"error": str(e)}
    except Exception as e:  # noqa: BLE001
        rep = {"error": f"Не вдалося прочитати файл: {type(e).__name__}"}
    finally:
        os.unlink(tmp.name)
    rep["kind"] = "turnover"
    rep["period"] = f"{a:%d.%m.%Y}–{b:%d.%m.%Y}"
    audit(db, "products.import_turnover", user=user, period=rep["period"], matched=rep.get("matched"))
    LAST_IMPORT[f"products-{user.id}"] = rep
    return RedirectResponse("/products?wh=all", 303)


@app.post("/products/bulk")
async def products_bulk(request: Request, user: User = Depends(require("products")), db: Session = Depends(db_dep)):
    form = await request.form()
    ids = [int(x) for x in form.getlist("ids") if str(x).isdigit()]
    field_, value = form.get("field"), (form.get("value") or "").strip()
    if field_ == "warehouse" and value not in catalog.WAREHOUSES:
        field_ = None
    if field_ in ("brand", "supplier", "warehouse") and ids:
        for p_ in db.scalars(select(Product).where(Product.id.in_(ids))):
            setattr(p_, field_, value[:200] or None)
        db.commit()
        audit(db, "product.bulk", user=user, field=field_, n=len(ids))
    back = str(form.get("back") or "/products")
    return RedirectResponse(back if back.startswith("/products") else "/products", 303)


STOCK_TABS = {"balances": "Залишки", "critical": "Критичні залишки", "supply": "Поставки", "order": "Замовлення",
              "transfer": "Переміщення", "writeoff": "Списання", "return": "Повернення", "inventory": "Інвентаризація",
              "suppliers": "Постачальники", "norms": "Норми списання", "moves": "Рух товарів"}
templates.env.globals["STOCK_TABS"] = STOCK_TABS
templates.env.globals["REPORT_TABS"] = __import__("app.services.reports", fromlist=["REPORTS"]).REPORTS


def _critical_rows(rows):
    from .services import stock as st
    return [r for r in rows if st.is_critical(r) or (r.deficit or 0) > 0]


@app.get("/stock", response_class=HTMLResponse)
def stock_page(request: Request, tab: str = "balances", only: str = "", wh: str = "", q: str = "", sup: str = "",
               user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .services import stock as st
    from .models import Supplier, ProcedureNorm, StockMove
    if only == "deficit":
        tab = "critical"
    tab = tab if tab in STOCK_TABS else "balances"
    if tab in ("supply", "order"):
        st.ensure_suppliers(db)
        _sup_names = sorted({x for (x,) in db.execute(select(Supplier.name))} |
                            {x for (x,) in db.execute(select(Product.supplier).where(Product.supplier.is_not(None)).distinct())},
                            key=str.lower)
    else:
        _sup_names = []
    from .integrations.checkbox import accounts as _cb_acc
    ctx = dict(tab=tab, TABS=STOCK_TABS, WAREHOUSES=catalog.WAREHOUSES, STATUS=st.STATUS, KINDS=st.DOC_KINDS, sup_names=_sup_names,
               fops=[a["name"] for a in _cb_acc(db)] if tab in ("supply", "order", "suppliers") else [],
               MOVE_KINDS=st.MOVE_KINDS, active=st.ledger_active(db), start=st.ledger_start(db), wh=wh, q=q, sup=sup)
    if tab in ("balances", "critical"):
        rows = [r for r in catalog.product_rows(db) if r.product.active is not False]
        per = st.balances(db) if ctx["active"] else {}
        if wh in catalog.WAREHOUSES:
            if ctx["active"]:
                rows = [r for r in rows if per.get((r.product.id, wh)) or (r.product.warehouse or "vitrina") == wh]
            else:
                rows = [r for r in rows if (r.product.warehouse or "vitrina") == wh]
        if q:
            words = set(norm(q).split())
            rows = [r for r in rows if words <= set(norm(" ".join(filter(None, [r.product.name, r.product.sku, r.product.brand]))).split())]
        if tab == "critical":
            rows = _critical_rows(rows)
            if sup:
                rows = [r for r in rows if (r.product.supplier or r.last_supplier or "") == sup]
            rows.sort(key=lambda r: ((r.product.supplier or r.last_supplier or "яяя").lower(), r.product.name))
        else:
            order = {k: i for i, k in enumerate(catalog.WAREHOUSES)}
            if not request.query_params.get("all"):                 # за замовчуванням — лише те, що є в наявності
                rows = [r for r in rows if (r.stock or 0) != 0]
            rows.sort(key=lambda r: (order.get(r.product.warehouse or "vitrina", 9), (r.product.brand or "яяя").lower(), r.product.name))
        wh_tot = defaultdict(lambda: {"n": 0, "value": 0.0})
        costs = {r.product.id: (r.cost_avg or r.cost_last or 0) for r in rows}
        if ctx["active"]:
            for (pid, w), qty in per.items():
                if qty and pid in costs:
                    wh_tot[w]["n"] += 1
                    wh_tot[w]["value"] += qty * costs[pid]
        ctx.update(rows=rows, per=per, wh_tot=dict(wh_tot), suggested=st.suggested_qty,
                   suppliers=sorted({r.product.supplier or r.last_supplier for r in rows if (r.product.supplier or r.last_supplier)}))
        lv = db.scalar(select(StockLevel).limit(1))
        ctx["cb_updated"] = lv.updated_at if lv else None
        ctx["snap"] = _balances_snapshot(db)
    elif tab in st.DOC_KINDS:
        arch = bool(request.query_params.get("archive"))
        ctx["docs"] = st.doc_rows(db, tab, archived=arch)
        ctx["archive"] = arch
        from .services.payments import MODE_LABELS
        ctx["MODES"] = MODE_LABELS
        ctx["archive_n"] = db.scalar(select(func.count(StockDoc.id)).where(StockDoc.kind == tab, StockDoc.archived_at.is_not(None))) or 0
        ctx["user_names"] = {u_.id: u_.name for u_ in db.scalars(select(User))} if arch else {}
    elif tab == "suppliers":
        from .models import SupplierInvoice
        st.ensure_suppliers(db)
        per = request.query_params.get("per", "365")
        d_from = None if per == "all" else datetime.now() - timedelta(days=int(per) if per.isdigit() else 365)
        ctx.update(sstats=st.supplier_stats(db, d_from), per=per, sel=request.query_params.get("s", ""),
                   debt_only=bool(request.query_params.get("debt")), imported=LAST_IMPORT.pop(f"inv-{user.id}", None))
        if ctx["sel"]:
            ctx["invoices"] = db.scalars(select(SupplierInvoice).where(SupplierInvoice.supplier == ctx["sel"])
                                         .order_by(SupplierInvoice.inv_date.desc()).limit(300)).all()
        elif ctx["debt_only"]:
            ctx["invoices"] = [x for x in db.scalars(select(SupplierInvoice).where(SupplierInvoice.status != "Чернетка")
                                                     .order_by(SupplierInvoice.inv_date.desc())) if x.debt > 1]
        agg = defaultdict(lambda: {"products": 0, "orders": 0, "supplies": 0, "sum": 0.0, "last": None})
        for p_ in db.scalars(select(Product).where(Product.supplier.is_not(None))):
            agg[p_.supplier.lower()]["products"] += 1
        for d in db.scalars(select(StockDoc).where(StockDoc.kind.in_(("supply", "order")), StockDoc.status != "cancelled")):
            a = agg[(d.supplier or "").lower()]
            a["orders" if d.kind == "order" else "supplies"] += 1
            if d.kind == "supply" and d.status == "posted":
                a["sum"] += sum((ln.qty or 0) * (ln.unit_cost or 0) for ln in d.lines)
                a["last"] = max(filter(None, [a["last"], d.doc_date]))
        ctx.update(suppliers=db.scalars(select(Supplier).order_by(Supplier.name)).all(), agg=agg)
        # оплати: monobank + накладні CRM
        from .models import BankTxn
        from .services import payments as pay
        from .integrations import monobank
        from .integrations.checkbox import accounts as _cb_acc
        sel_sup = next((x for x in ctx["suppliers"] if x.name == ctx["sel"]), None)
        ctx.update(sel_sup=sel_sup, mono=monobank.public(db), fops=[a["name"] for a in _cb_acc(db)],
                   txns=db.scalars(select(BankTxn).where(BankTxn.at >= datetime.now() - timedelta(days=60))
                                   .order_by(BankTxn.at.desc()).limit(200)).all(),
                   open_groups=pay.groups(db, only_open=True), PAY=pay.PAY_LABELS,
                   weak_ids={w.id for w in pay.weak_auto(db)},
                   crm_groups=pay.groups(db, ctx["sel"]) if ctx["sel"] else [],
                   sup_by_id={x.id: x for x in ctx["suppliers"]},
                   job=jobs.JOBS.get(request.query_params.get("job") or "") if request.query_params.get("job") else None)
    elif tab == "norms":
        from .models import CbService
        services = db.scalars(select(CbService).order_by(CbService.category, CbService.name)).all()
        norms = defaultdict(list)
        for n in db.scalars(select(ProcedureNorm)):
            norms[n.service_id].append(n)
        if q:
            words = set(norm(q).split())
            services = [x for x in services if words <= set(norm(" ".join(filter(None, (
                x.name, x.category, x.techcard, x.direction)))).split())]
        elif request.query_params.get("only") != "all":
            services = [x for x in services if x.id in norms] or services[:60]
        # групи: одна техкартка = одна норма на всі її послуги; послуги без техкартки — окремо
        groups: dict[str, dict] = {}
        for x in services:
            key = f"t:{x.techcard}" if x.techcard else f"s:{x.id}"
            g_ = groups.setdefault(key, {"key": key, "title": x.techcard or x.name, "techcard": x.techcard,
                                         "services": [], "norms": {}, "direction": x.direction})
            g_["services"].append(x)
        all_by_card = defaultdict(list)
        if any(g_["techcard"] for g_ in groups.values()):
            for x in db.scalars(select(CbService).where(CbService.techcard.is_not(None))):
                all_by_card[x.techcard].append(x)
        for g_ in groups.values():
            if g_["techcard"]:
                g_["services"] = all_by_card.get(g_["techcard"], g_["services"])
            n_svc = len(g_["services"])
            for x in g_["services"]:
                for n_ in norms.get(x.id, []):
                    k_ = (n_.product_id, n_.qty, n_.warehouse)
                    item = g_["norms"].setdefault(k_, {"n": n_, "count": 0})
                    item["count"] += 1
            g_["norms"] = [dict(v, total=n_svc) for v in g_["norms"].values()]
        glist = sorted(groups.values(), key=lambda g_: ((g_["direction"] or "яяя"), g_["title"].lower()))
        ctx.update(services=services, groups=glist, norms=norms, products=_product_options(db, consumables=True))
    elif tab == "moves":
        mq = select(StockMove).order_by(StockMove.at.desc(), StockMove.id.desc()).limit(300)
        if wh in catalog.WAREHOUSES:
            mq = mq.where(StockMove.warehouse == wh)
        ctx["moves"] = db.scalars(mq).all()
        ctx["pnames"] = {p_.id: p_ for p_ in db.scalars(select(Product))}
    return render(request, "stock.html", **ctx)


def _product_options(db: Session, consumables: bool = False):
    q = select(Product).where(Product.active.is_not(False)).order_by(Product.name)
    return [(p_.id, f"{p_.name}{(' · арт. ' + p_.sku) if p_.sku else ''} [#{p_.id}]", p_.warehouse) for p_ in db.scalars(q)]


def _pick_product(db: Session, ref: str) -> Product | None:
    import re as _re
    m = _re.search(r"\[#(\d+)\]\s*$", ref or "")
    if m:
        return db.get(Product, int(m.group(1)))
    ref = (ref or "").strip()
    if not ref:
        return None
    return db.scalar(select(Product).where((Product.sku == ref) | (Product.barcode == ref))) or catalog.match_product(db, ref)


@app.post("/stock/start")
def stock_start(from_file: str = Form(""), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .services import stock as st
    snap = _balances_snapshot(db) if from_file else None
    if from_file and not snap:
        return RedirectResponse("/stock?msg=Спершу+завантажте+звіт+«Залишки+по+складах»", 303)
    try:
        doc = st.start_ledger(db, user.id, snapshot=snap)
    except st.StockError as e:
        return RedirectResponse(f"/stock?msg={e}", 303)
    audit(db, "stock.start", "stock_doc", doc.id, user=user, lines=len(doc.lines), from_file=bool(snap))
    if snap:
        _balances_snapshot(db, clear=True)
    return RedirectResponse("/stock?msg=Облік+у+CRM+розпочато", 303)


def _balances_snapshot(db: Session, rep: dict | None = None, clear: bool = False) -> dict | None:
    """Знімок останнього імпортованого звіту «Залишки по складах» (SyncState wh_snapshot)."""
    import json
    stt = db.get(SyncState, "wh_snapshot")
    if clear:
        if stt:
            db.delete(stt)
            db.commit()
        return None
    if rep is not None:
        stt = stt or SyncState(key="wh_snapshot")
        stt.value, stt.updated_at = json.dumps(rep, ensure_ascii=False), datetime.utcnow()
        db.merge(stt)
        db.commit()
        return rep
    try:
        return json.loads(stt.value) if stt and stt.value else None
    except ValueError:
        return None


async def _import_balances_file(db: Session, path: str) -> dict:
    rep = await asyncio.to_thread(catalog.parse_balances, db, path)
    return _balances_snapshot(db, rep)


@app.post("/stock/balances/import")
async def stock_balances_import(file: UploadFile = File(...), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Звіт CleverBox «Залишки по складах» → собівартість залишку + знімок для старту обліку / звірки."""
    import tempfile
    suffix = ".xls" if (file.filename or "").lower().endswith(".xls") else ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
    msg = ""
    try:
        kind = await asyncio.to_thread(catalog.detect_export, tmp.name)
        if kind != "balances":
            raise ValueError(catalog.WRONG_FILE[kind])
        rep = await _import_balances_file(db, tmp.name)
        audit(db, "stock.balances_import", user=user, rows=rep["rows"], created=rep["created"], value=rep["value"])
    except ValueError as e:
        msg = str(e)
    except Exception as e:  # noqa: BLE001
        msg = f"Не вдалося прочитати файл: {type(e).__name__}"
    finally:
        os.unlink(tmp.name)
    return RedirectResponse("/stock?tab=balances" + (f"&msg={msg}" if msg else "#snap"), 303)


@app.post("/stock/balances/reconcile")
def stock_balances_reconcile(user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .services import stock as st
    snap = _balances_snapshot(db)
    if not snap:
        return RedirectResponse("/stock?msg=Спершу+завантажте+звіт+«Залишки+по+складах»", 303)
    try:
        docs = st.reconcile_from_snapshot(db, snap, user.id)
    except st.StockError as e:
        return RedirectResponse(f"/stock?msg={e}", 303)
    audit(db, "stock.balances_reconcile", user=user, docs=[d.id for d in docs])
    _balances_snapshot(db, clear=True)
    return RedirectResponse(f"/stock?tab=inventory&msg=Створено+{len(docs)}+чернеток+інвентаризації+—+перевірте+різницю+і+проведіть", 303)


@app.post("/stock/balances/dismiss")
def stock_balances_dismiss(user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    _balances_snapshot(db, clear=True)
    return RedirectResponse("/stock", 303)


@app.post("/stock/docs")
def stock_doc_create(kind: str = Form(...), wh: str = Form(""), supplier: str = Form(""),
                     user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .services import stock as st
    kw = {}
    if kind in ("transfer", "writeoff", "inventory"):
        kw["wh_from"] = wh if wh in catalog.WAREHOUSES else "vitrina"
    if kind in ("supply", "order"):
        kw["wh_to"] = wh if wh in catalog.WAREHOUSES else "vitrina"
        kw["supplier"] = supplier.strip() or None
    if kind == "return":
        kw.update(reason="supplier", wh_from="vitrina")
    if kind == "writeoff":
        kw["reason"] = st.WRITEOFF_REASONS[0]
    doc = st.new_doc(db, kind, user_id=user.id, **kw)
    if kind == "inventory":
        st.inventory_prefill(db, doc)
    db.commit()
    audit(db, "stock.create", "stock_doc", doc.id, user=user, kind=kind)
    return RedirectResponse(f"/stock/docs/{doc.id}", 303)


@app.get("/stock/lookup/products")
def stock_lookup_products(q: str = "", wh: str = "", user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Пошук товару для вікна «Нова поставка»: назва / артикул / штрихкод. Ціна й к-сть — в одиницях приходу."""
    q = (q or "").strip()
    if len(q) < 2:
        return JSONResponse([])
    words = set(norm(q).split())
    ql = q.lower()
    rows = []
    for r in catalog.product_rows(db, []):
        p_ = r.product
        if p_.active is False:
            continue
        exact = ql in ((p_.sku or "").lower(), (p_.barcode or "").lower())
        if not exact and not words <= set(norm(" ".join(filter(None, [p_.name, p_.sku, p_.brand, p_.barcode]))).split()):
            continue
        k = p_.coef if p_.coef and p_.coef > 0 else 1
        rows.append((0 if exact else (1 if (p_.warehouse or "vitrina") == wh else 2), {
            "id": p_.id, "name": p_.name, "sku": p_.sku or "", "barcode": p_.barcode or "", "brand": p_.brand or "",
            "unit": "шт" if k == 1 else f"уп. ({k:g} {p_.unit or 'од.'})", "coef": k,
            "price": r.cost_last or r.cost_avg or p_.cost or 0, "stock": r.stock,
            "wh": catalog.WAREHOUSES.get(p_.warehouse or "vitrina"), "supplier": p_.supplier or r.last_supplier or ""}))
    rows.sort(key=lambda x: (x[0], x[1]["name"].lower()))
    return JSONResponse([x[1] for x in rows[:25]])


@app.post("/stock/recognize")
async def stock_recognize(file: UploadFile = File(...), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Файл накладної → постачальник, № , дата, товари (зіставлені з каталогом), склад — для вікна «Нова поставка»."""
    from .services import invoice_ai
    data = await file.read()
    if not data or len(data) > 20 * 1024 * 1024:
        return JSONResponse({"error": "Файл порожній або більший за 20 МБ"}, status_code=400)
    try:
        parsed, src = await asyncio.to_thread(invoice_ai.extract, db, data, file.filename or "file", file.content_type)
        res = invoice_ai.match(db, parsed)
    except ValueError as e:
        return JSONResponse({"error": str(e)}, status_code=400)
    except Exception as e:  # noqa: BLE001
        log.warning("recognize: %s", e)
        return JSONResponse({"error": f"Не вдалося розпізнати: {str(e)[:160]}"}, status_code=400)
    res["source"] = src
    audit(db, "stock.recognize", user=user, file=file.filename, source=src, lines=len(res["lines"]), matched=res["matched"])
    return JSONResponse(res)


@app.post("/stock/docs/quick")
async def stock_doc_quick(request: Request, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Вікно «Нова поставка / замовлення»: шапка + рядки одним JSON. Рядки на різні склади → окремі документи
    з однією групою оплати (одна накладна). paid_by + paid=true — відмітити оплаченим цим ФОП. post=true — провести."""
    import uuid
    from .services import stock as st, payments as pay
    from .models import StockDocLine, Supplier
    data = await request.json()
    kind = data.get("kind") if data.get("kind") in ("supply", "order") else "supply"
    wh0 = data.get("wh") if data.get("wh") in catalog.WAREHOUSES else "vitrina"
    lines = [x for x in (data.get("lines") or []) if x.get("product_id") and (_num(str(x.get("qty"))) or 0) > 0]
    if not lines:
        return JSONResponse({"error": "Додайте хоча б один товар з кількістю"}, status_code=400)
    by_wh: dict[str, list] = {}
    for x in lines:
        by_wh.setdefault(x.get("wh") if x.get("wh") in catalog.WAREHOUSES else wh0, []).append(x)
    d = _day(str(data.get("date") or "")) or date.today()
    supplier = (data.get("supplier") or "").strip()[:200] or None
    paid_by = (data.get("paid_by") or "").strip()[:100] or None
    group = uuid.uuid4().hex[:12] if len(by_wh) > 1 else None
    docs = []
    for wh, ls in by_wh.items():
        doc = st.new_doc(db, kind, user_id=user.id, wh_to=wh, supplier=supplier, doc_date=d,
                         ext_no=(data.get("ext_no") or "").strip()[:100] or None,
                         note=(data.get("note") or "").strip()[:2000] or None)
        doc.pay_group, doc.paid_by = group, paid_by
        for x in ls:
            p_ = db.get(Product, int(x["product_id"]))
            if not p_:
                continue
            k = p_.coef if p_.coef and p_.coef > 0 else 1        # у вікні — одиниці приходу; в обліку — одиниці обліку
            qty, price = _num(str(x.get("qty"))) or 0, _num(str(x.get("price") or ""))
            doc.lines.append(StockDocLine(product_id=p_.id, qty=round(qty * k, 4),
                                          unit_cost=round(price / k, 4) if price is not None else None,
                                          expiry_date=_exp_day(x.get("expiry"))))
        db.flush()
        docs.append(doc)
    if supplier and not any(n.lower() == supplier.lower() for (n,) in db.execute(select(Supplier.name))):
        db.add(Supplier(name=supplier))
    db.commit()
    req = data.get("sup_req") if isinstance(data.get("sup_req"), dict) else {}
    if supplier and (req.get("edrpou") or req.get("iban")):          # реквізити з розпізнаної накладної → юр. особа
        from .services import suppliers as sp
        sx = next((x for x in db.scalars(select(Supplier)) if x.name.lower() == supplier.lower()), None)
        hit = sp.find(db, req.get("iban"), req.get("edrpou"))
        if sx and not (hit and hit[0].id == sx.id and hit[1]):
            try:
                sp.save_entity(db, sx, {"legal_name": str(req.get("name") or "")[:300], "edrpou": str(req.get("edrpou") or ""),
                                        "iban": str(req.get("iban") or "")})
            except ValueError:
                db.rollback()
    if kind == "supply" and data.get("paid") and paid_by:          # оплачено вже зараз (готівка / інший банк)
        g = pay.find_group(db, pay.group_key(docs[0]))
        if g:
            amt_ = g.left
            pay.allocate(db, g, amt_, fop=paid_by, user_id=user.id)
            db.commit()
            audit(db, "stock.pay", "stock_doc", docs[0].id, user=user, op="paid", fop=paid_by, amount=round(amt_, 2))
    for x_ in docs:
        audit(db, "stock.quick_doc", "stock_doc", x_.id, user=user, kind=kind, docs=[x.id for x in docs],
              lines=len(x_.lines))
    warn = None
    if data.get("post") and kind == "supply":
        if not st.ledger_active(db):
            warn = "Збережено чернеткою: облік складу ще не розпочато"
        else:
            for doc in docs:
                try:
                    st.post(db, doc)
                    audit(db, "stock.post", "stock_doc", doc.id, user=user, kind=kind)
                except st.StockError as e:
                    warn = f"Збережено чернеткою: {e}"
    if len(docs) > 1:
        warn = (warn + ". " if warn else "") + f"Накладну розділено на {len(docs)} склади: " + \
            ", ".join(f"{x.number} ({catalog.WAREHOUSES[x.wh_to]})" for x in docs)
    return JSONResponse({"id": docs[0].id, "ids": [x.id for x in docs], "url": f"/stock/docs/{docs[0].id}",
                         "number": docs[0].number, "warning": warn})


@app.post("/stock/docs/{did:int}/files")
async def stock_doc_files(did: int, files: list[UploadFile] = File(...), user: User = Depends(require("stock")),
                          db: Session = Depends(db_dep)):
    """Фото / скани накладної. Документи однієї накладної (кілька складів) отримують ті самі файли."""
    from .services import ingest
    from .models import StockDocFile
    doc = db.get(StockDoc, did) or (_ for _ in ()).throw(HTTPException(404))
    targets = db.scalars(select(StockDoc).where(StockDoc.pay_group == doc.pay_group)).all() if doc.pay_group else [doc]
    n = 0
    for f in files:
        data = await f.read()
        if not data or len(data) > 25 * 1024 * 1024:
            continue
        mime = f.content_type or ""
        if not (mime.startswith("image/") or mime == "application/pdf" or (f.filename or "").lower().endswith((".pdf", ".jpg", ".jpeg", ".png", ".heic", ".webp"))):
            continue
        path = ingest.store(data, f.filename or "file")
        for t in targets:
            db.add(StockDocFile(doc_id=t.id, path=path, name=(f.filename or "файл")[:200], mime=mime[:100] or None))
        n += 1
    db.commit()
    audit(db, "stock.files", "stock_doc", did, user=user, files=n)
    return JSONResponse({"ok": True, "files": n})


@app.get("/stock/files/{fid:int}")
def stock_file(fid: int, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .models import StockDocFile
    f = db.get(StockDocFile, fid) or (_ for _ in ()).throw(HTTPException(404))
    return FileResponse(f.path, media_type=f.mime or None, filename=f.name,
                        content_disposition_type="inline" if (f.mime or "").startswith(("image/", "application/pdf")) else "attachment")


@app.post("/stock/files/{fid:int}/delete")
def stock_file_delete(fid: int, del_password: str = Form(""), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .models import StockDocFile
    f = db.get(StockDocFile, fid)
    did = f.doc_id if f else 0
    err = _delpw(db, user, del_password, "stock_file", "stock_doc", did)
    if err:
        return RedirectResponse(_with_err(f"/stock/docs/{did}#files", err), 303)
    if f:
        name = f.name
        db.delete(f)
        db.commit()
        audit(db, "stock.file_delete", "stock_doc", did, user=user, name=name)
    return RedirectResponse(f"/stock/docs/{did}#files", 303)


@app.post("/stock/docs/{did:int}/pay")
def stock_doc_pay(did: int, action: str = Form("paid"), fop: str = Form(""), amount: str = Form(""),
                  user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Оплата накладної вручну (усі склади групи): action=paid — оплатити залишок / суму; clear — зняти ручні відмітки."""
    from .services import payments as pay
    doc = db.get(StockDoc, did) or (_ for _ in ()).throw(HTTPException(404))
    g = pay.find_group(db, pay.group_key(doc))
    amt = None
    if g:
        if action == "clear":
            pay.clear_manual(db, g)
        else:
            amt = _num(amount) or g.left
            pay.allocate(db, g, amt, fop=fop.strip() or None, user_id=user.id)
            for d_ in g.docs:
                d_.paid_by = fop.strip() or d_.paid_by
            db.commit()
        for d_ in g.docs:
            audit(db, "stock.pay", "stock_doc", d_.id, user=user, op=action, fop=fop or None,
                  amount=round(amt, 2) if amt and d_.id == did else None)
    return RedirectResponse(f"/stock/docs/{did}#pay", 303)


@app.post("/stock/order-from-critical")
def stock_order_from_critical(supplier: str = Form(""), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .services import stock as st
    rows = _critical_rows(catalog.product_rows(db))
    doc = st.order_from_critical(db, rows, supplier.strip() or None, user.id)
    if not doc:
        return RedirectResponse("/stock?tab=critical&msg=Немає+що+замовляти", 303)
    return RedirectResponse(f"/stock/docs/{doc.id}", 303)


@app.get("/stock/docs/{did:int}", response_class=HTMLResponse)
def stock_doc_page(did: int, request: Request, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .services import stock as st
    from .models import Supplier
    doc = db.get(StockDoc, did) or (_ for _ in ()).throw(HTTPException(404))
    bal = st.balances(db) if st.ledger_active(db) else {}
    lot_opts: dict[int, list] = {}                     # партії складу списання: товар → [(термін, к-сть)]
    src_wh = doc.wh_from or doc.wh_to
    if doc.kind in ("transfer", "writeoff", "return") and src_wh and bal:
        for (pid_, w_, e_), q_ in st.batch_balances(db, warehouse=src_wh).items():
            if e_ and q_ > 1e-6:
                lot_opts.setdefault(pid_, []).append((e_, round(q_, 3)))
        for v in lot_opts.values():
            v.sort()
    from .services import payments as pay
    from .models import StockDocFile, PaymentAlloc, BankTxn
    from .integrations.checkbox import accounts as _cb_acc
    grp = pay.find_group(db, pay.group_key(doc)) if doc.kind == "supply" else None
    allocs = db.scalars(select(PaymentAlloc).where(PaymentAlloc.doc_id == doc.id).order_by(PaymentAlloc.at)).all()
    txns = {t.id: t for t in db.scalars(select(BankTxn).where(BankTxn.id.in_([a.txn_id for a in allocs if a.txn_id] or [""])))}
    sup = next((x for x in db.scalars(select(Supplier)) if x.name.lower() == (doc.supplier or "").lower()), None) if doc.supplier else None
    return render(request, "stock_doc.html", d=doc, KINDS=st.DOC_KINDS, STATUS=st.STATUS, WAREHOUSES=catalog.WAREHOUSES,
                  files=db.scalars(select(StockDocFile).where(StockDocFile.doc_id == doc.id).order_by(StockDocFile.id)).all(),
                  grp=grp, allocs=allocs, txns=txns, PAY=pay.PAY_LABELS, MODES=pay.MODE_LABELS,
                  alloc_users={u_.id: u_.name for u_ in db.scalars(select(User))}, line_paid=pay.by_product(db, [doc.id]),
                  fops=[a["name"] for a in _cb_acc(db)], sup=sup, lot_opts=lot_opts, today=date.today(),
                  group_docs=[x for x in (grp.docs if grp else []) if x.id != doc.id],
                  REASONS=st.WRITEOFF_REASONS, products=_product_options(db), bal=bal, active=st.ledger_active(db),
                  suppliers=[x.name for x in db.scalars(select(Supplier).order_by(Supplier.name))],
                  history=__import__("app.services.doc_history", fromlist=["history"]).history(db, doc),
                  editable=doc.status in ("draft", "sent"))


@app.post("/stock/docs/{did:int}")
async def stock_doc_save(did: int, request: Request, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Шапка + усі рядки однією формою; action: save | post | add."""
    from .services import stock as st
    from .models import StockDocLine
    doc = db.get(StockDoc, did) or (_ for _ in ()).throw(HTTPException(404))
    f = await request.form()
    action = f.get("action", "save")
    from .services import doc_history as dh
    before = dh.snapshot(doc)
    if doc.status in ("draft", "sent"):
        doc.doc_date = _day(f.get("doc_date", "")) or doc.doc_date
        for k in ("supplier", "ext_no", "note", "reason"):
            if k in f:
                setattr(doc, k, (f.get(k) or "").strip()[:300] or None)
        for k in ("wh_from", "wh_to"):
            if k in f:
                setattr(doc, k, f.get(k) if f.get(k) in catalog.WAREHOUSES else None)
        for ln in list(doc.lines):
            if f.get(f"del_{ln.id}"):
                db.delete(ln)
                continue
            if f"qty_{ln.id}" in f:
                ln.qty = _num(f.get(f"qty_{ln.id}")) or 0
            if f"cost_{ln.id}" in f:
                ln.unit_cost = _num(f.get(f"cost_{ln.id}"))
            if f"exp_{ln.id}" in f:
                ln.expiry_date = _exp_day(f.get(f"exp_{ln.id}"))
        new_ref = (f.get("new_product") or "").strip()
        if new_ref:
            pr = _pick_product(db, new_ref)
            if not pr:
                db.commit()
                return RedirectResponse(f"/stock/docs/{did}?err=Товар+не+знайдено:+{new_ref[:60]}#add", 303)
            new_exp = _exp_day(f.get("new_exp"))
            line = next((ln for ln in doc.lines if ln.product_id == pr.id and ln.expiry_date == new_exp), None)   # інший термін — нова партія
            if line and doc.kind != "inventory":
                line.qty = (line.qty or 0) + (_num(f.get("new_qty")) or 1)
            elif not line:
                cost = _num(f.get("new_cost"))
                if cost is None and doc.kind in ("supply", "order"):
                    rr = next((r for r in catalog.product_rows(db, []) if r.product.id == pr.id), None)
                    cost = (rr.cost_last or rr.cost_avg) if rr else pr.cost
                doc.lines.append(StockDocLine(product_id=pr.id, qty=_num(f.get("new_qty")) or 1, unit_cost=cost,
                                              expiry_date=_exp_day(f.get("new_exp"))))
        db.commit()
        db.refresh(doc)
        changes = dh.diff(before, dh.snapshot(doc))
        if changes:
            audit(db, "stock.edit", "stock_doc", doc.id, user=user, changes=changes)
    try:
        if action == "post":
            st.post(db, doc)
            audit(db, "stock.post", "stock_doc", doc.id, user=user, kind=doc.kind)
            return RedirectResponse(f"/stock/docs/{did}?msg=Проведено", 303)
        if action == "send" and doc.kind == "order":
            doc.status = "sent"
            db.commit()
            audit(db, "stock.send", "stock_doc", doc.id, user=user)
            return RedirectResponse(f"/stock/docs/{did}?msg=Позначено+як+відправлене", 303)
        if action == "receive":
            sup = st.receive_order(db, doc, user.id)
            audit(db, "stock.receive", "stock_doc", sup.id, user=user, order=doc.number)
            return RedirectResponse(f"/stock/docs/{sup.id}?msg=Перевірте+кількості+й+ціни+та+проведіть+поставку", 303)
        if action == "cancel":
            st.cancel(db, doc)
            audit(db, "stock.cancel", "stock_doc", doc.id, user=user, kind=doc.kind)
            return RedirectResponse(f"/stock/docs/{did}?msg=Скасовано", 303)
        if action == "prefill" and doc.kind == "inventory":
            n = st.inventory_prefill(db, doc)
            return RedirectResponse(f"/stock/docs/{did}?msg=Додано+{n}+товарів", 303)
    except st.StockError as e:
        return RedirectResponse(f"/stock/docs/{did}?err={e}", 303)
    return RedirectResponse(f"/stock/docs/{did}" + ("#add" if f.get("new_product") else ""), 303)


@app.post("/stock/suppliers")
def stock_supplier_save(sid: str = Form(""), name: str = Form(...), contact: str = Form(""), phone: str = Form(""),
                        email: str = Form(""), terms: str = Form(""), note: str = Form(""),
                        legal_name: str = Form(""), edrpou: str = Form(""), iban: str = Form(""), bank: str = Form(""),
                        pay_purpose: str = Form(""), address: str = Form(""), contract: str = Form(""), vat: str = Form(""),
                        back: str = Form(""), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .models import Supplier
    x = db.get(Supplier, int(sid)) if sid.isdigit() else None
    if back and x:                                         # з вікна постачальника: лише контакти, реквізити — в юр. особах
        old = x.name
        x.name, x.contact, x.phone, x.email = name.strip()[:200] or x.name, contact.strip() or None, phone.strip() or None, email.strip() or None
        x.terms, x.note = terms.strip()[:300] or None, note.strip() or None
        if old != x.name:
            for p_ in db.scalars(select(Product).where(Product.supplier == old)):
                p_.supplier = x.name
            for d_ in db.scalars(select(StockDoc).where(StockDoc.supplier == old)):
                d_.supplier = x.name
        db.commit()
        return RedirectResponse(_card_back(back, x.id), 303)
    if not x:
        x = db.scalar(select(Supplier).where(Supplier.name == name.strip())) or Supplier(name=name.strip()[:200])
        db.add(x)
    old = x.name
    x.name, x.contact, x.phone, x.email = name.strip()[:200], contact.strip() or None, phone.strip() or None, email.strip() or None
    x.terms, x.note = terms.strip()[:300] or None, note.strip() or None
    import re as _re
    x.legal_name, x.bank = legal_name.strip()[:300] or None, bank.strip()[:200] or None
    x.edrpou = _re.sub(r"\D", "", edrpou)[:20] or None
    x.iban = _re.sub(r"\s", "", iban).upper()[:40] or None
    x.pay_purpose, x.address = pay_purpose.strip()[:300] or None, address.strip()[:300] or None
    x.contract, x.vat = contract.strip()[:200] or None, bool(vat) or None
    if x.iban and not _re.fullmatch(r"UA\d{27}", x.iban):
        db.commit()
        return RedirectResponse(f"/stock?tab=suppliers&s={_q(x.name)}&msg=IBAN+має+формат+UA+і+27+цифр+—+перевірте#sup", 303)
    if old and old != x.name:                              # перейменування — і в товарах
        for p_ in db.scalars(select(Product).where(Product.supplier == old)):
            p_.supplier = x.name
    db.commit()
    return RedirectResponse(f"/stock?tab=suppliers&s={_q(x.name)}&card={x.id}#sup", 303)


def _card_back(back: str, sid: int, msg: str = "") -> str:
    from urllib.parse import urlencode, urlsplit, parse_qsl
    u = urlsplit(back if back.startswith("/") and not back.startswith("//") else "/stock?tab=suppliers")
    qs = [(k, v) for k, v in parse_qsl(u.query) if k not in ("card", "msg")] + [("card", str(sid))] + ([("msg", msg)] if msg else [])
    return f"{u.path}?{urlencode(qs)}"


@app.get("/stock/suppliers/card", response_class=HTMLResponse)
def stock_supplier_card_by_name(request: Request, name: str = "", user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .models import Supplier
    key = name.strip().lower()
    x = next((s_ for s_ in db.scalars(select(Supplier)) if s_.name.lower() == key), None) if key else None
    if not x and key:
        x = Supplier(name=name.strip()[:200])
        db.add(x)
        db.commit()
    if not x:
        return HTMLResponse("Постачальника не знайдено", 404)
    return stock_supplier_card(request, x.id, user, db)


@app.get("/stock/suppliers/{sid}/card", response_class=HTMLResponse)
def stock_supplier_card(request: Request, sid: int, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .models import Supplier
    from .services import suppliers as sp, payments as pay
    x = db.get(Supplier, sid)
    if not x:
        return HTMLResponse("Постачальника не знайдено", 404)
    return templates.TemplateResponse(request, "_supplier_card.html", {"request": request, "user": user, "PAY": pay.PAY_LABELS,
                                                                        **sp.card(db, x)})


@app.post("/stock/suppliers/{sid}/entity")
async def stock_supplier_entity(request: Request, sid: int, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .models import Supplier
    from .services import suppliers as sp
    f = await request.form()
    x = db.get(Supplier, sid)
    back = str(f.get("back") or "/stock?tab=suppliers")
    if not x:
        return RedirectResponse(back, 303)
    eid = str(f.get("eid") or "")
    try:
        e = sp.save_entity(db, x, {k: str(v) for k, v in f.items()}, int(eid) if eid.isdigit() else None)
    except ValueError as err:
        db.rollback()
        return RedirectResponse(_card_back(back, sid, f"Не збережено: {err}"), 303)
    audit(db, "supplier.entity_save", "supplier", sid, user=user, entity_id_=e.id, kind=e.kind)
    return RedirectResponse(_card_back(back, sid), 303)


@app.post("/stock/suppliers/{sid}/entity/{eid}/delete")
def stock_supplier_entity_delete(sid: int, eid: int, back: str = Form(""), del_password: str = Form(""),
                                 user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .models import Supplier
    from .services import suppliers as sp
    err = _delpw(db, user, del_password, "supplier_entity", "supplier", sid)
    if err:
        return RedirectResponse(_card_back(back, sid, err), 303)
    x = db.get(Supplier, sid)
    if x:
        sp.delete_entity(db, x, eid)
        audit(db, "supplier.entity_delete", "supplier", sid, user=user, entity_id_=eid)
    return RedirectResponse(_card_back(back, sid), 303)


@app.post("/stock/docs/{did}/archive")
def stock_doc_archive(did: int, del_password: str = Form(""), note: str = Form(""), back: str = Form(""),
                      user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Видалення документа в архів — лише з паролем поточного користувача."""
    from .services import stock as st
    doc = db.get(StockDoc, did)
    back = back if back.startswith("/") and not back.startswith("//") else (f"/stock?tab={doc.kind}" if doc else "/stock")
    sep = "&" if "?" in back else "?"
    if not doc:
        return RedirectResponse(back, 303)
    if len(note.strip()) < 3:
        return RedirectResponse(f"{back}{sep}err={_q('Вкажіть причину видалення')}", 303)
    err = _delpw(db, user, del_password, "stock_doc", "stock_doc", did)
    if err:
        return RedirectResponse(f"{back}{sep}err={_q(err)}", 303)
    try:
        num, prev = doc.number, doc.status
        st.archive(db, doc, user.id, note)
    except st.StockError as e:
        return RedirectResponse(f"{back}{sep}err={_q(str(e))}", 303)
    audit(db, "stock.archive", "stock_doc", did, user=user, kind=doc.kind, number=num, was=prev, note=note[:300])
    if back.startswith(f"/stock/docs/{did}"):
        back, sep = f"/stock?tab={doc.kind}", "&"
    return RedirectResponse(f"{back}{sep}msg={_q(f'Документ {num} видалено в архів')}", 303)


@app.post("/stock/docs/{did}/restore")
def stock_doc_restore(did: int, user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    from .services import stock as st
    doc = db.get(StockDoc, did)
    if not doc:
        return RedirectResponse("/stock", 303)
    st.restore(db, doc)
    audit(db, "stock.restore", "stock_doc", did, user=user, kind=doc.kind, number=doc.number)
    return RedirectResponse(f"/stock/docs/{did}?msg={_q('Відновлено з архіву' + (' як чернетку — перевірте й проведіть' if doc.status == 'draft' else ''))}", 303)


def _delpw(db: Session, user: User, pw: str, what: str = "", entity: str | None = None, eid: int | None = None) -> str | None:
    """Пароль видалення: None — можна; інакше текст помилки (і запис у журнал)."""
    from .security import delete_password_error
    err = delete_password_error(user, pw)
    if err:
        audit(db, "delete.denied", entity, eid, user=user, what=what, reason="no_pw" if not user.delete_pw_hash else "wrong")
    return err


def _with_err(url: str, err: str) -> str:
    base, _, frag = url.partition("#")
    return f"{base}{'&' if '?' in base else '?'}delerr={_q(err)}" + (f"#{frag}" if frag else "")


def _q(v: str) -> str:
    from urllib.parse import quote
    return quote(v or "")


# ---------------- monobank: оплати постачальникам ----------------
def _mono_job(job, days: int = 2):
    from .integrations import monobank
    with SessionLocal() as db:
        job.log("Завантажую виписку monobank…")
        return monobank.sync(db, days=days)


CB_SETUP = "https://molodist.cbox.online/new/#setup/"
# Налаштування: блоки як у CleverBox. Пункт: (назва, опис, посилання в CRM | None, право, шлях у CleverBox | None)
SETTINGS_BLOCKS = [
    ('Клініка', 'clinic', 'Загальні параметри, філії, команда й графік', [
        ('Загальні', 'Країна, часовий пояс, мова, валюта, ПДВ; налаштування звітів', '/settings/clinic', 'users', 'main_settings'),
        ('Філії', 'Локації: ресурси (апарати), режим роботи, робочі місця (кабінети)', '/settings/clinic#loc', 'users', 'salons'),
        ('Команда', 'Лікарі й співробітники: посади, KPI, привʼязка до CleverBox', '/specialists', 'specialists', 'sfields/staff'),
        ('Графік роботи', 'Зміни спеціалістів по днях і кабінетах', '/specialists/schedule', 'schedule', None),
        ('Онлайн-бронювання', 'Віджет запису на сайті', None, None, 'webwidget'),
    ]),
    ('Співробітники', 'team', 'Посади, оклади, категорії й нарахування', [
        ('Посади співробітників', 'Директор, адміністратори, лікарі…: довідник посад для картки спеціаліста', '/settings/positions', 'specialists_edit', 'dirs/prof'),
        ('Налаштування окладів', None, None, None, 'salary'),
        ('Категорії спеціалістів', 'Ціни за категоріями', None, None, 'cmarket/customprice/main'),
        ('Статті для нарахувань', None, None, None, 'dirs/stype'),
        ('Застосунок для спеціаліста', None, None, None, 'cmarket/specapp/main'),
    ]),
    ('Фінанси', 'finance', 'Каси, статті, оплати й фіскалізація', [
        ('Каси / рахунки', 'Готівка, рахунки ФОП, термінал, онлайн-оплати: залишок після кожної операції, локація, тип', '/settings/cashboxes', 'finance', None),
        ('Статті фінансових операцій', 'Довідник статей доходів і витрат: реклама, податки, персонал, закупки…', '/settings/articles', 'finance', 'dirs/ftype'),
        ('Знижки й бонусна програма', 'Правила знижок, % нарахування бонусів, термін дії, персональні знижки', '/settings/loyalty', 'finance', None),
        ('Платіжні системи', 'Посилання на оплату для пацієнтів: monobank еквайринг, LiqPay, WayForPay', '/settings/payments', 'finance', 'payment_system'),
        ('Банк monobank', 'Виписки ФОП, оплати постачальникам і їх автоматичне рознесення', '/settings/keys?s=mono#mono', 'finance', None),
        ('Checkbox', 'Фіскалізація чеків CRM, каси (РРО), відправлення чеків; оборот ФОП — у «Звіти → Ліміти ФОП»', '/settings/checkbox', 'finance', 'cmarket/checkbox/main'),
    ]),
    ('Комунікації', 'chat', 'Канали звʼязку з пацієнтами', [
        ('Канали', 'SMS і Viber (TurboSMS), Telegram-бот, Instagram та Facebook', '/settings/channels', 'users', 'cmarket/channels/main'),
        ('Повідомлення пацієнтам', 'Розсилки й нагадування CRM; автоповідомлення — у CleverBox', '/campaigns', 'campaigns', 'messages'),
        ('IP-телефонія', 'Binotel, Stream Telecom, Ringostat: дзвінки → звернення', '/settings/telephony', 'users', 'cmarket/siphone/main'),
    ]),
    ('Звернення', 'inbox', 'Довідники для заявок і дзвінків', [
        ('Джерела звернень', None, '/settings/appeals?open=sources', 'users', 'dirs/appeal_type'),
        ('Мітки звернень', None, '/settings/appeals?open=marks', 'users', 'dirs/appeal_mark'),
        ('Статуси звернень', None, '/settings/appeals?open=statuses', 'users', 'dirs/appeal_status'),
        ('Причини відмов', None, '/settings/appeals?open=deny_reasons', 'users', 'dirs/appeal_deny'),
        ('Причини звернень', None, '/settings/appeals?open=causes', 'users', 'dirs/appeal_cause'),
    ]),
    ('Склад і товари', 'box', 'Списання, терміни, накладні, сайт', [
        ('Норми списання', 'Що списується зі складу після процедури', '/stock?tab=norms', 'stock', None),
        ('Сповіщення про терміни', 'Кому й коли надсилати попередження про терміни придатності', '/reports?r=expiry#notify', 'finance', None),
        ('Розпізнавання накладних', 'Ключ Claude API — фото/PDF накладної → документ поставки', '/settings/keys?s=ai#ai', 'finance', None),
        ('molodist.store', 'Shop-Express: посилання й фото товарів з сайту', '/settings/keys?s=shop#shop', 'finance', None),
    ]),
    ('Доступ і безпека', 'lock', 'Хто що бачить і може змінювати', [
        ('Права доступу', 'Ролі та їх дозволи, користувачі CRM, паролі входу й видалення', '/access/roles', 'users', 'roles'),
        ('Паролі для видалення', 'Окремий у кожного користувача — без нього нічого не видаляється', '/access', 'users', None),
        ('Журнал дій', 'Хто що і коли робив у CRM', '/access/audit', 'users', None),
        ('Мій кабінет', 'Пароль входу, Telegram для сповіщень', '/me', None, None),
    ]),
    ('Дані й інтеграції', 'data', 'Синхронізація, копії, API', [
        ('Дані й синхронізація', 'Імпорт / експорт, синхронізація з CleverBox', '/data', 'export_reports', None),
        ('Резервні копії', 'Щоночі о 03:30 і вручну', '/data#backup', 'export_reports', None),
        ('REST API', 'API-ключі CRM і вебхуки; токен CleverBox — у «Дані й синхронізація»', '/access/api', 'api', 'cmarket/rest/main'),
        ('Зіставлення послуг', 'Позиції плану ↔ послуги CleverBox', '/mapping', 'mapping', None),
    ]),
]


@app.get("/settings", response_class=HTMLResponse)
def settings_page(request: Request, user: User = Depends(current_user), db: Session = Depends(db_dep)):
    """Налаштування: усе в одному місці, блоками як у CleverBox. Що є в CRM — відкривається тут, решта — у CleverBox."""
    from .integrations import monobank, checkbox as chb, shop_express
    from .services import invoice_ai
    st_ = {}
    try:
        st_["/settings/keys?s=mono#mono"] = f"підключено {len(monobank.saved(db))} ФОП" if monobank.saved(db) else "не підключено"
        from .services import paylinks as _pls
        _pp = _pls.public(db)
        _ra = [a_ for a_ in _pp["accounts"] if a_["ready"]]
        st_["/settings/payments"] = f"підключено {len(_ra)}" if _ra else "не підключено"
        st_["/settings/keys?s=ai#ai"] = "ключ задано" if invoice_ai.api_key(db) else "не задано"
        st_["/settings/keys?s=shop#shop"] = "підключено" if shop_express.creds(db).get("login") else "не підключено"
        acc = chb.accounts(db)
        st_["/reports?r=fop"] = f"підключено {len(acc)} ФОП" if acc else "не підключено"
        n_keys = sum(bool(x) for x in (monobank.saved(db), acc, invoice_ai.api_key(db), shop_express.creds(db).get("login")))
        st_["/settings/keys"] = f"підключено {n_keys} з 4"
        users = db.scalars(select(User).where(User.active.is_(True))).all()
        st_["/access"] = f"паролі видалення: {sum(1 for u_ in users if u_.delete_pw_hash)} з {len(users)}"
    except Exception:  # noqa: BLE001
        pass
    blocks = []
    for title, icon, about, items in SETTINGS_BLOCKS:
        rows = [{"t": t, "d": d, "url": u if (u and (not p or can(user, p))) else None, "cb": (CB_SETUP + c) if c else None,
                 "status": st_.get(u) if u else None} for t, d, u, p, c in items]
        rows = [r for r in rows if r["url"] or r["cb"]]
        if rows:
            blocks.append({"title": title, "icon": icon, "about": about, "rows": rows, "ext": all(not r["url"] for r in rows)})
    return render(request, "settings.html", blocks=blocks)


# ---------------- Налаштування → Посади співробітників ----------------
@app.get("/settings/positions", response_class=HTMLResponse)
def settings_positions(request: Request, user: User = Depends(require("specialists_edit")), db: Session = Depends(db_dep)):
    from .services import positions as ps
    return render(request, "positions.html", nav="settings", items=ps.items(db), use=ps.usage(db))


@app.post("/settings/positions")
def settings_positions_save(pid: str = Form(""), name: str = Form(""), user: User = Depends(require("specialists_edit")),
                            db: Session = Depends(db_dep)):
    from .services import positions as ps
    err = ps.save(db, int(pid) if pid.isdigit() else None, name)
    if not err:
        audit(db, "positions.save", "position", int(pid) if pid.isdigit() else None, user=user, name=name)
    return RedirectResponse(_with_msg("/settings/positions", err or "Збережено", bool(err)), 303)


@app.post("/settings/positions/{pid:int}/toggle")
def settings_positions_toggle(pid: int, user: User = Depends(require("specialists_edit")), db: Session = Depends(db_dep)):
    from .services import positions as ps
    ps.toggle(db, pid)
    return RedirectResponse("/settings/positions", 303)


@app.post("/settings/positions/{pid:int}/delete")
def settings_positions_delete(pid: int, user: User = Depends(require("specialists_edit")), db: Session = Depends(db_dep)):
    from .services import positions as ps
    ps.delete(db, pid)
    audit(db, "positions.delete", "position", pid, user=user)
    return RedirectResponse(_with_msg("/settings/positions", "Посаду видалено"), 303)


@app.post("/settings/positions/order")
async def settings_positions_order(request: Request, user: User = Depends(require("specialists_edit")), db: Session = Depends(db_dep)):
    from .services import positions as ps
    js = await request.json()
    ps.reorder(db, [int(x) for x in js.get("ids") or [] if str(x).isdigit()])
    return JSONResponse({"ok": True})


# ---------------- Налаштування → Статті фінансових операцій ----------------
@app.get("/settings/articles", response_class=HTMLResponse)
def settings_articles(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    fa = _fa()
    xs = fa.items(db)
    use = fa.usage(db)
    groups: dict = {}
    for x in xs:
        groups.setdefault((x["kind"], fa.group(x["name"])), []).append({**x, "used": use.get(x["name"], 0)})
    return render(request, "fin_articles.html", nav="settings", groups=groups, KINDS=fa.KINDS, n=len(xs),
                  kind=request.query_params.get("kind", ""), q=request.query_params.get("q", ""))


@app.post("/settings/articles")
def settings_articles_save(aid: str = Form(""), name: str = Form(""), kind: str = Form(""), active: str = Form(""), salary: str = Form(""),
                           user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    try:
        x = _fa().save(db, int(aid) if aid.isdigit() else None, name, kind, bool(active), bool(salary))
    except ValueError as e:
        return RedirectResponse(_with_msg("/settings/articles", str(e), True), 303)
    audit(db, "fin.article_save", user=user, article=x["id"], name=x["name"], kind=x["kind"], active=x["active"])
    return RedirectResponse(_with_msg("/settings/articles", f"Статтю «{x['name']}» збережено"), 303)


@app.post("/settings/articles/{aid:int}/delete")
def settings_articles_delete(aid: int, del_password: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    err = _delpw(db, user, del_password, "fin_article", None, aid)
    if err:
        return RedirectResponse(_with_err("/settings/articles", err), 303)
    gone = _fa().delete(db, aid)
    audit(db, "fin.article_delete", user=user, article=aid, removed=gone)
    return RedirectResponse(_with_msg("/settings/articles", "Статтю видалено" if gone else "По статті є операції — її вимкнено (історія лишається)"), 303)


# ---------------- Налаштування → Каси / рахунки ----------------
@app.get("/settings/cashboxes", response_class=HTMLResponse)
def settings_cashboxes(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import cashboxes as cbx, clinic as cl, sales as sl
    boxes = cbx.overview(db)
    open_ = request.query_params.get("open", "")
    detail = next((b for b in boxes if str(b["id"]) == open_), None)
    return render(request, "cashboxes.html", nav="settings", boxes=boxes, TYPES=cbx.TYPES, locs=cl.get(db).get("locations") or [],
                  methods=sl.methods(db), open_=open_, detail=detail, dmoves=cbx.moves(db, detail, 100) if detail else [],
                  total=round(sum(b["balance"] for b in boxes if b.get("active")), 2))


@app.post("/settings/cashboxes")
async def settings_cashboxes_save(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import cashboxes as cbx
    f = await request.form()
    bid = int(f.get("id")) if str(f.get("id") or "").isdigit() else None
    data = {"name": f.get("name"), "type": f.get("type"), "active": bool(f.get("active")), "note": f.get("note"),
            "location_id": int(f.get("location_id")) if str(f.get("location_id") or "").isdigit() else None,
            "opening": _num(str(f.get("opening") or "0")) or 0, "opening_date": _day(str(f.get("opening_date") or "")) and str(f.get("opening_date")),
            "methods": f.getlist("methods"), "color": str(f.get("color") or "")[:9], "pos": bool(f.get("pos"))}
    try:
        x = cbx.save(db, bid, data)
    except ValueError as e:
        return RedirectResponse(_with_msg("/settings/cashboxes" + (f"?edit={bid}" if bid else ""), str(e), True), 303)
    audit(db, "cashbox.save", user=user, cashbox=x["id"], name=x["name"], type=x["type"], opening=x["opening"], methods=x["methods"])
    return RedirectResponse(_with_msg("/settings/cashboxes", f"Касу «{x['name']}» збережено"), 303)


@app.post("/settings/cashboxes/{bid:int}/delete")
def settings_cashboxes_delete(bid: int, del_password: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import cashboxes as cbx
    err = _delpw(db, user, del_password, "cashbox", None, bid)
    if err:
        return RedirectResponse(_with_err("/settings/cashboxes", err), 303)
    cbx.delete(db, bid)
    audit(db, "cashbox.delete", user=user, cashbox=bid)
    return RedirectResponse(_with_msg("/settings/cashboxes", "Касу видалено"), 303)


@app.post("/settings/cashboxes/{bid:int}/restore")
def settings_cashboxes_restore(bid: int, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import cashboxes as cbx
    cbx.restore(db, bid)
    audit(db, "cashbox.restore", user=user, cashbox=bid)
    return RedirectResponse(_with_msg("/settings/cashboxes", "Касу відновлено"), 303)


# ---------------- Налаштування → Платіжні системи (посилання на оплату) ----------------
@app.get("/settings/payments", response_class=HTMLResponse)
def settings_payments(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import paylinks as pls
    from .models import PayLink
    t_ = pls.public(db)
    links = db.scalars(select(PayLink).order_by(PayLink.id.desc()).limit(40)).all()
    pats = {p_.id: p_ for p_ in db.scalars(select(Patient).where(Patient.id.in_({l_.patient_id for l_ in links if l_.patient_id})))} if links else {}
    base = str(request.base_url).rstrip("/").replace("http://", "https://")
    from .services import cashboxes as _cbx
    return render(request, "paysys.html", nav="settings", t=t_, PROV=pls.PROVIDERS, links=links, pats=pats, SL=pls.STATUS_LABELS,
                  boxes=_cbx.names(db), accs={a_["id"]: a_["name"] for a_ in t_["accounts"]},
                  hook_url=lambda p: f"{base}/api/public/pay/{p}/{t_['hook']}")


@app.post("/settings/payments")
async def settings_payments_save(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import paylinks as pls
    f = await request.form()
    pls.save_common(db, int(f.get("validity_h")) if str(f.get("validity_h") or "").isdigit() else 24, str(f.get("text") or ""),
                    int(f.get("default")) if str(f.get("default") or "").isdigit() else None)
    audit(db, "paysys.save", user=user)
    return RedirectResponse(_with_msg("/settings/payments", "Збережено"), 303)


@app.post("/settings/payments/account")
async def settings_payments_account(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import paylinks as pls
    f = await request.form()
    aid = int(f.get("id")) if str(f.get("id") or "").isdigit() else None
    try:
        a_ = pls.save_account(db, aid, str(f.get("provider") or ""), str(f.get("name") or ""), {k: f.get(k) for k in f.keys()},
                              str(f.get("cashbox") or ""), bool(f.get("active")))
    except ValueError as e:
        return RedirectResponse(_with_msg("/settings/payments", str(e), True), 303)
    audit(db, "paysys.account_save", user=user, account=a_["id"], provider=a_["provider"], name=a_["name"])
    if f.get("test"):
        ok, msg = await asyncio.to_thread(pls.test, db, a_["id"])
        return RedirectResponse(_with_msg("/settings/payments", ("✅ " if ok else "❌ ") + msg, not ok), 303)
    return RedirectResponse(_with_msg("/settings/payments", f"Підключення «{a_['name']}» збережено"), 303)


@app.post("/settings/payments/account/{aid:int}/delete")
def settings_payments_account_delete(aid: int, del_password: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import paylinks as pls
    err = _delpw(db, user, del_password, "paysys", None, aid)
    if err:
        return RedirectResponse(_with_err("/settings/payments", err), 303)
    pls.delete_account(db, aid)
    audit(db, "paysys.account_delete", user=user, account=aid)
    return RedirectResponse(_with_msg("/settings/payments", "Підключення видалено"), 303)


@app.post("/paylinks")
async def paylink_create(request: Request, user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import paylinks as pls
    from .integrations.messaging import send_to_patient
    f = await request.form()
    back = str(f.get("back") or "/settings/payments")
    back = back if back.startswith("/") and not back.startswith("//") else "/"
    pid = int(f.get("patient_id")) if str(f.get("patient_id") or "").isdigit() else None
    aid = int(f.get("appointment_id")) if str(f.get("appointment_id") or "").isdigit() else None
    if pid:
        _patient_or_403(db, user, pid)
    base = str(request.base_url).rstrip("/").replace("http://", "https://")
    try:
        acc = int(f.get("account_id")) if str(f.get("account_id") or "").isdigit() else None
        pl = await asyncio.to_thread(pls.create, db, _num(str(f.get("amount") or "0")) or 0, str(f.get("purpose") or ""), base, pid, aid, user.id,
                                     None, acc)
    except pls.PayError as e:
        return RedirectResponse(_with_msg(back, str(e), True), 303)
    audit(db, "paylink.create", "paylink", pl.id, user=user, amount=pl.amount, patient=pid, provider=pl.provider)
    msg = f"Посилання на оплату №{pl.id} на {pl.amount:.2f} ₴ створено"
    if f.get("send") and pid:
        p_ = db.get(Patient, pid)
        res = await asyncio.to_thread(send_to_patient, p_, pls.message(db, pl))
        pl.info = {**(pl.info or {}), "sent": {"ok": res.ok, "channel": res.channel, "at": datetime.now().strftime("%d.%m.%Y %H:%M")}}
        db.commit()
        msg += f" і надіслано ({res.channel})" if res.ok else f", але не надіслано: {res.error}"
    return RedirectResponse(_with_msg(back, msg), 303)


@app.post("/paylinks/{lid:int}/check")
async def paylink_check(lid: int, back: str = Form("/settings/payments"), user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .services import paylinks as pls
    from .models import PayLink
    pl = db.get(PayLink, lid) or (_ for _ in ()).throw(HTTPException(404))
    back = back if back.startswith("/") and not back.startswith("//") else "/"
    try:
        st_ = await asyncio.to_thread(pls.check_status, db, pl)
    except (httpx.HTTPError, ValueError) as e:
        return RedirectResponse(_with_msg(back, f"Не вдалося перевірити: {e}", True), 303)
    return RedirectResponse(_with_msg(back, f"Посилання №{lid}: {pls.STATUS_LABELS.get(st_, st_)}"), 303)


@app.post("/paylinks/{lid:int}/send")
async def paylink_send(lid: int, ch: str = Form(...), back: str = Form("/settings/payments"), user: User = Depends(require("sales")),
                       db: Session = Depends(db_dep)):
    """Надіслати посилання на оплату пацієнту вибраним каналом: Telegram / Viber / SMS."""
    from .services import paylinks as pls
    from .models import PayLink
    from .integrations.messaging import send_to_patient
    pl = db.get(PayLink, lid) or (_ for _ in ()).throw(HTTPException(404))
    back = back if back.startswith("/") and not back.startswith("//") else "/"
    if ch not in ("telegram", "viber", "sms"):
        raise HTTPException(400)
    p_ = db.get(Patient, pl.patient_id) if pl.patient_id else None
    if not p_:
        return RedirectResponse(_with_msg(back, "У посилання немає пацієнта", True), 303)
    res = await asyncio.to_thread(send_to_patient, p_, pls.message(db, pl), [ch])
    lbl = {"telegram": "Telegram", "viber": "Viber", "sms": "SMS"}[ch]
    if res.ok:
        pl.info = {**(pl.info or {}), "sent": {"ok": True, "channel": res.channel, "at": datetime.now().strftime("%d.%m.%Y %H:%M")}}
        db.commit()
    audit(db, "paylink.send", "paylink", lid, user=user, channel=ch, ok=res.ok)
    return RedirectResponse(_with_msg(back, f"Посилання надіслано в {lbl}" if res.ok else f"{lbl}: не надіслано — {res.error}", not res.ok), 303)


@app.post("/paylinks/{lid:int}/cancel")
def paylink_cancel(lid: int, back: str = Form("/settings/payments"), user: User = Depends(require("sales")), db: Session = Depends(db_dep)):
    from .models import PayLink
    pl = db.get(PayLink, lid) or (_ for _ in ()).throw(HTTPException(404))
    back = back if back.startswith("/") and not back.startswith("//") else "/"
    if pl.status == "created":
        pl.status = "cancelled"
        db.commit()
        audit(db, "paylink.cancel", "paylink", lid, user=user)
    return RedirectResponse(_with_msg(back, f"Посилання №{lid} скасовано в CRM"), 303)


@app.post("/api/public/pay/{provider}/{hook}")
async def paylink_hook(request: Request, provider: str, hook: str, db: Session = Depends(db_dep)):
    """Вебхук платіжної системи про оплату посилання."""
    from .services import paylinks as pls
    import hmac as _hm
    if provider not in pls.PROVIDERS or not _hm.compare_digest(hook, pls.public(db)["hook"]):
        raise HTTPException(404)
    raw = await request.body()
    form = {}
    if provider != "mono" and raw and not raw.lstrip().startswith(b"{"):
        from urllib.parse import parse_qsl
        form = dict(parse_qsl(raw.decode("utf-8", "replace"), keep_blank_values=True))
    res = await asyncio.to_thread(pls.webhook, db, provider, raw, form)
    return JSONResponse(res)


# ---------------- Налаштування → IP-телефонія ----------------
@app.get("/settings/telephony", response_class=HTMLResponse)
def settings_telephony(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import telephony as tel, appeals as ap_
    from .models import Call
    t_ = tel.public(db)
    calls = db.scalars(select(Call).order_by(Call.at.desc()).limit(30)).all()
    base = str(request.base_url).rstrip("/").replace("http://", "https://")
    return render(request, "telephony.html", nav="settings", t=t_, PROV=tel.PROVIDERS, calls=calls, sources=ap_.dirs(db)["sources"],
                  hook_url=lambda p: f"{base}/api/public/telephony/{p}/{t_['hook']}")


@app.post("/settings/telephony")
async def settings_telephony_save(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import telephony as tel
    f = await request.form()
    prov = str(f.get("provider") or "")
    try:
        tel.save(db, prov, {k: f.get(k) for k in f.keys()}, str(f.get("source") or ""), bool(f.get("create_appeal")), bool(f.get("missed_only")))
    except ValueError as e:
        return RedirectResponse(_with_msg("/settings/telephony", str(e), True), 303)
    audit(db, "telephony.save", user=user, provider=prov)
    if f.get("test"):
        ok, msg = await asyncio.to_thread(tel.test, db)
        return RedirectResponse(_with_msg("/settings/telephony", ("✅ " if ok else "❌ ") + msg, not ok), 303)
    return RedirectResponse(_with_msg("/settings/telephony", "Збережено"), 303)


@app.post("/api/public/telephony/{provider}/{hook}")
async def telephony_hook(request: Request, provider: str, hook: str, db: Session = Depends(db_dep)):
    """Вебхук дзвінків провайдера IP-телефонії → журнал дзвінків і звернення."""
    from .services import telephony as tel
    import hmac as _hm
    t_ = tel.public(db)
    if provider not in tel.PROVIDERS or not _hm.compare_digest(hook, t_["hook"]):
        raise HTTPException(404)
    raw = await request.body()
    try:
        data = json.loads(raw or b"{}")
    except ValueError:
        from urllib.parse import parse_qsl
        data = dict(parse_qsl(raw.decode("utf-8", "replace")))
    if not isinstance(data, dict):
        data = {}
    c = tel.ingest(db, provider, data)
    return {"ok": True, "call": c.id, "appeal": c.appeal_id, "status": "success"}


# ---------------- Налаштування → Checkbox (фіскалізація) ----------------
@app.get("/settings/checkbox", response_class=HTMLResponse)
def settings_checkbox(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis, sales as sl
    from .models import Sale
    last = db.scalars(select(Sale).where(Sale.fiscal.is_not(None)).order_by(Sale.id.desc()).limit(10)).all()
    from .services import cashboxes as _cbx, clinic as _cl
    boxes = [{"name": x["name"], "cash": x.get("type") == "cash"} for x in _cbx.items(db)
             if x.get("active") and not x.get("deleted") and x.get("type") in ("cash", "card", "other")]
    return render(request, "checkbox_settings.html", nav="settings", c=_fis.cfg(db), rros=_fis.rros(db), methods=sl.methods(db),
                  boxes=boxes, locs=_cl.get(db).get("locations") or [], term=_fis.terminals(_fis.cfg(db)),
                  last=last, open_=request.query_params.get("open", ""))


@app.post("/settings/checkbox/receipt")
async def settings_checkbox_receipt(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis, sales as sl
    f = await request.form()
    c = _fis.cfg(db)
    c["auto"] = bool(f.get("auto"))
    c["default_rro"] = int(f.get("default_rro")) if str(f.get("default_rro") or "").isdigit() else None
    c["method_modes"] = {m: str(f.get(f"mode_{i}") or "skip") for i, m in enumerate(sl.methods(db))
                         if str(f.get(f"mode_{i}") or "skip") in ("cash", "card", "skip")}
    _fis.save_cfg(db, c)
    audit(db, "fiscal.cfg", user=user, auto=c["auto"], rro=c["default_rro"], modes=c["method_modes"])
    return RedirectResponse(_with_msg("/settings/checkbox", "Налаштування фіскального чека збережено"), 303)


@app.post("/settings/checkbox/send")
async def settings_checkbox_send(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    f = await request.form()
    c = _fis.cfg(db)
    c["send"] = {"on": bool(f.get("on")), "channels": [x for x in ("telegram", "viber", "sms") if f.get(f"ch_{x}")],
                 "text": str(f.get("text") or "")[:500] or _fis.DEFAULT["send"]["text"]}
    _fis.save_cfg(db, c)
    audit(db, "fiscal.send_cfg", user=user, on=c["send"]["on"], channels=c["send"]["channels"])
    return RedirectResponse(_with_msg("/settings/checkbox", "Відправлення чеків збережено"), 303)


@app.post("/settings/checkbox/extra")
async def settings_checkbox_extra(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    f = await request.form()
    c = _fis.cfg(db)
    if "_paylink" in f:
        c["paylink"] = bool(f.get("paylink"))
    if "_widget" in f:
        c["widget"] = bool(f.get("widget"))
    _fis.save_cfg(db, c)
    return RedirectResponse(_with_msg("/settings/checkbox", "Збережено"), 303)


@app.post("/settings/checkbox/rro")
async def settings_checkbox_rro(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    f = await request.form()
    rid = int(f.get("id")) if str(f.get("id") or "").isdigit() else None
    if not str(f.get("name") or "").strip():
        return RedirectResponse(_with_msg("/settings/checkbox?open=rro", "Вкажіть назву каси", True), 303)
    methods = {}
    for k in f.keys():
        if k.startswith("m_on::"):
            m_ = k[6:]
            methods[m_] = str(f.get("m_lbl::" + m_) or "cash")
    data = {k: f.get(k) for k in ("name", "fop", "login", "license_key", "password", "pin", "test")} | {"on": bool(f.get("on"))}
    if f.get("_methods"):
        data["methods"] = methods
        data["branches"] = f.getlist("branches")
    rid = _fis.save_rro(db, rid, data)
    audit(db, "fiscal.rro_save", user=user, rro=rid, test=bool(f.get("test")), methods=sorted(methods))
    if f.get("check"):
        ok, msg = await asyncio.to_thread(_fis.connect, db, rid)
        return RedirectResponse(_with_msg("/settings/checkbox?open=rro", ("✅ " if ok else "❌ ") + msg, not ok), 303)
    return RedirectResponse(_with_msg("/settings/checkbox?open=rro", "Касу збережено"), 303)


@app.post("/settings/checkbox/terminals")
async def settings_checkbox_terminals(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Термінали Checkbox PayLink: адреса утиліти PayLink і термінали (ID пристрою ↔ спосіб оплати каси)."""
    from .services import fiscal as _fis
    f = await request.form()
    c = _fis.cfg(db)
    t = _fis.terminals(c)
    act, tid = str(f.get("act") or "save"), str(f.get("tid") or "")
    if act == "conn":
        t["host"] = str(f.get("host") or "127.0.0.1").strip()[:120]
        t["port"] = int(f.get("port")) if str(f.get("port") or "").isdigit() else 9020
    elif act in ("on", "off", "del"):
        for x in t["items"]:
            if str(x["id"]) == tid:
                x["on"] = act == "on"
        if act == "del":
            t["items"] = [x for x in t["items"] if str(x["id"]) != tid]
    else:
        name, dev = str(f.get("name") or "").strip()[:80], str(f.get("device_id") or "").strip()[:80]
        if not name or not dev:
            return RedirectResponse(_with_msg("/settings/checkbox?open=term", "Вкажіть назву й ID термінала з PayLink", True), 303)
        x = next((x for x in t["items"] if str(x["id"]) == tid), None)
        if not x:
            x = {"id": max([int(y["id"]) for y in t["items"]] + [0]) + 1, "on": True}
            t["items"].append(x)
        x.update(name=name, device_id=dev, method=str(f.get("method") or "")[:60])
    c["terminals"] = t
    c["paylink"] = any(x.get("on") for x in t["items"])
    _fis.save_cfg(db, c)
    audit(db, "fiscal.terminals", user=user, act=act, terminal=tid)
    return RedirectResponse(_with_msg("/settings/checkbox?open=term", {"del": "Термінал видалено", "on": "Термінал підключено", "off": "Термінал відключено"}.get(act, "Збережено")), 303)


@app.post("/settings/checkbox/rro/{rid:int}/off")
def settings_checkbox_rro_off(rid: int, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    _fis.disconnect(db, rid)
    audit(db, "fiscal.rro_off", user=user, rro=rid)
    return RedirectResponse(_with_msg("/settings/checkbox?open=rro", "Касу відключено"), 303)


@app.get("/settings/checkbox/rro/{rid:int}/{what}")
async def settings_checkbox_rro_data(rid: int, what: str, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    if what not in ("shifts", "taxes"):
        raise HTTPException(404)
    try:
        rows = await asyncio.to_thread(_fis.shifts if what == "shifts" else _fis.taxes, db, rid)
    except (_fis.FiscalError, httpx.HTTPError) as e:
        return JSONResponse({"error": str(e)[:200]}, status_code=400)
    return JSONResponse({"rows": rows})


@app.post("/settings/checkbox/rro/{rid:int}/delete")
def settings_checkbox_rro_delete(rid: int, del_password: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis
    err = _delpw(db, user, del_password, "rro", None, rid)
    if err:
        return RedirectResponse(_with_err("/settings/checkbox?open=rro", err), 303)
    _fis.delete_rro(db, rid)
    audit(db, "fiscal.rro_delete", user=user, rro=rid)
    return RedirectResponse(_with_msg("/settings/checkbox?open=rro", "Касу видалено"), 303)


@app.post("/settings/checkbox/rro/{rid:int}/close-shift")
async def settings_checkbox_close_shift(rid: int, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import fiscal as _fis

    def _close():
        k = _fis.Kasa(_fis._rro(db, rid))
        if not k.shift():
            return "Зміна вже закрита"
        k.close_shift()
        return "Зміну закрито (Z-звіт)"
    try:
        msg = await asyncio.to_thread(_close)
    except (_fis.FiscalError, httpx.HTTPError) as e:
        return RedirectResponse(_with_msg("/settings/checkbox?open=rro", f"Checkbox: {e}", True), 303)
    audit(db, "fiscal.close_shift", user=user, rro=rid)
    return RedirectResponse(_with_msg("/settings/checkbox?open=rro", msg), 303)


# ---------------- Налаштування → Канали ----------------
@app.get("/settings/channels", response_class=HTMLResponse)
def settings_channels(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import channels as ch
    return render(request, "channels.html", nav="settings", chs=ch.channels(db), T=ch.TYPES, open_=request.query_params.get("open", ""),
                  dlr_url=str(request.base_url).rstrip("/").replace("http://", "https://") + "/api/public/turbosms/dlr")


@app.post("/settings/channels/{cid:int}")
async def settings_channel_save(request: Request, cid: int, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import channels as ch
    f = await request.form()
    params = {}
    for k in ("sender", "allowed_from", "allowed_to", "login", "username", "ttl"):
        if k in f:
            params[k] = str(f.get(k) or "").strip()[:80]
    if "ttl" in params:
        params["ttl"] = int(params["ttl"]) if params["ttl"].isdigit() else 3600
    if "fallback_sms" in f or f.get("_has_fallback"):
        params["fallback_sms"] = bool(f.get("fallback_sms"))
    secrets = {k: str(f.get(k) or "").strip() for k in ("token", "webhook_secret", "password", "page_token") if f.get(k)}
    secrets.update({f"clear_{k}": "1" for k in ("token", "webhook_secret", "password", "page_token") if f.get(f"clear_{k}")})
    try:
        ch.update(db, cid, str(f.get("name") or "").strip(), bool(f.get("on")), params, secrets)
    except ValueError as e:
        return RedirectResponse(_with_msg("/settings/channels", str(e), True), 303)
    audit(db, "channels.save", user=user, channel=cid, on=bool(f.get("on")), params=params, secrets=sorted(secrets))
    if f.get("test"):
        ok, msg = await asyncio.to_thread(ch.test, db, cid)
        return RedirectResponse(_with_msg(f"/settings/channels?open={cid}", ("✅ " if ok else "❌ ") + msg, not ok), 303)
    return RedirectResponse(_with_msg("/settings/channels", "Канал збережено"), 303)


@app.post("/settings/channels/add")
def settings_channel_add(type_: str = Form(..., alias="type"), name: str = Form(""), user: User = Depends(require("users")),
                         db: Session = Depends(db_dep)):
    from .services import channels as ch
    if type_ not in ch.TYPES:
        raise HTTPException(400)
    c = ch.add(db, type_, name.strip())
    audit(db, "channels.add", user=user, channel=c["id"], type=type_)
    return RedirectResponse(f"/settings/channels?open={c['id']}", 303)


@app.post("/settings/channels/{cid:int}/delete")
def settings_channel_delete(cid: int, del_password: str = Form(""), user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import channels as ch
    err = _delpw(db, user, del_password, "channel", None, cid)
    if err:
        return RedirectResponse(_with_err("/settings/channels", err), 303)
    ch.remove(db, cid)
    audit(db, "channels.delete", user=user, channel=cid)
    return RedirectResponse(_with_msg("/settings/channels", "Канал видалено"), 303)


@app.post("/api/public/turbosms/dlr")
async def turbosms_dlr(request: Request, db: Session = Depends(db_dep)):
    """Вебхук TurboSMS: статуси доставки SMS / Viber → журнал повідомлень і розсилки."""
    from .services import channels as ch
    from .integrations import messaging as _m
    from .models import MessageLog, CampaignRecipient
    raw = await request.body()
    if not ch.verify_dlr(db, raw, request.headers.get("X-Signature") or request.headers.get("Signature")):
        raise HTTPException(403, "bad signature")
    try:
        data = json.loads(raw or b"{}")
    except ValueError:
        data = dict((await request.form()).items())
    rows = data if isinstance(data, list) else (data.get("messages") or data.get("response_result") or [data])
    n = 0
    for x in rows if isinstance(rows, list) else []:
        mid, st_ = str(x.get("message_id") or ""), _m.STATUS_MAP.get(str(x.get("status") or "").lower())
        if not mid or not st_:
            continue
        for m in db.scalars(select(MessageLog).where(MessageLog.message_id == mid)):
            m.status = st_
        for r in db.scalars(select(CampaignRecipient).where(CampaignRecipient.message_id == mid)):
            r.status = st_
        n += 1
    db.commit()
    return {"ok": True, "updated": n}


# ---------------- Налаштування → Основні (загальні, звіти, локації) ----------------
@app.get("/settings/clinic", response_class=HTMLResponse)
def settings_clinic(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import clinic as cl
    from .models import CbFinOp
    data = cl.get(db)
    arts = sorted({a for (a,) in db.execute(select(CbFinOp.article).where(CbFinOp.kind == "income", CbFinOp.article.is_not(None)).distinct()) if a}
                  | set(_fa().names(db, "income", False)) | set(data["reports"]["exclude_income"]))
    return render(request, "clinic.html", nav="settings", c=data, cl=cl, income_articles=arts,
                  open_=request.query_params.get("open", ""))


def _clinic_back(open_: str = "", msg: str = "Збережено") -> RedirectResponse:
    return RedirectResponse(f"/settings/clinic?msg={_q(msg)}" + (f"&open={open_}" if open_ else ""), 303)


@app.post("/settings/clinic/general")
def settings_clinic_general(country: str = Form(""), timezone_: str = Form("Europe/Kyiv", alias="timezone"), language: str = Form(""),
                            currency: str = Form("UAH"), vat: str = Form(""), user: User = Depends(require("users")),
                            db: Session = Depends(db_dep)):
    from zoneinfo import ZoneInfo
    from .services import clinic as cl
    try:
        ZoneInfo(timezone_)
    except Exception:  # noqa: BLE001
        timezone_ = "Europe/Kyiv"
    data = cl.get(db)
    data["general"] = {"country": country[:60], "timezone": timezone_, "language": language[:40], "currency": currency[:6], "vat": bool(vat)}
    cl.save(db, data)
    audit(db, "settings.general", user=user, **data["general"])
    return _clinic_back()


@app.post("/settings/clinic/reports")
async def settings_clinic_reports(request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import clinic as cl
    f = await request.form()
    items = [str(x).strip()[:120] for x in f.getlist("exclude_income") if str(x).strip()]
    extra = str(f.get("extra") or "").strip()[:120]
    if extra:
        items.append(extra)
    data = cl.get(db)
    data["reports"]["exclude_income"] = sorted(set(items))
    cl.save(db, data)
    audit(db, "settings.reports", user=user, exclude_income=data["reports"]["exclude_income"])
    return _clinic_back()


@app.post("/settings/clinic/location")
def settings_clinic_location(lid: str = Form(""), name: str = Form(...), address: str = Form(""), phone: str = Form(""),
                             active: str = Form(""), user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    from .services import clinic as cl
    data = cl.get(db)
    loc = cl.location(data, int(lid)) if lid.isdigit() else None
    if not loc:
        nid = max([int(x.get("id", 0)) for x in data["locations"]] + [0]) + 1
        loc = {"id": nid, "hours": {str(i): {"on": i < 5, "from": "10:00", "to": "20:00"} for i in range(7)}, "resources": [], "places": []}
        data["locations"].append(loc)
    loc.update(name=name.strip()[:120] or "Локація", address=address.strip()[:200], phone=phone.strip()[:40], active=bool(active))
    cl.save(db, data)
    audit(db, "settings.location", user=user, location=loc["id"], name=loc["name"])
    return _clinic_back()


@app.post("/settings/clinic/location/{lid:int}/hours")
async def settings_clinic_hours(lid: int, request: Request, user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    import re as _re
    from .services import clinic as cl
    f = await request.form()
    data = cl.get(db)
    loc = cl.location(data, lid)
    if loc:
        hrs = {}
        for i in range(7):
            fr, to = str(f.get(f"from_{i}") or "10:00"), str(f.get(f"to_{i}") or "20:00")
            ok = _re.fullmatch(r"\d{2}:\d{2}", fr) and _re.fullmatch(r"\d{2}:\d{2}", to) and fr < to
            hrs[str(i)] = {"on": bool(f.get(f"on_{i}")) and bool(ok), "from": fr if ok else "10:00", "to": to if ok else "20:00"}
        loc["hours"] = hrs
        cl.save(db, data)
        audit(db, "settings.hours", user=user, location=lid)
    return _clinic_back(f"hours{lid}")


@app.post("/settings/clinic/location/{lid:int}/{what}")
def settings_clinic_list(lid: int, what: str, action: str = Form(...), idx: str = Form(""), name: str = Form(""),
                         del_password: str = Form(""), user: User = Depends(require("users")), db: Session = Depends(db_dep)):
    """Ресурси (апарати) і робочі місця (кабінети): add · rename · toggle · up · down · delete (з паролем видалення)."""
    from .services import clinic as cl
    if what not in ("resources", "places"):
        raise HTTPException(404)
    data = cl.get(db)
    loc = cl.location(data, lid)
    if not loc:
        return _clinic_back()
    items = loc.setdefault(what, [])
    i = int(idx) if idx.isdigit() and int(idx) < len(items) else None
    opn = f"{what}{lid}"
    if action == "add" and name.strip():
        items.append({"name": name.strip()[:120], "active": True})
    elif i is not None and action == "rename" and name.strip():
        items[i]["name"] = name.strip()[:120]
    elif i is not None and action == "toggle":
        items[i]["active"] = not items[i].get("active", True)
    elif i is not None and action in ("up", "down"):
        j = i - 1 if action == "up" else i + 1
        if 0 <= j < len(items):
            items[i], items[j] = items[j], items[i]
    elif i is not None and action == "delete":
        err = _delpw(db, user, del_password, f"settings_{what}", "location", lid)
        if err:
            return RedirectResponse(_with_err(f"/settings/clinic?open={opn}", err), 303)
        items.pop(i)
    cl.save(db, data)
    audit(db, f"settings.{what}", user=user, location=lid, op=action, name=name or None)
    return _clinic_back(opn)


@app.get("/settings/keys", response_class=HTMLResponse)
def settings_keys(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Одне місце для ключів інтеграцій: monobank (по ФОП), Checkbox, CleverBox. Самі ключі ніколи не показуються."""
    from .integrations import monobank, checkbox as chb
    from .models import BankTxn
    fops = [a["name"] for a in chb.accounts(db)]
    mono = {m["fop"]: m for m in monobank.public(db)}
    for n in mono:
        if n not in fops:
            fops.append(n)
    cnt = {f: db.scalar(select(func.count()).select_from(BankTxn).where(BankTxn.fop == f)) or 0 for f in fops}
    jid = request.query_params.get("job")
    return render(request, "keys.html", nav="keys", fops=fops, mono=mono, cnt=cnt,
                  cb=[{"name": a["name"], "configured": a["configured"]} for a in chb.accounts(db)], cb_status=chb.statuses(db),
                  cleverbox=bool(settings.cleverbox_token), job=jobs.JOBS.get(jid) if jid else None,
                  ai_key=bool(__import__("app.services.invoice_ai", fromlist=["x"]).api_key(db)),
                  shop=__import__("app.integrations.shop_express", fromlist=["x"]).creds(db),
                  shop_rep=__import__("app.integrations.shop_express", fromlist=["x"]).last_report(db))


def _shop_job(job):
    from .integrations import shop_express as sx
    with SessionLocal() as db:
        job.log("Завантажую каталог molodist.store…")
        rep = sx.sync(db)
        return {k: rep[k] for k in ("rows", "matched", "urls", "photos")}


@app.post("/settings/keys/shop")
def settings_keys_shop(site: str = Form(""), login: str = Form(""), password: str = Form(""), action: str = Form("save"),
                       user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Shop-Express: API-користувач магазину → перевірка входу → синхронізація посилань і фото."""
    from .integrations import shop_express as sx
    if action == "clear":
        sx.clear(db)
        audit(db, "keys.shop", user=user, op="clear")
        return RedirectResponse("/settings/keys#shop", 303)
    if action == "save":
        sx.save_creds(db, site, login, password)
        c = sx.creds(db)
        try:
            sx.ShopExpress(c["site"], c["login"], c["password"])
        except Exception as e:  # noqa: BLE001
            audit(db, "keys.shop", user=user, op="save", ok=False)
            return RedirectResponse(f"/settings/keys?shoperr={_q(str(e)[:200])}#shop", 303)
        audit(db, "keys.shop", user=user, op="save", ok=True)
    job = jobs.start("Синхронізація molodist.store", _shop_job)
    return RedirectResponse(f"/settings/keys?job={job.id}#shop", 303)


@app.post("/settings/keys/ai")
def settings_keys_ai(key: str = Form(""), action: str = Form("save"), user: User = Depends(require("finance")),
                     db: Session = Depends(db_dep)):
    from .services import invoice_ai
    invoice_ai.save_key(db, "" if action == "clear" else key)
    audit(db, "keys.ai", user=user, op=action)
    return RedirectResponse("/settings/keys?saved=1#ai", 303)


@app.post("/stock/mono/connect")
def stock_mono_connect(fop: str = Form(...), token: str = Form(""), back: str = Form(""),
                       user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .integrations import monobank
    err = monobank.connect(db, fop.strip()[:100], token)
    audit(db, "mono.connect", user=user, fop=fop, ok=err is None)            # токен у журнал не пишемо
    if back == "keys":
        return RedirectResponse(f"/settings/keys?mono={'err' if err else 'ok'}&fop={_q(fop)}#mono", 303)
    return RedirectResponse(f"/stock?tab=suppliers&mono={'err' if err else 'ok'}#mono", 303)


@app.post("/stock/mono/watch")
async def stock_mono_watch(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .integrations import monobank
    f = await request.form()
    monobank.set_watch(db, f.get("fop", ""), f.getlist("acc"))
    return RedirectResponse("/settings/keys?saved=1#mono" if f.get("back") == "keys" else "/stock?tab=suppliers#mono", 303)


@app.post("/stock/mono/disconnect")
def stock_mono_disconnect(fop: str = Form(...), back: str = Form(""), user: User = Depends(require("finance")),
                          db: Session = Depends(db_dep)):
    from .integrations import monobank
    monobank.disconnect(db, fop)
    audit(db, "mono.disconnect", user=user, fop=fop)
    return RedirectResponse("/settings/keys#mono" if back == "keys" else "/stock?tab=suppliers#mono", 303)


@app.post("/stock/mono/sync")
def stock_mono_sync(days: int = Form(2), back: str = Form(""), user: User = Depends(require("finance")),
                    db: Session = Depends(db_dep)):
    job = jobs.start("Виписка monobank", _mono_job, max(1, min(days, 31)))
    if back == "keys":
        return RedirectResponse(f"/settings/keys?job={job.id}#mono", 303)
    return RedirectResponse(f"/stock?tab=suppliers&job={job.id}#mono", 303)


@app.post("/stock/mono/recheck")
def stock_mono_recheck(user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Зняти автоматичні рознесення, де постачальника впізнано лише за назвою (могли промахнутись)."""
    from .models import PaymentAlloc
    from .services import payments as pay
    n_ = 0
    for w in pay.weak_auto(db):
        docs_ = {a.doc_id for a in db.scalars(select(PaymentAlloc).where(PaymentAlloc.txn_id == w.id))}
        pay.unallocate_txn(db, w)
        w.supplier_id = None
        db.commit()
        for d_ in docs_:
            audit(db, "stock.pay_unassign", "stock_doc", d_, user=user, amount=round(-w.amount, 2), fop=w.fop, txn=w.id,
                  why="упізнано лише за назвою")
        n_ += 1
    pay.rematch_all(db)                               # частково рознесені доносять залишок
    audit(db, "mono.recheck_weak", user=user, n=n_)
    return RedirectResponse(f"/stock?tab=suppliers&msg={_q(f'Знято сумнівних автоматичних рознесень: {n_}. Вони чекають ручного рішення нижче.')}#mono", 303)


@app.post("/stock/mono/txn/{tid}")
def stock_mono_txn(tid: str, action: str = Form(...), group: str = Form(""), supplier_id: str = Form(""), back: str = Form(""),
                   user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Ручне рознесення платежу: assign (на накладну) · unassign · ignore · auto (спробувати ще раз)."""
    from .models import BankTxn, Supplier
    from .services import payments as pay
    t = db.get(BankTxn, tid) or (_ for _ in ()).throw(HTTPException(404))
    if supplier_id.isdigit():
        t.supplier_id = int(supplier_id)
    if action == "assign" and group:
        g = pay.find_group(db, group)
        if g:
            done = pay.allocate(db, g, -t.amount, txn=t, fop=t.fop, user_id=user.id)
            sup_ = next((x for x in db.scalars(select(Supplier)) if x.name.lower() == (g.supplier or "").lower()), None)
            if sup_:
                t.supplier_id = sup_.id
                pay.learn_entity(db, t, sup_)                    # наступного разу — автоматично за IBAN
            t.status, t.note = ("matched" if done >= -t.amount - pay.TOL else "partial"), f"→ {g.label}"
            for d_ in g.docs:
                audit(db, "stock.pay_mono", "stock_doc", d_.id, user=user, amount=round(done, 2), fop=t.fop, txn=tid)
    elif action == "unassign":
        from .models import PaymentAlloc
        docs_ = {a.doc_id for a in db.scalars(select(PaymentAlloc).where(PaymentAlloc.txn_id == t.id))}
        pay.unallocate_txn(db, t)
        for d_ in docs_:
            audit(db, "stock.pay_unassign", "stock_doc", d_, user=user, amount=round(-t.amount, 2), fop=t.fop, txn=tid)
        t.status = "ignored_once"                     # цей платіж не розносити знову одразу на те саме
        db.commit()
        pay.rematch_all(db)                           # інші платежі (напр. частково рознесені) доносять залишок
        t.status = "new"
    elif action == "ignore":
        t.status = "ignored"
    elif action == "auto":
        pay.match_txn(db, t)
    db.commit()
    audit(db, "mono.txn", user=user, txn=tid, op=action)
    if back.startswith("/stock/docs/"):
        return RedirectResponse(back, 303)
    return RedirectResponse("/stock?tab=suppliers#mono", 303)


@app.post("/stock/suppliers/import")
async def stock_invoices_import(file: UploadFile = File(...), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Excel «Склад → Поставки» з CleverBox → оборот і борг по постачальниках."""
    import tempfile
    from .services import stock as st
    suffix = ".xls" if (file.filename or "").lower().endswith(".xls") else ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
    try:
        kind = await asyncio.to_thread(catalog.detect_export, tmp.name)
        if kind == "products":                       # переплутали форму — товари теж приймаємо
            rep = await asyncio.to_thread(catalog.import_products, db, tmp.name)
            rep = {"error": f"Це був файл «Товари» — імпортовано в Продукти ({rep.get('rows', 0)} товарів)."}
        elif kind != "supply":
            raise ValueError(catalog.WRONG_FILE[kind])
        else:
            rep = await asyncio.to_thread(st.import_supplier_invoices, db, tmp.name)
    except ValueError as e:
        rep = {"error": str(e)}
    except Exception as e:  # noqa: BLE001
        rep = {"error": f"Не вдалося прочитати файл: {type(e).__name__}"}
    finally:
        os.unlink(tmp.name)
    audit(db, "stock.invoices_import", user=user, **rep)
    LAST_IMPORT[f"inv-{user.id}"] = rep
    return RedirectResponse("/stock?tab=suppliers", 303)


@app.post("/stock/invoices/{iid:int}/paid")
def stock_invoice_paid(iid: int, back: str = Form("/stock?tab=suppliers"), note: str = Form(""),
                       user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Власник відмічає поставку оплаченою (якщо в CleverBox оплату не провели)."""
    from .models import SupplierInvoice
    x = db.get(SupplierInvoice, iid)
    if x:
        x.paid = x.amount
        x.paid_note = (note.strip() or f"оплачено, відмітив {user.name} {date.today():%d.%m.%Y}")[:300]
        db.commit()
        audit(db, "stock.invoice_paid", user=user, invoice=x.cb_number, amount=x.amount)
    return RedirectResponse(back if back.startswith("/stock") else "/stock?tab=suppliers", 303)


@app.post("/stock/norms")
def stock_norm_add(service_id: int = Form(...), product: str = Form(...), qty: str = Form("1"), warehouse: str = Form(""),
                   techcard: str = Form(""), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Норма на послугу; якщо передано техкартку — на всі послуги цієї техкартки (без дублів)."""
    from .models import ProcedureNorm, CbService
    pr = _pick_product(db, product)
    if not pr:
        return RedirectResponse(f"/stock?tab=norms&msg=Товар+не+знайдено#s{service_id}", 303)
    ids = [service_id]
    if techcard:
        ids = [x for (x,) in db.execute(select(CbService.id).where(CbService.techcard == techcard))] or ids
    have = {sid for (sid,) in db.execute(select(ProcedureNorm.service_id).where(
        ProcedureNorm.product_id == pr.id, ProcedureNorm.service_id.in_(ids)))}
    for sid in ids:
        if sid not in have:
            db.add(ProcedureNorm(service_id=sid, product_id=pr.id, qty=_num(qty) or 1,
                                 warehouse=warehouse if warehouse in catalog.WAREHOUSES else None))
    db.commit()
    audit(db, "stock.norm_add", user=user, service=service_id, techcard=techcard or None, services=len(ids), product=pr.id)
    return RedirectResponse(f"/stock?tab=norms#s{service_id}", 303)


@app.post("/stock/norms/{nid:int}/delete")
def stock_norm_delete(nid: int, techcard: str = Form(""), del_password: str = Form(""), user: User = Depends(require("stock")),
                      db: Session = Depends(db_dep)):
    from .models import ProcedureNorm, CbService
    n = db.get(ProcedureNorm, nid)
    sid = n.service_id if n else ""
    err = _delpw(db, user, del_password, "norm", "norm", nid)
    if err:
        return RedirectResponse(_with_err(f"/stock?tab=norms#s{sid}", err), 303)
    if n:
        if techcard:      # прибрати матеріал з усієї техкартки
            ids = [x for (x,) in db.execute(select(CbService.id).where(CbService.techcard == techcard))]
            for m in db.scalars(select(ProcedureNorm).where(ProcedureNorm.product_id == n.product_id,
                                                            ProcedureNorm.service_id.in_(ids or [sid]))):
                db.delete(m)
        else:
            db.delete(n)
        db.commit()
    return RedirectResponse(f"/stock?tab=norms#s{sid}", 303)


@app.get("/finance", response_class=HTMLResponse)
def finance_page(request: Request, df: str | None = None, dt_: str | None = None,
                 user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    today = date.today()
    d_from = date.fromisoformat(df) if df else today.replace(day=1)
    d_to = date.fromisoformat(dt_) if dt_ else today
    fs = fin_svc.summary(db, d_from, d_to)
    fc = fin_svc.plan_forecast(db)
    sts = all_statuses(db, d_from, d_to)
    total, _ = kpi_summary(sts)
    m_max = max([m["revenue"] for m in fs.months] + [r["total"] for r in fc] + [1])
    return render(request, "finance.html", fs=fs, fc=fc, total=total, d_from=d_from, d_to=d_to, m_max=m_max)


# ---------------- Сертифікати й абонементи ----------------
@app.get("/gifts", response_class=HTMLResponse)
def gifts_page(request: Request, tab: str = "certs", st: str = "", q: str = "", seller: str = "",
               user: User = Depends(require("patients_all")), db: Session = Depends(db_dep)):
    from .models import GiftCertificate, Subscription
    from .services import certificates as cs
    tab = tab if tab in ("certs", "subs") else "certs"
    today = date.today()
    model, states = (GiftCertificate, cs.STATES) if tab == "certs" else (Subscription, cs.SUB_STATES)
    items = db.scalars(select(model).order_by(model.sold_at.desc())).all()
    counts = defaultdict(lambda: {"n": 0, "sum": 0.0})
    for x in items:
        k = x.state(today)
        counts[k]["n"] += 1
        counts[k]["sum"] += (x.balance or 0) if tab == "certs" else x.value_left
    sellers = sorted({x.seller for x in items if x.seller})
    if st in states:
        items = [x for x in items if x.state(today) == st]
    elif not q and not seller:
        items = [x for x in items if x.state(today) in ("active", "expiring", "unpaid")] or items[:100]
    if seller:
        items = [x for x in items if x.seller == seller]
    if q:
        words = set(norm(q).split())
        items = [x for x in items if words <= set(norm(" ".join(filter(None, (
            x.client_name, x.code, str(x.id), x.type_name, x.seller)))).split())]
    if st == "expiring" or not st:
        items.sort(key=lambda x: (x.state(today) not in ("expiring",), x.valid_to or date.max))
    other = Subscription if tab == "certs" else GiftCertificate
    return render(request, "gifts.html", nav="gifts", tab=tab, items=items[:500], total=len(items), states=states,
                  counts=dict(counts), st=st, q=q, seller=seller, sellers=sellers, today=today,
                  other_n=db.scalar(select(func.count()).select_from(other)) or 0, last=LAST_IMPORT.pop(f"gifts-{user.id}", None))


@app.post("/gifts/import")
async def gifts_import(file: UploadFile = File(...), user: User = Depends(require("import")), db: Session = Depends(db_dep)):
    """Excel CleverBox «Сертифікати» або «Абонементи» — тип визначається за колонками."""
    from .services import certificates as cs
    tab = "certs"
    try:
        sheets = importer.read_sheets(await file.read(), file.filename or "data.xls")
        head, rows = (sheets[0][1][0], sheets[0][1][1:]) if sheets and sheets[0][1] else ([], [])
        if cs.is_certificates_export(head):
            rep = await asyncio.to_thread(cs.import_certificates, db, rows, head)
        elif cs.is_subscriptions_export(head):
            rep, tab = await asyncio.to_thread(cs.import_subscriptions, db, rows, head), "subs"
        else:
            rep = {"error": "Це не вивантаження CleverBox «Сертифікати» чи «Абонементи». Колонки: "
                            + ", ".join(str(h) for h in head[:12])}
    except Exception as e:  # noqa: BLE001
        rep = {"error": f"Не вдалося прочитати файл: {type(e).__name__}"}
    rep["sheet"] = file.filename
    audit(db, "gifts.import", user=user, **{k: v for k, v in rep.items() if k in ("type", "rows", "updated", "error", "sheet")})
    LAST_IMPORT[f"gifts-{user.id}"] = rep
    return RedirectResponse(f"/gifts?tab={tab}", 303)


# ---------------- Звіти ----------------
def _report_period(period: str, df: str | None, dt_: str | None) -> tuple[date, date, str]:
    today = date.today()
    m0 = today.replace(day=1)
    if df or dt_:
        a = _day(df) or m0
        b = _day(dt_) or today
        return (a, b, "custom") if a <= b else (b, a, "custom")
    if period == "prev_month":
        b = m0 - timedelta(days=1)
        return b.replace(day=1), b, period
    if period == "quarter":
        q0 = date(today.year, 3 * ((today.month - 1) // 3) + 1, 1)
        return q0, today, period
    if period == "year":
        return date(today.year, 1, 1), today, period
    if period == "30d":
        return today - timedelta(days=29), today, period
    if period == "12m":
        y, m = (today.year - 1, today.month + 1) if today.month < 12 else (today.year, 1)
        return date(y, m, 1), today, period
    return m0, today, "month"


def _expiry_page(request: Request, db: Session):
    from .services import expiry as ex, reports as rp
    wh = request.query_params.get("wh", "")
    rep = ex.report(db)
    cfg = ex.config(db)
    users = db.scalars(select(User).where(User.active.is_(True)).order_by(User.role, User.name)).all()
    return render(request, "expiry.html", nav="reports", REPORTS=rp.REPORTS, key="expiry", rep=rep, wh=wh, cfg=cfg,
                  users=users, LABELS=ex.LABELS, WEEKDAYS=ex.WEEKDAYS, admin_chat=bool(settings.telegram_admin_chat_id),
                  bot=bool(settings.telegram_bot_token), recips=ex.recipients(db, cfg), ROLES=ROLES)


@app.post("/reports/expiry/writeoff")
def expiry_writeoff(wh: str = Form(...), user: User = Depends(require("stock")), db: Session = Depends(db_dep)):
    """Чернетка списання всіх прострочених партій складу — перевірити й провести."""
    from .services import stock as st, expiry as ex
    from .models import StockDocLine
    rows = [x for x in ex.rows(db) if x["wh"] == wh and x["days"] is not None and x["days"] < 0]
    if not rows or not st.ledger_active(db):
        return RedirectResponse(f"/reports?r=expiry&wh={wh}", 303)
    reason = next((r for r in st.WRITEOFF_REASONS if "простроч" in r.lower()), st.WRITEOFF_REASONS[0])
    doc = st.new_doc(db, "writeoff", user_id=user.id, wh_from=wh, reason=reason, note="Прострочені партії (зі звіту термінів)")
    for x in rows:
        doc.lines.append(StockDocLine(product_id=x["product"].id, qty=x["qty"], expiry_date=x["expiry"]))
    db.commit()
    audit(db, "expiry.writeoff_draft", "stock_doc", doc.id, user=user, wh=wh, lines=len(rows))
    return RedirectResponse(f"/stock/docs/{doc.id}?msg=Чернетка+списання+прострочених+партій+—+перевірте+і+проведіть", 303)


@app.post("/reports/expiry/notify")
async def expiry_notify(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import expiry as ex
    f = await request.form()
    cfg = {"users": [int(x) for x in f.getlist("users") if str(x).isdigit()], "admin_chat": bool(f.get("admin_chat")),
           "weekday": int(f.get("weekday") or 0), "days": int(f.get("days") or 90), "alerts": bool(f.get("alerts"))}
    ex.save_config(db, cfg)
    audit(db, "expiry.notify", user=user, **cfg)
    if f.get("test"):
        n = ex.send(db, ex.summary_text(db, cfg["days"]) or "🗓 Терміни придатності: прострочених і близьких партій немає ✅", cfg)
        return RedirectResponse(f"/reports?r=expiry&sent={n}#notify", 303)
    return RedirectResponse("/reports?r=expiry&saved=1#notify", 303)


def _fop_page(request: Request, db: Session):
    from .services import fop as fop_svc
    from .services import reports as rp
    from .integrations.checkbox import accounts
    from .models import FopMonth
    years = sorted({y for (y,) in db.execute(select(FopMonth.year).distinct())} | {date.today().year}, reverse=True)
    try:
        year = int(request.query_params.get("year") or date.today().year)
    except ValueError:
        year = date.today().year
    st_ = db.get(SyncState, "checkbox_last")
    jid = request.query_params.get("job")
    f_ = fop_svc.build(db, year)
    qp = request.query_params
    for i, x in enumerate(f_["fops"]):               # ручний період у картці кожного ФОП: ?p0_from=2025-01&p0_to=2026-03
        a = fop_svc.parse_ym(qp.get(f"p{i}_from")) or (year, 1)
        b = fop_svc.parse_ym(qp.get(f"p{i}_to")) or (year, 12)
        x["period"] = fop_svc.period(db, x["name"], a, b)
        x["period"]["custom"] = bool(qp.get(f"p{i}_from") or qp.get(f"p{i}_to"))
    return render(request, "fop.html", nav="reports", REPORTS=rp.REPORTS, key="fop", f=f_,
                  years=years, year=year, cb_status=__import__("app.integrations.checkbox", fromlist=["x"]).statuses(db),
                  accs=[{"name": a["name"], "configured": a["configured"], "login": a["login"], "source": a["source"]}
                        for a in accounts(db)],
                  cb_last=st_.value if st_ else None, job=jobs.JOBS.get(jid) if jid else None,
                  last=LAST_IMPORT.pop(f"fop-{request.state.user.id}", None))


def _fop_job(job, start: str | None = None):
    from .integrations import checkbox as chb
    with SessionLocal() as db:
        s0 = date.fromisoformat(start) if start else None
        job.log("Завантажую Z-звіти з Checkbox…")
        return chb.sync(db, s0, date.today())


@app.post("/reports/fop/sync")
def fop_sync(start: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    job = jobs.start("Оборот ФОП з Checkbox", _fop_job, start or None)
    audit(db, "fop.sync", user=user, start=start or None, job=job.id)
    return RedirectResponse(f"/reports?r=fop&job={job.id}", 303)


@app.post("/reports/fop/connect")
def fop_connect(name: str = Form(...), login: str = Form(...), password: str = Form(""),
                user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Підключити кабінет Checkbox ФОП: логін/пароль касира (пароль — зашифровано в базі)."""
    from .integrations import checkbox as chb
    if not name.strip() or not login.strip():
        return RedirectResponse("/reports?r=fop&cmsg=Вкажіть+назву+ФОП+і+логін+касира#connect", 303)
    err = chb.connect(db, name, login, password)
    audit(db, "fop.connect", user=user, fop=name.strip(), ok=err is None)       # без пароля
    if err:
        return RedirectResponse(f"/reports?r=fop&cerr={name.strip()}#connect", 303)
    job = jobs.start("Оборот ФОП з Checkbox", _fop_job, f"{date.today().year}-01-01")
    return RedirectResponse(f"/reports?r=fop&job={job.id}&cok={name.strip()}#connect", 303)


@app.post("/reports/fop/disconnect")
def fop_disconnect(name: str = Form(...), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .integrations import checkbox as chb
    chb.disconnect(db, name)
    audit(db, "fop.disconnect", user=user, fop=name)
    return RedirectResponse("/reports?r=fop#connect", 303)


@app.post("/reports/fop/import")
async def fop_import(file: UploadFile = File(...), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Google-таблиця «Ліміти ФОП» (Файл → Завантажити → .xlsx) — аркуші «<ФОП> <рік>»."""
    from .services import fop as fop_svc
    try:
        sheets = importer.read_sheets(await file.read(), file.filename or "fop.xlsx")
        if not fop_svc.is_fop_workbook(sheets):
            raise ValueError("Не знайдено аркушів «<ФОП> <рік>» з рядками «Готівка» / «Безготівка»")
        rep = fop_svc.import_workbook(db, sheets)
    except ValueError as e:
        rep = {"error": str(e)}
    except Exception as e:  # noqa: BLE001
        rep = {"error": f"Не вдалося прочитати файл: {type(e).__name__}"}
    audit(db, "fop.import", user=user, rows=rep.get("rows"), error=rep.get("error"))
    LAST_IMPORT[f"fop-{user.id}"] = rep
    return RedirectResponse("/reports?r=fop", 303)


@app.post("/reports/fop/limit")
def fop_limit(year: int = Form(...), amount: str = Form(...), user: User = Depends(require("finance")),
              db: Session = Depends(db_dep)):
    from .services import fop as fop_svc
    v = _num(amount)
    if v:
        fop_svc.set_limit(db, year, v)
        audit(db, "fop.limit", user=user, year=year, amount=v)
    return RedirectResponse(f"/reports?r=fop&year={year}", 303)


def _cashflow_page(request: Request, db: Session, period: str, df: str | None):
    from .services import cashflow as cf
    from .services import reports as rp
    years = sorted({d.year for (d,) in db.execute(select(CashDay.day))} | {date.today().year})
    try:
        year = int(request.query_params.get("year") or date.today().year)
    except ValueError:
        year = date.today().year
    import json as _json
    probe_raw = db.get(SyncState, "cleverbox_finance_probe")
    try:
        probe = _json.loads(probe_raw.value) if probe_raw and probe_raw.value else None
    except ValueError:
        probe = None
    meth = db.get(SyncState, "cleverbox_finance_method")
    last = db.get(SyncState, "cleverbox_finance_last")
    jid = request.query_params.get("job")
    return render(request, "cashflow.html", nav="reports", REPORTS=rp.REPORTS, key="cashflow", cf=cf.build(db, year),
                  probe=probe, cb_method=meth.value if meth else None, cb_last=last.value if last else None,
                  job=jobs.JOBS.get(jid) if jid else None,
                  years=years, year=year, last=LAST_IMPORT.pop(f"cash-{request.state.user.id}", None))


@app.post("/reports/cashflow/import")
async def cashflow_import(file: UploadFile = File(...), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Excel CleverBox «Прибуток» (Profit_*.xls) за будь-який період — дні перезаписуються."""
    from .services import cashflow as cf
    try:
        sheets = importer.read_sheets(await file.read(), file.filename or "profit.xls")
        head, rows = (sheets[0][1][0], sheets[0][1][1:]) if sheets and sheets[0][1] else ([], [])
        if not cf.is_profit_export(head):
            raise ValueError("Це не звіт CleverBox «Прибуток» (потрібні колонки Дата / Надходження / Витрата)")
        rep = await asyncio.to_thread(cf.import_profit, db, rows, head)
        cf.auto_fill(db)
    except ValueError as e:
        rep = {"error": str(e)}
    except Exception as e:  # noqa: BLE001
        rep = {"error": f"Не вдалося прочитати файл: {type(e).__name__}"}
    audit(db, "cashflow.import", user=user, **{k: v for k, v in rep.items() if k in ("rows", "updated", "period", "error")})
    LAST_IMPORT[f"cash-{user.id}"] = rep
    return RedirectResponse("/reports?r=cashflow", 303)


def _cf_job(job, kind: str, start: str | None = None):
    from .integrations import cb_finance
    with SessionLocal() as db:
        if kind == "probe":
            job.log("Перевіряю фінансові методи API CleverBox…")
            return cb_finance.probe(db)
        s0 = date.fromisoformat(start) if start else date(date.today().year, 1, 1)
        job.log(f"Завантажую рух коштів з {s0:%d.%m.%Y}…")
        rep = cb_finance.sync_cash(db, s0, date.today())
        from .services import cashflow as cf
        cf.auto_fill(db)
        return rep


@app.post("/reports/cashflow/cleverbox")
def cashflow_cleverbox(action: str = Form("probe"), start: str = Form(""), user: User = Depends(require("finance")),
                       db: Session = Depends(db_dep)):
    """Кнопки на сторінці Cashflow: перевірити API CleverBox / дозавантажити історію."""
    if action == "reset":
        for k in ("cleverbox_finance_method", "cleverbox_finance_probe_at"):
            st_ = db.get(SyncState, k)
            if st_:
                db.delete(st_)
        db.commit()
        action = "probe"
    job = jobs.start("Cashflow з CleverBox" if action != "probe" else "Пошук фінансових методів CleverBox",
                     _cf_job, action, start or None)
    audit(db, "cashflow.cleverbox", user=user, op=action, job=job.id)
    return RedirectResponse(f"/reports?r=cashflow&job={job.id}", 303)


@app.post("/reports/cashflow/opening")
def cashflow_opening(d: str = Form(...), amount: str = Form(...), user: User = Depends(require("finance")),
                     db: Session = Depends(db_dep)):
    from .services import cashflow as cf
    day, val = _day(d), _num(amount)
    if day and val is not None:
        cf.set_opening(db, day, val)
        audit(db, "cashflow.opening", user=user, date=str(day), amount=val)
    return RedirectResponse("/reports?r=cashflow", 303)


REPORT_PERIODS = {"month": "Цей місяць", "prev_month": "Минулий місяць", "30d": "30 днів", "quarter": "Квартал",
                  "year": "З початку року", "12m": "12 місяців"}


@app.get("/reports", response_class=HTMLResponse)
def reports_page(request: Request, r: str = "revenue", period: str = "month", df: str | None = None, dt_: str | None = None,
                 fmt: str = "", user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import reports as rp
    key = r if r in rp.REPORTS else "revenue"
    if key == "cashflow":
        return _cashflow_page(request, db, period, df)
    if key == "fop":
        return _fop_page(request, db)
    if key == "expiry":
        return _expiry_page(request, db)
    d_from, d_to, period = _report_period(period, df, dt_)
    rep = rp.build(db, key, d_from, d_to)
    if fmt == "xlsx":
        audit(db, "reports.export", user=user, report=key, d_from=str(d_from), d_to=str(d_to))
        return Response(rp.to_xlsx(rep, d_from, d_to),
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition": f'attachment; filename="report-{key}-{d_from}-{d_to}.xlsx"'})
    p_from, p_to = rp.prev_period(d_from, d_to)
    return render(request, "reports.html", rep=rep, REPORTS=rp.REPORTS, PERIODS=REPORT_PERIODS, key=key, period=period,
                  d_from=d_from, d_to=d_to, p_from=p_from, p_to=p_to,
                  c_max=max([v for _, v in rep.chart] + [1]))


# ---------------- Фінанси → Планування ціни реалізації ----------------
def _pricing_rows(db: Session, wh: str, brand: str, q: str, only: str):
    from .services import pricing as pr_svc
    rt = pr_svc.rates(db)
    rows = [r for r in catalog.product_rows(db, []) if r.product.active is not False]
    if wh in catalog.WAREHOUSES:
        rows = [r for r in rows if (r.product.warehouse or "vitrina") == wh]
    if brand:
        rows = [r for r in rows if (r.product.brand or "") == ("" if brand == "—" else brand)]
    if q:
        words = set(norm(q).split())
        rows = [r for r in rows if words <= set(norm(" ".join(filter(None, [r.product.name, r.product.sku, r.product.brand]))).split())]
    out = []
    for r in rows:
        p_ = r.product
        cur = p_.cost_currency or "UAH"
        pr = pr_svc.PriceRow(p_, r.cost_last or r.cost_avg, rt.get(cur, 1.0))   # для ціноутворення — остання закупівля
        out.append(pr)
    if only == "nofx":
        out = [x for x in out if not x.product.cost_fx]
    elif only == "noplan":
        out = [x for x in out if not x.product.planned_price]
    elif only == "diff":
        out = [x for x in out if x.product.planned_price and x.price_now and abs(x.product.planned_price - x.price_now) >= 1]
    elif only == "lowmargin":
        out = [x for x in out if (pr_svc.margin(x.price_now, x.cost_uah) or 100) < 40]
    out.sort(key=lambda x: ((x.product.brand or "яяя").lower(), x.product.name))
    return out, rt


# ---------------- Фінанси: розділи як у CleverBox ----------------
def _period(df: str, dt_: str, default_days: int | None = None) -> tuple[date, date]:
    today = date.today()
    d_to = _day(dt_) or today
    d_from = _day(df) or (d_to - timedelta(days=default_days) if default_days is not None else today.replace(day=1))
    return (d_from, d_to) if d_from <= d_to else (d_to, d_from)


def _fa():
    from .services import fin_articles
    return fin_articles


def _cashboxes(db: Session) -> list[str]:
    from .models import CbFinOp
    from .services import cashboxes as _cbx
    base = _cbx.names(db)                                             # «Налаштування → Каси / рахунки»
    cb = [a for (a,) in db.execute(select(CbFinOp.account).where(CbFinOp.account.is_not(None)).distinct()) if a]
    return base + sorted(set(cb) - set(base)) + ["Інше"]               # + рахунки з журналу CleverBox


def _month(m: str) -> str:
    import re as _re
    return m if _re.fullmatch(r"\d{4}-\d{2}", m or "") else fo_month_key(date.today())


def fo_month_key(d: date) -> str:
    return f"{d.year}-{d.month:02d}"


def _journal(db: Session, request: Request, d_from: date, d_to: date):
    from .services import finops as fo
    qp = request.query_params
    users = {u_.id: u_.name for u_ in db.scalars(select(User))}
    rows = fo.journal(db, d_from, d_to, users, archive=bool(qp.get("archive")))
    opts = {"accounts": sorted({r.account for r in rows if r.account}), "articles": sorted({r.article for r in rows if r.article}),
            "authors": sorted({r.author for r in rows if r.author})}
    f = {k: qp.get(k, "") for k in ("src", "kind", "account", "article", "author", "q")}
    if f["src"]:
        rows = [r for r in rows if r.src == f["src"]]
    if f["kind"]:
        rows = [r for r in rows if (r.is_storno if f["kind"] == "storno" else (r.kind == f["kind"] and not r.is_storno))]
    for k in ("account", "article", "author"):
        if f[k]:
            rows = [r for r in rows if getattr(r, k) == f[k]]
    if f["q"].strip():
        ql = f["q"].strip().lower().lstrip("#")
        rows = [r for r in rows if ql in f"{r.num} {r.who} {r.basis} {r.comment} {r.article}".lower()]
    return rows, opts, f


@app.get("/finance/ops", response_class=HTMLResponse)
def finance_ops(request: Request, df: str = "", dt_: str = "", user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import finops as fo
    from .models import CbFinOp
    d_from, d_to = _period(df, dt_)
    rows, opts, f = _journal(db, request, d_from, d_to)
    last_cb = db.scalar(select(func.max(CbFinOp.updated_at)))
    from .services import clinic as _cl
    return render(request, "fin_ops.html", nav="finance", rows=rows[:2000], n_all=len(rows),
                  tot=fo.journal_totals(rows, _cl.excluded_income(db)),
                  d_from=d_from, d_to=d_to, f=f, opts=opts, archive=bool(request.query_params.get("archive")),
                  EXP=_fa().names(db, "expense"), INC=_fa().names(db, "income"), SAL=_fa().salary_names(db),
                  KL=fo.KIND_LABELS, boxes=_cashboxes(db), last_cb=last_cb,
                  cb_imported=LAST_IMPORT.pop(f"cbfin-{user.id}", None),
                  job=jobs.JOBS.get(request.query_params.get("job") or ""),
                  doctors=db.scalars(select(Doctor).where(Doctor.active.is_(True)).order_by(Doctor.name)).all(),
                  today=date.today(), this_month=fo_month_key(date.today()))


@app.get("/finance/ops/export.xlsx")
def finance_ops_export(request: Request, df: str = "", dt_: str = "", user: User = Depends(require("finance")),
                       db: Session = Depends(db_dep)):
    from openpyxl import Workbook
    from .services import finops as fo
    import io
    d_from, d_to = _period(df, dt_)
    rows, _, _ = _journal(db, request, d_from, d_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "Фіноперації"
    ws.append(["Дата", "Тип", "№", "Джерело", "Від кого / кому", "Підстава", "Стаття", "Рахунок", "Вид / на рахунок", "Автор", "Сума", "Коментар"])
    for r in rows:
        ws.append([r.at, "Сторно" if r.is_storno else fo.KIND_LABELS.get(r.kind, r.kind), r.num, "CRM" if r.src == "crm" else "CleverBox",
                   r.who, r.basis, r.article, r.account, r.vid, r.author, round(r.sign * r.amount if r.sign else r.amount, 2), r.comment])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(ws.max_row, 11).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJKL", (17, 12, 10, 11, 28, 22, 26, 22, 22, 20, 13, 30)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="finops_{d_from}_{d_to}.xlsx"'})


@app.post("/finance/ops/import")
async def finance_ops_import(file: UploadFile = File(...), back: str = Form("/finance/ops"), user: User = Depends(require("finance")),
                             db: Session = Depends(db_dep)):
    """Великий файл (десятки тисяч операцій) — у фоні, сторінка показує хід."""
    import tempfile
    from .services import cb_finops_import as cfi
    suffix = ".xls" if (file.filename or "").lower().endswith(".xls") else ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(await file.read())
    if not await asyncio.to_thread(cfi.is_journal, tmp.name):
        LAST_IMPORT[f"cbfin-{user.id}"] = {"error": "Це не журнал фіноперацій CleverBox: потрібні колонки «Стаття» і «Сума»."}
        return RedirectResponse(back if back.startswith("/finance") else "/finance/ops", 303)
    uid, fname = user.id, file.filename

    def _run(job, path):
        with SessionLocal() as s_:
            rep = cfi.import_file(s_, path, job)
            u_ = s_.get(User, uid)
            audit(s_, "fin.cb_import", user=u_, file=fname, **rep)
            return rep
    job = jobs.start("Імпорт журналу фіноперацій CleverBox", _run, tmp.name)
    from urllib.parse import urlsplit, parse_qsl, urlencode
    u = urlsplit(back if back.startswith("/finance") else "/finance/ops")
    qs = [(k, v) for k, v in parse_qsl(u.query) if k not in ("msg", "err", "job")] + [("job", job.id)]
    return RedirectResponse(f"{u.path}?{urlencode(qs)}", 303)


@app.post("/finance/ops")
def finance_op_add(day: str = Form(""), kind: str = Form("expense"), category: str = Form(""), amount: str = Form(...),
                   cashbox: str = Form(""), cashbox_to: str = Form(""), counterparty: str = Form(""), doctor_id: str = Form(""),
                   salary_month: str = Form(""), comment: str = Form(""),
                   back: str = Form("/finance/ops"), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .models import FinOp
    amt = _num(amount)
    back = back if back.startswith("/finance") else "/finance/ops"
    if not amt or amt <= 0:
        return RedirectResponse(_with_msg(back, "Вкажіть суму більше 0", err=True), 303)
    kind = kind if kind in ("income", "expense", "transfer") else "expense"
    if kind == "transfer":
        if not cashbox or not cashbox_to or cashbox == cashbox_to:
            return RedirectResponse(_with_msg(back, "Переміщення: оберіть два різні рахунки", err=True), 303)
        category = category or "Переміщення між рахунками"
    if not category.strip():
        return RedirectResponse(_with_msg(back, "Оберіть статтю", err=True), 303)
    x = FinOp(day=_day(day) or date.today(), kind=kind, category=category.strip()[:60],
              amount=round(amt, 2), cashbox=cashbox.strip()[:100] or None, cashbox_to=(cashbox_to.strip()[:100] or None) if kind == "transfer" else None,
              counterparty=counterparty.strip()[:200] or None, comment=comment.strip()[:500] or None,
              doctor_id=int(doctor_id) if doctor_id.isdigit() else None,
              salary_month=_month(salary_month) if salary_month else None, created_by=user.id)
    if x.category in _fa().salary_names(db) and x.doctor_id and not x.salary_month:
        x.salary_month = fo_month_key(x.day)
    db.add(x)
    db.commit()
    audit(db, "fin.op_add", "fin_op", x.id, user=user, kind=x.kind, category=x.category, amount=x.amount, cashbox=x.cashbox)
    return RedirectResponse(_with_msg(back, f"Записано CRM-{x.id}: {x.category} {x.amount:,.2f} ₴".replace(",", " ")), 303)


@app.post("/finance/ops/storno")
def finance_op_storno(num: str = Form(...), comment: str = Form(""), back: str = Form("/finance/ops"),
                      user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    """Сторно операції CRM: зворотна операція на ту саму суму й рахунок; оригінал позначається «сторновано»."""
    import re as _re
    from .models import FinOp
    back = back if back.startswith("/finance") else "/finance/ops"
    m = _re.search(r"(\d+)", num or "")
    x = db.get(FinOp, int(m[1])) if m else None
    if not x or x.deleted_at:
        return RedirectResponse(_with_msg(back, "Операцію CRM з таким № не знайдено (операції CleverBox сторнуються в CleverBox)", err=True), 303)
    if x.storno_of or db.scalar(select(FinOp.id).where(FinOp.storno_of == x.id, FinOp.deleted_at.is_(None))):
        return RedirectResponse(_with_msg(back, f"CRM-{x.id} уже сторновано або сама є сторно", err=True), 303)
    if len(comment.strip()) < 3:
        return RedirectResponse(_with_msg(back, "Вкажіть причину сторно", err=True), 303)
    kind = {"income": "expense", "expense": "income", "transfer": "transfer"}[x.kind]
    y = FinOp(day=date.today(), kind=kind, category=x.category, amount=x.amount,
              cashbox=x.cashbox_to if x.kind == "transfer" else x.cashbox, cashbox_to=x.cashbox if x.kind == "transfer" else None,
              counterparty=x.counterparty, doctor_id=x.doctor_id, salary_month=x.salary_month, storno_of=x.id,
              comment=comment.strip()[:500], created_by=user.id)
    db.add(y)
    db.commit()
    audit(db, "fin.op_storno", "fin_op", x.id, user=user, storno=y.id, amount=x.amount, comment=y.comment)
    return RedirectResponse(_with_msg(back, f"Сторно CRM-{x.id} записано як CRM-{y.id}"), 303)


@app.post("/finance/ops/{oid:int}/delete")
def finance_op_delete(oid: int, del_password: str = Form(""), back: str = Form("/finance/ops"),
                      user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .models import FinOp
    back = back if back.startswith("/finance") else "/finance/ops"
    err = _delpw(db, user, del_password, "fin_op", "fin_op", oid)
    if err:
        return RedirectResponse(_with_err(back, err), 303)
    x = db.get(FinOp, oid)
    if x and not x.deleted_at:
        x.deleted_at = datetime.utcnow()
        db.commit()
        audit(db, "fin.op_delete", "fin_op", oid, user=user, category=x.category, amount=x.amount)
    return RedirectResponse(_with_msg(back, "Операцію видалено"), 303)


def _with_msg(url: str, msg: str, err: bool = False) -> str:
    base, _, frag = url.partition("#")
    return f"{base}{'&' if '?' in base else '?'}{'err' if err else 'msg'}={_q(msg)}" + (f"#{frag}" if frag else "")


@app.get("/finance/day", response_class=HTMLResponse)
def finance_day(request: Request, d: str = "", user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import finops as fo
    day = _day(d) or date.today()
    cs = fo.all_checks(db, day, day)                         # чеки CleverBox + каса CRM
    prev = fo.sales_summary(fo.all_checks(db, day - timedelta(days=7), day - timedelta(days=7)))
    return render(request, "fin_day.html", nav="finance", day=day, s=fo.sales_summary(cs), rows=fo.sales_rows(db, cs),
                  prev=prev, timedelta=timedelta)


@app.get("/finance/sales", response_class=HTMLResponse)
def finance_sales(request: Request, df: str = "", dt_: str = "", q: str = "", staff: str = "",
                  user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import finops as fo
    d_from, d_to = _period(df, dt_)
    cs = fo.all_checks(db, d_from, d_to)
    rows = fo.sales_rows(db, cs, q, staff)
    staff_names = sorted({it.get("initiator_name") for c in cs for it in (c.items or []) if it.get("initiator_name")})
    return render(request, "fin_sales.html", nav="finance", d_from=d_from, d_to=d_to, q=q, staff=staff, staff_names=staff_names,
                  s=fo.sales_summary([r["check"] for r in rows]), rows=rows[:800], n=len(rows))


@app.get("/finance/payroll", response_class=HTMLResponse)
def finance_payroll(request: Request, m: str = "", user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import finops as fo
    month = _month(m)
    rows = fo.payroll(db, month)
    all_docs = db.scalars(select(Doctor).where(Doctor.active.is_(True)).order_by(Doctor.name)).all()
    return render(request, "fin_payroll.html", nav="finance", month=month, rows=rows, months=_month_list(),
                  total={k: round(sum(r["m"][k] for r in rows), 2) for k in ("services", "products", "accrued", "adj")},
                  no_rule=[d_ for d_ in all_docs if d_.id not in {r["doctor"].id for r in rows if r["rule"]}])


def _month_list(n: int = 24) -> list[str]:
    t = date.today().replace(day=1)
    out = []
    for _ in range(n):
        out.append(fo_month_key(t))
        t = (t - timedelta(days=1)).replace(day=1)
    return out


@app.post("/finance/payroll/rule")
def finance_payroll_rule(doctor_id: int = Form(...), base: str = Form(""), pct_services: str = Form(""), pct_products: str = Form(""),
                         note: str = Form(""), m: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .models import SalaryRule
    from .services import finops as fo
    if not db.get(SyncState, "payroll_start"):
        fo.set_payroll_start(db, _month(m))
    r = db.scalar(select(SalaryRule).where(SalaryRule.doctor_id == doctor_id)) or SalaryRule(doctor_id=doctor_id)
    r.base, r.pct_services, r.pct_products = _num(base), _num(pct_services), _num(pct_products)
    r.note, r.updated_at = note.strip()[:300] or None, datetime.utcnow()
    db.add(r)
    db.commit()
    audit(db, "fin.salary_rule", "doctor", doctor_id, user=user, base=r.base, pct_services=r.pct_services, pct_products=r.pct_products)
    return RedirectResponse(f"/finance/payroll?m={_month(m)}#d{doctor_id}", 303)


@app.post("/finance/payroll/adj")
def finance_payroll_adj(doctor_id: int = Form(...), month: str = Form(...), amount: str = Form(...), sign: str = Form("+"),
                        comment: str = Form(""), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .models import SalaryAdj
    amt = _num(amount)
    mk = _month(month)
    if amt:
        a = SalaryAdj(doctor_id=doctor_id, month=mk, amount=round(-abs(amt) if sign == "-" else abs(amt), 2),
                      comment=comment.strip()[:300] or None, created_by=user.id)
        db.add(a)
        db.commit()
        audit(db, "fin.salary_adj", "doctor", doctor_id, user=user, month=mk, amount=a.amount, comment=a.comment)
    return RedirectResponse(f"/finance/payroll?m={mk}#d{doctor_id}", 303)


@app.post("/finance/payroll/adj/{aid:int}/delete")
def finance_payroll_adj_delete(aid: int, del_password: str = Form(""), user: User = Depends(require("finance")),
                               db: Session = Depends(db_dep)):
    from .models import SalaryAdj
    a = db.get(SalaryAdj, aid)
    back = f"/finance/payroll?m={a.month}" if a else "/finance/payroll"
    err = _delpw(db, user, del_password, "salary_adj", "doctor", a.doctor_id if a else None)
    if err:
        return RedirectResponse(_with_err(back, err), 303)
    if a:
        audit(db, "fin.salary_adj_delete", "doctor", a.doctor_id, user=user, month=a.month, amount=a.amount)
        db.delete(a)
        db.commit()
    return RedirectResponse(back, 303)


@app.get("/finance/settle", response_class=HTMLResponse)
def finance_settle(request: Request, m: str = "", user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import finops as fo
    from .models import FinOp
    month = _month(m)
    rows = fo.settlements(db, month)
    a, b = fo.month_bounds(month)
    pays = db.scalars(select(FinOp).where(FinOp.deleted_at.is_(None), FinOp.category.in_(_fa().salary_names(db)), FinOp.doctor_id.is_not(None))
                      .order_by(FinOp.day.desc())).all()
    pays = [x for x in pays if (x.salary_month or fo_month_key(x.day)) == month]
    return render(request, "fin_settle.html", nav="finance", month=month, rows=rows, sal_cat=_fa().salary_name(db), months=_month_list(), start=fo.payroll_start(db),
                  boxes=_cashboxes(db), pays=pays, docs={d_.id: d_ for d_ in db.scalars(select(Doctor))}, today=date.today(),
                  total={k: round(sum(r[k] for r in rows), 2) for k in ("opening", "accrued", "paid", "closing")})


@app.post("/finance/settle/start")
def finance_settle_start(month: str = Form(...), user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import finops as fo
    fo.set_payroll_start(db, _month(month))
    audit(db, "fin.payroll_start", user=user, month=_month(month))
    return RedirectResponse(f"/finance/settle?m={_month(month)}", 303)


@app.get("/finance/balances", response_class=HTMLResponse)
def finance_balances(request: Request, q: str = "", user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import finops as fo
    rows = fo.balances(db, q)
    pos = round(sum(p.cb_balance for p in rows if p.cb_balance > 0), 2)
    neg = round(sum(p.cb_balance for p in rows if p.cb_balance < 0), 2)
    return render(request, "fin_balances.html", nav="finance", rows=rows, q=q, pos=pos, neg=neg,
                  synced=db.scalar(select(func.max(Patient.cb_synced_at))))


@app.get("/finance/bonuses", response_class=HTMLResponse)
def finance_bonuses(request: Request, df: str = "", dt_: str = "", q: str = "", user: User = Depends(require("finance")),
                    db: Session = Depends(db_dep)):
    from .services import finops as fo
    d_from, d_to = _period(df, dt_, default_days=90)
    from .services import loyalty as ly
    from .models import BonusTxn
    a_, b_ = datetime.combine(d_from, datetime.min.time()), datetime.combine(d_to, datetime.max.time())
    ltx = db.scalars(select(BonusTxn).where(BonusTxn.at >= a_, BonusTxn.at <= b_).order_by(BonusTxn.at.desc()).limit(1000)).all()
    lp = {p_.id: p_ for p_ in db.scalars(select(Patient).where(Patient.id.in_([t.patient_id for t in ltx] or [0])))}
    return render(request, "fin_bonuses.html", nav="finance", d_from=d_from, d_to=d_to, q=q, b=fo.bonuses(db, d_from, d_to, q),
                  ltx=ltx, lp=lp, KL=ly.KIND_LABELS, lyon=ly.settings(db).get("enabled"),
                  synced=db.scalar(select(func.max(Patient.cb_synced_at))))


@app.get("/finance/pricing", response_class=HTMLResponse)
def finance_pricing(request: Request, wh: str = "vitrina", brand: str = "", q: str = "", only: str = "",
                    user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import pricing as pr_svc
    rows, rt = _pricing_rows(db, wh, brand, q, only)
    brands = sorted({x for (x,) in db.execute(select(Product.brand).where(Product.brand.is_not(None)).distinct())}, key=str.lower)
    return render(request, "pricing.html", nav="finance", rows=rows[:600], total=len(rows), rt=rt, wh=wh, brand=brand, q=q,
                  only=only, brands=brands, WAREHOUSES=catalog.WAREHOUSES, CURRENCIES=pr_svc.CURRENCIES, MARGINS=pr_svc.MARGINS)


@app.post("/finance/rates")
async def finance_rates(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import pricing as pr_svc
    f = await request.form()
    back = str(f.get("back") or "/finance/pricing")
    if f.get("nbu"):
        try:
            nb = await asyncio.to_thread(pr_svc.nbu_rates)
            pr_svc.save_rates(db, nb.get("USD"), nb.get("EUR"), None, "НБУ")
        except Exception as e:  # noqa: BLE001
            log.warning("NBU rates: %s", e)
            return RedirectResponse(back + ("&" if "?" in back else "?") + "msg=Не+вдалося+отримати+курс+НБУ", 303)
    else:
        rnd = f.get("round")
        pr_svc.save_rates(db, _num(f.get("usd") or ""), _num(f.get("eur") or ""), int(rnd) if str(rnd).isdigit() else None)
    audit(db, "pricing.rates", user=user)
    return RedirectResponse(back if back.startswith("/finance") else "/finance/pricing", 303)


@app.post("/finance/pricing")
async def finance_pricing_save(request: Request, user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    from .services import pricing as pr_svc
    f = await request.form()
    ids = [int(x) for x in f.getlist("pid") if str(x).isdigit()]
    n = 0
    now = datetime.now()
    for p_ in db.scalars(select(Product).where(Product.id.in_(ids))):
        cur = f.get(f"cur_{p_.id}")
        cfx, rrp, plan = _num(f.get(f"cfx_{p_.id}") or ""), _num(f.get(f"rrp_{p_.id}") or ""), _num(f.get(f"plan_{p_.id}") or "")
        new = (cur if cur in pr_svc.CURRENCIES else p_.cost_currency, cfx, rrp, plan)
        if new != (p_.cost_currency, p_.cost_fx, p_.rrp_fx, p_.planned_price):
            if plan != p_.planned_price:
                p_.planned_at = now if plan else None
            p_.cost_currency, p_.cost_fx, p_.rrp_fx, p_.planned_price = new
            n += 1
    db.commit()
    audit(db, "pricing.save", user=user, changed=n)
    back = str(f.get("back") or "/finance/pricing")
    sep = "&" if "?" in back else "?"
    return RedirectResponse((back if back.startswith("/finance") else "/finance/pricing") + f"{sep}msg=Збережено:+{n}", 303)


@app.get("/finance/pricing.xlsx")
def finance_pricing_xlsx(wh: str = "vitrina", brand: str = "", q: str = "", only: str = "",
                         user: User = Depends(require("finance")), db: Session = Depends(db_dep)):
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from fastapi.responses import StreamingResponse
    from .services import pricing as pr_svc
    rows, rt = _pricing_rows(db, wh, brand, q, only)
    wb = Workbook(); ws = wb.active; ws.title = "Планування цін"
    ws.append([f"Курс USD {rt['USD']}  ·  EUR {rt['EUR']}  ·  {rt['source']} {rt['updated'] or ''}"])
    head = ["Бренд", "Товар", "Артикул", "Склад", "Валюта", "Закупівля (вал.)", "Закупівля, грн", "РРЦ постачальника (вал.)",
            "Ціна в програмі", "Валовий прибуток, грн", "Маржа, %", "Ціна при курсі", "Маржа при курсі, %"] + \
           [f"Ціна при маржі {m}%" for m in pr_svc.MARGINS] + ["Рекомендована ціна", "Різниця з поточною", "Націнка, %", "Маржа, %"]
    ws.append(head)
    for c in ws[2]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="3B3050")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    rnd = rt["round"]
    for x in rows:
        p_ = x.product; cost = x.cost_uah; now_ = x.price_now; at = x.price_at_rate; plan = p_.planned_price
        ws.append([p_.brand, p_.name, p_.sku, catalog.WAREHOUSES.get(p_.warehouse or "vitrina"), x.cur if p_.cost_fx else "UAH",
                   p_.cost_fx, cost, p_.rrp_fx, now_, (now_ - cost) if now_ and cost is not None else None,
                   pr_svc.margin(now_, cost), pr_svc.round_price(at, rnd), pr_svc.margin(at, cost)] +
                  [pr_svc.round_price(pr_svc.price_for_margin(cost, m), rnd) for m in pr_svc.MARGINS] +
                  [plan, (plan - now_) if plan and now_ else None, pr_svc.markup(plan, cost), pr_svc.margin(plan, cost)])
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTU", [14, 46, 12, 12, 8] + [13] * 16):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C3"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename=pricing-{date.today()}.xlsx"})


@app.get("/book/{item_id}", response_class=HTMLResponse)
def book_page(request: Request, item_id: int, day: str | None = None, user: User = Depends(require("book")),
              db: Session = Depends(db_dep)):
    from .services import booking
    item = db.get(PlanItem, item_id)
    if not item or item.kind != "procedure" or not plan_visible(user, item.plan):
        raise HTTPException(404, "Процедуру не знайдено")
    ps = compute_plan(db, item.plan)
    d = date.fromisoformat(day) if day else max(date.today(), (ps.due_date.date() if ps.due_date else date.today()))
    svc, slots, error = None, [], None
    if not settings.cleverbox_token:
        error = "CleverBox не підключено (CLEVERBOX_TOKEN)"
    else:
        try:
            svc, slots = booking.get_slots(db, item, d)
        except Exception as e:      # noqa: BLE001
            error = str(e)
    return render(request, "book.html", item=item, ps=ps, day=d, svc=svc, slots=slots, error=error,
                  prev=d - timedelta(days=1), next=d + timedelta(days=1))


@app.post("/book/{item_id}")
def book_save(item_id: int, day: str = Form(...), slot: str = Form(...), comment: str = Form(""),
              user: User = Depends(require("book")), db: Session = Depends(db_dep)):
    from .services import booking
    item = db.get(PlanItem, item_id)
    if not item or not plan_visible(user, item.plan):
        raise HTTPException(404, "Процедуру не знайдено")
    time_, master = slot.split("|")
    try:
        booking.book(db, user, item, date.fromisoformat(day), time_, int(master), comment)
    except Exception as e:          # noqa: BLE001
        raise HTTPException(400, f"CleverBox не прийняв запис: {e}")
    return RedirectResponse(f"/patients/{item.plan.patient_id}#plan{item.plan_id}", 303)

# ---------------- Скрипти, повідомлення пацієнтам, розсилки ----------------

from .models import ScriptTemplate, ProcedureNote, Campaign, CampaignRecipient, MessageLog
from .services import scripts as scripts_svc, campaigns as camp_svc
from .integrations import messaging


def _channels_ctx():
    return {"tg_on": messaging.telegram_available(), "ts_on": messaging.turbosms_available(),
            "viber_on": bool(settings.turbosms_sender_viber), "sms_on": bool(settings.turbosms_sender_sms)}


@app.post("/tasks/{tid}/send")
def task_send(tid: int, text: str = Form(...), channel: str = Form("auto"), close: str = Form(""),
              user: User = Depends(require("message")), db: Session = Depends(db_dep)):
    t = db.get(Task, tid)
    if not t or not t.patient_id or t not in db.scalars(tasks_svc.visible_query(user).where(Task.id == tid)).all():
        raise HTTPException(403, "Немає доступу до задачі")
    p = _patient_or_403(db, user, t.patient_id)
    res = camp_svc.send_single(db, p, text.strip(), channel, user, task_id=t.id)
    if not res.ok:
        raise HTTPException(400, f"Не надіслано ({res.channel}): {res.error}")
    note = f"Надіслано повідомлення ({res.channel}) {datetime.now():%d.%m %H:%M}"
    t.result = (t.result + "\n" if t.result else "") + note
    if close:
        t.status, t.done_at = "done", datetime.utcnow()
    db.commit()
    audit(db, "message.send", "patient", p.id, user=user, channel=res.channel, task=t.id)
    return RedirectResponse("/tasks", 303)


@app.post("/patients/{pid}/message")
def patient_message(pid: int, text: str = Form(...), channel: str = Form("auto"),
                    user: User = Depends(require("message")), db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    res = camp_svc.send_single(db, p, text.strip(), channel, user)
    if not res.ok:
        raise HTTPException(400, f"Не надіслано ({res.channel}): {res.error}")
    audit(db, "message.send", "patient", p.id, user=user, channel=res.channel)
    return RedirectResponse(f"/patients/{pid}#contact", 303)


@app.post("/patients/{pid}/optout")
def patient_optout(pid: int, value: str = Form(...), user: User = Depends(require("patient_edit")),
                   db: Session = Depends(db_dep)):
    p = _patient_or_403(db, user, pid)
    p.marketing_opt_out = value == "1"
    p.opt_out_at = datetime.now() if p.marketing_opt_out else None
    db.commit()
    audit(db, "patient.optout" if p.marketing_opt_out else "patient.optin", "patient", p.id, user=user)
    return RedirectResponse(f"/patients/{pid}#contact", 303)


@app.get("/scripts", response_class=HTMLResponse)
def scripts_page(request: Request, pid: int | None = None, user: User = Depends(require("scripts_edit")),
                 db: Session = Depends(db_dep)):
    tpls = db.scalars(select(ScriptTemplate).order_by(ScriptTemplate.kind, ScriptTemplate.channel)).all()
    notes = db.scalars(select(ProcedureNote).order_by(ProcedureNote.id)).all()
    sample = db.get(Patient, pid) if pid else next((p for p in db.scalars(select(Patient).join(BeautyPlan)).all()), None)
    previews = {}
    if sample:
        for kind in ("not_booked", "care_ending", "follow_up"):
            for sc in scripts_svc.for_patient(db, sample, kind, user.name):
                previews[sc.title] = sc.text
    return render(request, "scripts.html", tpls=tpls, notes=notes, sample=sample, previews=previews,
                  PLACEHOLDERS=scripts_svc.PLACEHOLDERS, KINDS=tasks_svc.KINDS)


@app.post("/scripts/templates")
async def scripts_tpl_save(request: Request, user: User = Depends(require("scripts_edit")), db: Session = Depends(db_dep)):
    f = await request.form()
    t = db.get(ScriptTemplate, int(f["id"])) if str(f.get("id", "")).isdigit() else ScriptTemplate()
    t.kind, t.channel = f.get("kind", "follow_up"), f.get("channel", "call")
    t.title, t.text = f.get("title", "").strip() or "Скрипт", f.get("text", "").strip()
    t.active = bool(f.get("active"))
    t.updated_at = datetime.utcnow()
    db.add(t)
    db.commit()
    audit(db, "script.update", user=user, template=t.title)
    return RedirectResponse("/scripts", 303)


@app.post("/scripts/notes")
async def scripts_note_save(request: Request, user: User = Depends(require("scripts_edit")), db: Session = Depends(db_dep)):
    f = await request.form()
    n = db.get(ProcedureNote, int(f["id"])) if str(f.get("id", "")).isdigit() else ProcedureNote()
    if f.get("delete") and n.id:
        err = _delpw(db, user, str(f.get("del_password") or ""), "script_note", "procedure_note", n.id)
        if err:
            return RedirectResponse(_with_err("/scripts", err), 303)
        db.delete(n)
    else:
        n.keyword, n.why, n.benefit = f.get("keyword", "").strip(), f.get("why", "").strip(), f.get("benefit", "").strip() or None
        if n.keyword and n.why:
            db.add(n)
    db.commit()
    audit(db, "script.note", user=user, keyword=f.get("keyword"))
    return RedirectResponse("/scripts#notes", 303)


@app.get("/campaigns", response_class=HTMLResponse)
def campaigns_page(request: Request, user: User = Depends(require("campaigns")), db: Session = Depends(db_dep)):
    cs = db.scalars(select(Campaign).order_by(Campaign.created_at.desc()).limit(100)).all()
    st = {c.id: camp_svc.stats(db, c) for c in cs if c.status in ("sending", "done", "scheduled")}
    balance = messaging.turbosms_balance() if messaging.turbosms_available() else None
    return render(request, "campaigns.html", cs=cs, st=st, balance=balance, SEGMENTS=clients_svc.SEGMENTS, settings=settings,
                  IDEAS=camp_svc.SEGMENT_IDEAS, doctors=db.scalars(select(Doctor).order_by(Doctor.name)).all(),
                  **_channels_ctx())


def _campaign_form(c: Campaign, f) -> None:
    c.name = (f.get("name") or "Розсилка").strip()
    c.segment = f.get("segment") if f.get("segment") in clients_svc.SEGMENTS else "sleeping"
    c.text = (f.get("text") or "").strip()
    c.channels = [ch for ch in ("telegram", "viber", "sms") if f.get(f"ch_{ch}")] or ["telegram"]
    c.filters = {k: f.get(k) for k in ("doctor_id", "source", "min_ltv", "max_ltv") if f.get(k)}
    if str(f.get("attribution_days", "")).isdigit():
        c.attribution_days = int(f["attribution_days"])


@app.post("/campaigns")
async def campaign_create(request: Request, user: User = Depends(require("campaigns")), db: Session = Depends(db_dep)):
    f = await request.form()
    c = Campaign(created_by=user.id, status="draft")
    _campaign_form(c, f)
    if not c.text:
        raise HTTPException(400, "Текст розсилки порожній")
    db.add(c)
    db.commit()
    audit(db, "campaign.create", "campaign", c.id, user=user, name=c.name, segment=c.segment)
    return RedirectResponse(f"/campaigns/{c.id}", 303)


@app.get("/campaigns/{cid}", response_class=HTMLResponse)
def campaign_page(request: Request, cid: int, user: User = Depends(require("campaigns")), db: Session = Depends(db_dep)):
    c = db.get(Campaign, cid) or (_ for _ in ()).throw(HTTPException(404, "Розсилку не знайдено"))
    aud = camp_svc.build_audience(db, c) if c.status == "draft" else None
    samples = [(p, ch, camp_svc.rendered_text(c, p)) for p, ch in (aud.rows[:5] if aud else [])]
    st = camp_svc.stats(db, c) if c.status != "draft" else None
    recips = c.recipients[:500] if c.status != "draft" else []
    return render(request, "campaign.html", c=c, aud=aud, samples=samples, st=st, recips=recips, settings=settings,
                  SEGMENTS=clients_svc.SEGMENTS, doctors=db.scalars(select(Doctor).order_by(Doctor.name)).all(),
                  parts=messaging.sms_parts(c.text + "\n" + settings.marketing_footer), quiet=messaging.is_quiet_time(),
                  **_channels_ctx())


@app.post("/campaigns/{cid}")
async def campaign_update(request: Request, cid: int, user: User = Depends(require("campaigns")),
                          db: Session = Depends(db_dep)):
    c = db.get(Campaign, cid)
    if not c or c.status != "draft":
        raise HTTPException(400, "Змінювати можна лише чернетку")
    _campaign_form(c, await request.form())
    db.commit()
    return RedirectResponse(f"/campaigns/{cid}", 303)


@app.post("/campaigns/{cid}/test")
def campaign_test(cid: int, to: str = Form(...), user: User = Depends(require("campaigns")), db: Session = Depends(db_dep)):
    """Тест собі: на свій телефон (Viber/SMS) або Telegram chat ID у боті пацієнтів."""
    c = db.get(Campaign, cid)
    to = to.strip()
    text = "[ТЕСТ] " + camp_svc.rendered_text(c, Patient(full_name=user.name))
    if messaging.normalize_phone(to):          # номер телефону → Viber/SMS через TurboSMS
        res = messaging.send_turbosms(to, text + "\n" + settings.marketing_footer,
                                      [ch for ch in ("viber", "sms") if ch in c.channels] or ["sms"])
    else:                                      # інакше — Telegram chat ID у боті пацієнтів
        res = messaging.send_telegram(to, text)
    audit(db, "campaign.test", "campaign", c.id, user=user, ok=res.ok, channel=res.channel)
    if not res.ok:
        raise HTTPException(400, f"Тест не надіслано ({res.channel}): {res.error}")
    return RedirectResponse(f"/campaigns/{cid}", 303)


@app.post("/campaigns/{cid}/launch")
def campaign_launch(cid: int, when: str = Form(""), confirm: str = Form(""),
                    user: User = Depends(require("campaigns_send")), db: Session = Depends(db_dep)):
    c = db.get(Campaign, cid)
    if not c or c.status != "draft":
        raise HTTPException(400, "Запустити можна лише чернетку")
    if confirm != "1":
        raise HTTPException(400, "Підтвердіть запуск")
    if when:
        c.scheduled_at = datetime.fromisoformat(when)
        c.status = "scheduled"
        db.commit()
        audit(db, "campaign.schedule", "campaign", c.id, user=user, at=when)
        return RedirectResponse(f"/campaigns/{cid}", 303)
    n = camp_svc.start(db, c)
    audit(db, "campaign.launch", "campaign", c.id, user=user, recipients=n)
    if messaging.is_quiet_time():
        c.status = "scheduled"
        db.commit()
    else:
        jobs.start(f"Розсилка #{c.id}", camp_svc.deliver, campaign_id=c.id)
    return RedirectResponse(f"/campaigns/{cid}", 303)


@app.post("/campaigns/{cid}/cancel")
def campaign_cancel(cid: int, user: User = Depends(require("campaigns_send")), db: Session = Depends(db_dep)):
    c = db.get(Campaign, cid)
    if c and c.status in ("draft", "scheduled", "sending"):
        c.status = "cancelled"
        for r in c.recipients:
            if r.status == "queued":
                r.status, r.skip_reason = "skipped", "розсилку скасовано"
        db.commit()
        audit(db, "campaign.cancel", "campaign", c.id, user=user)
    return RedirectResponse(f"/campaigns/{cid}", 303)


@app.post("/campaigns/{cid}/copy")
def campaign_copy(cid: int, user: User = Depends(require("campaigns")), db: Session = Depends(db_dep)):
    c = db.get(Campaign, cid)
    n = Campaign(name=c.name + " (копія)", segment=c.segment, filters=c.filters, channels=c.channels, text=c.text,
                 attribution_days=c.attribution_days, created_by=user.id)
    db.add(n)
    db.commit()
    return RedirectResponse(f"/campaigns/{n.id}", 303)


@app.get("/patients")
def patients_redirect(request: Request):
    return RedirectResponse("/clients" + (("?" + request.url.query) if request.url.query else ""), 303)


_VERSION_FILE = Path(__file__).with_name("VERSION")
APP_VERSION = _VERSION_FILE.read_text().strip() if _VERSION_FILE.exists() else "dev"


@app.get("/health")
def health():
    return {"ok": True, "version": APP_VERSION, "last_sync": str(state["last_sync"][0]) if state["last_sync"] else None}


# ---------------- Спеціалісти ----------------
from .services import specialists as spec_svc   # noqa: E402

COLOR_RE = __import__("re").compile(r"^#[0-9a-fA-F]{6}$")


@app.get("/specialists", response_class=HTMLResponse)
def specialists_page(request: Request, show: str = "active", user: User = Depends(require("specialists")),
                     db: Session = Depends(db_dep)):
    cs = spec_svc.cards(db)
    active = [c for c in cs if c.doctor.active]
    fired = [c for c in cs if not c.doctor.active]
    return render(request, "specialists.html", cards=fired if show == "fired" else active, show=show,
                  n_active=len(active), n_fired=len(fired), staff_fields=spec_svc.cb_staff_fields(db),
                  palette=spec_svc.PALETTE)


@app.post("/specialists")
def specialist_create(name: str = Form(...), position: str = Form(""), user: User = Depends(require("specialists_edit")),
                      db: Session = Depends(db_dep)):
    name = name.strip()
    if not name:
        raise HTTPException(400, "Вкажіть імʼя")
    if db.scalar(select(Doctor).where(Doctor.name == name)):
        raise HTTPException(400, "Спеціаліст з таким імʼям уже є")
    d = Doctor(name=name, position=position.strip() or None, source="manual", active=True, in_calendar=True)
    db.add(d)
    db.commit()
    audit(db, "specialist.create", "doctor", d.id, user=user)
    return RedirectResponse(f"/specialists/{d.id}", 303)


# ---------------- Графік спеціалістів ----------------
def _places(db: Session) -> list[str]:
    from .services import clinic as cl
    return [p["name"] for l_ in cl.get(db)["locations"] if l_.get("active") for p in l_.get("places", []) if p.get("active", True)]


@app.get("/specialists/schedule", response_class=HTMLResponse)
def schedule_page(request: Request, start: str = "", span: int = 14, doc: int = 0, user: User = Depends(require("specialists")),
                  db: Session = Depends(db_dep)):
    from .services import schedule as sch
    today = date.today()
    s = _day(start) or (today - timedelta(days=today.weekday()))
    span = span if span in (7, 14, 31) else 14
    days = [s + timedelta(days=i) for i in range(span)]
    docs = [d_ for d_ in sch.doctors(db) if not doc or d_.id == doc]
    sh = sch.shifts(db, days[0], days[-1])
    return render(request, "schedule.html", nav="specialists", days=days, docs=docs, all_docs=sch.doctors(db), sh=sh, span=span,
                  start=s, doc=doc, today=today, totals=sch.totals(sh, [d_.id for d_ in docs]), places=_places(db),
                  timedelta=timedelta, WD=["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Нд"], color=spec_svc.color)


def _sched_back(back: str, msg: str = "", err: bool = False) -> RedirectResponse:
    back = back if back.startswith("/specialists/schedule") else "/specialists/schedule"
    return RedirectResponse(_with_msg(back, msg, err) if msg else back, 303)


@app.post("/specialists/schedule/cell")
def schedule_cell(doctor_id: int = Form(...), day: str = Form(...), start: str = Form(""), end: str = Form(""), place: str = Form(""),
                  note: str = Form(""), off: str = Form(""), back: str = Form(""), user: User = Depends(require("schedule")),
                  db: Session = Depends(db_dep)):
    from .services import schedule as sch
    d_ = _day(day)
    if not d_:
        return _sched_back(back, "Невірна дата", True)
    try:
        sch.set_shift(db, doctor_id, d_, None if off else start, None if off else end, place.strip()[:120] or None,
                      note.strip()[:200] or None, user.id)
    except ValueError as e:
        return _sched_back(back, str(e), True)
    audit(db, "schedule.cell", "doctor", doctor_id, user=user, day=day, start=None if off else start, end=None if off else end, place=place or None)
    return _sched_back(back)


@app.post("/specialists/schedule/fill")
async def schedule_fill(request: Request, user: User = Depends(require("schedule")), db: Session = Depends(db_dep)):
    from .services import schedule as sch
    f = await request.form()
    back = str(f.get("back") or "")
    d_from, d_to = _day(str(f.get("d_from") or "")), _day(str(f.get("d_to") or ""))
    if not d_from or not d_to or d_to < d_from or (d_to - d_from).days > 370:
        return _sched_back(back, "Вкажіть період (до року)", True)
    ids = [int(x) for x in f.getlist("doctor_id") if str(x).isdigit()]
    if not ids:
        return _sched_back(back, "Оберіть спеціаліста", True)
    wd = {int(x) for x in f.getlist("wd") if str(x).isdigit()}
    n = 0
    try:
        for did in ids:
            n += sch.fill(db, did, d_from, d_to, wd, str(f.get("start") or ""), str(f.get("end") or ""),
                          str(f.get("place") or "").strip()[:120] or None, str(f.get("mode") or "weekdays"), user.id,
                          overwrite=bool(f.get("overwrite")))
    except ValueError as e:
        return _sched_back(back, str(e), True)
    audit(db, "schedule.fill", user=user, doctors=ids, d_from=str(d_from), d_to=str(d_to), mode=f.get("mode"), shifts=n)
    return _sched_back(back, f"Заповнено змін: {n}")


@app.post("/specialists/schedule/copy")
def schedule_copy(src: str = Form(...), weeks: int = Form(1), back: str = Form(""), user: User = Depends(require("schedule")),
                  db: Session = Depends(db_dep)):
    from .services import schedule as sch
    s = _day(src)
    if not s:
        return _sched_back(back, "Невірна дата", True)
    s = s - timedelta(days=s.weekday())
    n = sum(sch.copy_week(db, s, s + timedelta(days=7 * k), user.id) for k in range(1, max(1, min(weeks, 12)) + 1))
    audit(db, "schedule.copy", user=user, src=str(s), weeks=weeks, shifts=n)
    return _sched_back(back, f"Скопійовано змін: {n}")


@app.post("/specialists/schedule/clear")
def schedule_clear(doctor_id: str = Form(""), d_from: str = Form(...), d_to: str = Form(...), del_password: str = Form(""),
                   back: str = Form(""), user: User = Depends(require("schedule")), db: Session = Depends(db_dep)):
    from .services import schedule as sch
    err = _delpw(db, user, del_password, "schedule", "doctor", int(doctor_id) if doctor_id.isdigit() else None)
    if err:
        b = back if back.startswith("/specialists/schedule") else "/specialists/schedule"
        return RedirectResponse(_with_err(b, err), 303)
    a, b_ = _day(d_from), _day(d_to)
    n = sch.clear(db, int(doctor_id) if doctor_id.isdigit() else None, a, b_) if a and b_ else 0
    audit(db, "schedule.clear", user=user, doctor=doctor_id or "all", d_from=d_from, d_to=d_to, shifts=n)
    return _sched_back(back, f"Прибрано змін: {n}")


@app.get("/specialists/{did:int}", response_class=HTMLResponse)
def specialist_page(did: int, request: Request, user: User = Depends(require("specialists")),
                    db: Session = Depends(db_dep)):
    d = db.get(Doctor, did)
    if not d:
        raise HTTPException(404)
    card = spec_svc.cards(db, [did])[0]
    now = datetime.now()
    idx = spec_svc.DoctorIndex(spec_svc.all_doctors(db))
    upcoming = [v for v in db.scalars(select(Visit).where(Visit.status == "booked", Visit.starts_at > now,
                                                          Visit.starts_at <= now + timedelta(days=14),
                                                          Visit.master_name.is_not(None)).order_by(Visit.starts_at))
                if (m := idx.find(v.master_name)) and m.id == did][:30]
    patients = db.scalars(select(Patient).where(Patient.doctor_id == did).order_by(Patient.full_name).limit(200)).all()
    from .services import positions as _ps
    return render(request, "specialist.html", d=d, c=card, upcoming=upcoming, patients=patients,
                  palette=spec_svc.PALETTE, positions=_ps.names(db))


@app.post("/specialists/{did:int}")
def specialist_save(did: int, name: str = Form(...), position: str = Form(""), phone: str = Form(""),
                    email: str = Form(""), color: str = Form(""), bio: str = Form(""), work_days: str = Form(""),
                    telegram_chat_id: str = Form(""), photo_url: str = Form(""), status: str = Form("active"),
                    in_calendar: str = Form(""), sale_initiator: str = Form("1"), user: User = Depends(require("specialists_edit")),
                    db: Session = Depends(db_dep)):
    d = db.get(Doctor, did)
    if not d:
        raise HTTPException(404)
    name = name.strip()
    if name and name != d.name and not db.scalar(select(Doctor).where(Doctor.name == name)):
        d.name = name
    d.position, d.phone, d.email = position.strip() or None, phone.strip() or None, email.strip() or None
    d.bio, d.work_days = bio.strip() or None, work_days.strip() or None
    d.telegram_chat_id = telegram_chat_id.strip() or None
    photo_url = photo_url.strip()
    d.photo_url = photo_url if photo_url.startswith("https://") else (None if not photo_url else d.photo_url)
    d.color = color if COLOR_RE.match(color or "") else None
    want_active = status != "fired"
    if want_active != bool(d.active):
        d.active, d.status_locked = want_active, True
        d.fired_at = None if want_active else date.today()
    d.in_calendar = in_calendar == "1"
    d.sale_initiator = sale_initiator == "1"
    db.commit()
    audit(db, "specialist.update", "doctor", d.id, user=user)
    return RedirectResponse(f"/specialists/{did}?saved=1", 303)


@app.post("/specialists/sync")
async def specialists_sync(user: User = Depends(require("sync")), db: Session = Depends(db_dep)):
    """Оновити спеціалістів з CleverBox зараз (не чекаючи щоденного синку)."""
    if not settings.cleverbox_token:
        return RedirectResponse("/specialists?msg=Спершу+збережіть+токен+CleverBox+у+розділі+Дані", 303)
    from .integrations.cleverbox import CleverBox, sync_staff
    try:
        n = await asyncio.to_thread(lambda: sync_staff(db, CleverBox()))
        audit(db, "specialists.sync", user=user, linked=n)
        return RedirectResponse(f"/specialists?msg=Оновлено+з+CleverBox", 303)
    except Exception as e:  # noqa: BLE001
        log.warning("specialists sync: %s", type(e).__name__)
        return RedirectResponse("/specialists?msg=CleverBox+не+відповів,+спробуйте+пізніше", 303)


@app.post("/specialists/{did:int}/calendar")
def specialist_calendar_toggle(did: int, request: Request, enabled: str = Form(""),
                               user: User = Depends(require("specialists_edit")), db: Session = Depends(db_dep)):
    """Галочка «Показувати в календарі» на картці спеціаліста."""
    d = db.get(Doctor, did)
    if not d:
        raise HTTPException(404)
    d.in_calendar = enabled == "1"
    db.commit()
    audit(db, "specialist.calendar", "doctor", d.id, user=user, in_calendar=d.in_calendar)
    if "application/json" in request.headers.get("accept", ""):
        return {"ok": True, "in_calendar": d.in_calendar}
    return RedirectResponse("/specialists", 303)


@app.post("/specialists/{did:int}/unlock")
def specialist_unlock(did: int, user: User = Depends(require("specialists_edit")), db: Session = Depends(db_dep)):
    d = db.get(Doctor, did)
    if d:
        d.status_locked = False
        db.commit()
        audit(db, "specialist.status_auto", "doctor", d.id, user=user)
    return RedirectResponse(f"/specialists/{did}", 303)


# ---------------- Вебхуки CleverBox (події в реальному часі) ----------------
from .integrations import cb_webhooks   # noqa: E402


@app.post("/hooks/cleverbox/{secret}")
async def cleverbox_hook(secret: str, request: Request, db: Session = Depends(db_dep)):
    if not cb_webhooks.check_secret(db, secret):
        raise HTTPException(404)
    raw = await request.body()
    if len(raw) > 512_000:
        raise HTTPException(413)
    try:
        payload = json.loads(raw or b"{}")
    except ValueError:
        form = await request.form()
        payload = {k: v for k, v in form.items()}
        if "data" in payload and isinstance(payload["data"], str):
            try:
                payload["data"] = json.loads(payload["data"])
            except ValueError:
                pass
    ev = await asyncio.to_thread(cb_webhooks.handle, payload)
    return {"ok": True, "event": ev["event"]}


@app.post("/settings/cleverbox-webhooks")
async def cleverbox_webhooks_subscribe(user: User = Depends(require("migrate")), db: Session = Depends(db_dep)):
    if not settings.cleverbox_token:
        return RedirectResponse("/data?cb_notice=Спершу+збережіть+токен#migrate", 303)
    if not settings.public_url.startswith("https://"):
        return RedirectResponse("/data?cb_notice=PUBLIC_URL+у+.env+має+бути+https://crm.molodist.com.ua#migrate", 303)
    res = await asyncio.to_thread(cb_webhooks.subscribe, db)
    audit(db, "cleverbox.webhooks.subscribe", user=user, ok=sum(r["ok"] for r in res))
    ok = all(r["ok"] for r in res)
    msg = "Вебхуки+CleverBox+підключено" if ok else "CleverBox+відхилив+підписку+—+деталі+нижче"
    return RedirectResponse(f"/data?cb_notice={msg}#webhooks", 303)


# ---------------- Бюті-плани: дошка ----------------
PLAN_COLUMNS = [("new", "Нові", "ще не стартували, без запису"),
                ("need", "Потребують запису", "почали, але наступна процедура не записана"),
                ("booked", "Записані", "є запис наперед"),
                ("done", "Завершені", "усі процедури виконано")]


def plan_column(ps) -> str:
    if ps.plan.status == "completed" or ps.is_complete:
        return "done"
    if ps.is_booked:
        return "booked"
    if ps.realized_sum > 0 or ps.proc_done > 0:
        return "need"
    return "new"


@app.get("/plans", response_class=HTMLResponse)
def plans_board(request: Request, doc: str = "", q: str = "", plan: int | None = None,
                user: User = Depends(current_user), db: Session = Depends(db_dep)):
    from .services import specialists as spec_svc
    from .services.kpi import all_statuses
    doc_id = int(doc) if doc.isdigit() else None
    statuses = [ps for ps in all_statuses(db, doctor_id=doctor_scope(user)) if ps.plan.status != "cancelled"]
    if doc_id:
        statuses = [ps for ps in statuses if ps.plan.doctor_id == doc_id]
    if q:
        words = set(norm(q).split())
        statuses = [ps for ps in statuses if words <= set(norm(ps.plan.patient.full_name + " " + (ps.plan.patient.phone or "")).split())
                    or (ps.plan.patient.phone and q.strip() in ps.plan.patient.phone)]
    cols = {k: [] for k, *_ in PLAN_COLUMNS}
    for ps in statuses:
        cols[plan_column(ps)].append(ps)
    far = datetime(2100, 1, 1)
    cols["new"].sort(key=lambda ps: ps.plan.session_at or far, reverse=True)
    cols["need"].sort(key=lambda ps: ({"red": 0, "yellow": 1}.get(ps.flag, 2), ps.due_date or far))
    cols["booked"].sort(key=lambda ps: ps.next_booking or far)
    cols["done"].sort(key=lambda ps: ps.last_activity or datetime.min, reverse=True)
    sel = next((ps for ps in statuses if ps.plan.id == plan), None) if plan else None
    doctors = [d for d in spec_svc.all_doctors(db) if any(ps.plan.doctor_id == d.id for ps in statuses) or d.active]
    return render(request, "plans.html", cols=cols, COLS=PLAN_COLUMNS, sel=sel, doc=doc_id, q=q, total=len(statuses),
                  doctors=doctors, color=spec_svc.color, initials=spec_svc.initials, now=datetime.now(),
                  sums={k: sum(ps.planned_sum for ps in v) for k, v in cols.items()})
